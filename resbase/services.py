"""Bazaga yig'ish (narxlangan smetadan kod + narx), median narxlar va statistika.

Faqat ishonchli qatorlar olinadi: holati EXACT (RES bilan aniq moslangan) yoki smetachi qo'lda qo'ygan narx
(``manual`` belgisi bilan, medianaga kirmaydi). Shubhali moslik, «ТОПИЛМАДИ», nol narx va kodsiz qatorlar olinmaydi."""
import logging
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from statistics import median

from django.core.cache import cache
from django.db import transaction
from django.utils import timezone

from core.normalize import norm_code, norm_name, unit_key  # noqa: F401 — panel ham shu yerdan oladi
from jobs.models import Batch, Job, Project

from .models import Observation, Resource, ResourceName, Source, counted_q
from .regions import guess_region

log = logging.getLogger(__name__)

_CACHE_VER = "resbase:ver"
SKIP_LABELS = {
    "no_code": "кодсиз (маҳаллий материал, иш сарлавҳаси)",
    "not_exact": "аниқ мослашмаган (тахминий, топилмади)",
    "zero_price": "нарх йўқ ёки 0",
    "duplicate": "шу сметада такрор",
}


def bump():
    """Bazada nimadir o'zgardi — median keshi yangilanadi."""
    try:
        cache.incr(_CACHE_VER)
    except ValueError:
        cache.set(_CACHE_VER, 2, None)


def _ver():
    return cache.get_or_set(_CACHE_VER, 1, None)


# ===================== ish (papka / obyekt / smeta) =====================
def work_of(kind, work_id):
    model = {"folder": Project, "object": Batch, "smeta": Job}[kind]
    return model.objects.get(pk=work_id)


def jobs_of(kind, work) -> list:
    if kind == "folder":
        qs = Job.objects.filter(batch__project=work)
    elif kind == "object":
        qs = Job.objects.filter(batch=work)
    else:
        qs = Job.objects.filter(pk=work.pk)
    return list(qs.filter(status=Job.Status.DONE))


def title_of(kind, work) -> str:
    return work.local_name if kind == "smeta" else work.name


def region_suggestion(kind, work) -> tuple:
    """Viloyat taxmini: ish nomi, papkalar va bir nechta smetaning sarlavha/svodka matni."""
    texts = [title_of(kind, work)]
    if kind == "folder":
        texts += [f.get("path", "") for f in work.files[:400]]
    jobs = jobs_of(kind, work)[:4]
    for j in jobs:
        texts += [j.local_name, j.batch_label]
        texts += _head_texts(j)
    return guess_region(texts)


def _head_texts(job, rows=14) -> list:
    """Lokal faylning hamma listlari boshidagi matn (obyekt nomi, «… туманида жойлашган»)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(job.abs(job.local_path), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — fayl o'chirilgan bo'lishi mumkin
        return []
    out = []
    try:
        for ws in wb.worksheets[:6]:
            for row in ws.iter_rows(max_row=rows, values_only=True):
                out += [str(v) for v in row if isinstance(v, str)]
    finally:
        wb.close()
    return out


# ===================== yig'ish =====================
def _num(v):
    try:
        d = Decimal(str(v))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return d if d.is_finite() else None


def rows_from_job(job) -> tuple:
    """(olinadigan [{code, name, unit, price, manual, row, category}], tashlab ketilganlar Counter, jami resurs qatori)."""
    taken, skipped, seen, total = [], Counter(), set(), 0
    for r in (job.result or {}).get("rows") or []:
        status = r.get("status")
        if status in ("WORK_HEADER",) or not r.get("unit"):
            continue
        total += 1
        code = norm_code(r.get("code"))
        if not code:
            skipped["no_code"] += 1
            continue
        manual = bool(r.get("manual")) or status == "MANUAL"
        if status != "EXACT" and not manual:
            skipped["not_exact"] += 1
            continue
        price = _num(r.get("g"))
        if price is None or price <= 0:
            skipped["zero_price"] += 1
            continue
        if code in seen:
            skipped["duplicate"] += 1
            continue
        seen.add(code)
        taken.append({"code": code, "name": str(r.get("name") or "").strip()[:500], "unit": str(r.get("unit")).strip()[:40],
                      "price": price.quantize(Decimal("0.0001")), "manual": manual, "row": int(r.get("row") or 0),
                      "category": (r.get("default_category") or r.get("category") or "")[:1]})
    return taken, skipped, total


@transaction.atomic
def harvest(source: Source) -> Source:
    """Manbaning narxlarini (qayta) yig'adi: eski kuzatuvlar o'chiriladi, smetalardan yangidan olinadi."""
    work = work_of(source.kind, source.work_id)
    touched = set(source.observations.values_list("resource_id", flat=True))
    source.observations.all().delete()
    skipped, rows, taken_n, jobs = Counter(), 0, 0, jobs_of(source.kind, work)
    resources = {r.code: r for r in Resource.objects.all()}
    names = Counter()
    for job in jobs:
        taken, sk, total = rows_from_job(job)
        skipped.update(sk)
        rows += total
        for t in taken:
            res = resources.get(t["code"])
            if res is None:
                res = resources[t["code"]] = Resource.objects.create(
                    code=t["code"], name=t["name"], unit=t["unit"], unit_key=unit_key(t["unit"]), category=t["category"])
            elif not res.category and t["category"] and not res.edited:
                res.category = t["category"]
                res.save(update_fields=["category", "updated"])
            Observation.objects.create(source=source, resource=res, job_id=job.id, row=t["row"], name=t["name"],
                                       unit=t["unit"], price=t["price"], manual=t["manual"])
            names[(res.pk, t["name"], t["unit"])] += 1
            taken_n += 1
    for (rid, name, unit), n in names.items():
        obj, created = ResourceName.objects.get_or_create(resource_id=rid, name=name, unit=unit, defaults={"seen": n})
        if not created:
            obj.seen += n
            obj.save(update_fields=["seen"])
    source.title = title_of(source.kind, work)[:500]
    source.owner_id = work.owner_id
    source.stats = {"jobs": len(jobs), "rows": rows, "taken": taken_n, "skipped": dict(skipped)}
    source.harvested = timezone.now()
    source.save()
    recheck(touched | set(source.observations.values_list("resource_id", flat=True)))
    transaction.on_commit(bump)
    log.info("resbase: %s «%s» — %s ta narx (%s qator)", source.kind, source.title, taken_n, rows)
    return source


def add_source(kind, work_id, *, region="", year=None, quarter=None, user=None) -> Source:
    work = work_of(kind, work_id)
    hint_region, hint_word = region_suggestion(kind, work) if not region else (region, "")
    source, _ = Source.objects.update_or_create(kind=kind, work_id=work.pk, defaults={
        "title": title_of(kind, work)[:500], "owner_id": work.owner_id, "region": region or hint_region,
        "region_hint": hint_word, "year": year, "quarter": quarter, "added_by": user})
    return harvest(source)


def delete_source(source: Source):
    touched = set(source.observations.values_list("resource_id", flat=True))
    source.delete()
    recheck(touched)
    Resource.objects.filter(observations__isnull=True).delete()  # hech qayerda uchramaydigan kodlar qolmasin
    bump()


def wipe_all():
    Observation.objects.all().delete()
    Source.objects.all().delete()
    Resource.objects.all().delete()
    bump()


# ===================== narxlar =====================
def price_table(region, year, quarter) -> dict:
    """{kod: (mediana Decimal, kuzatuvlar soni)} — faqat shu viloyat va davr; qo'lda va istisno qilinganlar kirmaydi."""
    if not (region and year and quarter):
        return {}
    key = f"resbase:prices:{_ver()}:{region}:{year}:{quarter}"
    table = cache.get(key)
    if table is None:
        prices = defaultdict(list)
        qs = Observation.objects.filter(counted_q(), source__region=region, source__year=year,
                                        source__quarter=quarter).values_list("resource__code", "price")
        for code, price in qs:
            prices[code].append(price)
        table = {c: (Decimal(median(ps)).quantize(Decimal("0.01")), len(ps)) for c, ps in prices.items()}
        cache.set(key, table, 3600)
    return table


def resource_medians(resource) -> list:
    """Resurs sahifasi uchun: [{region, year, quarter, median, n, manual, min, max}] — viloyat va davr bo'yicha."""
    groups = defaultdict(list)
    for o in resource.observations.select_related("source").filter(excluded=False).exclude(review=Observation.Review.PENDING):
        s = o.source
        groups[(s.region, s.year, s.quarter)].append(o)
    out = []
    for (region, year, quarter), obs in groups.items():
        auto = [o.price for o in obs if not o.manual]
        out.append({"region": region, "year": year, "quarter": quarter, "n": len(auto), "manual": len(obs) - len(auto),
                    "median": Decimal(median(auto)).quantize(Decimal("0.01")) if auto else None,
                    "min": min(auto) if auto else None, "max": max(auto) if auto else None})
    out.sort(key=lambda g: (-(g["year"] or 0), -(g["quarter"] or 0), g["region"]))
    return out


def matrix():
    """Viloyat × davr: nechta kodning narxi bor (medianaga kiradigan kuzatuvlar bo'yicha)."""
    from django.db.models import Count
    rows = (Observation.objects.filter(counted_q())
            .values("source__region", "source__year", "source__quarter")
            .annotate(codes=Count("resource", distinct=True), obs=Count("id")))
    periods = sorted({(r["source__year"], r["source__quarter"]) for r in rows},
                     key=lambda p: (p[0] is None, -(p[0] or 0), -(p[1] or 0)))
    cells = {(r["source__region"], r["source__year"], r["source__quarter"]): r for r in rows}
    return periods, cells


# ===================== қамров =====================
COVER_KINDS = ("exact", "other", "none", "no_code")


def coverage(job, region, year, quarter, *, own=False) -> dict:
    """Baza bu smetani qanchalik narxlay oladi: har resurs qatori — shu viloyat/davr narxi bor (exact), faqat boshqa
    davr yoki viloyatda bor (other), kodi bor-u narxi yo'q (none), kodsiz (no_code). ``own=False`` — smetaning
    o'zidan yig'ilgan narxlar hisobga olinmaydi (aks holda o'z narxi bilan «qoplanib» qoladi)."""
    rows = [r for r in (job.result or {}).get("rows") or [] if r.get("status") != "WORK_HEADER" and r.get("unit")]
    codes = {norm_code(r.get("code")) for r in rows} - {""}
    cells = defaultdict(lambda: defaultdict(list))  # kod → (viloyat, yil, chorak) → narxlar
    qs = Observation.objects.filter(counted_q(), resource__code__in=codes)
    if not own:
        qs = qs.exclude(job_id=job.id)
    for code, r, y, q, price in qs.values_list("resource__code", "source__region", "source__year", "source__quarter",
                                                "price"):
        cells[code][(r, y, q)].append(price)
    want = (region, year, quarter) if region and year and quarter else None
    counts, missing, other, items = Counter(), {}, {}, []
    for r in rows:
        code = norm_code(r.get("code"))
        median_ = n = None
        if not code:
            kind = "no_code"
        elif want and want in cells.get(code, {}):
            ps = cells[code][want]
            kind, median_, n = "exact", Decimal(median(ps)).quantize(Decimal("0.01")), len(ps)
        elif cells.get(code):
            kind = "other"
            other.setdefault(code, {"code": code, "name": r.get("name") or "", "unit": r.get("unit"), "rows": 0,
                                    "where": sorted(cells[code], key=lambda c: (c[1] is None, -(c[1] or 0), -(c[2] or 0)))})
            other[code]["rows"] += 1
        else:
            kind = "none"
            missing.setdefault(code, {"code": code, "name": r.get("name") or "", "unit": r.get("unit"), "rows": 0})
            missing[code]["rows"] += 1
        counts[kind] += 1
        items.append({"row": r.get("row"), "code": code or str(r.get("code") or ""), "name": r.get("name") or "",
                      "unit": r.get("unit"), "kind": kind, "median": median_, "n": n, "price": _num(r.get("g"))})
    total = len(rows)
    return {"total": total, "counts": [(k, counts[k], round(100 * counts[k] / total) if total else 0) for k in COVER_KINDS],
            "missing": sorted(missing.values(), key=lambda m: (-m["rows"], int(m["code"]))),
            "other": sorted(other.values(), key=lambda m: (-m["rows"], int(m["code"]))), "items": items}


# ===================== нархлаш учун =====================
def period_complete(period) -> bool:
    return bool(period and period.get("region") and period.get("year") and period.get("quarter"))


def period_label(period) -> str:
    from .regions import REGION_LABELS
    if not period_complete(period):
        return ""
    return f"{REGION_LABELS.get(period['region'], period['region'])} · {period['year']}-{period['quarter']}"


def price_map(period) -> dict:
    """Narxlash uchun (``core.pipeline.Options.base_prices``): {kod: {price, n, unit_key, unit, category}} — faqat
    tanlangan viloyat va davr medianasi. Baza o'chirilgan yoki davr to'liq emas — {}."""
    from .models import BaseConfig
    if not period_complete(period) or not BaseConfig.get().enabled:
        return {}
    table = price_table(period["region"], int(period["year"]), int(period["quarter"]))
    meta = {c: (k, u, cat) for c, k, u, cat in
            Resource.objects.filter(code__in=table).values_list("code", "unit_key", "unit", "category")}
    return {code: {"price": float(price), "n": n, "unit_key": meta[code][0], "unit": meta[code][1],
                   "category": meta[code][2]} for code, (price, n) in table.items() if code in meta}


def parse_period(data) -> dict:
    """Формадан {region, year, quarter} — нотўғри ёки бўш қиймат = номаълум."""
    from .regions import REGION_LABELS
    region = data.get("region", "")
    try:
        year = int(data.get("year") or 0) or None
        quarter = int(data.get("quarter") or 0) or None
    except (TypeError, ValueError):
        year = quarter = None
    if year and not 2000 <= year <= 2100:
        year = None
    return {"region": region if region in REGION_LABELS else "", "year": year,
            "quarter": quarter if quarter in (1, 2, 3, 4) else None}


# ===================== текширув навбати =====================
PRICE_RATIO = 3  # medianadan shuncha baravardan ko'p farq — shubhali
NAME_MIN = 0.5   # nom o'xshashligi shundan past — shubhali
_WORD = re.compile(r"[0-9a-zа-яўқғҳ]+")


def _words(name) -> list:
    return _WORD.findall(norm_name(name))


def name_similarity(a, b) -> float:
    """0..1: qisqa nomning so'zlaridan qanchasi uzunida bor (raqamlar bilan) — «БУЛЬДОЗЕРЫ 79 КВТ» va
    «БУЛЬДОЗЕРЫ ПРИ РАБОТЕ … 79 (108) КВТ (Л.С.)» o'xshash; «ВОДА» va «ГВОЗДИ» — yo'q."""
    wa, wb = _words(a), _words(b)
    if not wa or not wb:
        return 0.0
    short, long_ = (wa, set(wb)) if len(wa) <= len(wb) else (wb, set(wa))
    hit = sum(1 for w in short if w in long_ or (len(w) >= 5 and any(x.startswith(w[:5]) for x in long_)))
    return hit / len(short)


def recheck(resource_ids=None):
    """Shubhali kuzatuvlarni qayta belgilaydi (``resource_ids`` — faqat shu resurslar, None — hammasi):
    * name — smetadagi nom resurs nomiga (va qabul qilingan nomlarga) o'xshamaydi;
    * unit — birlik boshqa (qabul qilingan birliklardan ham);
    * price — shu viloyat va davr narxlari medianasidan (3 tadan kam bo'lsa — hamma narxlar) 3 baravardan ko'p farq.
    Belgisi bor kuzatuv admin qaror qilguncha ``pending`` (medianaga kirmaydi); qabul/rad qarori saqlanadi."""
    R = Observation.Review
    qs = Observation.objects.select_related("source", "resource")
    if resource_ids is not None:
        qs = qs.filter(resource_id__in=list(resource_ids))
    by_res = defaultdict(list)
    for o in qs:
        by_res[o.resource_id].append(o)
    changed = []
    for obs in by_res.values():
        res = obs[0].resource
        ok_names = [res.name] + [o.name for o in obs if o.review == R.ACCEPTED]
        ok_units = {res.unit_key} | {unit_key(o.unit) for o in obs if o.review == R.ACCEPTED}
        flags = {}
        for o in obs:
            f = []
            if not o.manual:
                if max(name_similarity(o.name, n) for n in ok_names) < NAME_MIN:
                    f.append("name")
                if unit_key(o.unit) not in ok_units:
                    f.append("unit")
            flags[o.pk] = f
        # narx solishtiriladigan asos: ishonchli (nomi va birligi mos yoki qabul qilingan) kuzatuvlar
        base = [o for o in obs if not o.manual and not o.excluded and o.review != R.REJECTED
                and (not flags[o.pk] or o.review == R.ACCEPTED)]
        cells = defaultdict(list)
        for o in base:
            cells[(o.source.region, o.source.year, o.source.quarter)].append(o)
        for o in obs:
            if o.manual or o.excluded or o.price <= 0:
                continue
            # o'zi bilan birga medianasi (kamida 3 ta narx): 2 tadan qaysi biri xato ekanini bilib bo'lmaydi
            group = [x.price for x in cells[(o.source.region, o.source.year, o.source.quarter)] if x.pk != o.pk] + [o.price]
            if len(group) < 3:
                group = [x.price for x in base if x.pk != o.pk] + [o.price]
                if len(group) < 3:
                    continue
            m = median(group)
            if m > 0 and max(o.price / m, m / o.price) > PRICE_RATIO:
                flags[o.pk].append("price")
        for o in obs:
            new = ",".join(flags[o.pk])
            if o.review in (R.REJECTED, R.ACCEPTED):
                review = o.review  # admin qarori saqlanadi (qabul qilingan nom/birlik boshqalar uchun ham asos)
            else:
                review = R.PENDING if new else R.NONE
            if (o.flags, o.review) != (new, review):
                o.flags, o.review = new, review
                changed.append(o)
    Observation.objects.bulk_update(changed, ["flags", "review"], batch_size=500)
    if changed:
        transaction.on_commit(bump)
    return len(changed)


def decide(obs_ids, accept: bool) -> int:
    """Admin qarori: «қабул» — medianaga kiradi (shu nom va birlik keyin shubhali bo'lmaydi), «рад» — istisno."""
    R = Observation.Review
    qs = Observation.objects.filter(pk__in=list(obs_ids))
    rids = set(qs.values_list("resource_id", flat=True))
    n = qs.update(review=R.ACCEPTED, excluded=False) if accept else qs.update(review=R.REJECTED, excluded=True)
    recheck(rids)
    bump()
    return n
