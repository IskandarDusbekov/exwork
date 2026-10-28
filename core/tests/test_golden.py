"""Etalon testlar (reja §6): eski dastur natijasi bilan solishtirish.

``golden/<nom>/`` ichida: ``input_local.*``, ``input_res.*`` (ustun qo'shilgan),
``expected_narxlandi.xlsx`` (eski dastur chiqargan) va ixtiyoriy ``options.json``
(``{"fuzzy": 95}``). Papka bo'sh bo'lsa testlar skip qilinadi.
"""
import datetime as dt
import re

import pytest
from openpyxl import load_workbook

from core.log_sheet import LOG_TITLE
from core.pipeline import Options, run
from core.svodka import XlError

from .helpers import golden_cases, soffice_available

CASES = golden_cases()
if not CASES:
    pytestmark = pytest.mark.skip(reason="Etalon fayllar yo'q (core/tests/golden/<nom>/)")

# Eski dastur LOG'ga str(...) yozgan; Excel «1-6» kabi matnni o'zi sanaga aylantiradi.
_DATE_LIKE = re.compile(r"^\s*\d{1,4}[-./]\d{1,2}([-./]\d{1,4})?\s*$")


@pytest.fixture(scope="module", params=CASES, ids=lambda c: c.name)
def golden(request, tmp_path_factory):
    case = request.param
    if case.needs_soffice() and not soffice_available():
        pytest.skip("LibreOffice yo'q: .xls kirishni konvertatsiya qilib bo'lmaydi")
    out = tmp_path_factory.mktemp(case.name) / "out.xlsx"
    opts = Options(
        output_path=out,
        fuzzy_percent=case.options.get("fuzzy", 95),
        local_sheet=case.options.get("local_sheet"),
        res_sheet=case.options.get("res_sheet"),
    )
    result = run(case.local, case.res, opts)
    exp_wb = load_workbook(case.expected)
    exp_vals = load_workbook(case.expected, data_only=True)
    out_wb = load_workbook(out)
    return case, result, exp_wb, exp_vals, out_wb


def _main(wb, title=None):
    if title:
        return wb[title]
    return next(ws for ws in wb.worksheets if ws.title != LOG_TITLE)


def _same(a, b) -> bool:
    if a == "":
        a = None
    if b == "":
        b = None
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return a == b


def _column(ws, col, last):
    return [ws[f"{col}{r}"].value for r in range(1, last + 1)]


def test_sheet_order(golden):
    _, result, exp_wb, _, out_wb = golden
    assert out_wb.sheetnames == exp_wb.sheetnames
    assert out_wb.active.title == exp_wb.active.title


def test_g_column(golden):
    _, result, exp_wb, _, out_wb = golden
    exp, out = _main(exp_wb, result.local_sheet), _main(out_wb, result.local_sheet)
    last = max(exp.max_row, out.max_row)
    diffs = [(r, e, o) for r, (e, o) in enumerate(zip(_column(exp, "G", last), _column(out, "G", last)), 1)
             if not _same(e, o)]
    assert not diffs, f"G farqlari (qator, kutilgan, chiqqan): {diffs[:20]}"


@pytest.mark.parametrize("col", ["H", "I", "J", "K"])
def test_formula_columns(golden, col):
    _, result, exp_wb, _, out_wb = golden
    exp, out = _main(exp_wb, result.local_sheet), _main(out_wb, result.local_sheet)
    last = max(exp.max_row, out.max_row)
    diffs = [(r, e, o) for r, (e, o) in enumerate(zip(_column(exp, col, last), _column(out, col, last)), 1)
             if not _same(e, o)]
    assert not diffs, f"{col} farqlari: {diffs[:20]}"


def test_log_rows(golden):
    _, _, exp_wb, _, out_wb = golden
    exp = [list(r) for r in exp_wb[LOG_TITLE].iter_rows(values_only=True)]
    out = [list(r) for r in out_wb[LOG_TITLE].iter_rows(values_only=True)]
    assert len(exp) == len(out), f"LOG qatorlar soni: {len(exp)} != {len(out)}"
    diffs = []
    for i, (e_row, o_row) in enumerate(zip(exp, out), start=1):
        for j, (e, o) in enumerate(zip(e_row, o_row), start=1):
            if isinstance(e, dt.datetime) and isinstance(o, str) and _DATE_LIKE.match(o):
                continue  # Excel'ning avto-sana o'zgartirishi
            if not _same(e, o):
                diffs.append((i, j, e, o))
    assert not diffs, f"LOG farqlari (qator, ustun, kutilgan, chiqqan): {diffs[:20]}"


def test_totals_match_excel(golden):
    """Python hisoblagan svodka = Excel saqlagan qiymat (±1 so'm), xatolar ham bir xil."""
    _, result, _, exp_vals, _ = golden
    ws = _main(exp_vals, result.local_sheet)
    s = result.svodka_start
    for i, (key, value) in enumerate(result.totals.items()):
        cached = ws[f"H{s + i}"].value
        if cached is None:
            pytest.skip("Etalon faylda hisoblangan qiymatlar saqlanmagan")
        if isinstance(value, XlError):
            assert cached == str(value), f"{key}: {cached!r} != {value}"
        else:
            assert isinstance(cached, (int, float)), f"{key}: Excel {cached!r}, Python {value!r}"
            assert abs(cached - value) <= 1, f"{key}: {cached} != {value}"
