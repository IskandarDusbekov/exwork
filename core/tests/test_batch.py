"""Bir/ko'p faylli yuklash: list turi, nusxalar, juftlash (haqiqiy obyekt fayllarida)."""
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.batch import LOCAL, RESOURCE, UNKNOWN, apply_roles, inspect_workbook, mark_duplicates, plan_pairs

TESTS = Path(__file__).parent
OBJ_LOCAL = TESTS / "batch/obekt_lokal_7list.xlsx"   # 400/410/420 БВ+БР, 450_БВ
OBJ_RES = TESTS / "batch/obekt_res_7list.xlsx"       # o'sha listlar + 450_БР (C oldiga ustun qo'shilgan)


def _infos(*paths):
    infos = []
    for i, p in enumerate(paths):
        infos += inspect_workbook(load_workbook(p), i, Path(p).name)
    return mark_duplicates(infos)


def test_roles_numbers_and_names_from_content():
    infos = {i.sheet: i for i in _infos(OBJ_LOCAL)}
    assert {s: i.role for s, i in infos.items()} == {
        "202611400_БВ": LOCAL, "202611400_БР": RESOURCE, "202611410_БВ": LOCAL, "202611410_БР": RESOURCE,
        "202611420_БВ": LOCAL, "202611420_БР": RESOURCE, "202611450_БВ": LOCAL}
    m = infos["202611450_БВ"]
    assert (m.number, m.obj) == ("1-6", "МОСТ НА ПК 75+34")
    assert m.stroyka.startswith('IL "URGUT TUMANIDAGI')
    assert infos["202611410_БВ"].obj == "YANGI-DARGOM KANALI РВТ 80-100"


def test_one_object_file_pairs_by_content_not_only_number():
    # 400 va 410 ikkalasi ham «№ 1-1» — mazmun bo'yicha ajratiladi
    pairs, unused = plan_pairs(_infos(OBJ_LOCAL))
    got = {p.local: p.res for p in pairs}
    assert got == {"0:202611400_БВ": "0:202611400_БР", "0:202611410_БВ": "0:202611410_БР",
                   "0:202611420_БВ": "0:202611420_БР", "0:202611450_БВ": None}
    assert unused == []


def test_two_files_duplicates_and_cross_file_pair():
    infos = _infos(OBJ_LOCAL, OBJ_RES)
    dups = {i.key: i.duplicate_of for i in infos if i.duplicate_of}
    assert len(dups) == 6 and dups["1:202611400_БВ"] == "0:202611400_БВ"
    pairs, unused = plan_pairs(infos)
    assert len(pairs) == 4 and unused == []
    most = next(p for p in pairs if p.local == "0:202611450_БВ")
    assert most.res == "1:202611450_БР" and most.overlap > 0.9
    assert most.label == "МОСТ НА ПК 75+34 · № 1-6"
    assert any("смета № 1-6" in r for r in most.reasons)


def test_names_do_not_matter_lrv_res():
    d = TESTS / "excel_calc/bolnitsa7_kr_blok2"
    infos = _infos(d / "input_local.xlsx", d / "input_res.xlsx")   # listlar: «LRV», «RES»
    assert [i.role for i in infos] == [LOCAL, RESOURCE]
    assert infos[0].stroyka == "7-БОЛНИЦА 2019" and infos[0].obj == "БЛОК 1,2,3,4,5"
    pairs, _ = plan_pairs(infos)
    assert [(p.local, p.res) for p in pairs] == [("0:LRV", "1:RES")]
    assert pairs[0].label == "КОНСТРУКТИВНЫЕ РЕШЕНИЕ (КР) БЛОК 2 · № С-03"


def test_unknown_sheet_and_manual_roles(tmp_path):
    wb = Workbook()
    wb.active.title = "Свод"
    wb.active["A1"] = "СВОДНЫЙ СМЕТНЫЙ РАСЧЕТ"
    wb.save(tmp_path / "s.xlsx")
    infos = _infos(OBJ_LOCAL, tmp_path / "s.xlsx")
    svod = next(i for i in infos if i.sheet == "Свод")
    assert svod.role == UNKNOWN
    apply_roles(infos, {"0:202611400_БР": UNKNOWN})
    pairs, _ = plan_pairs(infos)
    assert next(p for p in pairs if p.local == "0:202611400_БВ").res != "0:202611400_БР"


# ---------------- tasnif: resurs/lokal ko'rinishidagi, lekin smeta bo'lmagan listlar ----------------
def _sheet(wb, title, header, rows, top=()):
    ws = wb.create_sheet(title)
    for r, text in enumerate(top, start=1):
        ws.cell(r, 1, text)
    h = len(top) + 2
    for c, text in enumerate(header, start=1):
        ws.cell(h, c, text)
        ws.cell(h + 1, c, c)  # «1 2 3 …» raqamlash qatori
    for i, row in enumerate(rows, start=h + 2):
        for c, v in enumerate(row, start=1):
            ws.cell(i, c, v)
    return ws


def test_catalog_priced_lrv_act_and_summary_are_skipped():
    wb = Workbook()
    wb.remove(wb.active)
    items = [(n, f"МАТЕРИАЛ {n}", "М3") for n in range(1, 9)]
    _sheet(wb, "РС", ["№ п/п", "Наименование ресурсов", "Ед. изм.", "Количество", "Цена"],
           [(n, name, unit, 2, 1000) for n, name, unit in items], top=["Форма N 6a", "ЛОКАЛЬНЫЙ РЕСУРСНЫЙ СМЕТНЫЙ РАСЧЕТ"])
    _sheet(wb, "Каталог", ["№ п/п", "Наименование выпускаемой продукции", "Ед. изм", "Отпускная текущая цена"],
           [(n, name, unit, 1000) for n, name, unit in items])
    lrv = []
    for n in range(1, 5):  # ish + uning resurslari «1.1, 1.2», narx yozilgan
        lrv += [(str(n), f"ИШ {n}", "100 М2", 1, 1), (f"{n}.1", "ЗАТРАТЫ ТРУДА РАБОЧИХ", "ЧЕЛ-Ч", 2, 3),
                (f"{n}.2", f"МАТЕРИАЛ {n}", "М3", 1, 1000)]
    _sheet(wb, "ЛРВ нархланган", ["№ п/п", "Наименование работ и затрат", "Единица измерения", "Количество",
                                   "Сметная стоимость"], lrv, top=["ЛОКАЛЬНО-РЕСУРСНАЯ ВЕДОМОСТЬ"])
    _sheet(wb, "Дефектный акт", ["№ п/п", "Наименование работ", "Ед. изм.", "Количество"],
           [(n, f"ИШ {n}", "100 М2", 1) for n in range(1, 6)], top=["ДЕФЕКТНЫЙ АКТ"])
    wb.create_sheet("СВОД").cell(2, 2, "СВОДНЫЙ СМЕТНЫЙ РАСЧЕТ СТОИМОСТИ СТРОИТЕЛЬСТВА")
    wb.create_sheet("Бўш")
    infos = {i.sheet: i for i in mark_duplicates(inspect_workbook(wb))}
    assert infos["РС"].role == RESOURCE
    assert {s: (infos[s].role, infos[s].kind) for s in ("Каталог", "ЛРВ нархланган", "Дефектный акт", "СВОД", "Бўш")} == {
        "Каталог": (UNKNOWN, "catalog"), "ЛРВ нархланган": (UNKNOWN, "priced"), "Дефектный акт": (UNKNOWN, "act"),
        "СВОД": (UNKNOWN, "summary"), "Бўш": (UNKNOWN, "empty")}


def test_visible_copy_preferred_over_hidden_original(tmp_path):
    """Smetachi asl listlarni yashirib, nusxasi bilan ishlagan: ko'rinadigan nusxa olinadi va juftlanadi."""
    wb = load_workbook(OBJ_LOCAL)
    src = wb["202611410_БВ"]
    src.sheet_state = "hidden"
    wb.copy_worksheet(src).title = "Лист5"
    wb.active = wb["202611400_БВ"]
    p = tmp_path / "hidden.xlsx"
    wb.save(p)
    infos = _infos(p)
    by = {i.sheet: i for i in infos}
    assert by["202611410_БВ"].hidden and by["202611410_БВ"].duplicate_of == "0:Лист5"
    assert by["Лист5"].duplicate_of is None
    pairs, _unused = plan_pairs(infos)
    got = {p.local: p.res for p in pairs}
    assert got["0:Лист5"] == "0:202611410_БР"
    assert "0:202611410_БВ" not in got
    # qo'lda tanlash uchun: har bir resurs listi mosligi bilan, eng mosi birinchi
    opts = next(p for p in pairs if p.local == "0:Лист5").options
    assert opts[0]["res"] == "0:202611410_БР" and opts[0]["same_file"]
    assert len(opts) == sum(1 for i in infos if i.role == RESOURCE and not i.duplicate_of)
