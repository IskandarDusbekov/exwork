import pytest
from openpyxl import Workbook, load_workbook

from core import convert

from .helpers import soffice_available


def test_xlsx_passthrough(tmp_path):
    p = tmp_path / "a.xlsx"
    Workbook().save(p)
    assert convert.ensure_xlsx(p, tmp_path / "out") == p


def test_unsupported_extension(tmp_path):
    with pytest.raises(convert.ConversionError, match="qo'llanmaydi"):
        convert.ensure_xlsx(tmp_path / "a.csv", tmp_path)


def test_configured_path_must_exist(monkeypatch, tmp_path):
    monkeypatch.setenv(convert.ENV_VAR, str(tmp_path / "yoq" / "soffice.exe"))
    with pytest.raises(convert.ConversionError, match="topilmadi"):
        convert.find_soffice()


def test_explicit_path_wins(monkeypatch, tmp_path):
    fake = tmp_path / "soffice.exe"
    fake.write_text("")
    monkeypatch.setenv(convert.ENV_VAR, str(tmp_path / "boshqa.exe"))
    assert convert.find_soffice(str(fake)) == str(fake)


def test_no_soffice_gives_clear_error(monkeypatch, tmp_path):
    monkeypatch.delenv(convert.ENV_VAR, raising=False)
    monkeypatch.setenv(convert.CONVERTER_ENV, "libreoffice")
    monkeypatch.setattr(convert.shutil, "which", lambda name: None)
    monkeypatch.setattr(convert, "_WINDOWS_PATHS", ())
    (tmp_path / "a.xls").write_bytes(b"")
    with pytest.raises(convert.ConversionError, match=convert.ENV_VAR):
        convert.ensure_xlsx(tmp_path / "a.xls", tmp_path / "out")


def test_auto_falls_back_to_excel(monkeypatch, tmp_path):
    monkeypatch.delenv(convert.ENV_VAR, raising=False)
    monkeypatch.delenv(convert.CONVERTER_ENV, raising=False)
    monkeypatch.setattr(convert.shutil, "which", lambda name: None)
    monkeypatch.setattr(convert, "_WINDOWS_PATHS", ())
    monkeypatch.setattr(convert, "excel_available", lambda: True)
    called = []
    monkeypatch.setattr(convert, "_convert_excel", lambda p, d: called.append(p) or d / "a.xlsx")
    (tmp_path / "a.xls").write_bytes(b"")
    assert convert.ensure_xlsx(tmp_path / "a.xls", tmp_path).name == "a.xlsx"
    assert called


def test_excel_opens_with_macros_disabled(monkeypatch, tmp_path):
    """Begona .xls Excel orqali ochilganda makros va hodisalar ishlamasligi kerak — bu Open'dan OLDIN o'rnatiladi."""
    import sys
    import types

    state = {}

    class Workbook:
        def SaveAs(self, path, FileFormat):  # noqa: N802, N803 — COM nomlari
            open(path, "wb").close()

        def Close(self, SaveChanges):  # noqa: N802, N803
            pass

    class Books:
        def Open(self, path, **kw):  # noqa: N802
            state["open"] = dict(kw, security=state.get("AutomationSecurity"), events=state.get("EnableEvents"))
            return Workbook()

    books = Books()

    class Excel:
        @property
        def Workbooks(self):  # noqa: N802
            return books

        def __setattr__(self, name, value):
            state[name] = value

        def Quit(self):  # noqa: N802
            pass

    client = types.SimpleNamespace(DispatchEx=lambda name: Excel())
    monkeypatch.setitem(sys.modules, "pythoncom", types.SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None))
    monkeypatch.setitem(sys.modules, "win32com", types.SimpleNamespace(client=client))
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(convert, "excel_available", lambda: True)
    monkeypatch.setattr(convert, "_excel_pid", lambda excel: None)
    (tmp_path / "a.xls").write_bytes(b"")

    out = convert._convert_excel(tmp_path / "a.xls", tmp_path)
    assert out.name == "a.xlsx"
    assert state["open"]["security"] == convert.MSO_AUTOMATION_FORCE_DISABLE
    assert state["open"]["events"] is False
    assert state["open"]["UpdateLinks"] == 0 and state["open"]["ReadOnly"] is True


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice o'rnatilmagan")
def test_real_conversion(tmp_path):
    src = tmp_path / "src.xlsx"
    wb = Workbook()
    wb.active["C5"] = "ВОДА"
    wb.save(src)
    # xlsx -> xls -> xlsx aylanmasi uchun avval xls yasab olamiz
    import subprocess
    exe = convert.find_soffice()
    subprocess.run([exe, "--headless", "--convert-to", "xls", "--outdir", str(tmp_path), str(src)],
                   check=True, capture_output=True, timeout=180)
    out = convert.ensure_xlsx(tmp_path / "src.xls", tmp_path / "conv")
    assert load_workbook(out).active["C5"].value == "ВОДА"
