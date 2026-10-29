import pytest
from openpyxl import load_workbook

from core.log_sheet import LOG_TITLE, excel_value
from core.pipeline import Options, default_output_path, main, run
from core.svodka import VALUE

from .conftest import FIRST_DATA_ROW, build_local, build_res


def test_run_sample(sample_files, tmp_path):
    local, res = sample_files
    out = tmp_path / "out.xlsx"
    result = run(local, res, Options(output_path=out))

    # Sarlavha qatorlari ham narxlanadi (compat): oraliq «Наименование…» bilan boshlanadi
    assert result.start_row == 4  # C4 «ЛОКАЛЬНАЯ РЕСУРСНАЯ ВЕДОМОСТЬ»
    assert result.end_row == FIRST_DATA_ROW + 7
    assert result.sum_row == result.end_row + 1
    assert result.svodka_start == result.sum_row + 1

    wb = load_workbook(out)
    assert wb.sheetnames == [LOG_TITLE, "202611450_БВ"]
    assert wb.active.title == LOG_TITLE
    assert [ws.sheet_view.tabSelected for ws in wb.worksheets] == [True, False]
    ws = wb["202611450_БВ"]

    r = FIRST_DATA_ROW
    assert ws[f"G{r}"].value == 0                      # nol qoidasi: ish sarlavhasi
    assert ws[f"G{r + 1}"].value == 29818.17           # ishchilar
    assert ws[f"G{r + 2}"].value == "--"               # машинистлар (eski dastur xatosi)
    assert ws[f"G{r + 4}"].value == 485994             # fuzzy
    assert ws[f"G{r + 7}"].value == "ТОПИЛМАДИ"
    assert ws[f"H{r + 1}"].value == f'=IF(F{r + 1}="",E{r + 1}*G{r + 1},F{r + 1}*G{r + 1})'
    assert (ws[f"I{r + 1}"].value, ws[f"J{r + 1}"].value, ws[f"K{r + 1}"].value) == (f"=H{r + 1}", None, None)
    assert (ws[f"I{r + 2}"].value, ws[f"J{r + 2}"].value) == (None, f"=H{r + 2}")
    assert (ws[f"J{r + 3}"].value, ws[f"K{r + 4}"].value) == (f"=H{r + 3}", f"=H{r + 4}")
    # birlashmaning chap-yuqori bo'lmagan katagiga yozilmaydi (E:F birlashgan)
    assert ws[f"F{r}"].value is None

    s, sm = result.svodka_start, result.sum_row
    assert ws[f"H{sm}"].value == f"=SUM(H{result.start_row}:H{result.end_row})"
    assert ws[f"Q{sm}"].value == f"=SUM(Q{result.start_row}:Q{result.end_row})"
    assert ws[f"C{s}"].value == "Зарплата" and ws[f"C{s}"].font.b
    assert ws[f"C{s + 12}"].value == "Всего с НДС      12 %"
    assert ws[f"D{s + 9}"].value == 0.19 and ws[f"D{s + 9}"].number_format == "0.00%"
    assert ws[f"H{s + 12}"].value == f"=H{s + 11}*1.12"
    assert ws[f"A{s}"].fill.fgColor.rgb == "FFD9D9D9"

    log = wb[LOG_TITLE]
    assert log.freeze_panes == "A2" and log.auto_filter.ref == "A1:J1"
    statuses = {row[0]: row[3] for row in log.iter_rows(min_row=2, values_only=True)}
    assert statuses[r + 4] == "FUZZY" and statuses[r + 7] == "NOT_FOUND"

    assert result.stats["FUZZY"] == 1
    assert result.stats["ZERO_RULE"] == 1
    assert {p["row"] for p in result.problems} >= {r + 2, r + 7}
    # «--» va sarlavhalardagi «ТОПИЛМАДИ» tufayli Excel'dagidek #VALUE!
    assert result.totals["mashina"] == VALUE
    assert result.totals["s_nds"] == VALUE


def test_totals_numeric(tmp_path):
    local = build_local(tmp_path / "local.xlsx", numeric_only=True)
    res = build_res(tmp_path / "res.xlsx")
    result = run(local, res, Options(output_path=tmp_path / "out.xlsx"))
    assert result.start_row == FIRST_DATA_ROW
    t = result.totals

    workers = 39.7 * 29818.17
    machines = 0.248 * 244250
    materials = 0 + 2.5375 * 485994 + 4.375 * 1000 + 0.5675 * 105000
    assert t["zarplata"] == pytest.approx(workers)
    assert t["mashina"] == pytest.approx(machines)
    assert t["materialy"] == pytest.approx(materials)
    direct = workers + machines + materials * 1.03
    assert t["pryamye"] == pytest.approx(direct)
    prochie = direct * 0.19
    bez = (prochie + direct) * 0.0032 + prochie + direct
    assert t["bez_nds"] == pytest.approx(bez)
    assert t["s_nds"] == pytest.approx(bez * 1.12)
    assert result.problems == []


def test_default_output_path(tmp_path):
    assert default_output_path(tmp_path / "Локал 3.xls").name == "Локал 3_narxlandi.xlsx"
    assert default_output_path(tmp_path / "a.xlsm").name == "a_narxlandi.xlsm"


def test_existing_log_is_replaced(sample_files, tmp_path):
    local, res = sample_files
    first = run(local, res, Options(output_path=tmp_path / "1.xlsx")).output_path
    second = run(first, res, Options(output_path=tmp_path / "2.xlsx", local_sheet="202611450_БВ")).output_path
    assert load_workbook(second).sheetnames.count(LOG_TITLE) == 1


def test_missing_sheet(sample_files, tmp_path):
    local, res = sample_files
    with pytest.raises(ValueError, match="topilmadi"):
        run(local, res, Options(output_path=tmp_path / "o.xlsx", local_sheet="yo'q"))


def test_cli(sample_files, tmp_path, capsys):
    local, res = sample_files
    out = tmp_path / "cli.xlsx"
    assert main([str(local), str(res), "-o", str(out), "--fuzzy", "95", "--mode", "compat"]) == 0
    assert out.is_file()
    text = capsys.readouterr().out
    assert "Tayyor" in text and "s_nds" in text and "#VALUE!" in text

    assert main([str(local), str(res), "-o", str(out)]) == 0  # standart: smart
    text = capsys.readouterr().out
    assert "rejim: smart" in text and "#VALUE!" not in text


@pytest.mark.parametrize("raw, value", [
    ("29818.17", 29818.17), ("153724", 153724), ("1.0", 1), (1.0, 1), ("--", "--"),
    ("None", "None"), ("", None), (" 12 ", 12), ("ТОПИЛМАДИ", "ТОПИЛМАДИ"),
])
def test_log_excel_value(raw, value):
    assert excel_value(raw) == value


def _add_broken_external_link(path):
    """LibreOffice o'girgan ayrim .xls'lardagidek: workbook boshqa faylga havola qiladi, havola XML'i buzuq."""
    import zipfile
    src = path.read_bytes()
    with zipfile.ZipFile(path.with_suffix(".tmp"), "w", zipfile.ZIP_DEFLATED) as zo, \
            zipfile.ZipFile(__import__("io").BytesIO(src)) as zi:
        for item in zi.infolist():
            data = zi.read(item)
            if item.filename == "xl/workbook.xml":
                data = data.replace(b"</sheets>", b'</sheets><externalReferences><externalReference r:id="rIdX"/>'
                                    b"</externalReferences>")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats'
                                    b'.org/officeDocument/2006/relationships/externalLink" '
                                    b'Target="externalLinks/externalLink1.xml"/></Relationships>')
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/externalLinks/externalLink1.xml" '
                                    b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
                                    b'externalLink+xml"/></Types>')
            zo.writestr(item, data)
        zo.writestr("xl/externalLinks/externalLink1.xml", b'<?xml version="1.0"?>\n<externalLink>\x01</externalLink>')
    path.with_suffix(".tmp").replace(path)


def test_broken_external_link(sample_files, tmp_path):
    import xml.etree.ElementTree as ET
    from jobs.services import sheet_names
    local, res = sample_files
    for p in (local, res):
        _add_broken_external_link(p)
    with pytest.raises(ET.ParseError):  # openpyxl o'zi (havolalar bilan) ocholmaydi
        load_workbook(local)
    assert "202611450_БВ" in sheet_names(local)
    result = run(local, res, Options(output_path=tmp_path / "out.xlsx"))
    assert result.stats["FUZZY"] == 1
