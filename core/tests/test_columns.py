"""Lokal ustunlarini aniqlash (sarlavha har xil yozilgan / yo'q / noto'g'ri) va narxlashdan keyingi tekshiruv."""
import pytest
from openpyxl import Workbook, load_workbook

from core.columns import LayoutError, detect_local_columns, header_role, num, prepare_local
from core.pipeline import Options, run

from .conftest import LOCAL_ROWS, build_local, build_res

# Ish (hajmi E da) va uning resurslari: birlikka × hajm = jami
ROWS = [
    ("1", "Е01-01", "РАЗРАБОТКА ГРУНТА", "1000 М3", "3,655", None),
    ("1.1", "00003", "ЗАТРАТЫ ТРУДА МАШИНИСТОВ", "ЧЕЛ-ЧАС", "25,96", "94,8838"),
    ("1.2", "02263", "ЭКСКАВАТОРЫ ОДНОКОВШОВЫЕ", "МАШ-ЧАС", "25,96", "94,8838"),
    ("1.3", "4001", "ЩЕБЕНЬ", "М3", "2", "7,31"),
    ("2", "Е01-02", "ДОРАБОТКА ГРУНТА ВРУЧНУЮ", "100 М3", "3,248", None),
    ("2.1", "00001", "ЗАТРАТЫ ТРУДА РАБОЧИХ-СТРОИТЕЛЕЙ", "ЧЕЛ-ЧАС", "97,2", "315,7056"),
    ("2.2", "4002", "ПЕСОК", "М3", "1,5", "4,872"),
]


def _sheet(head=("№ п/п", "Шифр", "Наименование работ и затрат", "Единица измерения"), qty=("на единицу измерения", "по проектным данным"),
           total_col=6, merged=False, header=True):
    """Lokal: nom C, birlik D; birlikka — E, jami — ``total_col``. ``merged`` — E:F va G:H birlashgan («1. НВ В2»)."""
    ws = Workbook().active
    r = 1
    if header:
        ws.append(list(head) + ["Количество"])
        ws.append([None] * 4 + [qty[0]] + [None] * (total_col - 6) + [qty[1]])
        ws.append(["1", "2", "3", "4", "5"] + [None] * (total_col - 6) + ["6"])
        r = 4
    for row in ROWS:
        values = list(row[:5]) + [None] * (total_col - 6) + [row[5]]
        ws.append(values)
    if merged:
        ws.merge_cells("E1:H1")
        for rr in range(2, ws.max_row + 1):
            ws.merge_cells(f"E{rr}:F{rr}")
            ws.merge_cells(f"G{rr}:H{rr}")
    return ws, r


def _prepare(ws):
    cols = detect_local_columns(ws)
    assert cols is not None
    return cols, prepare_local(ws, ws, cols)


def test_num_reads_text_numbers():
    assert num("94,8838") == pytest.approx(94.8838)
    assert num("1 234,5") == 1234.5 and num(" 7 ") == 7 and num(3) == 3.0
    assert num("ЧЕЛ-ЧАС") is None and num(None) is None and num(True) is None and num("=E5*2") is None


@pytest.mark.parametrize("text, role", [
    ("количество | на единицу измерения", "qty_unit"),
    ("КОЛ-ВО | НА ЕД.", "qty_unit"),
    ("Кол-во | по проекту", "qty_total"),
    ("КОЛИЧЕСТВО | ВСЕГО", "qty_total"),
    ("Потребность", "qty_total"),
    ("Кол.", "qty"),
    ("Обьем", "qty"),
    ("Цена", "price"),
    ("сметная стоимость | на единицу измерения", "price"),
    ("СМЕТНАЯ СТОИМОСТЬ | ОБЩАЯ", "amount"),
    ("Сумма", "amount"),
    ("Примечание", None),
])
def test_header_role(text, role):
    assert header_role(text) == role


def test_standard_layout_is_left_as_is():
    ws, first = _sheet()
    cols, notes = _prepare(ws)
    assert (cols.name, cols.unit, cols.qty_unit, cols.qty_total) == (3, 4, 5, 6)
    assert notes == [] and ws.cell(first + 1, 6).value == "94,8838"  # ҳеч нарса кўчирилмади


def test_work_header_merges_do_not_hide_total_column():
    """Фақат иш сарлавҳаларида (ҳажм) E:F бирлашган, ресурс қаторларида F тўла — F жами миқдор бўлиб қолади."""
    ws, first = _sheet()
    for i, row in enumerate(ROWS):
        if row[5] is None:
            ws.merge_cells(f"E{first + i}:F{first + i}")
    cols, notes = _prepare(ws)
    assert (cols.qty_unit, cols.qty_total) == (5, 6) and notes == []


def test_upper_case_and_abbreviated_headers():
    ws, first = _sheet(head=("№", "ОБОСН.", "НАИМ. РАБОТ И РЕСУРСОВ", "ЕД."), qty=("КОЛ-ВО НА ЕД.", "ВСЕГО"))
    cols = detect_local_columns(ws)
    assert (cols.name, cols.unit, cols.qty_unit, cols.qty_total, cols.source) == (3, 4, 5, 6, "header")


def test_merged_quantity_columns_are_split():
    """«1. НВ В2»: E:F — на единицу, G:H — по проектным данным (нарх устуни йўқ)."""
    ws, first = _sheet(total_col=7, merged=True)
    cols, notes = _prepare(ws)
    assert (cols.qty_unit, cols.qty_total) == (5, 7)
    row = first + 1
    assert [ws.cell(row, c).value for c in range(5, 9)] == [pytest.approx(25.96), pytest.approx(94.8838), None, None]
    assert ws.cell(2, 6).value == "по проектным данным" and ws.cell(3, 6).value == "6"
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "E1:F1" in merged and not any(m.startswith("G") for m in merged)
    assert any("стандарт жойга" in n for n in notes)
    # иккинчи марта — ўзгармайди
    assert _prepare(ws)[1] == []


def test_total_in_g_without_merges():
    ws, first = _sheet(total_col=7)
    cols, _notes = _prepare(ws)
    assert (cols.qty_unit, cols.qty_total) == (5, 7)
    assert ws.cell(first + 1, 6).value == pytest.approx(94.8838) and ws.cell(first + 1, 7).value is None


def test_wrong_header_is_corrected_by_content():
    """Сарлавҳа F ни «по проекту» дейди, лекин F бўш — жами миқдор аслида G да."""
    ws, first = _sheet(total_col=7)
    ws.cell(2, 6, "по проекту")
    ws.cell(2, 7, "Примечание")
    cols = detect_local_columns(ws)
    assert (cols.qty_unit, cols.qty_total) == (5, 7)


def test_no_header_detected_by_content():
    ws, first = _sheet(header=False, total_col=7)
    cols, _notes = _prepare(ws)
    assert cols.source == "content"
    assert (cols.name, cols.unit, cols.qty_unit, cols.qty_total) == (3, 4, 5, 7)
    assert ws.cell(first + 5, 6).value == pytest.approx(315.7056)


def test_money_columns_are_recognised():
    ws, first = _sheet()
    ws.cell(1, 7, "Сметная стоимость")
    ws.merge_cells("G1:H1")
    ws.cell(2, 7, "на единицу")
    ws.cell(2, 8, "общая")
    cols, notes = _prepare(ws)
    assert (cols.qty_unit, cols.qty_total, cols.money) == (5, 6, [7, 8]) and notes == []


def test_name_too_far_right_is_refused():
    ws = Workbook().active
    ws.append(["№", "Шифр", "Обоснование", "Код", "Наименование", "Ед. изм.", "Кол-во"])
    for i, row in enumerate(ROWS):
        ws.append([row[0], row[1], None, None, row[2], row[3], row[4]])
    cols = detect_local_columns(ws)
    with pytest.raises(LayoutError, match="Номи E"):
        prepare_local(ws, ws, cols)


def _to_merged_layout(path):
    """conftest локалини «1. НВ В2» кўринишига: жами миқдор G:H да, E:F ва G:H бирлашган."""
    wb = load_workbook(path)
    ws = wb.active
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= 5 and m.min_col >= 5:
            ws.unmerge_cells(str(m))
    ws.merge_cells("E5:H5")  # «Количество»
    for r in range(6, ws.max_row + 1):
        if ws.cell(r, 5).value is None and ws.cell(r, 6).value is None:
            continue
        ws.cell(r, 7).value, ws.cell(r, 6).value = ws.cell(r, 6).value, None
        ws.merge_cells(f"E{r}:F{r}")
        ws.merge_cells(f"G{r}:H{r}")
    wb.save(path)
    return path


def test_pipeline_prices_merged_layout_like_standard(tmp_path):
    res = build_res(tmp_path / "res.xlsx")
    standard = run(build_local(tmp_path / "a.xlsx"), res, Options(mode="smart", output_path=tmp_path / "a_out.xlsx"))
    merged = run(_to_merged_layout(build_local(tmp_path / "b.xlsx")), res, Options(mode="smart", output_path=tmp_path / "b_out.xlsx"))
    assert merged.checks == [] and standard.checks == []
    assert merged.totals["s_nds"] == pytest.approx(standard.totals["s_nds"])
    assert merged.totals["s_nds"] > 0


def test_check_reports_priced_row_without_quantity(tmp_path):
    path = build_local(tmp_path / "l.xlsx")
    wb = load_workbook(path)
    ws = wb.active
    row = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 3).value == LOCAL_ROWS[1][2])
    ws.cell(row, 5).value = ws.cell(row, 6).value = "—"  # миқдор сон эмас
    wb.save(path)
    result = run(path, build_res(tmp_path / "r.xlsx"), Options(mode="smart", output_path=tmp_path / "o.xlsx"))
    assert any(c["row"] == row and "миқдори йўқ" in c["problem"] for c in result.checks)
    assert result.warnings[0].startswith("Текширувдан ўтмади")


def test_check_reports_resource_without_prices(tmp_path):
    res = build_res(tmp_path / "r.xlsx")
    wb = load_workbook(res)
    ws = wb.active
    for r in range(1, ws.max_row + 1):
        if isinstance(ws.cell(r, 7).value, (int, float)):
            ws.cell(r, 7).value = 0  # нархланмаган ресурс шаблони
    wb.save(res)
    result = run(build_local(tmp_path / "l.xlsx"), res, Options(mode="smart", output_path=tmp_path / "o.xlsx"))
    assert any("нархи 0" in c["problem"] for c in result.checks)


# ===================== smetachi ko'rsatgan ustunlar =====================
def test_manual_columns_without_header():
    """Sarlavhasiz va mazmundan ham topilmaydigan list: smetachi ko'rsatgan ustun/qatorlar olinadi."""
    ws = Workbook().active
    ws.append(["Ҳисоб-китоб"])
    for row in ROWS:
        ws.append(list(row[:5]) + [row[5]])
    assert detect_local_columns(ws) is None or detect_local_columns(ws).source == "content"
    cols = detect_local_columns(ws, {"name": 3, "unit": 4, "qty_unit": 5, "qty_total": 6, "start": 2, "end": 6})
    assert (cols.name, cols.unit, cols.qty_unit, cols.qty_total) == (3, 4, 5, 6)
    assert (cols.data_start, cols.data_stop, cols.source) == (2, 7, "manual")


def test_manual_columns_stop_at_total():
    """Oxirgi qator ko'rsatilmasa — «Итого по локальной смете» dan oldin tugaydi (u resurs emas)."""
    ws = Workbook().active
    ws.append(["Ҳисоб-китоб"])
    for row in ROWS:
        ws.append(list(row[:5]) + [row[5]])
    ws.append([])
    ws.append([None, "Итого по локальной смете"])
    cols = detect_local_columns(ws, {"name": 3, "unit": 4, "qty_unit": 5, "qty_total": 6, "start": 2})
    assert cols.data_stop == ws.max_row


def test_manual_qty_overrides_detected():
    ws, start = _sheet()
    cols = detect_local_columns(ws, {"qty_total": 5})  # фақат жами миқдор — E
    assert (cols.qty_unit, cols.qty_total, cols.data_start) == (None, 5, start)


def test_forced_local_on_priced_sheet():
    """Narxlangan ЛРВ («1.1, 1.2…», ЦЕНА bor) o'tkazib yuboriladi, smetachi «Локал» desa — lokal sifatida olinadi."""
    from core.batch import KIND_PRICED, LOCAL, UNKNOWN, inspect_sheet
    ws = Workbook().active
    ws.append(["№№", "Обоснование", "Наименование работ и ресурсов", "Ед.изм", "Кол-во", None, "Цена", "Сумма"])
    ws.append([None, None, None, None, "на единицу", "по проекту"])
    ws.append([1, 2, 3, 4, 5, 6, 7, 8])
    for row in ROWS:
        ws.append(list(row[:5]) + [row[5], 100, 200])
    auto = inspect_sheet(ws)
    assert (auto.role, auto.kind) == (UNKNOWN, KIND_PRICED)
    forced = inspect_sheet(ws, role=LOCAL)
    assert forced.role == LOCAL and forced.rows > 0
    assert (forced.columns["name"], forced.columns["unit"], forced.columns["qty_total"]) == (3, 4, 6)
    assert forced.headers["3"].startswith("наименование")


def test_manual_start_with_content_layout_stops_at_total():
    """Mazmundan topilgan jadval «Итого» ni hisobga olmasa ham, qo'lda boshlanish berilganda u chegara bo'ladi."""
    ws = Workbook().active
    ws.append(["Ҳисоб-китоб"])
    for row in ROWS:
        ws.append(list(row[:5]) + [row[5]])
    ws.append([None, "Итого по локальной смете"])
    cols = detect_local_columns(ws, {"start": 2})
    assert cols is None or cols.data_stop == ws.max_row
