"""Sintetik smeta fayllari: etalon fayllar bo'lmasa ham testlar ishlashi uchun.

Lokal list haqiqiy «Форма 5» tuzilishiga o'xshatilgan: sarlavhalar, birlashtirilgan
kataklar, qator uslubi (customFormat), maxsus kenglik/balandliklar, hair chiziqlar va
rus kodli header/footer (openpyxl uni o'qiy olmaydi — writer tiklashi kerak).
"""
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.writer import _insert_header_footer, _sheet_parts

RU_HEADER_FOOTER = (
    "<headerFooter><oddHeader>&amp;Л&amp;9ПРОГРАММНЫЙ КОМПЛЕКС&amp;Ц&amp;С&amp;П202611450</oddHeader>"
    "<oddFooter>&amp;ЦСтраниц - &amp;К</oddFooter></headerFooter>"
)

LOCAL_ROWS = [
    # (A, B, C nom, D birlik, E, F)
    ("1", "Е3704-001-07", "УСТРОЙСТВО ПРОЛЕТНЫХ СТРОЕНИЙ", "100М3", 0.025, None),
    ("1.1", "1", "ЗАТРАТЫ ТРУДА РАБОЧИХ-СТРОИТЕЛЕЙ", "ЧЕЛ.-Ч", 1588, 39.7),
    ("1.2", "3", "ЗАТРАТЫ ТРУДА МАШИНИСТОВ", "ЧЕЛ.-Ч", 110.15, 2.7538),
    ("1.3", "762", "КРАНЫ НА АВТОМОБИЛЬНОМ ХОДУ", "МАШ.-Ч", 9.92, 0.248),
    ("1.4", "6322", "БЕТОН ТЯЖЕЛЫЙ КЛАССА В15 /М-200/ ФРАКЦИИ 5-20 ММ", "М3", 101.5, 2.5375),
    ("1.5", "23469", "ВОДА", "М3", 175, 4.375),
    ("1.6", "44629", "ТКАНЬ МЕШОЧНАЯ", "10М2", 22.7, 0.5675),
    ("1.7", "99999", "НЕСУЩЕСТВУЮЩИЙ МАТЕРИАЛ", "ШТ", 1, 1),
]
FIRST_DATA_ROW = 8

RES_ROWS = [
    # (D nom, E birlik, G narx) — C oldiga bo'sh ustun qo'shilgan
    ("ЗАТРАТЫ ТРУДА РАБОЧИХ-СТРОИТЕЛЕЙ", "ЧЕЛ.-Ч", 29818.17),
    ("ЗАТРАТЫ ТРУДА МАШИНИСТОВ", "ЧЕЛ.-Ч", "--"),
    ("КРАНЫ НА АВТОМОБИЛЬНОМ ХОДУ", "МАШ.-Ч", 244250),
    ("БЕТОН ТЯЖЕЛЫЙ КЛАССА В15 /М-200/ ФРАКЦИИ 5-20ММ", "М3", 485994),  # fuzzy
    ("ВОДА", "М3", 1000),
    ("ТКАНЬ МЕШОЧНАЯ", "10М2", 105000),
]


def add_header_footer(path: Path, title: str, hf: str = RU_HEADER_FOOTER):
    with zipfile.ZipFile(path) as z:
        part = _sheet_parts(z)[title]
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zn:
        for info, data in items:
            if info.filename == part:
                data = _insert_header_footer(data.decode("utf-8"), hf).encode("utf-8")
            zn.writestr(info, data)
    tmp.replace(path)


def build_local(path: Path, numeric_only: bool = False) -> Path:
    """``numeric_only=True`` — sarlavhalarsiz, barcha narxlar topiladigan variant."""
    wb = Workbook()
    ws = wb.active
    ws.title = "202611450_БВ"
    base_font = Font(name="Times New Roman Cyr", sz=10)
    hair = Side(style="hair")

    widths = {"A": 5.43, "B": 13.57, "C": 82.86, "D": 10.14, "E": 10.14, "F": 10.14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].width = 9.14

    unpriced = ("НЕСУЩЕСТВУЮЩИЙ МАТЕРИАЛ", "ЗАТРАТЫ ТРУДА МАШИНИСТОВ")
    rows = [r for r in LOCAL_ROWS if not numeric_only or r[2] not in unpriced]
    last = FIRST_DATA_ROW + len(rows) - 1
    for r in range(1, last + 4):
        # Qator uslubi: Excel yangi katakka shuni beradi
        ws.row_dimensions[r].font = base_font
        ws.row_dimensions[r].alignment = Alignment(vertical="top")

    if not numeric_only:
        ws["F1"] = "Форма N 5"
        ws["B2"] = "IL \"URGUT TUMANIDAGI KANAL\""
        ws.merge_cells("B2:F2")
        # O'ng chiziq faqat oxirgi katakda: openpyxl yuklashda buni yo'qotadi (writer tiklashi kerak)
        ws["F2"].border = Border(right=Side(style="medium"), bottom=hair)
        ws["C2"].border = Border(bottom=hair)
        ws["C4"] = "ЛОКАЛЬНАЯ РЕСУРСНАЯ ВЕДОМОСТЬ  № "
        ws["D4"] = "1-6"
        ws.merge_cells("D4:F4")
        ws.row_dimensions[4].height = 15.75
        for col, text in zip("ABCDE", ["N п.п.", "Шифр", "Наименование работ и затрат", "Единица измерения", "Количество "]):
            c = ws[f"{col}5"]
            c.value = text
            c.font = Font(name="Times New Roman Cyr", sz=10, b=True)
            c.fill = PatternFill(fill_type="solid", fgColor="FFCCFFFF")
            c.border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells("A5:A6")
        ws.merge_cells("C5:C6")
        ws.merge_cells("E5:F5")
        ws["E6"] = "на. ед. измерения"
        ws["F6"] = "по проектным данным"
        for i, col in enumerate("ABCDEF", start=1):
            ws[f"{col}7"] = i

    for i, (a, b, c, d, e, f) in enumerate(rows):
        r = FIRST_DATA_ROW + i
        for col, v in zip("ABCDEF", (a, b, c, d, e, f)):
            cell = ws[f"{col}{r}"]
            cell.value = v
            cell.font = Font(name="Times New Roman Cyr", sz=10, b=(i == 0))
        ws[f"A{r}"].border = Border(left=hair, top=hair)
        ws[f"A{r}"].number_format = "@"
        ws[f"E{r}"].number_format = "#,##0.0######"
        ws[f"F{r}"].number_format = "#,##0.0######"
    ws.merge_cells(f"E{FIRST_DATA_ROW}:F{FIRST_DATA_ROW}")  # ish sarlavhasida E:F birlashgan
    ws.row_dimensions[FIRST_DATA_ROW + 2].height = 24
    ws[f"A{last + 1}"].border = Border(top=hair)  # pastki qo'shni: ramka uni olib tashlashi kerak

    wb.save(path)
    add_header_footer(path, ws.title)
    return path


def build_res(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "202611450_БР"
    ws["D1"] = "Наименование работ и затрат"
    ws["E1"] = "Единица измерения"
    ws["G1"] = "на.ед.изм."
    for i, (name, unit, price) in enumerate(RES_ROWS, start=3):
        ws[f"D{i}"], ws[f"E{i}"], ws[f"G{i}"] = name, unit, price
    wb.save(path)
    return path


@pytest.fixture
def sample_files(tmp_path):
    return build_local(tmp_path / "local.xlsx"), build_res(tmp_path / "res.xlsx")
