"""Smart rejim (reja §4, T1–T8).

Narx QAT'IY: faqat nom aynan mos kelsa va birlik bir xil bo'lsa. Qolgan hammasi narxlanmaydi
(G da matn, Excel'da qizil) — smetachi o'zi kiritadi.

Eng muhim tekshiruv — haqiqiy smetalar: I/J/K summalari resurs faylning o'zidagi
«ЗАРАБОТНАЯ ПЛАТА / ЭКСПЛУАТАЦИЯ МАШИН / МАТЕРИАЛЫ» summalariga teng bo'lishi kerak
(narxlanmagan qatorlarga smetachi taklif narxini kiritgan deb hisoblaganda).
"""
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.detect import detect_layout
from core.log_sheet import LOG_TITLE
from core.pipeline import CHECK_TEXT, MACHINIST, NO_PRICE, NO_PRICE_TEXT, UNIT_MISMATCH, WORK_HEADER, Options, run
from core.svodka import parse_text_number

from .conftest import FIRST_DATA_ROW, build_local, build_res

TESTS = Path(__file__).parent

# (papka, I zarplata, J mashinalar, K materiallar) — resurs fayldagi «ИТОГО» qatorlari
REAL_CASES = [
    ("excel_calc/bolnitsa7_kr_blok2", 238_735_601.25, 192_861_567.35, 705_961_343.89 + 404_493_302.15),
    ("excel_calc/most_202611450", 3_073_209.69105, 482_280.11371, 15_041_767.981375),
]


def _run(local, res, tmp_path, **kw):
    return run(local, res, Options(output_path=tmp_path / "out.xlsx", mode="smart", **kw))


@pytest.mark.parametrize("folder, salary, machines, materials", REAL_CASES, ids=lambda v: str(v).split("/")[-1])
def test_totals_match_resource_file(folder, salary, machines, materials, tmp_path):
    d = TESTS / folder
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    assert r.stats["NOT_FOUND"] == 0
    # resurs fayl summalarni qatorma-qator yaxlitlaydi: 1 so'mgacha farq bo'lishi mumkin
    assert r.totals["zarplata"] == pytest.approx(salary, abs=1)
    assert r.totals["mashina"] == pytest.approx(machines, abs=1)
    assert isinstance(r.totals["s_nds"], float)

    # Narxlanmagan qatorlar summaga kirmaydi; smetachi taklifni kiritsa — RES bilan teng
    ws = load_workbook(d / "input_local.xlsx")[r.local_sheet]
    added = 0.0
    for row in r.rows:
        if row["suggested_price"] is not None:
            assert row["h"] == 0 and row["kind"] == "other"
            f, e = ws.cell(row["row"], 6).value, ws.cell(row["row"], 5).value
            qty = parse_text_number(str(e if f in (None, "") else f))
            added += qty * row["suggested_price"]
    assert r.totals["materialy"] + added == pytest.approx(materials, abs=1)
    assert r.stats["NOT_PRICED"] == sum(1 for row in r.rows if row["suggested_price"] is not None)


def test_range_stops_before_itogo(tmp_path):
    """T3: «ИТОГО ПО ЛОКАЛЬНОЙ…» bo'limi narxlanmaydi, sarlavha qatorlari ham."""
    d = TESTS / "golden/202611450"
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    assert (r.start_row, r.end_row) == (15, 44)
    ws = load_workbook(r.output_path)["202611450_БВ"]
    assert ws["G46"].value is None
    # «ИТОГО ПО ЛОКАЛЬНОЙ…» qismi (46..77, «СОСТАВИЛ» bilan) olib tashlandi — svodka sum qatoridan keyin
    assert r.sum_row == 45 and r.tail_removed == 32
    assert ws[f"C{r.svodka_start}"].value == "Зарплата" and r.svodka_start == 46
    assert all(ws.cell(row, 2).value != "СОСТАВИЛ" for row in range(1, ws.max_row + 1))
    assert ws.max_row == 46 + 12


def test_price_and_amount_headers_and_filter(tmp_path):
    """Narx/summa ustunlari tepasida «Цена»/«Сумма» (jadval sarlavhasi uslubida), raqamlash
    qatori Q gacha davom etadi (7 … 17); yuqori sarlavha rangi summa ustunida tugaydi; filtr oxirgi resurs qatorigacha."""
    from core.writer import AMOUNT_COL, PRICE_COL
    d = TESTS / "golden/202611450"
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    ws = load_workbook(r.output_path)["202611450_БВ"]
    layout = detect_layout(load_workbook(d / "input_local.xlsx")["202611450_БВ"])
    head, num = layout.header_row, layout.numbering_row
    assert ws[f"{PRICE_COL}{head}"].value == "Цена" and ws[f"{AMOUNT_COL}{head}"].value == "Сумма"
    assert ws[f"{PRICE_COL}{head}"].fill.fgColor.rgb == ws[f"E{head}"].fill.fgColor.rgb == ws[f"H{head}"].fill.fgColor.rgb
    assert f"G{head}:G{num - 1}" in {str(m) for m in ws.merged_cells.ranges}
    assert ws[f"G{head + 1}"].value is None and ws[f"I{head}"].fill.fill_type is None  # sarlavha rangi summa ustunida tugaydi
    assert [ws[f"{c}{num}"].value for c in "FGHIQ"] == [6, 7, 8, 9, 17]  # raqamlash (filtr) qatori oxirigacha
    assert ws[f"Q{num}"].fill.fgColor.rgb == ws[f"G{num}"].fill.fgColor.rgb
    assert ws.auto_filter.ref == f"A{num}:Q{r.end_row}"  # filtr esa hamma ustunlar bo'ylab

    renamed = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path, column_names={"G": "Нарх"})
    assert load_workbook(renamed.output_path)["202611450_БВ"][f"G{head}"].value == "Нарх"


def test_rows_below_svodka_hidden(tmp_path):
    """Svodkadan keyin 5 bo'sh qator ko'rinadi, pasti yashirin (zeroHeight): ichidagi bo'sh
    qatorlar standart balandlikda qoladi, yozuvli qatorlarniki esa Excel'ning o'ziga qoldiriladi."""
    from core.pipeline import EMPTY_ROWS_AFTER
    d = TESTS / "golden/202611450"
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    ws = load_workbook(r.output_path)["202611450_БВ"]
    last = r.svodka_start + 12
    assert ws.sheet_format.zeroHeight
    assert all(ws.row_dimensions[row].height for row in range(last + 1, last + EMPTY_ROWS_AFTER + 1))
    assert last + EMPTY_ROWS_AFTER + 1 not in ws.row_dimensions
    src = load_workbook(d / "input_local.xlsx")["202611450_БВ"]
    assert ws.row_dimensions[20].height == src.row_dimensions[20].height


def test_cleared_not_found(sample_files, tmp_path):
    """«ТОПИЛМАДИ» олиб ташланса: бирлиги бор қаторда G бўш, лекин қизил (шартли формат) ва ҳали
    нархланмаган; бирлиги йўқ қаторда G, H ва категория бўш, нархланмаганлар қаторига кирмайди."""
    from core.pipeline import CLEARED
    local, res = sample_files
    row = FIRST_DATA_ROW + 7
    before = _run(local, res, tmp_path).stats["NOT_PRICED"]
    r = _run(local, res, tmp_path, cleared=[row])
    ws = load_workbook(r.output_path)["202611450_БВ"]
    assert r.cleared == [row] and ws[f"G{row}"].value is None and ws[f"H{row}"].value is not None
    assert any(f"G{row}" in str(cf.sqref) and str(cf.sqref) != f"G{r.start_row}:G{r.end_row}"
               for cf in ws.conditional_formatting)
    assert r.stats["NOT_PRICED"] == before and row in [p["row"] for p in r.problems]

    wb = load_workbook(local)
    ws = wb["202611450_БВ"]
    ws[f"D{row}"] = None  # бирлиги йўқ (бўлим номи каби)
    # охирги эмас: ундан кейин ресурс бор (охиридаги бирликсиз қаторлар — имзолар, нархланмайди)
    for col in "ABCDEF":
        ws[f"{col}{row + 1}"] = ws[f"{col}{FIRST_DATA_ROW + 5}"].value  # «ВОДА»
    ws[f"A{row + 1}"] = "1.8"
    wb.save(tmp_path / "nounit.xlsx")
    r = _run(tmp_path / "nounit.xlsx", res, tmp_path, cleared=[row])
    ws = load_workbook(r.output_path)["202611450_БВ"]
    st = {x["row"]: x["status"] for x in r.rows}
    assert st[row] == CLEARED and r.stats["NOT_PRICED"] == before - 1
    assert [ws[f"{c}{row}"].value for c in "GHIJK"] == [None] * 5


def test_not_found_suggests_nearest_resource_price(sample_files, tmp_path):
    """Topilmagan (o'xshashlik chegarasidan past) qatorga RES dagi eng yaqin resursning narxi taklif:
    G ga yozilmaydi; smetachi tasdiqlasa (prices) — narx RES dan aynan."""
    local, res = sample_files
    row = FIRST_DATA_ROW + 7
    wb = load_workbook(local)
    ws = wb["202611450_БВ"]
    ws[f"C{row}"], ws[f"D{row}"] = "ТКАНЬ МЕШОЧНАЯ ДВОЙНАЯ", "10М2"
    wb.save(tmp_path / "near.xlsx")
    r = _run(tmp_path / "near.xlsx", res, tmp_path)
    x = {y["row"]: y for y in r.rows}[row]
    assert x["status"] == "NOT_FOUND" and x["g"] == "ТОПИЛМАДИ"
    assert x["suggested_price"] == 105000 and x["near"]["same_unit"] and 0.5 <= x["near"]["score"] < 0.95
    accepted = _run(tmp_path / "near.xlsx", res, tmp_path, prices={row: 105000})
    assert load_workbook(accepted.output_path)["202611450_БВ"][f"G{row}"].value == 105000


def test_fuzzy_row_has_near_and_can_be_cleared(sample_files, tmp_path):
    """«ТЕКШИРИНГ» (o'xshash) qator ham 1-bosqichda: RES dagi nomzodi (asl yozuvda) va yozuvni olib tashlash."""
    local, res = sample_files
    r = _run(local, res, tmp_path)
    fuzzy = next(x for x in r.rows if x["status"] == "FUZZY")
    assert fuzzy["g"] == CHECK_TEXT and fuzzy["near"]["price"] == 485994
    assert fuzzy["near"]["name"] == "БЕТОН ТЯЖЕЛЫЙ КЛАССА В15 /М-200/ ФРАКЦИИ 5-20ММ"
    c = _run(local, res, tmp_path, cleared=[fuzzy["row"]])
    ws = load_workbook(c.output_path)["202611450_БВ"]
    assert ws[f"G{fuzzy['row']}"].value is None and c.cleared == [fuzzy["row"]]
    assert fuzzy["row"] in [p["row"] for p in c.problems]  # narx hali kerak — qizil


def test_rules_on_sample(sample_files, tmp_path):
    local, res = sample_files
    r = _run(local, res, tmp_path)
    st = {row["row"]: row["status"] for row in r.rows}
    first = FIRST_DATA_ROW
    assert st[first] == WORK_HEADER                 # «1 Е3704…», ostida 1.1, 1.2 …
    assert st[first + 2] == MACHINIST               # «ЗАТРАТЫ ТРУДА МАШИНИСТОВ», narx «--»
    assert st[first + 7] == "NOT_FOUND"

    ws = load_workbook(r.output_path)["202611450_БВ"]
    assert ws[f"G{first}"].value == 0 and ws[f"G{first + 2}"].value == 0
    assert ws[f"G{first + 7}"].value == "ТОПИЛМАДИ"
    assert ws[f"H{first + 7}"].value.startswith("=IF(ISNUMBER(G")
    # «БЕТОН … 5-20 ММ» resursda «5-20ММ» — ўхшаш, lekin aniq emas: narxlanmaydi, taklif bor
    assert st[first + 4] == "FUZZY" and ws[f"G{first + 4}"].value == CHECK_TEXT
    # topilmagan qator svodkani buzmaydi (T6), lekin muammolar ro'yxatida turadi
    assert isinstance(r.totals["s_nds"], float)
    assert [p["row"] for p in r.problems] == [first + 4, first + 7]
    assert r.problems[0]["suggested_price"] == 485994
    # Excel'da qizil: shartli format G ustunida, narx yozilsa o'zi yo'qoladi
    rules = [(str(cf.sqref), rule.formula[0]) for cf in ws.conditional_formatting for rule in cf.rules]
    assert rules == [(f"G{r.start_row}:G{r.end_row}", f'AND(G{r.start_row}<>"",NOT(ISNUMBER(G{r.start_row})))')]


def test_log_is_last_and_smeta_stays_active(sample_files, tmp_path):
    local, res = sample_files
    r = _run(local, res, tmp_path)
    wb = load_workbook(r.output_path)
    res_ws = load_workbook(res).active
    # LOG dan keyin — resurs listi (qiymatlari bilan), smetachi narxni shu faylda solishtiradi
    assert wb.sheetnames == ["202611450_БВ", LOG_TITLE, res_ws.title]
    copied = wb[res_ws.title]
    assert [[c.value for c in row] for row in copied.iter_rows(max_row=12)] ==            [[c.value for c in row] for row in res_ws.iter_rows(max_row=12)]
    assert wb.active.title == "202611450_БВ"
    log = wb[LOG_TITLE]
    assert log["K1"].value == "note" and log.auto_filter.ref == "A1:K1"
    statuses = {row[0]: row[3] for row in log.iter_rows(min_row=2, values_only=True)}
    assert statuses[FIRST_DATA_ROW] == WORK_HEADER


def _res_without_blank_column(path: Path, rows):
    """Asl «Форма 6а»: nom C, birlik D, narx F — bo'sh ustun qo'shilmagan."""
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"], ws["F1"] = (
        "N п.п.", "Шифр", "Наименование работ и затрат", "Единица измерения", "Количество ", "Сметная стоимость")
    ws.merge_cells("F1:I1")
    ws["F2"], ws["H2"] = "в базисном уровне", "в текущем (прогнозном)"
    ws.merge_cells("F2:G2")
    ws.merge_cells("H2:I2")
    ws["F3"], ws["G3"], ws["H3"], ws["I3"] = "на.ед.изм.", "общая", "на.ед.изм.", "общая"
    for i, col in enumerate("ABCDEFGHI", start=1):
        ws[f"{col}4"] = i
    for r, (name, unit, price) in enumerate(rows, start=5):
        ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"], ws[f"H{r}"] = name, unit, price, "--"
    wb.save(path)
    return path


def test_resource_columns_detected_without_blank_column(tmp_path):
    """T1: foydalanuvchi bo'sh ustun qo'shmasa ham narx «в базисном уровне / на.ед.изм.» dan olinadi."""
    from .conftest import RES_ROWS
    res = _res_without_blank_column(tmp_path / "res.xlsx", RES_ROWS)
    layout = detect_layout(load_workbook(res).active, need_price=True, stop_at_total=False)
    assert (layout.name_col, layout.unit_col, layout.price_col, layout.data_start) == (3, 4, 6, 5)

    local = build_local(tmp_path / "local.xlsx")
    a = _run(local, res, tmp_path)
    b = run(local, build_res(tmp_path / "res_blank.xlsx"), Options(output_path=tmp_path / "b.xlsx", mode="smart"))
    assert a.columns["res_price"] == "F" and b.columns["res_price"] == "G"
    assert a.totals == b.totals


def _mini(tmp_path, local_rows, res_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "БВ"
    ws["A1"], ws["C1"], ws["D1"], ws["E1"] = "N п.п.", "Наименование работ и затрат", "Единица измерения", "Кол-во"
    for i, col in enumerate("ABCDEF", start=1):
        ws[f"{col}2"] = i
    for r, (a, name, unit, f) in enumerate(local_rows, start=3):
        ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = a, name, unit, f
    wb.save(tmp_path / "l.xlsx")
    wb = Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["G1"] = "НАИМЕНОВАНИЕ РЕСУРСА", "ЕД.ИЗМ", "ЦЕНА"
    for r, (name, unit, price) in enumerate(res_rows, start=2):
        ws[f"D{r}"], ws[f"E{r}"], ws[f"G{r}"] = name, unit, price
    wb.save(tmp_path / "r.xlsx")
    return _run(tmp_path / "l.xlsx", tmp_path / "r.xlsx", tmp_path)


def test_no_price_unit_mismatch_and_text_price(tmp_path):
    r = _mini(tmp_path, [
        ("1", "ГВОЗДИ СТРОИТЕЛЬНЫЕ", "Т", 2),
        ("2", "ВОДА", "Т", 3),                 # resursda faqat М3 da bor
        ("3", "ЭЛЕКТРОДЫ", "Т", 1),
    ], [
        ("ГВОЗДИ СТРОИТЕЛЬНЫЕ", "Т", "--"),
        ("ВОДА", "М3", "1 000,5"),             # raqamli matn narx
        ("ЭЛЕКТРОДЫ", "Т", 10),
    ])
    st = {row["row"]: row for row in r.rows}
    assert st[3]["status"] == NO_PRICE and st[3]["g"] == NO_PRICE_TEXT and st[3]["h"] == 0
    # nom aniq, lekin birlik boshqa: narx yozilmaydi, faqat taklif
    assert st[4]["status"] == UNIT_MISMATCH and st[4]["g"] == CHECK_TEXT and st[4]["suggested_price"] == "1 000,5"
    assert st[5]["h"] == 10
    assert r.stats[NO_PRICE] == 1 and r.stats[UNIT_MISMATCH] == 1
    assert {p["row"] for p in r.problems} == {3, 4}


def test_not_exact_is_never_priced(tmp_path):
    """Ўхшаш ва ШАБҲАЛИ — narx bir xil bo'lsa ham — yozilmaydi (taxmin qilinmaydi)."""
    r = _mini(tmp_path, [
        ("1", "АРМАТУРА А500 Д12", "Т", 1),
        ("2", "КРАСКА МА-15", "Т", 1),
    ], [
        ("АРМАТУРА А500 Д12А", "Т", 4_600_000),
        ("АРМАТУРА А500 Д12Б", "Т", 4_600_000),
        ("КРАСКА МА-15А", "Т", 1),
        ("КРАСКА МА-15Б", "Т", 2),
    ])
    st = {row["row"]: row for row in r.rows}
    for row in (3, 4):
        assert st[row]["status"] == "AMBIGUOUS" and st[row]["g"] == CHECK_TEXT and st[row]["h"] == 0
    assert st[3]["suggested_price"] == 4_600_000 and "2 та номзод" in st[3]["note"]
    assert isinstance(r.totals["s_nds"], float) and r.stats["NOT_PRICED"] == 2


def test_number_columns_widened_so_excel_shows_numbers(tmp_path):
    """Milliardli summalar «######» bo'lib qolmasin: G..K kengaytiriladi, A–F va L:XFD tegilmaydi."""
    from core.writer import column_width, number_text
    d = TESTS / "excel_calc/bolnitsa7_kr_blok2"
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    ws_in = load_workbook(d / "input_local.xlsx")["LRV"]
    ws = load_workbook(r.output_path)["LRV"]
    assert column_width(ws_in, 8) < 9  # kirishda H ~8.8 — «2 106 365 947» sig'maydi
    need = len(number_text(r.totals["s_nds"]))
    assert column_width(ws, 8) >= need
    assert r.columns["widened"].startswith("G, H")
    for col in range(1, 7):
        assert column_width(ws, col) == column_width(ws_in, col)
    assert column_width(ws, 30) == column_width(ws_in, 30)  # G:XFD guruhining qolgan qismi


def test_custom_percents(tmp_path):
    """Foizlar o'zgartiriladi; НДС formulasi D ga bog'langan — Excel'da ham o'zgartirish mumkin."""
    d = TESTS / "golden/202611450"
    base = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    r = run(d / "input_local.xlsx", d / "input_res.xlsx", Options(
        mode="smart", output_path=tmp_path / "p.xlsx", percents={"prochie": 0.1727, "s_nds": 0.20}))
    t = r.totals
    assert t["pryamye"] == pytest.approx(base.totals["pryamye"])
    assert t["prochie"] == pytest.approx((t["pryamye"] - t["transport_obor"] - t["oborudovanie"]) * 0.1727)
    assert t["s_nds"] == pytest.approx(t["bez_nds"] * 1.20)
    assert r.percents["prochie"] == 0.1727 and r.percents["transport_mat"] == 0.03

    ws = load_workbook(r.output_path)["202611450_БВ"]
    s = r.svodka_start
    assert ws[f"D{s + 9}"].value == 0.1727 and ws[f"D{s + 12}"].value == 0.2
    assert ws[f"H{s + 12}"].value == f"=H{s + 11}*(1+D{s + 12})"
    assert ws[f"C{s + 12}"].value == "Всего с НДС      20 %"

    with pytest.raises(ValueError, match="0 dan 100%"):
        run(d / "input_local.xlsx", d / "input_res.xlsx",
            Options(mode="smart", output_path=tmp_path / "x.xlsx", percents={"s_nds": 12}))


def test_compat_keeps_old_percents(tmp_path):
    d = TESTS / "golden/202611450"
    r = run(d / "input_local.xlsx", d / "input_res.xlsx",
            Options(output_path=tmp_path / "c.xlsx", percents={"prochie": 0.5}))
    ws = load_workbook(r.output_path)["202611450_БВ"]
    s = r.svodka_start
    assert ws[f"D{s + 9}"].value == 0.19 and ws[f"H{s + 12}"].value == f"=H{s + 11}*1.12"


# ===================== kategoriyalar: I/J/K/L/M =====================
def _equipment_case(tmp_path):
    """RES da bo'limlar: ТРУДОВЫЕ / МАТЕРИАЛЬНЫЕ / ОБОРУДОВАНИЕ / КАБЕЛЬНАЯ ПРОДУКЦИЯ."""
    wb = Workbook()
    ws = wb.active
    ws.title = "БВ"
    ws["A1"], ws["C1"], ws["D1"], ws["E1"] = "N п.п.", "Наименование работ и затрат", "Единица измерения", "Кол-во"
    for i, col in enumerate("ABCDEF", start=1):
        ws[f"{col}2"] = i
    local = [
        ("1", "ЗАТРАТЫ ТРУДА РАБОЧИХ-СТРОИТЕЛЕЙ", "ЧЕЛ-Ч", 10),  # 3 -> I
        ("2", "БЕТОН В20", "М3", 2),                           # 4 -> K
        ("3", "НАСОС ЦЕНТРОБЕЖНЫЙ К-50", "ШТ", 1),             # 5 -> K, RES: ОБОРУДОВАНИЕ -> taklif M
        ("4", "КАБЕЛЬ ВВГ 3Х2,5", "М", 100),                   # 6 -> K, RES: КАБЕЛЬ -> taklif L
        ("5", "ЛИФТ ПАССАЖИРСКИЙ", "ШТ", 1),                   # 7 -> K, RES da yo'q, nomdan -> taklif M
        ("6", "ПРОВОЛОКА ВЯЗАЛЬНАЯ", "Т", 0.5),                # 8 -> K, «проволока» kabel emas
    ]
    for r, (a, name, unit, f) in enumerate(local, start=3):
        ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = a, name, unit, f
    wb.save(tmp_path / "l.xlsx")

    wb = Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["G1"] = "НАИМЕНОВАНИЕ РЕСУРСА", "ЕД.ИЗМ", "ЦЕНА"
    res = [
        ("ТРУДОВЫЕ РЕСУРСЫ", None, None),
        ("ЗАТРАТЫ ТРУДА РАБОЧИХ-СТРОИТЕЛЕЙ", "ЧЕЛ-Ч", 15000),
        ("СТРОИТЕЛЬНЫЕ МАТЕРИАЛЫ И КОНСТРУКЦИИ", None, None),
        ("БЕТОН В20", "М3", 400000),
        ("ПРОВОЛОКА ВЯЗАЛЬНАЯ", "Т", 6000000),
        ("ИТОГО", "СУМ", 1),
        ("ОБОРУДОВАНИЕ", None, None),
        ("НАСОС ЦЕНТРОБЕЖНЫЙ К-50", "ШТ", 9000000),
        ("КАБЕЛЬНО-ПРОВОДНИКОВАЯ ПРОДУКЦИЯ", None, None),
        ("КАБЕЛЬ ВВГ 3Х2,5", "М", 12000),
    ]
    for r, (name, unit, price) in enumerate(res, start=2):
        if unit is None:
            ws[f"A{r}"] = name   # bo'lim sarlavhasi (birliksiz)
        else:
            ws[f"D{r}"], ws[f"E{r}"], ws[f"G{r}"] = name, unit, price
    wb.save(tmp_path / "r.xlsx")
    return tmp_path / "l.xlsx", tmp_path / "r.xlsx"


def test_equipment_and_cable_only_suggested_from_res_sections(tmp_path):
    """I/J/K RES bo'limi bo'yicha o'zi joylashadi, L/M esa faqat TAKLIF: НАСОС (ОБОРУДОВАНИЕ) va
    КАБЕЛЬ (КАБЕЛЬНАЯ ПРОДУКЦИЯ) K da qoladi, «M/L ga o'tkazamizmi?» deb so'raladi."""
    local, res = _equipment_case(tmp_path)
    r = _run(local, res, tmp_path)
    rows = {x["row"]: x for x in r.rows}
    assert [rows[n]["category"] for n in range(3, 9)] == ["I", "K", "K", "K", "K", "K"]
    assert rows[5]["suggested_category"] == "M" and "ОБОРУДОВАНИЕ" in rows[5]["suggest_reason"]
    assert rows[6]["suggested_category"] == "L" and "КАБЕЛЬ" in rows[6]["suggest_reason"]
    # RES da yo'q (лифт) — faqat nomdan taklif
    assert rows[7]["category"] == "K" and rows[7]["suggested_category"] == "M" and "лифт" in rows[7]["suggest_reason"]
    assert rows[8]["suggested_category"] is None and rows[4]["suggested_category"] is None
    assert r.totals["oborudovanie"] == pytest.approx(0)
    assert r.totals["kabel"] == pytest.approx(0)
    assert r.totals["materialy"] == pytest.approx(2 * 400_000 + 9_000_000 + 100 * 12_000 + 0.5 * 6_000_000)


def test_rows_carry_res_section_order(tmp_path):
    """Sayt jadvali RES tartibida guruhlanishi uchun: har qatorda RES bo'limi va RES qatori."""
    local, res = _equipment_case(tmp_path)
    r = _run(local, res, tmp_path)
    assert [(s["title"], s["column"]) for s in r.res_sections] == [
        ("ТРУДОВЫЕ РЕСУРСЫ", "I"), ("СТРОИТЕЛЬНЫЕ МАТЕРИАЛЫ И КОНСТРУКЦИИ", "K"),
        ("ОБОРУДОВАНИЕ", "M"), ("КАБЕЛЬНО-ПРОВОДНИКОВАЯ ПРОДУКЦИЯ", "L")]
    rows = {x["row"]: x for x in r.rows}
    assert [rows[n]["res_section"] for n in range(3, 9)] == [0, 1, 2, 3, None, 1]
    assert rows[4]["res_row"] == 5 and rows[8]["res_row"] == 6


def test_section_titles_marked_not_found(tmp_path):
    """«ВЕНТИЛЯЦИЯ», «ФАСОННЫЕ ЧАСТИ» (№ va birliksiz) — oldingidek ТОПИЛМАДИ (qizil): smetachi o'zi ko'radi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "БВ"
    ws["A1"], ws["C1"], ws["D1"], ws["E1"] = "N п.п.", "Наименование работ и затрат", "Единица измерения", "Кол-во"
    for i, col in enumerate("ABCDEF", start=1):
        ws[f"{col}2"] = i
    ws["C3"] = "ВЕНТИЛЯЦИЯ"
    ws["A4"], ws["C4"], ws["D4"], ws["F4"] = "1", "БЕТОН В20", "М3", 2
    ws["C5"] = "ФАСОННЫЕ ЧАСТИ"
    ws["A6"], ws["C6"], ws["D6"], ws["F6"] = "2", "ПРОВОЛОКА ВЯЗАЛЬНАЯ", "Т", 0.5
    wb.save(tmp_path / "l.xlsx")
    wb = Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["G1"] = "НАИМЕНОВАНИЕ РЕСУРСА", "ЕД.ИЗМ", "ЦЕНА"
    ws["D2"], ws["E2"], ws["G2"] = "БЕТОН В20", "М3", 400000
    ws["D3"], ws["E3"], ws["G3"] = "ПРОВОЛОКА ВЯЗАЛЬНАЯ", "Т", 6000000
    wb.save(tmp_path / "r.xlsx")

    r = _run(tmp_path / "l.xlsx", tmp_path / "r.xlsx", tmp_path)
    assert r.start_row == 3 and r.stats["NOT_FOUND"] == 2
    assert [x[0] for x in r.log] == [3, 4, 5, 6]
    ws = load_workbook(r.output_path)["БВ"]
    assert ws["G3"].value == "ТОПИЛМАДИ" and ws["G5"].value == "ТОПИЛМАДИ"
    assert ws["A5"].fill.fill_type is None  # sariq bo'yalmaydi
    assert ws["H4"].value and ws["K6"].value == "=H6"


def test_manual_categories_override_auto_sections(tmp_path):
    """Qo'lda belgilangan ustun ustun turadi: smetachi taklifni tasdiqlaydi (НАСОС -> M, КАБЕЛЬ -> L)
    va o'zi ham belgilaydi (БЕТОН -> M, ПРОВОЛОКА -> L)."""
    local, res = _equipment_case(tmp_path)
    r = run(local, res, Options(mode="smart", output_path=tmp_path / "o.xlsx",
                                categories={4: "M", "8": "L", 5: "M", 6: "L"}))
    rows = {x["row"]: x for x in r.rows}
    assert rows[4]["category"] == "M" and rows[4]["override"] and rows[4]["default_category"] == "K"
    assert rows[8]["category"] == "L" and rows[8]["override"] and rows[8]["default_category"] == "K"
    # tasdiqlangan takliflar: endi taklif qolmaydi
    assert rows[5]["category"] == "M" and rows[5]["override"] and rows[5]["suggested_category"] is None
    assert rows[6]["category"] == "L" and rows[6]["override"] and rows[6]["suggested_category"] is None
    assert r.categories == {4: "M", 8: "L", 5: "M", 6: "L"}

    ws = load_workbook(r.output_path)["БВ"]
    assert (ws["K4"].value, ws["M4"].value) == (None, "=H4")
    assert (ws["K5"].value, ws["M5"].value) == (None, "=H5")
    assert (ws["K6"].value, ws["L6"].value) == (None, "=H6")
    assert (ws["K8"].value, ws["L8"].value) == (None, "=H8")

    t = r.totals
    # ОБОРУДОВАНИЕ: НАСОС + БЕТОН
    assert t["oborudovanie"] == pytest.approx(9_000_000 + 2 * 400_000)
    assert t["transport_obor"] == pytest.approx((9_000_000 + 2 * 400_000) * 0.02)
    # КАБЕЛЬ: КАБЕЛЬ + ПРОВОЛОКА
    assert t["kabel"] == pytest.approx(100 * 12_000 + 0.5 * 6_000_000)
    assert t["transport_kabel"] == pytest.approx((100 * 12_000 + 0.5 * 6_000_000) * 0.02)
    # K da faqat narxsiz ЛИФТ qoldi (H=0) — material summasi nol
    assert t["materialy"] == pytest.approx(0)


def test_manual_category_validation_and_compat(tmp_path):
    local, res = _equipment_case(tmp_path)
    with pytest.raises(ValueError, match="kategoriya"):
        run(local, res, Options(mode="smart", output_path=tmp_path / "x.xlsx", categories={5: "Z"}))
    c = run(local, res, Options(output_path=tmp_path / "c.xlsx", categories={5: "M"}))  # compat e'tiborsiz
    assert all(x["category"] in "IJK" for x in c.rows) and c.categories == {}


def test_manual_prices_and_accepted_suggestions(tmp_path):
    """Qizil qatorga narx kiritish yoki taklifni tasdiqlash: narx yoziladi, qizil yo'qoladi."""
    d = TESTS / "excel_calc/bolnitsa7_kr_blok2"
    base = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    red = [x for x in base.rows if x["suggested_price"] is not None]
    assert len(red) == 12
    accept = {x["row"]: parse_text_number(str(x["suggested_price"])) for x in red}
    accept[red[0]["row"]] = 4_500_000  # bittasini smetachi boshqa narx bilan to'g'irlaydi

    r = run(d / "input_local.xlsx", d / "input_res.xlsx",
            Options(mode="smart", output_path=tmp_path / "m.xlsx", prices=accept))
    rows = {x["row"]: x for x in r.rows}
    assert r.stats["NOT_PRICED"] == 0 and r.stats["MANUAL"] == 12 and r.problems == []
    first = rows[red[0]["row"]]
    assert first["status"] == "MANUAL" and first["g"] == 4_500_000 and first["orig_status"] == "AMBIGUOUS"
    assert first["note"].startswith("нарх қўлда киритилди")
    assert rows[red[1]["row"]]["note"].startswith("таклиф тасдиқланди")
    assert r.prices == accept
    wb = load_workbook(r.output_path)
    second_row = red[1]["row"]
    assert wb["LRV"][f"G{second_row}"].value == 4_600_000
    log_status = {row[0]: row[3] for row in wb["LOG"].iter_rows(min_row=2, values_only=True)}
    assert log_status[second_row] == "MANUAL"

    # Hammasi taklif bo'yicha bo'lganda — RES dagi materiallar summasiga teng
    all_ok = run(d / "input_local.xlsx", d / "input_res.xlsx", Options(
        mode="smart", output_path=tmp_path / "a.xlsx",
        prices={x["row"]: parse_text_number(str(x["suggested_price"])) for x in red}))
    assert all_ok.totals["materialy"] == pytest.approx(705_961_343.89 + 404_493_302.15, abs=1)

    with pytest.raises(ValueError, match="manfiy"):
        run(d / "input_local.xlsx", d / "input_res.xlsx",
            Options(mode="smart", output_path=tmp_path / "x.xlsx", prices={red[0]["row"]: -1}))


def test_run_many_prices_every_local_sheet_in_one_file(tmp_path):
    """Bitta obyekt faylida 4 ta lokal: hammasi bitta faylda narxlanadi, har biriga o'z LOG'i."""
    from core.pipeline import run_many
    b = TESTS / "batch"
    local, res2 = b / "obekt_lokal_7list.xlsx", b / "obekt_res_7list.xlsx"
    steps = [(local, Options(mode="smart", local_sheet=f"2026114{n}_БВ", res_sheet=f"2026114{n}_БР"))
             for n in ("00", "10", "20")]
    steps.append((res2, Options(mode="smart", local_sheet="202611450_БВ", res_sheet="202611450_БР")))
    results = run_many(local, steps, tmp_path / "hammasi.xlsx")
    wb = load_workbook(tmp_path / "hammasi.xlsx")
    assert [s for s in wb.sheetnames if s.startswith("LOG")] == [
        "LOG 202611400_БВ", "LOG 202611410_БВ", "LOG 202611420_БВ", "LOG 202611450_БВ"]
    for r in results:
        ws = wb[r.local_sheet]
        assert ws[f"C{r.svodka_start}"].value == "Зарплата"
        assert ws[f"H{r.start_row + 1}"].value.startswith("=IF(ISNUMBER(")
    # МОСТ smetasi alohida narxlangandagi bilan bir xil
    alone = _run(TESTS / "golden/202611450/input_local.xlsx", TESTS / "golden/202611450/input_res.xlsx", tmp_path)
    assert results[-1].totals["s_nds"] == pytest.approx(alone.totals["s_nds"])


def test_formula_results_are_stored_for_protected_view(tmp_path):
    """Brauzerdan yuklangan fayl «Защищённый просмотр»da ochiladi va qayta hisoblanmaydi —
    shuning uchun har bir formulaning natijasi faylda saqlangan bo'lishi kerak."""
    d = TESTS / "excel_calc/bolnitsa7_kr_blok2"
    r = _run(d / "input_local.xlsx", d / "input_res.xlsx", tmp_path)
    ws = load_workbook(r.output_path, data_only=True)["LRV"]
    for row in r.rows:
        if row["h"] is not None:
            assert ws[f"H{row['row']}"].value == pytest.approx(row["h"]), row["row"]
            col = row["category"]
            assert ws[f"{col}{row['row']}"].value == pytest.approx(row["h"])
    s = r.svodka_start
    for i, (key, value) in enumerate(r.totals.items()):
        assert ws[f"H{s + i}"].value == pytest.approx(value), key
    assert ws[f"I{r.sum_row}"].value == pytest.approx(r.totals["zarplata"])
    # formulalar o'zi ham joyida (Excel tahrirlashda qayta hisoblaydi)
    assert load_workbook(r.output_path)["LRV"][f"H{s + 12}"].value == f"=H{s + 11}*(1+D{s + 12})"


def test_run_many_keeps_results_of_every_sheet(tmp_path):
    from core.pipeline import run_many
    f = TESTS / "batch/obekt_lokal_7list.xlsx"
    steps = [(f, Options(mode="smart", local_sheet=f"2026114{n}_БВ", res_sheet=f"2026114{n}_БР")) for n in ("00", "10", "20")]
    results = run_many(f, steps, tmp_path / "all.xlsx")
    wb = load_workbook(tmp_path / "all.xlsx", data_only=True)
    for r in results:
        cell = wb[r.local_sheet][f"H{r.svodka_start + 12}"].value
        assert cell == pytest.approx(r.totals["s_nds"]), r.local_sheet


def test_output_opens_on_the_priced_sheet(tmp_path):
    """Manba fayl «202611450_БВ» da ochiladigan qilib saqlangan; 410 qismining fayli esa
    narxlangan «202611410_БВ» da ochilishi kerak (aks holda smetachi bo'sh listni ko'radi)."""
    from core.pipeline import run_many
    f = TESTS / "batch/obekt_lokal_7list.xlsx"
    assert load_workbook(f).active.title == "202611450_БВ"
    r = run(f, f, Options(mode="smart", local_sheet="202611410_БВ", res_sheet="202611410_БР",
                          output_path=tmp_path / "p.xlsx"))
    wb = load_workbook(r.output_path)
    assert wb.active.title == "202611410_БВ"
    assert [s.title for s in wb.worksheets if s.sheet_view.tabSelected] == ["202611410_БВ"]

    steps = [(f, Options(mode="smart", local_sheet=f"2026114{n}_БВ", res_sheet=f"2026114{n}_БР")) for n in ("10", "20")]
    run_many(f, steps, tmp_path / "all.xlsx")
    assert load_workbook(tmp_path / "all.xlsx").active.title == "202611410_БВ"


def test_custom_columns_named_and_in_svodka(tmp_path):
    """Smetachi N ustunini qo'shib nomlaydi, qatorlarni unga o'tkazadi: H shu ustunga tushadi,
    sum qatori va svodkada alohida qator (Всего прямых затрат ga qo'shiladi)."""
    local, res = _equipment_case(tmp_path)
    names = {"N": "Монтаж ишлари", "K": "Материаллар"}
    r = run(local, res, Options(mode="smart", output_path=tmp_path / "o.xlsx", column_names=names,
                                categories={4: "N", 8: "N"}))
    assert r.extra_columns == [("N", "Монтаж ишлари")]
    rows = {x["row"]: x for x in r.rows}
    assert rows[4]["category"] == "N" and rows[8]["category"] == "N"
    ws = load_workbook(r.output_path)["БВ"]
    assert (ws["K4"].value, ws["N4"].value) == (None, "=H4")
    assert ws[f"N{r.sum_row}"].value == f"=SUM(N{r.start_row}:N{r.end_row})"
    assert ws["N1"].value == "Монтаж ишлари" and ws["K1"].value == "Материаллар"  # jadval sarlavhasi
    n = 2 * 400_000 + 0.5 * 6_000_000
    assert r.totals["extra_N"] == pytest.approx(n)
    labels = [ws.cell(row, 3).value for row in range(r.svodka_start, r.svodka_start + 14)]
    assert labels.index("Монтаж ишлари") == labels.index("Всего прямых  затрат") - 1
    without = run(local, res, Options(mode="smart", output_path=tmp_path / "p.xlsx"))
    assert r.totals["pryamye"] == pytest.approx(without.totals["pryamye"]
                                                 - n * 0.03)  # N da qoladi, faqat material transporti (3%) tushadi
    assert ws[f"G4"].number_format == '#,##0;-#,##0;"-"'  # kasrsiz, 0 o'rniga «-»
    with pytest.raises(ValueError, match="kategoriya"):
        run(local, res, Options(mode="smart", output_path=tmp_path / "x.xlsx", categories={4: "O"}))


def test_hidden_local_sheet_is_priced_and_shown(tmp_path):
    """Smetachi asl listni yashirib, nusxasi bilan ishlagan bo'lishi mumkin: yashirin list narxlansa
    yiqilmasin («Only visible sheets can be made active») va natijada ko'rinadigan bo'lsin."""
    f = TESTS / "batch/obekt_lokal_7list.xlsx"
    wb = load_workbook(f)
    wb.active = wb["202611450_БВ"]
    wb["202611410_БВ"].sheet_state = "hidden"
    src = tmp_path / "hidden.xlsx"
    wb.save(src)
    r = run(src, src, Options(mode="smart", local_sheet="202611410_БВ", res_sheet="202611410_БР",
                              output_path=tmp_path / "p.xlsx"))
    out = load_workbook(r.output_path)
    assert out["202611410_БВ"].sheet_state == "visible"
    assert out.active.title == "202611410_БВ"


def _coded_local(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "БВ"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = "N п.п.", "Шифр", "Наименование работ и затрат", "Единица измерения", "Кол-во"
    for i, col in enumerate("ABCDEF", start=1):
        ws[f"{col}2"] = i
    for r, (code, name, unit, f) in enumerate(rows, start=3):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = str(r - 2), code, name, unit, f
    wb.save(path)
    return path


BASE = {"258": {"price": 275_000.0, "n": 3, "unit_key": "машч", "unit": "МАШ.-Ч", "category": "J"},
        "43113": {"price": 150_000.0, "n": 1, "unit_key": "м3", "unit": "М3", "category": "K"},
        "9219": {"price": 5_000.0, "n": 2, "unit_key": "т", "unit": "Т", "category": "K"}}


def test_base_prices_rows_res_could_not(tmp_path):
    from core.pipeline import BASE as BASE_ST
    local = _coded_local(tmp_path / "l.xlsx", [
        ("000258", "БУЛЬДОЗЕРЫ 79 КВТ", "МАШ.-Ч", 2),   # RES'да йўқ — базадан
        ("43113", "ЩЕБЕНЬ ФРАКЦИИ 20-40", "М3", 3),     # RES нархламади (ном аниқ эмас) — базадан
        ("1", "ЗАТРАТЫ ТРУДА РАБОЧИХ", "ЧЕЛ.-Ч", 10),    # RES'да аниқ — RES нархи
        ("9219", "ВОДА", "М3", 4),                      # базада бор, лекин бирлик бошқа — нархланмайди
        ("С1", "АРМАТУРА", "Т", 1),                     # кодсиз
    ])
    wb = Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["G1"] = "НАИМЕНОВАНИЕ РЕСУРСА", "ЕД.ИЗМ", "ЦЕНА"
    ws["D2"], ws["E2"], ws["G2"] = "ЗАТРАТЫ ТРУДА РАБОЧИХ", "ЧЕЛ.-Ч", 36_000
    ws["D3"], ws["E3"], ws["G3"] = "ЩЕБЕНЬ ФРАКЦИИ 20-40 ММ", "М3", 140_000
    wb.save(tmp_path / "r.xlsx")
    r = _run(local, tmp_path / "r.xlsx", tmp_path, base_prices=BASE, base_label="Тошкент ш. · 2026-4")
    st = {row["row"]: row for row in r.rows}
    assert st[3]["status"] == BASE_ST and st[3]["g"] == 275_000 and st[3]["h"] == 550_000 and st[3]["category"] == "J"
    assert st[3]["orig_status"] == "NOT_FOUND" and "Тошкент ш. · 2026-4, 3 та смета" in st[3]["note"]
    assert st[4]["status"] == BASE_ST and st[4]["g"] == 150_000 and st[4]["orig_status"] != "EXACT"
    assert st[5]["status"] == "EXACT" and st[5]["g"] == 36_000
    assert st[6]["status"] == "NOT_FOUND" and "бирлик бошқа" in st[6]["note"]
    assert st[7]["status"] == "NOT_FOUND"
    assert r.stats[BASE_ST] == 2 and r.stats["NOT_PRICED"] == 2
    # қўлда нарх базадан устун
    r = _run(local, tmp_path / "r.xlsx", tmp_path, base_prices=BASE, prices={3: 1})
    assert {row["row"]: row for row in r.rows}[3]["status"] == "MANUAL"


def test_without_res_only_base(tmp_path):
    from core.pipeline import BASE as BASE_ST
    local = _coded_local(tmp_path / "l.xlsx", [("258", "БУЛЬДОЗЕРЫ", "МАШ.-Ч", 2), ("777", "НОМАЪЛУМ", "Т", 1)])
    r = _run(local, None, tmp_path, base_prices=BASE)
    st = {row["row"]: row for row in r.rows}
    assert st[3]["status"] == BASE_ST and st[3]["h"] == 550_000
    assert st[4]["status"] == "NOT_FOUND"
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["БВ", LOG_TITLE]  # RES листи қўшилмайди
    with pytest.raises(ValueError):
        run(local, None, Options(output_path=tmp_path / "o2.xlsx", mode="compat"))


def test_signature_rows_after_last_resource_not_priced(tmp_path):
    """Smeta oxiridagi imzolar («ГИП», «СОСТАВИЛ») birliksiz — narxlanmaydi, «ТОПИЛМАДИ» yozilmaydi;
    o'rtadagi birliksiz bo'lim nomi esa avvalgidek qoladi."""
    r = _mini(tmp_path, [
        ("", "РАЗДЕЛ 1. ЗЕМЛЯНЫЕ РАБОТЫ", None, None),
        ("1", "ЭЛЕКТРОДЫ", "Т", 1),
        (None, None, None, None),
        (None, "ГИП                     Ш. Жумаев", None, None),
        (None, "СОСТАВИЛ                К. Карабеков", None, None),
    ], [("ЭЛЕКТРОДЫ", "Т", 10)])
    st = {row["row"]: row for row in r.rows}
    assert r.end_row == 4 and set(st) == {3, 4}
    assert st[4]["status"] == "EXACT"
    assert st[3]["status"] == "NOT_FOUND"  # бўлим номи — смета­чи 1-босқичда тозалайди
    assert r.sum_row == 5
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb["БВ"]["G6"].value is None and wb["БВ"]["G7"].value is None
    assert wb["БВ"]["C6"].value.startswith("ГИП")
