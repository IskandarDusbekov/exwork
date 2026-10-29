from openpyxl import Workbook

from core.matcher import AMBIGUOUS, EXACT, FUZZY, NOT_FOUND, ResourceIndex, best_two_matches
from core.normalize import norm_name, norm_unit


def _index(rows):
    """rows: (nom, birlik, narx) — resurs listining D, E, G ustunlari."""
    wb = Workbook()
    ws = wb.active
    for i, (name, unit, price) in enumerate(rows, start=1):
        ws.cell(row=i, column=4, value=name)
        ws.cell(row=i, column=5, value=unit)
        ws.cell(row=i, column=7, value=price)
    return ResourceIndex.from_sheet(ws, "D", "G")


def _match(index, name, unit, fuzzy=95):
    return index.match(norm_name(name), norm_unit(unit), fuzzy / 100)


def test_index_skips_empty_and_keeps_first_duplicate():
    idx = _index([
        ("Наименование", None, "x"),
        (None, "М3", 1),
        ("ВОДА", "М3", 1000),
        ("вода", "м3", 2000),
    ])
    assert idx.by_unit == {"м3": {"вода": 1000}}


def test_exact():
    idx = _index([("ВОДА", "М3", 1000)])
    m = _match(idx, " Вода ", "м3")
    assert (m.status, m.price, m.unit, m.score) == (EXACT, 1000, "м3", 1.0)


def test_exact_in_other_unit_wins():
    idx = _index([("ВОДА", "Т", 7)])
    m = _match(idx, "ВОДА", "М3")
    assert (m.status, m.unit, m.price) == (EXACT, "т", 7)


def test_fuzzy():
    idx = _index([("БЕТОН ТЯЖЕЛЫЙ КЛАССА В15 /М-200/ ФРАКЦИИ 5-20 ММ", "М3", 485994)])
    m = _match(idx, "БЕТОН ТЯЖЕЛЫЙ КЛАССА В15 /М-200/ ФРАКЦИИ 5-20ММ", "М3")
    assert m.status == FUZZY and m.price == 485994 and 0.95 <= m.score < 1


def test_ambiguous():
    idx = _index([
        ("КРАСКА МАСЛЯНАЯ МА-15А", "Т", 1),
        ("КРАСКА МАСЛЯНАЯ МА-15Б", "Т", 2),
    ])
    m = _match(idx, "КРАСКА МАСЛЯНАЯ МА-15", "Т")
    assert m.status == AMBIGUOUS and {m.price, m.second_price} == {1, 2}
    assert m.second_name is not None and m.score - m.second_score < 0.01


def test_not_found_below_threshold():
    idx = _index([("ВОДА", "М3", 1000)])
    assert _match(idx, "ГВОЗДИ СТРОИТЕЛЬНЫЕ", "Т").status == NOT_FOUND


def test_best_two_matches_order():
    best, bs, second, ss = best_two_matches("abcd", ["abcx", "abcd", "zzzz"])
    assert best == "abcd" and bs == 1.0 and second == "abcx" and ss == 0.75


def test_not_found_keeps_nearest_for_suggestion():
    """Chegaradan past bo'lsa ham eng yaqin resurs eslab qolinadi; birligi bir xili ustun,
    juda o'xshamasa (< 50%) — taklif yo'q."""
    idx = _index([("ГАЙКА М20 ОЦИНКОВАННАЯ", "ШТ", 1500), ("ГАЙКА М22", "КГ", 99000), ("ЦЕМЕНТ", "Т", 900000)])
    m = _match(idx, "ГАЙКА М22 ОЦИНК", "ШТ")
    assert m.status == NOT_FOUND
    assert (m.near_name, m.near_unit, m.near_price) == (norm_name("ГАЙКА М20 ОЦИНКОВАННАЯ"), norm_unit("ШТ"), 1500)
    assert 0.5 <= m.near_score < 0.95
    assert _match(idx, "ПЕСОК КАРЬЕРНЫЙ", "М3").near_name is None


def test_fast_best_two_equals_reference():
    """Тезлаштирилган қидирув (тайёр SequenceMatcher + юқори чегара) — эталон билан айнан бир хил натижа."""
    import random
    from difflib import SequenceMatcher

    from core.matcher import _best_two_prepared

    rnd = random.Random(7)
    alphabet = "АБВГДЕЖЗИКЛМНОПРСТУФ 0123-.,/"
    words = ["".join(rnd.choice(alphabet) for _ in range(rnd.randint(2, 40))) for _ in range(300)]
    words += ["ТРУБА СТАЛЬНАЯ 57Х3", "ТРУБА СТАЛЬНАЯ 57Х3,5", "ТРУБЫ СТАЛЬНЫЕ 57Х3", "БЕТОН В15", "БЕТОН В20", "А" * 250]
    prepared = [(c, SequenceMatcher(None, "", c)) for c in words]
    for key in words[:120] + ["ТРУБА СТАЛЬНАЯ 57Х4", "БЕТОН", "", "А" * 249]:
        assert _best_two_prepared(key, prepared) == best_two_matches(key, words), key


def test_fast_best_two_small_alphabet_and_autojunk():
    """LCS чегараси: ҳарфлари кам (ўхшаш номлар кўп) ва 200+ белгили (difflib «autojunk») номларда ҳам эталон."""
    import random
    from difflib import SequenceMatcher

    from core.matcher import _best_two_prepared

    rnd = random.Random(11)
    for alphabet, lengths in (("АБВ ", (1, 12)), ("АБ1-", (180, 260)), ("АБВГДЕ 12", (5, 60))):
        words = ["".join(rnd.choice(alphabet) for _ in range(rnd.randint(*lengths))) for _ in range(150)]
        prepared = [(c, SequenceMatcher(None, "", c)) for c in words]
        for key in words[:50] + ["".join(rnd.choice(alphabet) for _ in range(rnd.randint(*lengths))) for _ in range(30)]:
            assert _best_two_prepared(key, prepared) == best_two_matches(key, words), key


def test_match_cache_returns_independent_copies():
    idx = _index([("ВОДА", "М3", 1000)])
    a = _match(idx, "ВОДА", "М3")
    a.price = 0
    assert _match(idx, "ВОДА", "М3").price == 1000


def test_fast_best_two_ties_keep_original_order():
    """Тенг баллда асл тартиб (олдин келгани) — тез қидирув номзодларни бошқа тартибда кўрса ҳам."""
    import itertools
    from difflib import SequenceMatcher

    from core.matcher import _best_two_prepared

    words = ["".join(p) for p in itertools.permutations("АБВГ")] + ["АБ", "БА", "АБВГДЕ", "ГВБА", "ЖЖЖ"]
    words = list(dict.fromkeys(words))
    prepared = [(c, SequenceMatcher(None, "", c)) for c in words]
    for key in ["АБВГ", "АБ", "ВГ", "Ж", "АБВГД", "БАГВ", ""]:
        assert _best_two_prepared(key, prepared) == best_two_matches(key, words), key
        rev = list(reversed(words))
        assert _best_two_prepared(key, [(c, SequenceMatcher(None, "", c)) for c in rev]) == best_two_matches(key, rev)


def test_floor_pruning_gives_reference_matches():
    """Паст ўхшашликни ҳисобламаслик (NEAR_MIN_SCORE) — ҳар қандай чегарада эталон билан бир хил Match."""
    import random
    from unittest import mock

    import core.matcher as m

    rnd = random.Random(3)
    alphabet = "АБВГДЕЖЗИКЛМН 0123"
    rows = [("".join(rnd.choice(alphabet) for _ in range(rnd.randint(3, 30))), rnd.choice(["ШТ", "М", "Т"]), i)
            for i in range(250)]
    rows += [("ТРУБА СТАЛЬНАЯ 57Х3", "М", 1), ("ТРУБА СТАЛЬНАЯ 57Х3,5", "М", 2), ("ТРУБЫ СТАЛЬНЫЕ 57Х3", "Т", 3)]
    keys = [r[0] for r in rows[:60]] + ["ТРУБА СТАЛЬНАЯ 57Х4", "ТРУБА СТАЛ", "БЕТОН", "ЖЖЖЖ"]

    def reference(self, name, unit, threshold, delta=0.01):
        with mock.patch.object(m, "_best_two_prepared", lambda k, prepared, floor=0.0: m.best_two_matches(
                k, [c for c, _ in prepared])):
            return self._match(name, unit, threshold, delta)

    for threshold in (0.3, 0.6, 0.8, 0.95):
        fast, ref = _index(rows), _index(rows)
        for key in keys:
            for unit in ("ШТ", "М", "Т", "КГ"):
                k = norm_name(key)
                assert fast.match(k, unit, threshold) == reference(ref, k, unit, threshold), (key, unit, threshold)


def test_top_k_equals_brute_force():
    """«Батафсил» таклифлари: энг ўхшаш 3 таси — ҳамма ресурсни difflib билан солиштирганда чиқадиган баллар."""
    import random
    from difflib import SequenceMatcher

    from core.matcher import NEAR_MIN_SCORE

    rnd = random.Random(5)
    alphabet = "АБВГДЕЖЗИКЛМН 0123-"
    rows = [("".join(rnd.choice(alphabet) for _ in range(rnd.randint(3, 30))), rnd.choice(["ШТ", "М", "Т"]), i)
            for i in range(300)]
    idx = _index(rows)
    for key in [r[0] for r in rows[:40]] + ["АБВГ", "ЖЖЖЖЖ", "А"]:
        brute = sorted((SequenceMatcher(None, key, n).ratio() for u in idx.units for n in idx._names[u]), reverse=True)
        expect = [s for s in brute if s >= NEAR_MIN_SCORE][:3]
        got = idx.top(key, 3)
        assert [s for s, _u, _n in got] == expect, key
        for s, u, n in got:
            assert SequenceMatcher(None, key, n).ratio() == s and n in idx._names[u]
