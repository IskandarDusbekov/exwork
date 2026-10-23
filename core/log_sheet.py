"""LOG listi (reja §2.10).

Eski dastur bilan farq qilmasligi uchun ikki nozik joy takrorlanadi:
* LOG oxiriga emas, faol listdan **oldin** qo'shiladi va faol bo'ladi (pywin32'da
  ``Worksheets.Add(After=...)`` ishlamagani uchun eski dastur shunday qiladi). Smart rejimda
  esa rejadagidek oxirgi list bo'ladi.
* Eski dastur qiymatlarni ``str(...)`` qilib yozgan, Excel esa raqamga o'xshash matnni
  songa aylantirgan. Biz ham raqamga o'xshash matnni son qilib yozamiz.
"""
import re
from typing import List, Sequence

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LOG_TITLE = "LOG"
HEADER = [
    "row_index", "local_name_raw", "local_unit_raw",
    "status", "picked_unit", "picked_name",
    "best_score", "second_name", "second_score",
    "written",
]
HEADER_FILL = "FFD9E1F2"
NOTE_HEADER = "note"  # faqat smart rejimda
STATUS_FILLS = {
    "EXACT": "FFC6EFCE",
    "FUZZY": "FFFFEB9C",
    "AMBIGUOUS": "FFF8CBAD",
    "NOT_FOUND": "FFFFC7CE",
    # smart rejim
    "NO_PRICE": "FFFFC7CE",
    "WORK_HEADER": "FFE7E6E6",
    "MACHINIST": "FFE7E6E6",
    "MANUAL": "FFDDEBF7",
}
MAX_WIDTH = 60
# Excel ColumnWidth=N faylga N + 0.7109375 bo'lib yoziladi (Calibri 11 uchun padding)
_EXCEL_PAD = 0.7109375

_NUMERIC = re.compile(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?")


def excel_value(v):
    """COM ``Range.Value = str`` qanday saqlasa, shunday qiymat."""
    if v is None or v == "":
        return None
    if isinstance(v, str):
        s = v.strip()
        if _NUMERIC.fullmatch(s):
            f = float(s)
            return int(f) if f.is_integer() and abs(f) < 1e15 else f
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# smart rejimda «aniq emas» = qizil (narx yozilmagan), «0 qoidasi» = kulrang
SMART_FILLS = dict(STATUS_FILLS, FUZZY="FFFFC7CE", AMBIGUOUS="FFFFC7CE", UNIT_MISMATCH="FFFFC7CE")


def write_log(wb, rows: Sequence[Sequence], smart: bool = False, title: str = LOG_TITLE) -> "object":
    """compat: LOG faol listdan oldin va faol bo'ladi (eski dastur kabi).
    smart: LOG oxirgi list, izoh ustuni bilan; faol list o'zgarmaydi (fayl smetaning o'zida ochiladi)."""
    if title in wb.sheetnames:
        del wb[title]

    header = HEADER + [NOTE_HEADER] if smart else HEADER
    if smart:
        active = wb.active
        ws = wb.create_sheet(title)
    else:
        index = wb.index(wb.active) if wb.active is not None else 0
        ws = wb.create_sheet(title, index)

    for c, name in enumerate(header, start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=excel_value(v))

    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Calibri", sz=11, b=True, family=2, scheme="minor")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"

    for r in range(2, len(rows) + 2):
        status = ws.cell(row=r, column=4).value
        color = (SMART_FILLS if smart else STATUS_FILLS).get(str(status).strip().upper()) if status else None
        if color:
            fill = PatternFill(fill_type="solid", fgColor=color)
            for c in range(1, len(header) + 1):
                ws.cell(row=r, column=c).fill = fill

    _autofit(ws, len(rows) + 1, len(header))

    if smart:
        wb.active = active
        return ws
    for other in wb.worksheets:
        other.sheet_view.tabSelected = other is ws
    wb.active = ws
    return ws


def _autofit(ws, last_row: int, ncols: int = len(HEADER)):
    """Excel AutoFit'ga yaqin: belgilar soni bo'yicha, B–J ≤ 60, A = 10."""
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        width = 0.0
        for r in range(1, last_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            n = len(str(v))
            width = max(width, n * (1.25 if r == 1 else 1.1) + 1)
        if letter == "A":
            width = 10
        else:
            width = min(width, MAX_WIDTH)
        ws.column_dimensions[letter].width = round(width, 2) + _EXCEL_PAD


def log_row(r: int, name_raw, unit_raw, match, written) -> List:
    """Eski dasturdagi LOG qatori bilan bir xil tartib va ko'rinish."""
    if match.status == "EXACT":
        return [r, str(name_raw), str(unit_raw), "EXACT", match.unit, match.name, 1.0, "", "", str(written)]
    if match.status == "NOT_FOUND":
        return [r, str(name_raw), str(unit_raw), "NOT_FOUND", "", "", "", "", "", written]
    return [
        r, str(name_raw), str(unit_raw), match.status, match.unit, match.name,
        round(match.score, 4), match.second_name or "",
        round(match.second_score, 4) if match.second_name else "",
        str(written),
    ]
