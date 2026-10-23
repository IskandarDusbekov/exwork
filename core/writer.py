"""Lokal listga yozish (reja §2.8): G narx, H summa, I/J/K ajratish, ramka, sum qatori.

openpyxl faqat kerakli kataklarga yozadi; mavjud uslub, birlashtirish, ustun kengligi va
qator balandligiga tegmaydi. Birlashtirilgan katakka faqat chap-yuqori katakka yoziladi.
"""
import re
import shutil
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

from .classify import KIND_COLUMN

NUM_FMT = "#,##0"
PERCENT_FMT = "0.00%"
GRAY = "FFD9D9D9"
SUM_COLS = ["H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]

_THIN = Side(style="thin")


class SheetWriter:
    """Listga yozadi va nimalar yozilganini eslab qoladi (svodka hisoblash uchun)."""

    def __init__(self, ws):
        self.ws = ws
        self.written: Dict[str, object] = {}

    def cell(self, addr: str):
        """Katakni qaytaradi. Faylda yo'q katak yaratilsa, Excel kabi qator (yoki ustun)
        uslubini meros qiladi — aks holda openpyxl standart shriftni qo'yib yuboradi."""
        row, col = coordinate_to_tuple(addr)
        exists = (row, col) in self.ws._cells
        cell = self.ws.cell(row=row, column=col)
        if not exists:
            style = self._inherited_style(row, col)
            if style is not None:
                cell._style = copy(style)
        return cell

    def _inherited_style(self, row: int, col: int):
        rd = self.ws.row_dimensions.get(row)
        if rd is not None and rd.customFormat:
            return rd._style
        for cd in self.ws.column_dimensions.values():
            if cd.min and cd.max and cd.min <= col <= cd.max and cd.has_style:
                return cd._style
        return None

    def set(self, addr: str, value) -> bool:
        """Qiymat yoki formula yozadi; birlashmaning chap-yuqori bo'lmagan katagini o'tkazib yuboradi."""
        cell = self.cell(addr)
        if isinstance(cell, MergedCell):
            return False
        cell.value = value
        self.written[addr] = value
        return True

    def number_format(self, col: str, start: int, end: int, fmt: str = NUM_FMT):
        for r in range(start, end + 1):
            self.cell(f"{col}{r}").number_format = fmt

    def bold(self, addr: str):
        cell = self.cell(addr)
        font = copy(cell.font)
        font.b = True
        cell.font = font

    def _cells(self, start_row: int, end_row: int, start_col: str, end_col: str):
        c1, c2 = column_index_from_string(start_col), column_index_from_string(end_col)
        for r in range(start_row, end_row + 1):
            for c in range(c1, c2 + 1):
                yield self.cell(f"{get_column_letter(c)}{r}")

    def fill(self, start_row: int, end_row: int, start_col: str = "A", end_col: str = "H", color: str = GRAY):
        for cell in self._cells(start_row, end_row, start_col, end_col):
            cell.fill = PatternFill(fill_type="solid", fgColor=color)

    def frame(self, start_row: int, end_row: int, start_col: str = "A", end_col: str = "H"):
        """Tashqi va ichki barcha chiziqlar yupqa (Excel Borders 7..12).

        Excel umumiy chegarani bitta katakda saqlaydi: oraliq chetiga chiziq qo'yilganda
        qo'shni katakning qarama-qarshi chizig'i olib tashlanadi. Shuni takrorlaymiz.
        """
        for cell in self._cells(start_row, end_row, start_col, end_col):
            old = cell.border
            cell.border = Border(
                left=_THIN, right=_THIN, top=_THIN, bottom=_THIN,
                diagonal=old.diagonal, diagonalUp=old.diagonalUp, diagonalDown=old.diagonalDown,
                vertical=old.vertical, horizontal=old.horizontal, outline=old.outline,
            )
        c1, c2 = column_index_from_string(start_col), column_index_from_string(end_col)
        for c in range(c1, c2 + 1):
            self._clear_side(start_row - 1, c, "bottom")
            self._clear_side(end_row + 1, c, "top")
        for r in range(start_row, end_row + 1):
            self._clear_side(r, c1 - 1, "right")
            self._clear_side(r, c2 + 1, "left")

    def _clear_side(self, row: int, col: int, side: str):
        if row < 1 or col < 1 or (row, col) not in self.ws._cells:
            return
        cell = self.ws._cells[(row, col)]
        current = getattr(cell.border, side)
        if current is None or current.style is None:
            return
        border = copy(cell.border)
        setattr(border, side, Side())
        cell.border = border


def amount_formula(r: int) -> str:
    return f'=IF(F{r}="",E{r}*G{r},F{r}*G{r})'


def safe_amount_formula(r: int) -> str:
    """Smart (T6): G da son bo'lmasa 0 — bitta qator butun svodkani #ЗНАЧ! qilmaydi."""
    return f'=IF(ISNUMBER(G{r}),IF(F{r}="",E{r}*G{r},F{r}*G{r}),0)'


def write_priced_rows(w: SheetWriter, start: int, end: int, g_values: Dict[int, object], kinds: Dict[int, str],
                      formula=amount_formula):
    """G..K ustunlarini ``start..end`` oralig'ida to'ldiradi va A–H ramka chizadi."""
    for r in range(start, end + 1):
        w.set(f"G{r}", g_values.get(r))
    w.number_format("G", start, end)

    for r in range(start, end + 1):
        w.set(f"H{r}", formula(r))
    w.number_format("H", start, end)

    for r in range(start, end + 1):
        col = KIND_COLUMN.get(kinds[r], "K")
        ref = f"=H{r}"
        for c in "IJK":
            w.set(f"{c}{r}", ref if col == c else None)
        if col in "LM":
            # L/M (kabel, uskuna) — faqat shu qatorga; boshqa qatorlardagi qo'lda kiritilganlarga tegilmaydi
            w.set(f"{col}{r}", ref)
            w.cell(f"{col}{r}").number_format = NUM_FMT
    for col in "IJK":
        w.number_format(col, start, end)

    w.frame(start, end, "A", "H")


def write_sum_row(w: SheetWriter, sum_row: int, start: int, end: int):
    for col in SUM_COLS:
        w.set(f"{col}{sum_row}", f"=SUM({col}{start}:{col}{end})")
        w.cell(f"{col}{sum_row}").number_format = NUM_FMT


# ===================== narxsiz qatorlarni qizil qilish (smart) =====================
UNPRICED_FILL = "FFFFC7CE"
UNPRICED_FONT = "FF9C0006"


def mark_unpriced(ws, start: int, end: int, col: str = "G"):
    """G da son bo'lmagan (matn) kataklar qizil. Shartli format: smetachi narx yozishi bilan
    qizil o'zi yo'qoladi; katakning o'z uslubiga tegilmaydi."""
    rng = f"{col}{start}:{col}{end}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({col}{start}<>"",NOT(ISNUMBER({col}{start})))'],
        fill=PatternFill(bgColor=UNPRICED_FILL), font=Font(color=UNPRICED_FONT, b=True), stopIfTrue=True,
    ))


# ===================== ustun kengligi (smart) =====================
# Excel son katakka sig'masa «######» ko'rsatadi. Smetalarda G dan oxirigacha (G:XFD) bitta
# umumiy kenglik (~9 belgi) bo'ladi, milliardli summalar esa 13+ belgi. Smart rejim biz
# yozadigan ustunlarni faqat kerak bo'lsa kengaytiradi (hech qachon toraytirmaydi).

def column_width(ws, col: int) -> float:
    for key, cd in ws.column_dimensions.items():
        lo = cd.min or column_index_from_string(key)
        hi = cd.max or lo
        if lo <= col <= hi and cd.width:
            return cd.width
    return 8.43  # Excel standarti


def set_column_width(ws, col: int, width: float):
    """Bitta ustun kengligini o'rnatadi; ustun guruhga kirsa (masalan G:XFD), guruh bo'linadi."""
    dims = ws.column_dimensions
    found = None
    for key, cd in list(dims.items()):
        lo = cd.min or column_index_from_string(key)
        hi = cd.max or lo
        if lo <= col <= hi:
            found = (key, cd, lo, hi)
            break
    if found is None:
        letter = get_column_letter(col)
        dims[letter] = ColumnDimension(ws, index=letter, width=width, min=col, max=col)
        return
    key, cd, lo, hi = found
    if lo == hi:
        cd.width = width
        return

    def part(a: int, b: int, w: float) -> ColumnDimension:
        d = ColumnDimension(ws, index=get_column_letter(a), width=w, bestFit=cd.bestFit, hidden=cd.hidden,
                            outlineLevel=cd.outlineLevel, collapsed=cd.collapsed, min=a, max=b)
        d._style = copy(cd._style)
        return d

    del dims[key]
    if lo < col:
        dims[get_column_letter(lo)] = part(lo, col - 1, cd.width)
    dims[get_column_letter(col)] = part(col, col, width)
    if col < hi:
        dims[get_column_letter(col + 1)] = part(col + 1, hi, cd.width)


def number_text(v) -> str:
    """``#,##0`` formatida Excel ko'rsatadigan matn uzunligi uchun."""
    if isinstance(v, bool) or v is None:
        return ""
    if isinstance(v, (int, float)):
        return f"{round(v):,}"
    return "#ЗНАЧ!" if str(v).startswith("#") else ""


def fit_number_columns(ws, values: Dict[str, object], first_col: str = "G", last_col: str = "Q") -> List[str]:
    """``values``: manzil -> hisoblangan qiymat. Sig'maydigan ustunlarni kengaytiradi, ro'yxatini qaytaradi."""
    normal = ws.parent._fonts[0].sz or 11
    need: Dict[int, float] = {}
    lo, hi = column_index_from_string(first_col), column_index_from_string(last_col)
    for addr, v in values.items():
        text = number_text(v)
        if not text:
            continue
        row, col = coordinate_to_tuple(addr)
        if not lo <= col <= hi:
            continue
        cell = ws._cells.get((row, col))
        size = (cell.font.sz if cell is not None and cell.font and cell.font.sz else normal)
        ratio = min(1.5, max(0.8, size / normal))
        need[col] = max(need.get(col, 0), len(text) * ratio * 1.1 + 2)
    widened = []
    for col, width in sorted(need.items()):
        if width > column_width(ws, col) + 0.01:
            set_column_width(ws, col, round(width, 2))
            widened.append(get_column_letter(col))
    return widened


# ===================== birlashgan kataklar uslubini tiklash =====================
# openpyxl yuklashda birlashmaning chap-yuqori bo'lmagan kataklarini yangi MergedCell bilan
# almashtiradi: asl uslub yo'qoladi, chegaraga faqat chap-yuqori katakdagi tomonlar ko'chadi.
# Masalan, «E15:F15» da o'ng chiziq faqat F15 da bo'lsa, u yo'qolardi. Asl uslub indeksini
# XML'dan o'qib qaytaramiz.

_S_ATTR_RE = re.compile(r'\bs="(\d+)"')


def _raw_style_ids(sheet_xml: str) -> Dict[tuple, int]:
    ids = {}
    for m in re.finditer(r"<c\b[^>]*>", sheet_xml):
        tag = m.group(0)
        ref = re.search(r'\br="([A-Z]+)(\d+)"', tag)
        if not ref:
            continue
        s = _S_ATTR_RE.search(tag)
        ids[(int(ref.group(2)), column_index_from_string(ref.group(1)))] = int(s.group(1)) if s else 0
    return ids


def restore_merged_styles(wb, path: Path) -> int:
    """``path`` dan yuklangan ``wb`` dagi MergedCell'larga asl uslubni qaytaradi."""
    restored = 0
    with zipfile.ZipFile(path) as z:
        parts = _sheet_parts(z)
        for ws in wb.worksheets:
            part = parts.get(ws.title)
            if not part or not ws.merged_cells.ranges:
                continue
            ids = _raw_style_ids(z.read(part).decode("utf-8"))
            for mcr in ws.merged_cells.ranges:
                for coord in mcr.cells:
                    cell = ws._cells.get(coord)
                    if not isinstance(cell, MergedCell):
                        continue
                    style_id = ids.get(coord, 0)
                    cell._style = copy(wb._cell_styles[style_id]) if style_id else StyleArray()
                    restored += 1
    return restored


# ===================== header/footer tiklash =====================
# openpyxl rus tilidagi kodli (&Л, &Ц ...) header/footer'ni o'qiy olmaydi va saqlashda
# tashlab yuboradi. Saqlangandan keyin uni asl fayldagi XML'dan qaytaramiz.

_HF_RE = re.compile(r"<headerFooter\b[^>]*/>|<headerFooter\b.*?</headerFooter>", re.S)
_AFTER_HF = (
    "rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors", "smartTags",
    "drawing", "legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects",
    "controls", "webPublishItems", "tableParts", "extLst",
)


def _sheet_parts(z: zipfile.ZipFile) -> Dict[str, str]:
    """List nomi -> zip ichidagi XML yo'li."""
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels_xml = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels_xml):
        tag = m.group(0)
        rid = re.search(r'\bId="([^"]+)"', tag)
        tgt = re.search(r'\bTarget="([^"]+)"', tag)
        if rid and tgt:
            t = tgt.group(1)
            targets[rid.group(1)] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    parts = {}
    for m in re.finditer(r"<sheet\b[^>]*>", wb_xml):
        tag = m.group(0)
        name = re.search(r'\bname="([^"]*)"', tag)
        rid = re.search(r'\br:id="([^"]+)"', tag) or re.search(r'\bid="([^"]+)"', tag)
        if name and rid and rid.group(1) in targets:
            parts[_xml_unescape(name.group(1))] = targets[rid.group(1)]
    return parts


def _xml_unescape(s: str) -> str:
    return s.replace("&quot;", '"').replace("&apos;", "'").replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")


def _insert_header_footer(sheet_xml: str, hf: str) -> str:
    if _HF_RE.search(sheet_xml):
        return _HF_RE.sub(lambda _: hf, sheet_xml, count=1)
    for tag in _AFTER_HF:
        m = re.search(rf"<{tag}\b", sheet_xml)
        if m:
            return sheet_xml[:m.start()] + hf + sheet_xml[m.start():]
    i = sheet_xml.rindex("</worksheet>")
    return sheet_xml[:i] + hf + sheet_xml[i:]


def restore_header_footer(src_path: Path, out_path: Path, sheet_titles: Iterable[str]) -> int:
    """``src_path`` dagi listlarning headerFooter'ini ``out_path`` ga ko'chiradi. Nechta tiklanganini qaytaradi."""
    with zipfile.ZipFile(src_path) as zs:
        src_parts = _sheet_parts(zs)
        hf_by_title: Dict[str, str] = {}
        for title in sheet_titles:
            part = src_parts.get(title)
            if not part:
                continue
            m = _HF_RE.search(zs.read(part).decode("utf-8"))
            if m:
                hf_by_title[title] = m.group(0)
    if not hf_by_title:
        return 0

    with zipfile.ZipFile(out_path) as zo:
        out_parts = _sheet_parts(zo)
        replace: Dict[str, bytes] = {}
        for title, hf in hf_by_title.items():
            part = out_parts.get(title)
            if part:
                xml = zo.read(part).decode("utf-8")
                replace[part] = _insert_header_footer(xml, hf).encode("utf-8")
        if not replace:
            return 0
        fd, tmp = tempfile.mkstemp(suffix=out_path.suffix, dir=out_path.parent)
        with open(fd, "wb") as fh, zipfile.ZipFile(fh, "w", zipfile.ZIP_DEFLATED) as zn:
            for item in zo.infolist():
                data = replace.get(item.filename)
                zn.writestr(item, data if data is not None else zo.read(item.filename))
    shutil.move(tmp, out_path)
    return len(replace)


def read_header_footer(path: Path, title: str) -> Optional[str]:
    with zipfile.ZipFile(path) as z:
        part = _sheet_parts(z).get(title)
        if not part:
            return None
        m = _HF_RE.search(z.read(part).decode("utf-8"))
        return m.group(0) if m else None


# ===================== formulalarga tayyor natija (cached value) =====================
# openpyxl formulaning faqat matnini yozadi (<f>…</f><v/>). Excel oddiy rejimda qayta
# hisoblaydi, lekin brauzerdan yuklangan fayl «Защищённый просмотр»da ochiladi va u yerda
# hisoblanmaydi — H, I/J/K va svodka bo'sh ko'rinadi. Shuning uchun natijani o'zimiz yozamiz
# (Python baholovchisi Excel bilan qatorma-qator bir xil ekani testlarda tekshirilgan).

_FORMULA_CELL_RE = re.compile(r'<c r="([A-Z]+[0-9]+)"((?:\s[^>]*?)?)>(<f>[^<]*</f>)<v\s*/>(?:</v>)?</c>')


def _cached_xml(ref: str, attrs: str, formula: str, value) -> Optional[str]:
    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
    if isinstance(value, bool):
        return f'<c r="{ref}"{attrs} t="b">{formula}<v>{int(value)}</v></c>'
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"{attrs}>{formula}<v>{repr(float(value)) if isinstance(value, float) else value}</v></c>'
    if isinstance(value, str) and value.startswith("#"):
        return f'<c r="{ref}"{attrs} t="e">{formula}<v>{value}</v></c>'
    return None


def write_cached_values(path: Path, values_by_sheet: Dict[str, Dict[str, object]]) -> int:
    """``{list: {manzil: qiymat}}`` — formula kataklariga natijani yozadi. Nechta katak yozilganini qaytaradi."""
    path = Path(path)
    count = 0
    with zipfile.ZipFile(path) as z:
        parts = _sheet_parts(z)
        replace: Dict[str, bytes] = {}
        for title, values in values_by_sheet.items():
            part = parts.get(title)
            if not part or not values:
                continue

            def fix(m):
                nonlocal count
                ref = m.group(1)
                if ref not in values:
                    return m.group(0)
                xml = _cached_xml(ref, m.group(2), m.group(3), values[ref])
                if xml is None:
                    return m.group(0)
                count += 1
                return xml

            replace[part] = _FORMULA_CELL_RE.sub(fix, z.read(part).decode("utf-8")).encode("utf-8")
        if not replace:
            return 0
        fd, tmp = tempfile.mkstemp(suffix=path.suffix, dir=path.parent)
        with open(fd, "wb") as fh, zipfile.ZipFile(fh, "w", zipfile.ZIP_DEFLATED) as zn:
            for item in z.infolist():
                data = replace.get(item.filename)
                zn.writestr(item, data if data is not None else z.read(item.filename))
    shutil.move(tmp, path)
    return count
