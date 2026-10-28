"""Resurs indeksi va moslashtirish: EXACT / FUZZY / AMBIGUOUS / NOT_FOUND (reja §2.3–2.4).

Moslik rejimida faqat ``difflib.SequenceMatcher`` ishlatiladi.
"""
from dataclasses import dataclass
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Union

from openpyxl.utils import column_index_from_string

from .classify import section_kind
from .normalize import norm_name, norm_unit

EXACT = "EXACT"
FUZZY = "FUZZY"
AMBIGUOUS = "AMBIGUOUS"
NOT_FOUND = "NOT_FOUND"
NEAR_MIN_SCORE = 0.5  # bundan past o'xshashlik — taklif ham qilinmaydi (tasodifiy nom)


@dataclass
class Match:
    status: str
    unit: str = ""
    name: str = ""
    price: Any = None
    score: Optional[float] = None
    second_name: Optional[str] = None
    second_score: Optional[float] = None
    second_price: Any = None
    # NOT_FOUND: o'xshashlik chegarasidan past bo'lsa ham eng yaqin resurs (faqat taklif uchun)
    near_name: Optional[str] = None
    near_unit: Optional[str] = None
    near_price: Any = None
    near_score: Optional[float] = None


def best_two_matches(key: str, candidates: List[str]):
    best = None
    second = None
    best_score = 0.0
    second_score = 0.0
    for c in candidates:
        score = SequenceMatcher(None, key, c).ratio()
        if score > best_score:
            second, second_score = best, best_score
            best, best_score = c, score
        elif score > second_score:
            second, second_score = c, score
    return best, best_score, second, second_score


def _col(c: Union[str, int]) -> int:
    return c if isinstance(c, int) else column_index_from_string(c)


class ResourceIndex:
    """``by_unit[birlik][nom] = narx``; takrorlangan nomdan birinchisi olinadi."""

    def __init__(self, by_unit: Dict[str, Dict[str, Any]], sections: Optional[Dict[tuple, tuple]] = None,
                 section_list: Optional[List[tuple]] = None, positions: Optional[Dict[tuple, tuple]] = None):
        self.by_unit = by_unit
        self.units = list(by_unit.keys())
        self._names = {u: list(d.keys()) for u, d in by_unit.items()}
        self.sections = sections or {}  # (birlik, nom) -> (kategoriya, bo'lim sarlavhasi)
        self.section_list = section_list or []  # RES dagi tartibda: [(kategoriya, sarlavha), ...]
        self.positions = positions or {}  # (birlik, nom) -> (bo'lim tartib raqami yoki -1, RES qatori)
        self.raw = {}  # (birlik, nom) -> (RES dagi asl nom, asl birlik) — smetachiga ko'rsatish uchun

    def raw_of(self, unit: str, name: str) -> tuple:
        return self.raw.get((unit, name), (name, unit))

    def section_of(self, unit: str, name: str) -> Optional[tuple]:
        return self.sections.get((unit, name))

    def position_of(self, unit: str, name: str) -> Optional[tuple]:
        """Resurs RES da qayerda: ``(bo'lim raqami (section_list da) yoki -1, RES qatori)``."""
        return self.positions.get((unit, name))

    @classmethod
    def from_sheet(cls, ws, name_col: Union[str, int] = "D", price_col: Union[str, int] = "G",
                   unit_col: Union[str, int, None] = None, start_row: int = 1,
                   stop_row: Optional[int] = None) -> "ResourceIndex":
        """Ustunlar harf yoki raqam; ``unit_col`` berilmasa — nomdan keyingi ustun (eski dastur kabi)."""
        rn = _col(name_col)
        ru = _col(unit_col) if unit_col is not None else rn + 1
        rp = _col(price_col)
        by_unit: Dict[str, Dict[str, Any]] = {}
        sections: Dict[tuple, tuple] = {}
        section_list: List[tuple] = []
        positions: Dict[tuple, tuple] = {}
        raw: Dict[tuple, tuple] = {}
        current = None  # oxirgi ko'rilgan bo'lim: (kategoriya, sarlavha)
        for r in range(start_row, min(stop_row or ws.max_row + 1, ws.max_row + 1)):
            unit = norm_unit(ws.cell(row=r, column=ru).value)
            name = norm_name(ws.cell(row=r, column=rn).value)
            if not unit:
                # birliksiz qator — bo'lim sarlavhasi bo'lishi mumkin («ОБОРУДОВАНИЕ», «ТРУДОВЫЕ РЕСУРСЫ»)
                for c in range(1, rn + 1):
                    text = ws.cell(row=r, column=c).value
                    kind = section_kind(text)
                    if kind:
                        current = (kind, str(text).strip())
                        section_list.append(current)
                        break
                continue
            if not name:
                continue
            d = by_unit.setdefault(unit, {})
            if name not in d:
                d[name] = ws.cell(row=r, column=rp).value
                positions[(unit, name)] = (len(section_list) - 1, r)
                raw[(unit, name)] = (str(ws.cell(row=r, column=rn).value).strip(), str(ws.cell(row=r, column=ru).value).strip())
                if current:
                    sections[(unit, name)] = current
        index = cls(by_unit, sections, section_list, positions)
        index.raw = raw
        return index

    def match(self, local_name: str, local_unit: str, threshold: float, ambiguity_delta: float = 0.01) -> Match:
        """``local_name``/``local_unit`` allaqachon normallashtirilgan bo'lishi kerak."""
        best_fuzzy = None
        best_ambig = None
        near_same = near_any = None  # (o'xshashlik, birlik, nom): birligi bir xili ustun

        search_units = []
        if local_unit and local_unit in self.by_unit:
            search_units.append(local_unit)
        search_units.extend(u for u in self.units if u != local_unit)

        for u in search_units:
            udict = self.by_unit[u]
            if not udict:
                continue
            if local_name in udict:
                return Match(EXACT, u, local_name, udict[local_name], 1.0)

            b, bs, s, ss = best_two_matches(local_name, self._names[u])
            if not b:
                continue
            if local_unit and u == local_unit:
                near_same = (bs, u, b)
            if near_any is None or bs > near_any[0]:
                near_any = (bs, u, b)
            if bs >= threshold:
                if s and (bs - ss) < ambiguity_delta:
                    if best_ambig is None or bs > best_ambig.score:
                        # narx G ga yozilmaydi (compat); smart rejim ikkala narxni solishtiradi
                        best_ambig = Match(AMBIGUOUS, u, b, udict[b], bs, s, ss, udict[s])
                elif best_fuzzy is None or bs > best_fuzzy.score:
                    best_fuzzy = Match(FUZZY, u, b, udict[b], bs, s, ss)

        if best_fuzzy is not None:
            return best_fuzzy
        if best_ambig is not None:
            return best_ambig
        m = Match(NOT_FOUND)
        near = near_same if near_same and near_same[0] >= NEAR_MIN_SCORE else near_any
        if near and near[0] >= NEAR_MIN_SCORE:
            m.near_score, m.near_unit, m.near_name = near[0], near[1], near[2]
            m.near_price = self.by_unit[near[1]][near[2]]
        return m
