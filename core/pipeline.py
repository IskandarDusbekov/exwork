"""Narxlash jarayoni: ``run(local, res, opts) -> Result`` va CLI.

python -m core.pipeline local.xlsx res.xlsx -o out.xlsx --fuzzy 95 [--mode smart|compat]

Ikki rejim:
* ``compat`` — eski ``.pyw`` dastur bilan aynan bir xil natija (etalon testlar uchun);
* ``smart`` — reja §4 tuzatishlari: ustunlar va ma'lumot oralig'i sarlavha bo'yicha topiladi,
  ish sarlavhalari va mashinist mehnati 0, H formulasi xatoga chidamli, svodka smeta pastidagi
  «СОСТАВИЛ» qatorlaridan keyin yoziladi. Narx QAT'IY: faqat nom aynan mos kelsa va birlik bir
  xil bo'lsa. Ўхшаш / ШАБҲАЛИ / boshqa birlik / narxsiz / topilmagan qatorlarda G ga matn
  yoziladi va Excel'da qizil bo'ladi (taklif faqat LOG'da) — smetachi narxni o'zi kiritadi.
"""
import argparse
import re
import sys
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import convert
from .classify import (
    COLUMN_KIND, KIND_COLUMN, OTHER, is_chel_ch, is_machinists_labor, is_workers_row, name_hint, row_kind,
)
from .detect import cell_value, detect_layout
from .log_sheet import LOG_TITLE, log_row, write_log
from .matcher import AMBIGUOUS, EXACT, FUZZY, NOT_FOUND, Match, ResourceIndex
from .normalize import norm_name, norm_unit
from .svodka import (
    DEFAULT_PERCENTS, XlError, make_evaluator, parse_text_number, resolve_percents, write_svodka,
)
from .writer import (
    SheetWriter, amount_formula, fit_number_columns, mark_unpriced, restore_header_footer, write_cached_values, restore_merged_styles, safe_amount_formula,
    write_priced_rows, write_sum_row,
)

NOT_FOUND_TEXT = "ТОПИЛМАДИ"
AMBIGUOUS_TEXT = "ШАБҲАЛИ"
CHECK_TEXT = "ТЕКШИРИНГ"   # smart: aniq mos kelmadi, smetachi narxni o'zi kiritadi
NO_PRICE_TEXT = "НАРХ ЙЎҚ"

# smart rejimdagi qo'shimcha statuslar
NO_PRICE = "NO_PRICE"          # resursda bor, lekin narx raqam emas («--»)
WORK_HEADER = "WORK_HEADER"    # ish sarlavhasi: ostida resurslar bor, o'zi narxlanmaydi
MACHINIST = "MACHINIST"        # mashinist mehnati: narxi mashina-soat narxi ichida
UNIT_MISMATCH = "UNIT_MISMATCH"  # nom aniq, lekin birlik boshqa — narx yozilmaydi
MANUAL = "MANUAL"              # smetachi narxni o'zi kiritgan yoki taklifni tasdiqlagan

MODES = ("smart", "compat")
LOCAL_NAME_COL = 3  # C, birlik D (compat)
RES_NAME_COL = "D"  # birlik E (compat: foydalanuvchi C oldiga bo'sh ustun qo'shadi)
RES_PRICE_COL = "G"


@dataclass
class Options:
    fuzzy_percent: int = 95
    ambiguity_delta: float = 0.01
    mode: str = "compat"            # "compat" | "smart"
    local_sheet: Optional[str] = None  # None -> 1-list (eski dastur kabi)
    res_sheet: Optional[str] = None
    soffice: Optional[str] = None      # LibreOffice yo'li (.xls uchun)
    output_path: Optional[Path] = None
    percents: Optional[Dict[str, float]] = None  # smart: svodka foizlari (ulush), None -> standart
    categories: Optional[Dict[int, str]] = None  # smart: qo'lda kategoriya {qator: "I".."M"}
    prices: Optional[Dict[int, float]] = None    # smart: qo'lda narx / tasdiqlangan taklif {qator: narx}
    log_title: Optional[str] = None              # LOG list nomi (bir faylda bir nechta lokal narxlanganda)
    active_sheet: Optional[str] = None           # smart: fayl qaysi listda ochilsin (standart: narxlangan lokal)


@dataclass
class Result:
    output_path: Path
    stats: Dict[str, int]
    problems: List[dict]
    totals: Dict[str, object]  # son yoki XlError ("#VALUE!" ...)
    local_sheet: str = ""
    start_row: int = 0
    end_row: int = 0
    sum_row: int = 0
    svodka_start: int = 0
    log: List[list] = field(default_factory=list)
    rows: List[dict] = field(default_factory=list)  # start..end: g, h (hisoblangan), kind, status
    mode: str = "compat"
    warnings: List[str] = field(default_factory=list)
    columns: Dict[str, str] = field(default_factory=dict)  # qaysi ustunlar ishlatildi
    percents: Dict[str, float] = field(default_factory=dict)  # svodkada ishlatilgan foizlar
    categories: Dict[int, str] = field(default_factory=dict)  # qo'llangan qo'lda kategoriyalar
    prices: Dict[int, float] = field(default_factory=dict)  # qo'llangan qo'lda narxlar
    cached_values: Dict[str, Dict[str, object]] = field(default_factory=dict)  # {list: {manzil: natija}}


@dataclass
class _Priced:
    g: Dict[int, object]
    log_rows: List[list]
    start: int
    end: int
    status: Dict[int, str]
    notes: Dict[int, str] = field(default_factory=dict)
    zero_rows: set = field(default_factory=set)
    unit_mismatch: int = 0
    suggestions: Dict[int, object] = field(default_factory=dict)  # smart: taklif narxi (yozilmaydi)
    sections: Dict[int, tuple] = field(default_factory=dict)  # smart: aniq topilgan resursning RES bo'limi
    orig_status: Dict[int, str] = field(default_factory=dict)  # smart: qo'lda narxdan oldingi status


def default_output_path(local_path: Path) -> Path:
    ext = ".xlsm" if local_path.suffix.lower() == ".xlsm" else ".xlsx"
    return local_path.with_name(f"{local_path.stem}_narxlandi{ext}")


def _pick_sheet(wb, name: Optional[str]):
    if name is None:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise ValueError(f"'{name}' listi topilmadi. Bor listlar: {', '.join(wb.sheetnames)}")
    return wb[name]


# Katak qiymatini o'qiydi, lekin ws.cell() kabi yangi katak yaratmaydi (aks holda writer
# uni «faylda bor» deb o'ylab, qator uslubini bermay qo'yadi).
_peek = cell_value


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _as_price(v) -> Optional[float]:
    """Smart: son yoki raqamli matn («29818,17») -> son; «--», bo'sh -> None."""
    if _is_number(v):
        return v
    if isinstance(v, str):
        return parse_text_number(v)
    return None


def run(local_path, res_path, opts: Optional[Options] = None) -> Result:
    opts = opts or Options()
    if opts.mode not in MODES:
        raise ValueError(f"Noma'lum rejim: {opts.mode}. Bor rejimlar: {', '.join(MODES)}")
    local_path, res_path = Path(local_path), Path(res_path)
    out_path = Path(opts.output_path) if opts.output_path else default_output_path(local_path)

    with tempfile.TemporaryDirectory(prefix="narxlash_") as tmp:
        local_xlsx = convert.ensure_xlsx(local_path, Path(tmp) / "local", opts.soffice)
        res_xlsx = convert.ensure_xlsx(res_path, Path(tmp) / "res", opts.soffice)
        return _run_xlsx(local_xlsx, res_xlsx, out_path, opts)


def _run_xlsx(local_xlsx: Path, res_xlsx: Path, out_path: Path, opts: Options) -> Result:
    smart = opts.mode == "smart"
    warnings: List[str] = []
    columns: Dict[str, str] = {}

    # --- resurs indeksi ---
    wb_res = load_workbook(res_xlsx, data_only=True)
    ws_res = _pick_sheet(wb_res, opts.res_sheet)
    res_layout = detect_layout(ws_res, need_price=True, stop_at_total=False) if smart else None
    if smart and (res_layout is None or res_layout.price_col is None):
        warnings.append("Resurs jadvalining sarlavhasi topilmadi — eski ustunlar olindi (nom D, birlik E, narx G)")
        res_layout = None
    if res_layout:
        warnings += res_layout.warnings
        index = ResourceIndex.from_sheet(ws_res, res_layout.name_col, res_layout.price_col,
                                         res_layout.unit_col, res_layout.data_start)
        columns.update(res_name=get_column_letter(res_layout.name_col), res_unit=get_column_letter(res_layout.unit_col),
                       res_price=get_column_letter(res_layout.price_col))
    else:
        index = ResourceIndex.from_sheet(ws_res, RES_NAME_COL, RES_PRICE_COL)
        columns.update(res_name=RES_NAME_COL, res_unit="E", res_price=RES_PRICE_COL)
    wb_res.close()

    # --- lokal ---
    keep_vba = local_xlsx.suffix.lower() == ".xlsm"
    wb = load_workbook(local_xlsx, keep_vba=keep_vba)
    restore_merged_styles(wb, local_xlsx)
    ws = _pick_sheet(wb, opts.local_sheet)
    sheet_titles = list(wb.sheetnames)
    threshold = int(max(80, min(100, opts.fuzzy_percent))) / 100.0

    if smart:
        priced, ln, lu = _price_smart(ws, index, threshold, opts.ambiguity_delta, warnings, opts.prices)
    else:
        ln, lu = LOCAL_NAME_COL, LOCAL_NAME_COL + 1
        priced = _price_compat(ws, index, threshold, opts.ambiguity_delta)
    columns.update(local_name=get_column_letter(ln), local_unit=get_column_letter(lu))
    start_r, end_r = priced.start, priced.end

    default_kinds = {r: row_kind(_peek(ws, r, ln), _peek(ws, r, lu)) for r in range(start_r, end_r + 1)}
    kinds = dict(default_kinds)
    applied_categories: Dict[int, str] = {}
    category_hints: Dict[int, tuple] = {}
    if smart:
        for r, col in (opts.categories or {}).items():
            r = int(r)
            if col not in COLUMN_KIND:
                raise ValueError(f"Noma'lum kategoriya: {col!r} (I, J, K, L, M bo'lishi kerak)")
            if start_r <= r <= end_r:
                kinds[r] = COLUMN_KIND[col]
                applied_categories[r] = col
        category_hints = _category_hints(ws, priced, default_kinds, ln)

    # --- yozish ---
    w = SheetWriter(ws)
    write_priced_rows(w, start_r, end_r, priced.g, kinds, safe_amount_formula if smart else amount_formula)
    sum_row = end_r + 1
    write_sum_row(w, sum_row, start_r, end_r)
    # compat: svodka darhol sum qatoridan keyin (pastki «СОСТАВИЛ» ustiga tushishi mumkin);
    # smart (T8): listdagi oxirgi to'ldirilgan qatordan keyin, bitta bo'sh qator qoldirib.
    svodka_start = (_last_content_row(ws) + 2) if smart else sum_row + 1
    percents = resolve_percents(opts.percents) if smart else dict(DEFAULT_PERCENTS)
    svodka_cells = write_svodka(w, svodka_start, sum_row, percents, smart=smart)
    if smart:
        mark_unpriced(ws, start_r, end_r)
    write_log(wb, priced.log_rows, smart=smart, title=opts.log_title or LOG_TITLE)
    if smart:
        # Fayl narxlangan listda ochilsin: manba fayl boshqa (narxlanmagan) listda saqlangan bo'lishi mumkin,
        # shunda smetachi ochganda «hammasi bo'sh» deb o'ylaydi.
        target = wb[opts.active_sheet] if opts.active_sheet in wb.sheetnames else ws
        for sheet in wb.worksheets:
            sheet.sheet_view.tabSelected = sheet is target
        wb.active = target

    # --- svodka summalari (Excel saqlagan qiymatlar + biz yozgan formulalar) ---
    wb_vals = load_workbook(local_xlsx, data_only=True)
    ws_vals = wb_vals[ws.title]
    ev = make_evaluator(w.written, lambda a: ws_vals[a].value)
    totals = {key: ev.value(addr) for key, addr in svodka_cells.items()}

    # Sahifada ko'rsatish uchun: har bir qator Excel'da qanday chiqishi
    log_by_row = {row[0]: row for row in priced.log_rows}
    rows = []
    for r in range(start_r, end_r + 1):
        log = log_by_row.get(r)
        status = priced.status.get(r, "")
        rows.append({
            "row": r,
            "name": _peek(ws, r, ln),
            "unit": _peek(ws, r, lu),
            "qty": _quantity(ws_vals, r),
            "g": w.written.get(f"G{r}"),
            "h": ev.value(f"H{r}") if f"H{r}" in w.written else None,
            "kind": kinds[r],
            "category": KIND_COLUMN[kinds[r]],
            "default_category": KIND_COLUMN[default_kinds[r]],
            "override": r in applied_categories,
            "suggested_category": None,
            "suggest_reason": "",
            "status": status,
            "picked_name": log[5] if log else "",
            "score": log[6] if log else "",
            "zero": r in priced.zero_rows or status in (WORK_HEADER, MACHINIST),
            "note": priced.notes.get(r, ""),
            "suggested_price": priced.suggestions.get(r),
            "manual": status == MANUAL,
            "orig_status": priced.orig_status.get(r, status),
        })
        hint = category_hints.get(r)
        if hint and KIND_COLUMN[hint[0]] != rows[-1]["category"]:
            rows[-1]["suggested_category"], rows[-1]["suggest_reason"] = KIND_COLUMN[hint[0]], hint[1]
    wb_vals.close()

    if smart:
        shown = {f"G{r['row']}": r["g"] for r in rows}
        for r in rows:
            for col in ("H", "I", "J", "K", "L", "M"):
                if f"{col}{r['row']}" in w.written and w.written[f"{col}{r['row']}"] is not None:
                    shown[f"{col}{r['row']}"] = r["h"]
        for col in ("H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"):
            shown[f"{col}{sum_row}"] = ev.value(f"{col}{sum_row}")
        shown.update({addr: totals[key] for key, addr in svodka_cells.items()})
        widened = fit_number_columns(ws, shown)
        if widened:
            columns["widened"] = ", ".join(widened)

    # formulalar natijasi — fayl «Защищённый просмотр»da ham to'liq ko'rinsin
    cached = {addr: ev.value(addr) for addr, v in w.written.items() if isinstance(v, str) and v.startswith("=")}

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
    restore_header_footer(local_xlsx, out_path, sheet_titles)
    write_cached_values(out_path, {ws.title: cached})

    # --- statistika va muammolar ---
    stats = Counter({EXACT: 0, FUZZY: 0, AMBIGUOUS: 0, NOT_FOUND: 0})
    if smart:
        stats.update({NO_PRICE: 0, UNIT_MISMATCH: 0, WORK_HEADER: 0, MACHINIST: 0, MANUAL: 0})
        stats.update(priced.status.values())
    else:
        stats.update(row[3] for row in priced.log_rows)
        stats["ZERO_RULE"] = len(priced.zero_rows)

    problems = []
    for row in priced.log_rows:
        r = row[0]
        final = priced.g.get(r)
        status = priced.status.get(r, row[3])
        note = priced.notes.get(r, "")
        if not _is_number(final):
            problems.append({
                "row": r, "name": row[1], "unit": row[2], "status": status,
                "picked_name": row[5], "score": row[6], "written": final, "note": note,
                "suggested_price": priced.suggestions.get(r),
            })
    stats["NON_NUMERIC"] = len(problems)
    if smart:
        stats["NOT_PRICED"] = len(problems)
        if problems:
            warnings.append(f"{len(problems)} та қатор нархланмади (Excel'да G қизил): нархни қўлда киритинг — "
                            "H ва сводка ўзи қайта ҳисобланади. Ҳозирги сумма тўлиқ эмас.")

    return Result(
        output_path=out_path, stats=dict(stats), problems=problems, totals=totals,
        local_sheet=ws.title, start_row=start_r, end_row=end_r, sum_row=sum_row,
        svodka_start=svodka_start, log=priced.log_rows, rows=rows, mode=opts.mode,
        warnings=warnings, columns=columns, percents=percents, categories=applied_categories,
        cached_values={ws.title: cached},
        prices={r: priced.g[r] for r, st in priced.status.items() if st == MANUAL},
    )


def _category_hints(ws, priced: "_Priced", default_kinds: Dict[int, str], ln: int) -> Dict[int, tuple]:
    """Kategoriya TAKLIFLARI (o'zi o'zgartirmaydi): avval aniq topilgan resursning RES dagi bo'limi,
    bo'lmasa K dagi qatorlar uchun nomdagi so'z («кабель», «насос» …). ``{qator: (tur, sabab)}``."""
    hints = {}
    for r, kind in default_kinds.items():
        if priced.status.get(r) in (WORK_HEADER, MACHINIST) or r not in priced.status:
            continue
        section = priced.sections.get(r)
        if section and section[0] != kind:
            hints[r] = (section[0], f"RES'да «{section[1]}» бўлимида")
        elif not section and kind == OTHER:
            hint = name_hint(_peek(ws, r, ln))
            if hint:
                hints[r] = (hint[0], f"номида «{hint[1]}…»")
    return hints


def _quantity(ws_vals, r: int):
    """H formulasidagi miqdor: F (loyiha bo'yicha), bo'sh bo'lsa E — Excel qanday ko'rsa, shunday."""
    f = ws_vals.cell(row=r, column=6).value
    return ws_vals.cell(row=r, column=5).value if f in (None, "") else f


def log_title_for(sheet: str) -> str:
    """«LOG 202611450_БВ» — Excel list nomi 31 belgidan oshmasin, taqiqlangan belgilar olib tashlanadi."""
    return ("LOG " + re.sub(r"[\\/*?:\[\]]", "", sheet))[:31]


def run_many(local_path, steps: List[Tuple[object, Options]], output_path) -> List[Result]:
    """Bitta lokal faylning bir nechta listini ketma-ket narxlaydi: har qadamning natijasi
    keyingisiga kirish bo'ladi. ``steps``: ``[(res_path, Options(local_sheet=…, …)), …]``.
    Har bir list o'z LOG'ini oladi («LOG <list>»). Oxirgi fayl ``output_path`` ga yoziladi."""
    output_path = Path(output_path)
    results = []
    with tempfile.TemporaryDirectory(prefix="narxlash_many_") as tmp:
        current = Path(local_path)
        for n, (res_path, opts) in enumerate(steps, start=1):
            last = n == len(steps)
            out = output_path if last else Path(tmp) / f"step{n}{output_path.suffix}"
            opts = Options(**{**opts.__dict__, "output_path": out,
                              "log_title": opts.log_title or log_title_for(opts.local_sheet or ""),
                              # umumiy fayl birinchi narxlangan listda ochilsin
                              "active_sheet": opts.active_sheet or steps[0][1].local_sheet})
            results.append(run(current, res_path, opts))
            current = out
    # har qadam oldingi listlarning natijalarini openpyxl orqali qayta saqlaganda yo'qotadi — oxirida hammasini yozamiz
    merged: Dict[str, Dict[str, object]] = {}
    for r in results:
        merged.update(r.cached_values)
    write_cached_values(output_path, merged)
    return results


def _last_content_row(ws) -> int:
    last = 0
    for (r, _c), cell in ws._cells.items():
        if r > last and cell.value is not None and str(cell.value).strip() != "":
            last = r
    return last


# ===================== compat: eski dastur mantig'i =====================
def _price_compat(ws, index: ResourceIndex, threshold: float, delta: float) -> _Priced:
    ln, lu = LOCAL_NAME_COL, LOCAL_NAME_COL + 1
    max_l = ws.max_row
    g_values: Dict[int, object] = {}
    log_rows: List[list] = []
    zero_rows = set()
    first_priced_row: Optional[int] = None
    last_data_row = 1

    for r in range(1, max_l + 1):
        name_raw = _peek(ws, r, ln)
        unit_raw = _peek(ws, r, lu)

        if name_raw is not None and str(name_raw).strip() != "":
            last_data_row = r
        # Nol qoidasi: ishchilar mehnati qatoridan oldingi qator (ish sarlavhasi) = 0
        if r > 1 and is_workers_row(name_raw, unit_raw):
            zero_rows.add(r - 1)

        local_name = norm_name(name_raw)
        if not local_name:
            continue
        m = index.match(local_name, norm_unit(unit_raw), threshold, delta)
        if m.status in (EXACT, FUZZY):
            value = m.price
        elif m.status == AMBIGUOUS:
            value = f"{AMBIGUOUS_TEXT} ({round(m.score, 3)})"
        else:
            value = NOT_FOUND_TEXT
        g_values[r] = value
        if first_priced_row is None:
            first_priced_row = r
        log_rows.append(log_row(r, name_raw, unit_raw, m, value))

    for rr in zero_rows:
        g_values[rr] = 0
        if first_priced_row is None:
            first_priced_row = rr

    if first_priced_row is None:
        first_priced_row = 2 if max_l >= 2 else 1
    start_r = first_priced_row
    end_r = last_data_row if last_data_row >= start_r else max_l
    status = {row[0]: row[3] for row in log_rows}
    return _Priced(g_values, log_rows, start_r, end_r, status, zero_rows=zero_rows)


# ===================== smart =====================
_TOP_ID = re.compile(r"\d+")


def _row_id(v) -> str:
    """№ ustuni: «1», «1.1», 1.0 -> «1»."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _work_headers(ws, start: int, stop: int, ln: int, lu: int) -> set:
    """Ish sarlavhalari (T5): № da nuqta yo'q («12») va keyingi raqamli qator «12.1…» bilan
    boshlanadi. Raqamlash bo'lmasa — eski qoida: ishchilar mehnatidan oldingi qator."""
    ids = [(r, _row_id(_peek(ws, r, 1))) for r in range(start, stop)]
    ids = [(r, i) for r, i in ids if i]
    headers = set()
    for (r, i), (_r2, i2) in zip(ids, ids[1:]):
        if _TOP_ID.fullmatch(i) and i2.startswith(i + "."):
            headers.add(r)
    for r in range(start + 1, stop):
        if is_workers_row(_peek(ws, r, ln), _peek(ws, r, lu)) and norm_name(_peek(ws, r - 1, ln)):
            headers.add(r - 1)
    return headers


def _smart_log_row(r, name_raw, unit_raw, m: Match, status: str, written, note: str) -> list:
    unit = "" if unit_raw is None else str(unit_raw)
    if m.status == NOT_FOUND:
        return [r, str(name_raw), unit, status, "", "", "", "", "", written, note]
    score = 1.0 if m.status == EXACT else round(m.score, 4)
    return [r, str(name_raw), unit, status, m.unit, m.name, score, m.second_name or "",
            round(m.second_score, 4) if m.second_name else "", written, note]


def _price_smart(ws, index: ResourceIndex, threshold: float, delta: float, warnings: List[str],
                 prices: Optional[Dict[int, float]] = None):
    layout = detect_layout(ws)
    if layout is None:
        warnings.append("Lokal jadval sarlavhasi («Наименование…», «Ед. изм.») topilmadi — butun list ko'rildi")
        ln, lu, start, stop = LOCAL_NAME_COL, LOCAL_NAME_COL + 1, 1, ws.max_row + 1
    else:
        warnings += layout.warnings
        ln, lu, start, stop = layout.name_col, layout.unit_col, layout.data_start, layout.data_stop
        if (ln, lu) != (LOCAL_NAME_COL, LOCAL_NAME_COL + 1):
            warnings.append(f"Lokalda nom {get_column_letter(ln)}, birlik {get_column_letter(lu)} ustunida; "
                            "H formulasi baribir E/F miqdorlardan foydalanadi")

    headers = _work_headers(ws, start, stop, ln, lu)
    g: Dict[int, object] = {}
    status: Dict[int, str] = {}
    notes: Dict[int, str] = {}
    log_rows: List[list] = []
    named = []
    unit_mismatch = 0
    suggestions: Dict[int, object] = {}
    sections: Dict[int, tuple] = {}
    orig_status: Dict[int, str] = {}
    manual_prices = {int(r): float(v) for r, v in (prices or {}).items()}
    for r, v in manual_prices.items():
        if v < 0:
            raise ValueError(f"{r}-qator: narx manfiy bo'lishi mumkin emas ({v})")

    for r in range(start, stop):
        name_raw, unit_raw = _peek(ws, r, ln), _peek(ws, r, lu)
        local_name = norm_name(name_raw)
        if not local_name:
            continue
        named.append(r)
        local_unit = norm_unit(unit_raw)
        m = index.match(local_name, local_unit, threshold, delta)
        machinist = is_chel_ch(unit_raw) and is_machinists_labor(name_raw)
        note = ""

        price = _as_price(m.price)
        suggestion = None
        # Qat'iy qoida: narx faqat nom AYNAN mos kelsa va birlik bir xil bo'lsa yoziladi.
        # Qolgan hammasi G da matn (Excel'da qizil) — smetachi narxni o'zi kiritadi.
        if r in headers:
            st, value, note = WORK_HEADER, 0, "иш сарлавҳаси: остидаги ресурслар нархланади"
        elif machinist and not (m.status == EXACT and price is not None):
            st, value, note = MACHINIST, 0, "машинист меҳнати машина нархи ичида"
        elif m.status == EXACT and local_unit and m.unit != local_unit:
            st, value, suggestion = UNIT_MISMATCH, CHECK_TEXT, m.price
            note = f"ном аниқ, лекин бирлик бошқа: локал «{unit_raw}», ресурс «{m.unit}» ({m.price})"
            unit_mismatch += 1
        elif m.status == EXACT and price is None:
            st, value, note = NO_PRICE, NO_PRICE_TEXT, f"ресурсда нарх рақам эмас: «{m.price}»"
        elif m.status == EXACT:
            st, value = EXACT, price
            section = index.section_of(m.unit, m.name)
            if section:
                sections[r] = section
        elif m.status == FUZZY:
            st, value, suggestion = FUZZY, CHECK_TEXT, m.price
            note = f"таклиф (аниқ эмас): «{m.name}» {round(m.score, 3)} — {m.price}"
        elif m.status == AMBIGUOUS:
            st, value, suggestion = AMBIGUOUS, CHECK_TEXT, m.price
            note = (f"2 та номзод: «{m.name}» ({m.price}) ва «{m.second_name}» ({m.second_price}), "
                    f"ўхшашлик {round(m.score, 3)}")
        else:
            st, value = NOT_FOUND, NOT_FOUND_TEXT

        if r in manual_prices:
            # Smetachi narxni o'zi kiritgan yoki taklifni tasdiqlagan — u har doim ustun
            manual = manual_prices[r]
            accepted = suggestion is not None and _as_price(suggestion) == manual
            orig_status[r] = st
            st, value = MANUAL, manual
            note = ("таклиф тасдиқланди" if accepted else "нарх қўлда киритилди") + (f" (асли: {note})" if note else "")

        g[r], status[r] = value, st
        if note:
            notes[r] = note
        if suggestion is not None:
            suggestions[r] = suggestion
        log_rows.append(_smart_log_row(r, name_raw, unit_raw, m, st, value, note))

    if not named:
        raise ValueError("Lokal listda narxlanadigan qator topilmadi")
    p = _Priced(g, log_rows, named[0], named[-1], status, notes)
    p.unit_mismatch = unit_mismatch
    p.suggestions = suggestions
    p.sections = sections
    p.orig_status = orig_status
    return p, ln, lu


# ===================== CLI =====================
def _fmt(v) -> str:
    if isinstance(v, XlError):
        return str(v)
    if _is_number(v):
        return f"{v:,.2f}".replace(",", " ")
    return repr(v)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="python -m core.pipeline", description="Смета нархлаш")
    p.add_argument("local", type=Path, help="Lokal (БВ) fayl: .xls/.xlsx/.xlsm")
    p.add_argument("res", type=Path, help="Resurs (БР) fayl: .xls/.xlsx/.xlsm")
    p.add_argument("-o", "--output", type=Path, help="Chiqish fayli (standart: <lokal>_narxlandi.xlsx)")
    p.add_argument("--fuzzy", type=int, default=95, help="Fuzzy chegarasi, 80–100 (standart 95)")
    p.add_argument("--mode", choices=MODES, default="smart",
                   help="smart (standart) yoki compat — eski dastur bilan aynan bir xil")
    p.add_argument("--local-sheet", help="Lokal list nomi (standart: 1-list)")
    p.add_argument("--res-sheet", help="Resurs list nomi (standart: 1-list)")
    p.add_argument("--soffice", help=f"LibreOffice yo'li (yoki {convert.ENV_VAR})")
    p.add_argument("--percent", action="append", default=[], metavar="KALIT=FOIZ",
                   help="smart: svodka foizi, masalan --percent prochie=17,27 --percent s_nds=20 "
                        f"(kalitlar: {', '.join(DEFAULT_PERCENTS)})")
    p.add_argument("--price", action="append", default=[], metavar="QATOR=NARX",
                   help="smart: qo'lda narx, masalan --price 183=4600000")
    args = p.parse_args(argv)
    try:
        prices = {int(k): parse_text_number(v) for k, v in (item.split("=", 1) for item in args.price)}
    except (ValueError, TypeError):
        print("Xato: --price QATOR=NARX ko'rinishida bo'lsin, masalan 183=4600000", file=sys.stderr)
        return 1
    try:
        percents = {k.strip(): parse_text_number(v) / 100 for k, v in (item.split("=", 1) for item in args.percent)}
    except (ValueError, TypeError):
        print("Xato: --percent KALIT=FOIZ ko'rinishida bo'lsin, masalan prochie=17,27", file=sys.stderr)
        return 1

    opts = Options(
        fuzzy_percent=args.fuzzy, mode=args.mode, local_sheet=args.local_sheet, res_sheet=args.res_sheet,
        soffice=args.soffice, output_path=args.output, percents=percents or None, prices=prices or None,
    )
    try:
        res = run(args.local, args.res, opts)
    except (convert.ConversionError, ValueError, FileNotFoundError) as e:
        print(f"Xato: {e}", file=sys.stderr)
        return 1

    print(f"Tayyor: {res.output_path}  (rejim: {res.mode})")
    print(f"List: {res.local_sheet}, qatorlar {res.start_row}..{res.end_row}, "
          f"sum qatori {res.sum_row}, svodka {res.svodka_start}")
    print("Ustunlar: " + ", ".join(f"{k}={v}" for k, v in res.columns.items()))
    for wmsg in res.warnings:
        print(f"Ogohlantirish: {wmsg}")
    print("Statistika: " + ", ".join(f"{k}={v}" for k, v in res.stats.items()))
    for key, value in res.totals.items():
        print(f"  {key:16} {_fmt(value)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
