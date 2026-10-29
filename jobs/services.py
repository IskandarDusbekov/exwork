"""Yuklash -> konvertatsiya -> narxlash. Django view'lari faqat shu funksiyalarni chaqiradi."""
import datetime as dt
import shutil
import time
from pathlib import Path
from typing import Optional

from django.conf import settings
from openpyxl import load_workbook

from core import convert
from core.columns import LayoutError
from core.batch import apply_roles, inspect_workbook, mark_duplicates, plan_pairs
from core.detect import pair_sheets
from core.pipeline import MODES, Options, run
from core.svodka import DEFAULT_PERCENTS, parse_text_number

from django.utils import timezone

from .models import Batch, Job, Project, Setting

ALLOWED_EXT = (".xls", ".xlsx", ".xlsm")


class UploadError(Exception):
    pass


def _check_upload(f, label: str) -> str:
    ext = Path(f.name).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise UploadError(f"{label}: фақат .xls, .xlsx, .xlsm файллар қабул қилинади")
    if f.size > settings.NARXLASH_MAX_UPLOAD:
        raise UploadError(f"{label}: файл жуда катта ({f.size // (1024 * 1024)} МБ)")
    return ext


def _save(f, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "wb") as out:
        for chunk in f.chunks():
            out.write(chunk)


def sheet_names(path: Path):
    wb = load_workbook(path, read_only=True, keep_links=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def remembered_percents(owner) -> dict:
    """Смета­чи охирги марта танлаган сводка фоизлари (бўш — стандарт)."""
    s = Setting.objects.filter(owner_id=getattr(owner, "pk", owner), key="percents").first()
    return dict(s.value) if s and isinstance(s.value, dict) else {}


def remember_percents(owner, percents: Optional[dict]):
    """Охирги танлов доим сақланади (стандартга қайтарилса — стандарт қийматлар ёзилади, бўш эмас)."""
    Setting.objects.update_or_create(owner_id=getattr(owner, "pk", owner), key="percents", defaults={"value": {**DEFAULT_PERCENTS, **(percents or {})}})


def base_period_of(owner) -> dict:
    """Смета­чининг ресурс базаси танлови (вилоят ва давр); танламаган бўлса — админ белгилагани.
    Базадан нархлаш ўчирилган бўлса — {} (смета базасиз нархланади)."""
    from resbase.models import BaseConfig
    cfg = BaseConfig.get()
    if not cfg.enabled:
        return {}
    s = Setting.objects.filter(owner_id=getattr(owner, "pk", owner), key="base_period").first()
    return dict(s.value) if s and isinstance(s.value, dict) and s.value else cfg.period


def remember_base_period(owner, period: dict):
    Setting.objects.update_or_create(owner_id=getattr(owner, "pk", owner), key="base_period", defaults={"value": period})


BASE_RES = "__base__"  # жуфтликда RES ўрнига «ресурс базаси» танланди


def base_ready(owner) -> bool:
    """RES'сиз нархлаш мумкинми: база ёқилган ва смета­чида (ёки сукут бўйича) вилоят ва давр тўлиқ."""
    from resbase.services import period_complete
    return period_complete(base_period_of(owner))


def pair_ready(p: dict) -> bool:
    """Жуфтлик нархланадими: белгиланган ва RES (ёки RES ўрнига ресурс базаси) танланган."""
    return bool(p.get("include") and (p.get("res") or p.get("base")))


def base_options(job: Job) -> dict:
    """``Options`` uchun: ресурс базаси narxlari (baza o'chiq yoki davr tanlanmagan — bo'sh)."""
    from resbase.services import period_label, price_map
    prices = price_map(job.base_period)
    return {"base_prices": prices or None, "base_label": period_label(job.base_period) if prices else ""}


def create_job(owner, local_file, res_file=None) -> Job:
    """Fayllarni saqlaydi, .xlsx ga o'giradi, listlarni o'qib standart juftlikni tanlaydi."""
    local_ext = _check_upload(local_file, "Локал файл")
    res_ext = _check_upload(res_file, "Ресурс файл") if res_file else local_ext

    job = Job(owner=owner, local_name=Path(local_file.name).name, res_name=Path((res_file or local_file).name).name,
              percents=remembered_percents(owner), base_period=base_period_of(owner))
    clock = FileClock()
    try:
        _save(local_file, job.abs(f"upload/local{local_ext}"))
        local_xlsx = clock.open(job.abs(f"upload/local{local_ext}"), job.abs("conv"), job.local_name, 0)
        if res_file:
            _save(res_file, job.abs(f"upload/res{res_ext}"))
            res_xlsx = clock.open(job.abs(f"upload/res{res_ext}"), job.abs("conv"), job.res_name, 1)
        job.local_sheets = clock.read(local_xlsx, 0)
        if res_file:
            job.res_sheets = clock.read(res_xlsx, 1)
        else:
            res_xlsx, job.res_sheets = local_xlsx, job.local_sheets
    except convert.ConversionError as e:
        shutil.rmtree(job.dir, ignore_errors=True)
        save_timings(clock, getattr(owner, "pk", owner))
        raise UploadError(str(e)) from e
    except Exception as e:
        shutil.rmtree(job.dir, ignore_errors=True)
        raise UploadError(f"Файлни очиб бўлмади: {e}") from e

    job.local_path = str(local_xlsx.relative_to(job.dir))
    job.res_path = str(res_xlsx.relative_to(job.dir))
    job.local_sheet, job.res_sheet = pair_sheets(job.local_sheets, job.res_sheets)
    job.save()
    save_timings(clock, job.owner_id, job=job)
    return job


def _jsonable(v):
    if v is None or isinstance(v, (bool, int, float)):
        return v
    if isinstance(v, str):
        return str(v)  # XlError ham oddiy "#VALUE!" satrga aylanadi
    if isinstance(v, (dt.date, dt.datetime, dt.time)):
        return v.isoformat()
    if isinstance(v, dict):
        return {str(k): _jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    return str(v)


def parse_percents(data, previous: Optional[dict] = None) -> dict:
    """Formadan ``pct_<kalit>`` maydonlari: «17,27» -> 0.1727. Bo'sh — standart;
    noto'g'ri yozilgan («abc», 150) — avvalgi qiymati qoladi."""
    previous = previous or {}
    out = {}
    for key in DEFAULT_PERCENTS:
        raw = (data.get(f"pct_{key}") or "").strip()
        if not raw:
            continue
        value = parse_text_number(raw.rstrip("%").strip())
        if value is not None and 0 <= value <= 100:
            out[key] = round(value / 100, 8)
        elif key in previous:
            out[key] = previous[key]
    return out


def run_job(job: Job, local_sheet: Optional[str], res_sheet: Optional[str], fuzzy: int,
            mode: Optional[str] = None, percents: Optional[dict] = None,
            categories: Optional[dict] = None, prices: Optional[dict] = None,
            cleared: Optional[list] = None) -> Job:
    """Sozlamalarni qo'llab, darhol (shu jarayonda) qayta narxlaydi — worker ichida. Sayt so'rovlari
    ``jobs.recalc.request`` ni chaqiradi: sozlama saqlanadi, hisob fonda."""
    apply_settings(job, local_sheet, res_sheet, fuzzy, mode, percents, categories, prices, cleared)
    recompute(job)
    job.save()
    return job


# sozlamalar (smetachi o'zgartiradi) — qayta hisoblash ularga tegmaydi, aks holda hisob davomidagi o'zgarish yo'qoladi
SETTING_FIELDS = ["local_sheet", "res_sheet", "fuzzy", "mode", "percents", "categories", "prices", "cleared",
                  "local_columns", "res_columns", "column_names", "base_period"]
# natija — faqat qayta hisoblash yozadi
RESULT_FIELDS = ["status", "error", "duration", "rows", "output_path", "result"]


def apply_settings(job: Job, local_sheet: Optional[str], res_sheet: Optional[str], fuzzy,
                   mode: Optional[str] = None, percents: Optional[dict] = None,
                   categories: Optional[dict] = None, prices: Optional[dict] = None,
                   cleared: Optional[list] = None) -> Job:
    """``percents``/``categories``/``prices``/``cleared``: None — avvalgilari qoladi, {} / [] — tozalanadi.
    Lokal list almashsa, qo'lda kategoriya, narxlar va tozalashlar bekor (qator raqamlari boshqa listga tegishli)."""
    if local_sheet and local_sheet != job.local_sheet:
        job.categories = {}
        job.prices = {}
        job.cleared = []
        job.local_columns = {}  # ustunlar boshqa listga tegishli
    if res_sheet and res_sheet != job.res_sheet:
        job.res_columns = {}
    if categories is not None:
        job.categories = categories
    if prices is not None:
        job.prices = prices
    if cleared is not None:
        job.cleared = sorted({int(r) for r in cleared})
    job.local_sheet = local_sheet or job.local_sheet
    job.res_sheet = res_sheet or job.res_sheet
    job.fuzzy = max(80, min(100, int(fuzzy)))
    job.mode = mode if mode in MODES else job.mode
    if percents is not None:
        job.percents = percents
    if not job.percents:
        # смета ўз фоизини олмаган — смета­чининг охирги танлови (у ҳеч қачон ўчмайди)
        job.percents = remembered_percents(job.owner_id)
    return job


def recompute(job: Job) -> Job:
    """Joriy sozlamalar bilan narxlaydi va natija maydonlarini (``RESULT_FIELDS``) to'ldiradi — saqlamaydi.
    Natija fayli avval vaqtinchalik nomga yoziladi: hisob davomida yuklab olinayotgan eski fayl buzilmaydi."""
    out = job.abs("out") / job.download_name
    tmp = out.with_name(f".{out.stem}.tmp{out.suffix}")
    t0 = time.monotonic()
    try:
        result = run(job.abs(job.local_path), job.abs(job.res_path) if job.res_path else None, Options(
            fuzzy_percent=job.fuzzy, mode=job.mode, local_sheet=job.local_sheet, res_sheet=job.res_sheet,
            output_path=tmp, percents=job.percents or None,
            categories={int(r): c for r, c in (job.categories or {}).items()} or None,
            prices={int(r): float(p) for r, p in (job.prices or {}).items()} or None,
            column_names=job.column_names or None, cleared=job.cleared or None,
            local_columns=job.local_columns or None, res_columns=job.res_columns or None, **base_options(job),
        ))
    except Exception as e:  # foydalanuvchiga ko'rsatamiz, server yiqilmasin
        job.status = Job.Status.FAILED
        job.error = str(e) if isinstance(e, LayoutError) else f"{type(e).__name__}: {e}"  # жадвал тузилиши — тушунарли сабаб
        job.duration = round(time.monotonic() - t0, 2)
        tmp.unlink(missing_ok=True)
        return job

    tmp.replace(out)
    job.status = Job.Status.DONE
    job.error = ""
    job.duration = round(time.monotonic() - t0, 2)
    job.rows = len(result.rows)
    job.output_path = str(out.relative_to(job.dir))
    job.result = {
        "stats": result.stats,
        "totals": {k: _jsonable(v) for k, v in result.totals.items()},
        "problems": [{k: _jsonable(v) for k, v in p.items()} for p in result.problems],
        "rows": [{k: _jsonable(v) for k, v in r.items()} for r in result.rows],
        "start_row": result.start_row,
        "end_row": result.end_row,
        "sum_row": result.sum_row,
        "svodka_start": result.svodka_start,
        "mode": result.mode,
        "warnings": result.warnings,
        "checks": result.checks,  # narxlashdan keyingi tekshiruv: bo'sh bo'lmasa — natijaga ishonib bo'lmaydi
        "columns": result.columns,
        "percents": result.percents,
        "categories": {str(r): c for r, c in result.categories.items()},
        "prices": {str(r): p for r, p in result.prices.items()},
        "res_sections": result.res_sections,
        "extra_columns": [list(c) for c in result.extra_columns],
        "tail_removed": result.tail_removed,
        "cleared": result.cleared,
    }
    return job


# ===================== obyekt: bir yoki bir nechta fayl =====================
MAX_FILES = 10


def object_name(batch: "Batch", names: list) -> str:
    """Объект номи — локал файл номидан (ресурс файли эмас): «Намуна локал» ёки «Намуна локал + 1 та локал».
    Локал аниқланмаган бўлса — биринчи файл номи."""
    count = {}  # файл → [локал листлар, ресурс листлар]
    for sh in batch.sheets:
        role = sh.get("role")
        if role in ("local", "resource"):
            count.setdefault(int(sh["key"].split(":", 1)[0]), [0, 0])[role == "resource"] += 1
    # асосан локалдан иборат файллар (битта файлда ҳар иккаласи бўлиши мумкин); бўлмаса — локали борлари
    local_files = (sorted(i for i, (lo, re_) in count.items() if lo > re_)
                   or sorted(i for i, (lo, _) in count.items() if lo))
    if not local_files:
        return names[0] if len(names) == 1 else f"{names[0]} + {len(names) - 1} та файл"
    first = Path(names[local_files[0]]).stem
    return first if len(local_files) == 1 else f"{first} + {len(local_files) - 1} та локал"


def create_batch(owner, files) -> "Batch":
    """Fayllarni saqlaydi, .xlsx ga o'giradi, har bir listni tahlil qilib juftliklarni taklif qiladi."""
    if not files:
        raise UploadError("Камида битта файл танланг")
    if len(files) > MAX_FILES:
        raise UploadError(f"Бир марта кўпи билан {MAX_FILES} та файл юкланг")
    names = [Path(f.name).name for f in files]
    batch = Batch(owner=owner, name=names[0] if len(names) == 1 else f"{names[0]} + {len(names) - 1} та файл")
    clock, seconds = FileClock(), {}
    try:
        for i, f in enumerate(files):
            ext = _check_upload(f, f.name)
            _save(f, batch.abs(f"upload/{i}{ext}"))
            xlsx = clock.open(batch.abs(f"upload/{i}{ext}"), batch.abs(f"conv/{i}"), names[i], i)
            batch.files.append({"name": names[i], "path": str(xlsx.relative_to(batch.dir)), "sheets": clock.read(xlsx, i)})
        analyze_batch(batch, timings=seconds)
        batch.name = object_name(batch, names)
        clock.analyzed(seconds)
    except convert.ConversionError as e:
        shutil.rmtree(batch.dir, ignore_errors=True)
        save_timings(clock, getattr(owner, "pk", owner))
        raise UploadError(str(e)) from e
    except UploadError:
        shutil.rmtree(batch.dir, ignore_errors=True)
        raise
    except Exception as e:
        shutil.rmtree(batch.dir, ignore_errors=True)
        raise UploadError(f"Файлни очиб бўлмади: {e}") from e
    batch.save()
    save_timings(clock, batch.owner_id, batch=batch)
    return batch


class FileClock:
    """Bitta faylni ochish vaqtini o'lchaydi: ``convert`` (.xls → .xlsx) va ``read`` (listlar ro'yxati).
    ``rows`` ga yig'iladi, keyin ``save_timings`` bilan bazaga."""

    def __init__(self):
        self.rows = []

    def open(self, src: Path, out_dir: Path, name: str, index: int):
        src = Path(src)
        row = {"index": index, "name": name[:255], "ext": src.suffix.lower()[:8],
               "size": src.stat().st_size if src.exists() else None,
               "method": convert.converter_name() if src.suffix.lower() == ".xls" else ""}
        self.rows.append(row)
        try:
            t = time.monotonic()
            xlsx = convert.ensure_xlsx(src, out_dir)
            if xlsx != src:
                row["convert_s"] = round(time.monotonic() - t, 3)
            return xlsx
        except Exception as e:
            row["error"] = f"{type(e).__name__}: {e}"[:300]
            raise

    def read(self, xlsx: Path, index: int):
        row = next(r for r in self.rows if r["index"] == index)
        t = time.monotonic()
        try:
            return sheet_names(xlsx)
        except Exception as e:
            row["error"] = f"{type(e).__name__}: {e}"[:300]
            raise
        finally:
            row["read_s"] = round(time.monotonic() - t, 3)

    def analyzed(self, seconds: dict):
        for r in self.rows:
            if r["index"] in seconds:
                r["analyze_s"] = round(seconds[r["index"]], 3)


def save_timings(clock: FileClock, owner_id, project_id=None, batch=None, job=None):
    """Panel uchun — yozilmasa ham asosiy ish to'xtamaydi."""
    from .models import FileTiming
    try:
        FileTiming.objects.bulk_create([FileTiming(
            owner_id=owner_id, project_id=project_id, batch=batch if batch and batch.pk else None,
            job=job if job and job.pk else None, **{k: v for k, v in r.items() if k != "index"})
            for r in clock.rows])
    except Exception:  # noqa: BLE001
        pass


def analyze_batch(batch: "Batch", roles: Optional[dict] = None, columns: Optional[dict] = None,
                  timings: Optional[dict] = None) -> "Batch":
    """Listlar turini aniqlaydi (``roles`` / ``columns`` — smetachining qo'lda tuzatishlari: list turi
    va ustunlar), nusxalarni ajratadi va juftliklarni taklif qiladi."""
    if roles is not None:
        batch.roles = roles
    if columns is not None:
        batch.columns = columns
    infos, prices = [], {}
    for i, f in enumerate(batch.files):
        t = time.monotonic()
        wb = load_workbook(batch.abs(f["path"]), keep_links=False)
        try:
            found = inspect_workbook(wb, i, f["name"], batch.roles, batch.columns)
            for info in found:
                prices[info.key] = count_prices(wb[info.sheet], info.columns or {})
            infos += found
        finally:
            wb.close()
            if timings is not None:
                timings[i] = time.monotonic() - t
    apply_roles(infos, batch.roles or {})
    mark_duplicates(infos)
    pairs, _unused = plan_pairs(infos)
    batch.sheets = [dict(i.to_dict(), prices=prices.get(i.key)) for i in infos]
    # RES топилмаган локал — база ёқилган ва давр танланган бўлса, ресурс базасидан нархланади
    base = base_ready(batch.owner_id)
    batch.pairs = [dict(p.to_dict(), include=p.res is not None or base, base=p.res is None and base) for p in pairs]
    return batch


def count_prices(ws, columns: dict):
    """Ресурс листида нечта қаторда нарх бор (0 дан катта сон). ``None`` — нарх устуни аниқланмаган.
    Нархланмаган RES (ЦЕНА ҳамма жойда 0) билан нархлаш бефойда — смета­чи олдиндан огоҳлантирилади."""
    col, start, end = columns.get("price"), columns.get("start"), columns.get("end")
    if not col or not start:
        return None
    n = 0
    for (v,) in ws.iter_rows(min_row=start, max_row=end or ws.max_row, min_col=col, max_col=col, values_only=True):
        if isinstance(v, str):
            v = parse_text_number(v)
        if isinstance(v, (int, float)) and v > 0:
            n += 1
    return n


def unpriced_res(batch: "Batch") -> set:
    """Нархи йўқ RES листлари (калитлари)."""
    return {s["key"] for s in batch.sheets if s.get("role") == "resource" and s.get("prices") == 0}


def apply_decisions(batch: "Batch", decisions: dict):
    """Смета­чи танлови (қайси локал нархлансин, қайси RES билан, номи) жуфтликларга ёзилади."""
    resources = {s["key"] for s in batch.sheets if s["role"] == "resource" and not s.get("duplicate_of")}
    for n, p in enumerate(batch.pairs):
        d = decisions.get(n, {})
        p["include"] = bool(d.get("include", p.get("include")))
        if d.get("res") == BASE_RES:
            p["res"], p["base"] = None, True
        elif "res" in d:
            p["res"] = d["res"] if d["res"] in resources else None
            p["base"] = False
        if d.get("label", "").strip():
            p["label"] = d["label"].strip()[:200]


def confirm_batch(batch: "Batch", decisions: dict) -> list:
    """``decisions``: ``{juftlik_raqami: {"include": bool, "res": kalit|"", "label": str}}``.
    Har bir tanlangan juftlik uchun Job yaratib, darhol narxlaydi. Avvalgi natijalar o'chiriladi."""
    for job in batch.jobs.all():
        shutil.rmtree(job.dir, ignore_errors=True)
        job.delete()
    apply_decisions(batch, decisions)

    jobs = []
    base_period = base_period_of(batch.owner_id)
    from resbase.services import period_complete
    if not period_complete(base_period):  # база ўчирилган ёки давр танланмаган — RES'сиз жуфтлик нархланмайди
        for p in batch.pairs:
            if not p.get("res"):
                p["base"] = False
    todo = sum(1 for p in batch.pairs if pair_ready(p))
    for n, p in enumerate(batch.pairs):
        if not pair_ready(p):
            continue
        if batch.project_id:  # papka qismi: panel va sahifada jonli jarayon
            Batch.objects.filter(pk=batch.pk).update(step=f"{len(jobs) + 1}/{todo} смета нархланяпти: {p['label']}"[:300])
            Project.objects.filter(pk=batch.project_id).update(heartbeat=timezone.now())
        loc, res = batch.sheet(p["local"]), batch.sheet(p["res"]) if p.get("res") else None  # None — ресурс базасидан
        lf, rf = batch.files[loc["file"]], batch.files[res["file"]] if res else None
        job = Job(owner_id=batch.owner_id, batch=batch, batch_label=p["label"], batch_file=loc["file"], batch_order=n,
                  local_name=p["label"], res_name=f'{rf["name"]} · {res["sheet"]}' if res else "ресурс базаси",
                  local_sheets=lf["sheets"], res_sheets=rf["sheets"] if rf else [],
                  local_sheet=loc["sheet"], res_sheet=res["sheet"] if res else "", mode="smart",
                  percents=remembered_percents(batch.owner_id), base_period=base_period,
                  local_columns=(batch.columns or {}).get(p["local"], {}),
                  res_columns=(batch.columns or {}).get(p["res"], {}) if res else {})
        local_dst = job.abs(f"conv/{loc['file']}.xlsx")
        local_dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(batch.abs(lf["path"]), local_dst)
        job.local_path = str(local_dst.relative_to(job.dir))
        job.res_path = ""
        if res:
            res_dst = job.abs(f"conv/{res['file']}.xlsx")
            if res_dst != local_dst:
                shutil.copy2(batch.abs(rf["path"]), res_dst)
            job.res_path = str(res_dst.relative_to(job.dir))
        job.save()
        run_job(job, None, None, job.fuzzy, job.mode)
        jobs.append(job)
    batch.status = Batch.Status.DONE
    batch.save()
    return jobs


def delete_batch(batch: "Batch"):
    """Obyekt (bir/ko'p fayl) butunlay: yuklangan fayllar va uning smetalari."""
    dirs = [batch.dir] + [j.dir for j in batch.jobs.all()]
    batch.delete()
    for d in dirs:
        shutil.rmtree(d, ignore_errors=True)


def delete_job(job: Job):
    """Alohida smeta: yuklangan fayllar va natija."""
    d = job.dir
    job.delete()
    shutil.rmtree(d, ignore_errors=True)


def batch_outputs(batch: "Batch") -> list:
    """Har bir lokal fayl uchun natija: «<fayl>_narxlandi.xlsx». Faylda bitta narxlangan lokal bo'lsa —
    o'sha qismning tayyor fayli (qo'lda o'zgarishlari bilan), bir nechta bo'lsa — hammasi bitta faylda:
    har bir lokal list narxlanadi va o'z «LOG …» listini oladi."""
    from collections import defaultdict

    from core.pipeline import run_many

    groups = defaultdict(list)
    for job in batch.jobs.filter(status=Job.Status.DONE).order_by("batch_order"):
        groups[job.batch_file].append(job)
    if not groups:
        raise UploadError("Нархланган қисм йўқ")
    out_dir = batch.abs("out")
    shutil.rmtree(out_dir, ignore_errors=True)
    out_dir.mkdir(parents=True)
    outputs = []
    for f, jobs in groups.items():
        out = out_dir / f"{Path(batch.files[f]['name']).stem}_narxlandi.xlsx"
        if len(jobs) == 1 and jobs[0].output_path and jobs[0].abs(jobs[0].output_path).is_file():
            shutil.copy2(jobs[0].abs(jobs[0].output_path), out)
        else:
            steps = [(job.abs(job.res_path) if job.res_path else None, Options(
                mode=job.mode, fuzzy_percent=job.fuzzy, local_sheet=job.local_sheet, res_sheet=job.res_sheet,
                percents=job.percents or None,
                categories={int(r): c for r, c in (job.categories or {}).items()} or None,
                prices={int(r): float(v) for r, v in (job.prices or {}).items()} or None,
                column_names=job.column_names or None, cleared=job.cleared or None,
                local_columns=job.local_columns or None, res_columns=job.res_columns or None, **base_options(job),
            )) for job in jobs]
            run_many(batch.abs(batch.files[f]["path"]), steps, out)
        outputs.append(out)
    return outputs


def build_combined(batch: "Batch"):
    """Obyektning hamma natijasi: bitta fayl bo'lsa — .xlsx, bir nechta bo'lsa — .zip."""
    import zipfile

    outputs = batch_outputs(batch)
    if len(outputs) == 1:
        return outputs[0]
    zpath = outputs[0].parent / f"{Path(batch.files[0]['name']).stem}_narxlandi.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for o in outputs:
            z.write(o, o.name)
    return zpath
