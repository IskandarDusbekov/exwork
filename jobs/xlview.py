"""Excel listini saytda Excel'dagidek ko'rsatish: kataklar, birlashtirishlar, ramka, rang, shrift,
ustun kengligi va qator balandligi, son formati. Smetachi faylni yuklab olib ochmasdan ko'radi.

Natija — tayyor HTML (``<table>``) va listlar ro'yxati; JSON qilib kesh papkaga yoziladi
(fayl o'zgarmaguncha qayta hisoblanmaydi).
"""
import datetime
import hashlib
import json
import re
from decimal import ROUND_HALF_UP, Decimal
from html import escape
from pathlib import Path
from typing import Dict, Optional

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

from .templatetags.narx import _EXCEL_ERRORS_RU

MAX_ROWS = 3000  # bundan ko'p qator — kesiladi (sahifa og'irlashmasin)
CACHE_VERSION = 2  # 2: shrift nomi va rang tozalanadi (eski keshda XSS qolmasin)
GUTTER_PX = 46
_RED_FILL, _RED_FONT = "#FFC7CE", "#9C0006"  # writer.mark_unpriced bilan bir xil

# Office mavzusi (fayl o'z mavzusini bermasa); tartib — Excel'dagi theme indeksi bo'yicha
_DEFAULT_THEME = ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD", "C0504D",
                  "9BBB59", "8064A2", "4BACC6", "F79646", "0000FF", "800080"]
_THEME_TAGS = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
               "accent5", "accent6", "hlink", "folHlink"]
_BORDERS = {
    "thin": "1px solid", "hair": "1px solid", "dotted": "1px dotted", "dashed": "1px dashed",
    "dashDot": "1px dashed", "dashDotDot": "1px dashed", "medium": "2px solid", "mediumDashed": "2px dashed",
    "mediumDashDot": "2px dashed", "mediumDashDotDot": "2px dashed", "slantDashDot": "2px dashed",
    "thick": "3px solid", "double": "3px double",
}


def _theme(wb) -> list:
    xml = wb.loaded_theme
    if not xml:
        return list(_DEFAULT_THEME)
    if isinstance(xml, bytes):
        xml = xml.decode("utf-8", "ignore")
    out = []
    for tag, default in zip(_THEME_TAGS, _DEFAULT_THEME):
        m = re.search(rf"<a:{tag}>(.*?)</a:{tag}>", xml, re.S)
        v = m and re.search(r'lastClr="([0-9A-Fa-f]{6})"|srgbClr val="([0-9A-Fa-f]{6})"', m.group(1))
        out.append((v.group(1) or v.group(2)).upper() if v else default)
    return out


def _tint(hex6: str, tint: float) -> str:
    rgb = [int(hex6[i:i + 2], 16) for i in (0, 2, 4)]
    if tint < 0:
        rgb = [round(c * (1 + tint)) for c in rgb]
    elif tint > 0:
        rgb = [round(c + (255 - c) * tint) for c in rgb]
    return "".join(f"{max(0, min(255, c)):02X}" for c in rgb)


def _number(v, fmt: str) -> str:
    """Son Excel'dagi kabi (rus lokali: bo'shliq — minglik, vergul — kasr)."""
    sections = (fmt or "General").split(";")
    sec, own_sign = sections[0], False
    if v < 0 and len(sections) > 1:
        sec, v, own_sign = sections[1], -v, True  # manfiy bo'lim ishorani o'zi yozadi
    elif v == 0 and len(sections) > 2:
        sec = sections[2]
    sec = re.sub(r"\[[^\]]*\]|_.|\*.", "", sec)
    literals = []

    def keep(m):
        literals.append(m.group(1) if m.group(1) is not None else m.group(2))
        return chr(0xE000 + len(literals) - 1)  # shaxsiy belgi: «0», «#» bilan adashmaydi

    code = re.sub(r'"([^"]*)"|\\(.)', keep, sec)

    def lit(s):
        return re.sub("[-]", lambda m: literals[ord(m.group(0)) - 0xE000], s)

    m = re.search(r"[#0?][#0?,. ]*", code)
    if m is None:
        if re.search(r"general|основной", code, re.I) or not code.strip():
            text = str(int(v)) if float(v).is_integer() and abs(v) < 1e15 else f"{v:.10g}"
            text = text.replace(".", ",").replace("e+", "E+").replace("e-", "E-")
            return lit(re.sub(r"general|основной", lambda _m: text, code, flags=re.I)) if code.strip() else text
        return lit(code)
    prefix, spec, suffix = code[:m.start()], m.group(0).rstrip(" ,"), code[m.end():]
    if "%" in prefix + suffix:
        v *= 100
    whole, _, frac = spec.partition(".")
    max_dec, min_dec = sum(ch in "0#?" for ch in frac), frac.count("0")
    q = Decimal(repr(float(v))).quantize(Decimal(1).scaleb(-max_dec), ROUND_HALF_UP)  # Excel kabi: 0,5 -> 1
    text = format(q, ",f" if "," in whole else "f")
    if max_dec > min_dec and "." in text:  # 0.## — ortiqcha nollarsiz
        head, frac_digits = text.split(".")
        frac_digits = frac_digits.rstrip("0").ljust(min_dec, "0")
        text = head + ("." + frac_digits if frac_digits else "")
    if text.startswith("-"):
        text = text[1:]
        if not own_sign and float(text.replace(",", "") or 0) != 0:
            prefix = "-" + prefix
    text = text.replace(",", " ").replace(".", ",")
    return lit(prefix) + text + lit(suffix)


def display(v, fmt: str = "General") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "ИСТИНА" if v else "ЛОЖЬ"
    if isinstance(v, (int, float)):
        return _number(v, fmt)
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y %H:%M" if (v.hour or v.minute) else "%d.%m.%Y")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.strftime("%d.%m.%Y" if isinstance(v, datetime.date) else "%H:%M")
    s = str(v)
    return _EXCEL_ERRORS_RU.get(s, s)


_HEX6 = re.compile(r"[0-9A-Fa-f]{6}")
_FONT_BAD = re.compile(r"[^\w .\-]")


def _safe_font(name) -> str:
    """Shrift nomi fayldan keladi va <style> ichiga yoziladi: «'», «}», «<» kabi belgilar bilan CSS'dan chiqib
    HTML/skript qo'shib bo'lardi (XSS — admin boshqaning faylini ochganda uning sessiyasida). Faqat harf, raqam,
    bo'shliq, nuqta, chiziqcha qoladi."""
    return _FONT_BAD.sub("", str(name or ""))[:64].strip()


class _Styles:
    """Katak uslubi -> CSS sinfi (bir xil uslubli kataklar bitta sinfni bo'lishadi)."""

    def __init__(self, wb):
        self.theme = _theme(wb)
        self.classes: Dict[tuple, str] = {}
        self.rules = []

    def color(self, c) -> Optional[str]:
        if c is None:
            return None
        v = None
        if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) >= 6:
            v = c.rgb[-6:].upper()
        elif c.type == "indexed" and isinstance(c.indexed, int) and c.indexed < 64:
            v = COLOR_INDEX[c.indexed][-6:]
        elif c.type == "theme" and isinstance(c.theme, int) and c.theme < len(self.theme):
            v = self.theme[c.theme]
        if v is None or not _HEX6.fullmatch(v):
            return None  # fayldan kelgan qiymat CSS'ga tushadi — faqat 6 xonali hex
        if c.tint:
            v = _tint(v, c.tint)
        return "#" + v

    def side(self, s) -> str:
        if s is None or not s.style or s.style not in _BORDERS:
            return ""
        return f"{_BORDERS[s.style]} {self.color(s.color) or '#000'}"

    def css(self, cell, right, bottom, kind: str, red: bool) -> str:
        key = (tuple(cell._style) if cell.has_style else (), right, bottom, kind, red)
        name = self.classes.get(key)
        if name:
            return name
        name = f"c{len(self.classes)}"
        self.classes[key] = name
        p = []
        f = cell.font
        if f is not None:
            if f.b:
                p.append("font-weight:700")
            if f.i:
                p.append("font-style:italic")
            deco = [d for d, on in (("underline", f.u), ("line-through", f.strike)) if on]
            if deco:
                p.append("text-decoration:" + " ".join(deco))
            if f.sz and float(f.sz) != 11:
                p.append(f"font-size:{float(f.sz):g}pt")
            font = _safe_font(f.name)
            if font and font not in ("Calibri", "Arial"):
                p.append(f"font-family:'{font}',Calibri,Arial,sans-serif")
            col = self.color(f.color)
            if col and col != "#000000":
                p.append(f"color:{col}")
        fill = cell.fill
        if fill is not None and fill.fill_type and fill.fill_type != "none":
            col = self.color(fill.fgColor)
            if col:
                p.append(f"background:{col}")
        if red:
            p.append(f"background:{_RED_FILL};color:{_RED_FONT};font-weight:700")
        a = cell.alignment
        h = a.horizontal if a is not None else None
        if h in (None, "general"):
            h = {"num": "right", "bool": "center"}.get(kind, "left")
        h = {"centerContinuous": "center", "fill": "left", "distributed": "justify"}.get(h, h)
        if h != "left":
            p.append(f"text-align:{h}")
        v = a.vertical if a is not None else None
        p.append("vertical-align:" + {"top": "top", "center": "middle", "justify": "middle",
                                      "distributed": "middle"}.get(v, "bottom"))
        if a is not None and a.wrap_text:
            p.append("white-space:pre-wrap;overflow-wrap:anywhere")
        if a is not None and a.indent:
            p.append(f"padding-left:{3 + int(a.indent) * 9}px")
        if a is not None and a.textRotation in (90, 255):
            p.append("writing-mode:vertical-rl" + (";transform:rotate(180deg)" if a.textRotation == 90 else ""))
        if right:
            p.append(f"border-right:{right}")
        if bottom:
            p.append(f"border-bottom:{bottom}")
        self.rules.append(f".xs .{name}{{{';'.join(p)}}}")
        return name


def _kind(v) -> str:
    if isinstance(v, bool) or (isinstance(v, str) and v in _EXCEL_ERRORS_RU):
        return "bool"
    if isinstance(v, (int, float, datetime.date, datetime.time)):
        return "num"
    return "text"


def _bounds(ws):
    last_r = last_c = 0
    for (r, c), cell in ws._cells.items():
        if cell.value is not None and cell.value != "":
            last_r, last_c = max(last_r, r), max(last_c, c)
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= last_r:
            last_c = max(last_c, rng.max_col)
    return last_r, last_c


def _missing_formulas(path: Path, title: str) -> set:
    """Formulasi bor, lekin saqlangan natijasi yo'q kataklar (openpyxl saqlaganda natija yo'qoladi)."""
    out = set()
    wb = load_workbook(path, read_only=True, keep_links=False)
    try:
        for row in wb[title].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    out.add((cell.row, cell.column))
    finally:
        wb.close()
    return out


def _values(path: Path, title: str, cells: set) -> dict:
    out = {}
    if not cells:
        return out
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if title in wb.sheetnames:
            for row in wb[title].iter_rows(max_row=max(r for r, _c in cells)):
                for cell in row:
                    if (getattr(cell, "row", None), getattr(cell, "column", None)) in cells:
                        out[(cell.row, cell.column)] = cell.value
    finally:
        wb.close()
    return out


def render(path: Path, sheet: Optional[str] = None, default: Optional[str] = None,
           fallback: Optional[Path] = None, red: frozenset = frozenset(), red_sheet: Optional[str] = None,
           max_rows: int = MAX_ROWS) -> dict:
    """``sheet`` → ``default`` → faol list. ``fallback`` — natijasi yo'q formulalar qiymati shu fayldan
    (bir xil list va katak). ``red`` — qizil kataklar (``"G15"``), ``red_sheet`` listida (None — faol list)."""
    wb = load_workbook(path, data_only=True, keep_links=False)
    try:
        names = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] or wb.sheetnames
        title = next((t for t in (sheet, default) if t and t in wb.sheetnames), wb.active.title)
        if red_sheet is None:
            red_sheet = wb.active.title
        ws = wb[title]
        styles = _Styles(wb)
        last_r, last_c = _bounds(ws)
        total_rows = last_r
        truncated = last_r > max_rows
        last_r = min(last_r, max_rows)

        fixes = {}
        if fallback is not None and last_r:
            empty = {rc for rc in _missing_formulas(path, title) if rc[0] <= last_r and ws._cells.get(rc) is not None
                     and ws._cells[rc].value is None}
            fixes = _values(fallback, title, empty)

        fmt = ws.sheet_format
        default_w = fmt.defaultColWidth or ((fmt.baseColWidth or 8) + 0.43)
        widths, hidden_c = {}, set()
        for dim in ws.column_dimensions.values():
            lo = dim.min or 0
            hi = dim.max or lo
            for c in range(lo, min(hi, last_c) + 1):
                if dim.hidden:
                    hidden_c.add(c)
                if dim.width:
                    widths[c] = dim.width
        cols = [c for c in range(1, last_c + 1) if c not in hidden_c]
        rows = [r for r in range(1, last_r + 1)
                if not (ws.row_dimensions.get(r) is not None and ws.row_dimensions[r].hidden)]
        default_h = fmt.defaultRowHeight or 15

        covered, spans = set(), {}
        col_set, row_set = set(cols), set(rows)
        for rng in ws.merged_cells.ranges:
            if rng.min_row > last_r:
                continue
            spans[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
            for r in range(rng.min_row, min(rng.max_row, last_r) + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        covered.add((r, c))

        def value(r, c):
            cell = ws._cells.get((r, c))
            if cell is None or isinstance(cell, MergedCell):
                return None
            v = cell.value
            return fixes.get((r, c), v) if v is None else v

        def side(r, c, attr):
            cell = ws._cells.get((r, c))
            b = cell.border if cell is not None and cell.has_style else None
            return getattr(b, attr, None) if b is not None else None

        col_px = {c: round(widths.get(c, default_w) * 7 + 5) for c in cols}
        out = [f'<table class="xs" style="width:{GUTTER_PX + sum(col_px.values())}px"><colgroup>'
               f'<col style="width:{GUTTER_PX}px">']
        out += [f'<col style="width:{col_px[c]}px">' for c in cols]
        out.append('</colgroup><thead><tr><th class="xs-corner"></th>')
        out += [f"<th>{get_column_letter(c)}</th>" for c in cols]
        out.append("</tr></thead><tbody>")
        for r in rows:
            rd = ws.row_dimensions.get(r)
            ht = rd.ht if rd is not None and rd.ht else default_h
            out.append(f'<tr style="height:{round(ht * 4 / 3)}px"><th>{r}</th>')
            for i, c in enumerate(cols):
                if (r, c) in covered:
                    continue
                r2, c2 = spans.get((r, c), (r, c))
                v = value(r, c)
                text = display(v, ws._cells[(r, c)].number_format if (r, c) in ws._cells else "General")
                # o'ng/past ramka: o'zining yoki qo'shnisining (chap/yuqori) chizig'i
                right = styles.side(side(r, c2, "right")) or styles.side(side(r, c2 + 1, "left"))
                bottom = styles.side(side(r2, c, "bottom")) or styles.side(side(r2 + 1, c, "top"))
                cell = ws._cells.get((r, c))
                is_red = title == red_sheet and f"{get_column_letter(c)}{r}" in red
                cls = styles.css(cell, right, bottom, _kind(v), is_red) if cell is not None else (
                    styles.css(ws.cell(r, c), right, bottom, "text", is_red) if (right or bottom or is_red) else "")
                attrs = f' class="{cls}"' if cls else ""
                cs = sum(1 for x in range(c, c2 + 1) if x in col_set)
                rs = sum(1 for x in range(r, min(r2, last_r) + 1) if x in row_set)
                if cs > 1:
                    attrs += f' colspan="{cs}"'
                if rs > 1:
                    attrs += f' rowspan="{rs}"'
                if text and (r, c) not in spans and _kind(v) == "text":
                    # matn qo'shni bo'sh kataklarga toshadi (Excel kabi); qo'shni to'la bo'lsa — kesiladi
                    nxt = cols[i + 1] if i + 1 < len(cols) else None
                    if nxt is None or value(r, nxt) is not None or (r, nxt) in spans or (r, nxt) in covered:
                        attrs += ' data-clip="1"'
                out.append(f"<td{attrs}>{escape(text)}</td>" if text else f"<td{attrs}></td>")
            out.append("</tr>")
        out.append("</tbody></table>")
        return {
            "sheets": names, "sheet": title, "rows": len(rows), "cols": len(cols), "truncated": truncated,
            "total_rows": total_rows,
            "html": f"<style>{''.join(styles.rules)}</style>" + "".join(out),
        }
    finally:
        wb.close()


def cached(cache_dir: Path, key: str, path: Path, build) -> dict:
    """``build()`` natijasini fayl o'zgarmaguncha (mtime) saqlaydi."""
    stamp = f"{CACHE_VERSION}|{key}|{path.stat().st_mtime_ns}"
    target = cache_dir / (hashlib.sha1(stamp.encode()).hexdigest()[:20] + ".json")
    if target.is_file():
        try:
            return json.loads(target.read_text("utf-8"))
        except (OSError, ValueError):
            pass
    data = build()
    try:
        cache_dir.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    except OSError:
        pass
    return data
