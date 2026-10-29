"""Навбат ва worker: эгаллаш (lease), эскирган ишни қайта олиш, ``run_worker --once`` билан бутун папка."""
import io
import json
import shutil
import tempfile
from datetime import timedelta
from pathlib import Path
from unittest import mock

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import User

from . import folders, recalc
from .models import Batch, Job, Project, Worker

TMP_MEDIA = tempfile.mkdtemp(prefix="narxlash_worker_media_")
GOLDEN = Path(__file__).resolve().parent.parent / "core/tests/golden/202611450"


@override_settings(MEDIA_ROOT=TMP_MEDIA, NARXLASH_TASKS="worker")
class WorkerQueueTests(TransactionTestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="Ali", last_name="V")
        self.client.force_login(self.user)

    def _project(self, **kw):
        return Project.objects.create(owner=self.user, name="P", status=Project.Status.PROCESSING, **kw)

    def test_claim_is_exclusive(self):
        p = self._project()
        self.assertEqual(folders.claim_next("w1"), p.id)
        self.assertIsNone(folders.claim_next("w2"))  # бошқаси ололмайди
        p.refresh_from_db()
        self.assertEqual(p.worker, "w1")
        self.assertTrue(folders.is_working(p))

    def test_dead_worker_lease_is_taken_over(self):
        p = self._project(worker="dead:1#0", heartbeat=timezone.now() - folders.STALE_AFTER - timedelta(seconds=1))
        self.assertFalse(folders.is_working(p))
        self.assertEqual(folders.claim_next("w2"), p.id)
        fresh = self._project(worker="alive:1#0", heartbeat=timezone.now())
        self.assertIsNone(folders.claim_next("w3"))  # тирик worker'ники — тегилмайди
        self.assertTrue(folders.is_working(fresh))

    def test_finished_or_other_status_not_claimed(self):
        Project.objects.create(owner=self.user, name="D", status=Project.Status.DONE)
        self.assertIsNone(folders.claim_next("w1"))

    def test_start_does_not_spawn_thread_in_worker_mode(self):
        p = self._project()
        with mock.patch.object(folders.threading, "Thread") as thread:
            folders.start(p)
        thread.assert_not_called()

    def test_stop_waits_for_active_worker(self):
        p = self._project(worker="w1", heartbeat=timezone.now())
        Batch.objects.create(owner=self.user, project=p, folder="A", name="A", stage=Batch.Stage.QUEUED)
        folders.stop_project(p)
        p.refresh_from_db()
        self.assertTrue(p.stop_requested)  # worker кейинги файлдан олдин ўзи тўхтайди
        self.assertEqual(p.parts.get().stage, Batch.Stage.QUEUED)
        idle = self._project()
        Batch.objects.create(owner=self.user, project=idle, folder="B", name="B", stage=Batch.Stage.QUEUED)
        folders.stop_project(idle)  # ҳеч ким ишламаяпти — дарҳол
        idle.refresh_from_db()
        self.assertFalse(idle.stop_requested)
        self.assertEqual(idle.parts.get().stage, Batch.Stage.STOPPED)

    def test_run_worker_processes_uploaded_folder(self):
        local, res = (GOLDEN / "input_local.xlsx").read_bytes(), (GOLDEN / "input_res.xlsx").read_bytes()
        files = {"МТТ/А/a_lokal.xlsx": local, "МТТ/А/a_res.xlsx": res,
                 "МТТ/Б/b_lokal.xlsx": local, "МТТ/Б/b_res.xlsx": res}
        r = self.client.post(reverse("project_create"), {
            "files": [SimpleUploadedFile(p.split("/")[-1], b) for p, b in files.items()], "paths": list(files)})
        self.assertEqual(r.status_code, 200, r.content)
        project = Project.objects.get()
        # сайт ўзи ишламайди — навбатда турибди
        self.assertEqual(project.status, Project.Status.PROCESSING)
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.QUEUED})

        out = io.StringIO()
        call_command("run_worker", once=True, threads=2, poll=0.1, stdout=out, stderr=io.StringIO())
        project.refresh_from_db()
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.READY}, out.getvalue())
        self.assertEqual(project.status, Project.Status.READY)
        self.assertEqual(project.worker, "")  # lease бўшатилди
        self.assertTrue(all(p.durations.get("converting") is not None for p in project.parts.all()))
        self.assertFalse(Worker.objects.exists())  # чиққанда рўйхатдан ўчди
        # ҳар бир файл очиш вақти ёзилди (.xlsx — ўгирилмайди)
        from .models import FileTiming
        timings = list(FileTiming.objects.filter(project=project))
        self.assertEqual(sorted(t.name for t in timings), sorted(p.split("/")[-1] for p in files))
        self.assertTrue(all(t.convert_s is None and t.read_s is not None and t.analyze_s is not None for t in timings))
        self.assertTrue(all(t.size and t.owner_id == self.user.pk and not t.error for t in timings))

        # «Ҳисоблашни бошлаш» → яна навбатга, worker нархлайди
        self.client.post(reverse("project_price", args=[project.id]))
        project.refresh_from_db()
        self.assertEqual(project.status, Project.Status.PROCESSING)
        call_command("run_worker", once=True, threads=1, poll=0.1, stdout=io.StringIO(), stderr=io.StringIO())
        project.refresh_from_db()
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.DONE})
        self.assertEqual(project.status, Project.Status.DONE)
        job = project.parts.first().jobs.first()
        self.assertTrue(job.rows and job.duration is not None)  # тезлик: қатор сони ва вақт

    def test_lost_lease_stops_without_overwriting(self):
        """Worker узоқ жавоб бермай иш бошқасига ўтган бўлса — эскиси қолган қисмларни олмайди,
        ҳолатни ва янги эгасининг lease'ини ўзгартирмайди."""
        p = self._project(worker="new:1#0", heartbeat=timezone.now())
        part = Batch.objects.create(owner=self.user, project=p, folder="A", name="A", stage=Batch.Stage.QUEUED)
        with mock.patch.object(folders, "analyze_part") as analyze:
            folders._run(p.id, worker="old:1#0")
        analyze.assert_not_called()
        p.refresh_from_db()
        part.refresh_from_db()
        self.assertEqual((p.status, p.worker), (Project.Status.PROCESSING, "new:1#0"))
        self.assertEqual(part.stage, Batch.Stage.QUEUED)

    def test_lease_lost_mid_run(self):
        p = self._project()
        self.assertEqual(folders.claim_next("old:1#0"), p.id)
        a = Batch.objects.create(owner=self.user, project=p, folder="A", name="A", stage=Batch.Stage.QUEUED)
        b = Batch.objects.create(owner=self.user, project=p, folder="B", name="B", stage=Batch.Stage.QUEUED)

        def taken_over(part):  # биринчи қисм ишланаётганда lease бошқа worker'га ўтди
            Batch.objects.filter(id=part.id).update(stage=Batch.Stage.READY)
            Project.objects.filter(id=p.id).update(worker="new:1#0", heartbeat=timezone.now())

        with mock.patch.object(folders, "analyze_part", side_effect=taken_over) as analyze:
            folders._run(p.id, worker="old:1#0")
        self.assertEqual(analyze.call_count, 1)
        a.refresh_from_db()
        b.refresh_from_db()
        p.refresh_from_db()
        self.assertEqual((a.stage, b.stage), (Batch.Stage.READY, Batch.Stage.QUEUED))  # B — янги эгасида
        self.assertEqual((p.status, p.worker), (Project.Status.PROCESSING, "new:1#0"))

    def test_settle_keeps_part_queued_meanwhile(self):
        """settle() ҳолатни ҳисоблаётганда смета­чи «Ҳисоблашни бошлаш»ни босди — лойиҳа навбатдан чиқмайди."""
        p = self._project()
        part = Batch.objects.create(owner=self.user, project=p, folder="A", name="A", stage=Batch.Stage.READY)

        def clicked(*a, **kw):  # текширувдан кейин, ёзувдан олдин
            Batch.objects.filter(id=part.id).update(stage=Batch.Stage.PRICING)
            return False

        with mock.patch.object(folders, "Job") as job:
            job.objects.filter.return_value.exists.side_effect = clicked
            folders.settle(p)
        p.refresh_from_db()
        self.assertEqual(p.status, Project.Status.PROCESSING)
        self.assertEqual(folders.claim_next("w1"), p.id)  # worker уни олади

    def test_settle_without_queue(self):
        p = self._project()
        Batch.objects.create(owner=self.user, project=p, folder="A", name="A", stage=Batch.Stage.READY)
        folders.settle(p)
        p.refresh_from_db()
        self.assertEqual(p.status, Project.Status.READY)


@override_settings(MEDIA_ROOT=TMP_MEDIA, NARXLASH_TASKS="thread")
class ThreadModeLeaseTests(TransactionTestCase):
    """Сайт ичидаги оқим ҳам lease олади: ёнида worker ҳам ишласа — битта папкани иккаласи олмайди."""

    def test_thread_skips_project_held_by_worker(self):
        user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="A", last_name="V")
        p = Project.objects.create(owner=user, name="P", status=Project.Status.PROCESSING,
                                   worker="srv:1#0", heartbeat=timezone.now())
        with mock.patch.object(folders.threading, "Thread") as thread:
            folders.start(p)
        thread.assert_not_called()
        self.assertNotIn(p.id, folders._running)

    def test_thread_takes_lease(self):
        user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="A", last_name="V")
        p = Project.objects.create(owner=user, name="P", status=Project.Status.PROCESSING)
        with mock.patch.object(folders.threading, "Thread") as thread:
            folders.start(p)
        thread.assert_called_once()
        folders._running.discard(p.id)
        p.refresh_from_db()
        self.assertTrue(p.worker.startswith("thread:"))
        self.assertIsNone(folders.claim_next("w1"))  # worker уни ололмайди


class UnpricedResTests(TransactionTestCase):
    """RES'да нарх умуман йўқ бўлса (ЦЕНА ҳамма жойда 0) — нархлашдан олдин айтилади."""

    def test_count_prices(self):
        from openpyxl import Workbook

        from .services import count_prices
        ws = Workbook().active
        for r, v in enumerate([None, 0, "0", "0,00", 0.0], start=2):
            ws.cell(r, 3, v)
        cols = {"price": 3, "start": 2, "end": 6}
        self.assertEqual(count_prices(ws, cols), 0)
        ws.cell(4, 3, "12 500,50")
        self.assertEqual(count_prices(ws, cols), 1)
        self.assertIsNone(count_prices(ws, {"name": 2}))  # нарх устуни аниқланмаган — номаълум

    def test_ready_stage_warns_before_pricing(self):
        part = Batch(pairs=[{"local": "0:L", "res": "0:R", "include": True, "label": "A"}],
                     sheets=[{"key": "0:L", "role": "local"}, {"key": "0:R", "role": "resource", "prices": 0}])
        stage, msg = folders.ready_stage(part)
        self.assertEqual(stage, Batch.Stage.REVIEW)
        self.assertIn("нарх йўқ", msg)
        part.sheets[1]["prices"] = 40
        self.assertEqual(folders.ready_stage(part), (Batch.Stage.READY, ""))
        part.sheets[1].pop("prices")  # эски таҳлил (маълумот йўқ) — огоҳлантирилмайди
        self.assertEqual(folders.ready_stage(part)[0], Batch.Stage.READY)


@override_settings(MEDIA_ROOT=TMP_MEDIA, NARXLASH_TASKS="worker")
class PartQueueTests(TransactionTestCase):
    def setUp(self):
        self.user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="Ali", last_name="V")
        self.client.force_login(self.user)

    def test_part_confirm_goes_to_queue_not_request(self):
        local, res = (GOLDEN / "input_local.xlsx").read_bytes(), (GOLDEN / "input_res.xlsx").read_bytes()
        files = {"МТТ/А/a_lokal.xlsx": local, "МТТ/А/a_res.xlsx": res}
        self.client.post(reverse("project_create"), {
            "files": [SimpleUploadedFile(p.split("/")[-1], b) for p, b in files.items()], "paths": list(files)})
        call_command("run_worker", once=True, threads=1, poll=0.1, stdout=io.StringIO(), stderr=io.StringIO())
        project = Project.objects.get()
        part = project.parts.get()
        self.assertEqual(part.stage, Batch.Stage.READY)
        self.assertTrue(part.sheets and any(s.get("prices") for s in part.sheets))  # RES'да нарх бор

        with mock.patch("jobs.folders.price_part") as price_part:
            r = self.client.post(reverse("batch_confirm", args=[part.id]),
                                 {"include_0": "1", "res_0": part.pairs[0]["res"], "label_0": "Янги ном"})
        price_part.assert_not_called()  # сўров ичида нархланмайди
        self.assertRedirects(r, reverse("project", args=[project.id]), fetch_redirect_response=False)
        part.refresh_from_db()
        project.refresh_from_db()
        self.assertEqual((part.stage, project.status), (Batch.Stage.PRICING, Project.Status.PROCESSING))
        self.assertEqual(part.pairs[0]["label"], "Янги ном")

        call_command("run_worker", once=True, threads=1, poll=0.1, stdout=io.StringIO(), stderr=io.StringIO())
        part.refresh_from_db()
        self.assertEqual(part.stage, Batch.Stage.DONE)
        self.assertEqual(part.jobs.get().local_name, "Янги ном")

    def test_interrupted_part_is_requeued(self):
        project = Project.objects.create(owner=self.user, name="P", status=Project.Status.READY)
        Batch.objects.create(owner=self.user, project=project, folder="A", name="A", stage=Batch.Stage.PRICING)
        folders.ensure_running(project)  # иш узилган: қисм «нархланяпти», лекин ҳеч ким ишламаяпти
        project.refresh_from_db()
        self.assertEqual(project.status, Project.Status.PROCESSING)
        self.assertEqual(folders.claim_next("w1"), project.id)


class DownloadNameTests(TransactionTestCase):
    def test_newline_in_label_is_removed(self):
        """Лист номида қатор ташланиши бўлса — Windows'да натижа файли яратилмасди (OSError)."""
        name = Job(batch_label="ДАННЫМ ЛОКАЛЬНЫХ СМЕТ \nНА ЦЕНТРАЛЬНЫЙ СКЛАД", local_name="x.xlsx").download_name
        self.assertEqual(name, "ДАННЫМ ЛОКАЛЬНЫХ СМЕТ НА ЦЕНТРАЛЬНЫЙ СКЛАД_narxlandi.xlsx")


@pytest.mark.recalc_queue
@override_settings(MEDIA_ROOT=TMP_MEDIA, NARXLASH_TASKS="worker")
class RecalcQueueTests(TransactionTestCase):
    """Смета ўзгарди (L/M устун, нарх, фоиз…) — сўров ичида қайта нархланмайди: navbat → worker."""

    def setUp(self):
        self.user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="Ali", last_name="V")
        self.client.force_login(self.user)
        local, res = (GOLDEN / "input_local.xlsx").read_bytes(), (GOLDEN / "input_res.xlsx").read_bytes()
        files = {"МТТ/А/a_lokal.xlsx": local, "МТТ/А/a_res.xlsx": res}
        self.client.post(reverse("project_create"), {
            "files": [SimpleUploadedFile(p.split("/")[-1], b) for p, b in files.items()], "paths": list(files)})
        self._work()
        part = Project.objects.get().parts.get()
        self.client.post(reverse("batch_confirm", args=[part.id]), {"include_0": "1", "res_0": part.pairs[0]["res"]})
        self._work()
        self.job = Job.objects.get()
        self.assertEqual((self.job.status, self.job.mode, self.job.recalc_pending), (Job.Status.DONE, "smart", False))
        self.row = next(r["row"] for r in self.job.result["rows"]
                        if r.get("status") == "EXACT" and r.get("category", "K") == "K")

    def _work(self):
        call_command("run_worker", once=True, threads=1, poll=0.1, stdout=io.StringIO(), stderr=io.StringIO())

    def _move(self, category="M"):
        return self.client.post(reverse("job_categories", args=[self.job.id]),
                                json.dumps({"rows": [self.row], "category": category}), content_type="application/json")

    def _category(self):
        self.job.refresh_from_db()
        return next(r for r in self.job.result["rows"] if r["row"] == self.row).get("category")

    def test_edit_is_queued_and_worker_recalculates(self):
        with mock.patch.object(recalc, "recompute") as recompute:
            r = self._move()
        recompute.assert_not_called()  # сўров ичида ҳисобланмайди
        self.assertEqual(r.json()["queued"], True)
        self.job.refresh_from_db()
        self.assertTrue(self.job.recalc_pending)
        self.assertEqual(self.job.categories, {str(self.row): "M"})  # созлама дарҳол сақланди
        self.assertEqual(self._category(), "K")  # натижа ҳали эски
        state = self.client.get(reverse("job_state", args=[self.job.id])).json()
        self.assertEqual((state["pending"], state["running"]), (True, False))

        self._work()
        self.assertEqual(self._category(), "M")
        self.assertFalse(self.job.recalc_pending)
        state = self.client.get(reverse("job_state", args=[self.job.id])).json()
        self.assertFalse(state["pending"])
        self.assertIn("rows", state["svodka"])

    def test_repeated_edits_are_merged_into_one_run(self):
        self._move("M")
        self._move("L")
        self.job.refresh_from_db()
        self.assertEqual(self.job.recalc_seq - self.job.recalc_done, 2)
        with mock.patch.object(recalc, "recompute", wraps=recalc.recompute) as recompute:
            self._work()
        self.assertEqual(recompute.call_count, 1)  # иккита ўзгариш — битта ҳисоб
        self.assertEqual(self._category(), "L")

    def test_edit_during_run_keeps_job_queued(self):
        self._move("M")
        self.assertEqual(recalc.claim_next("w1"), self.job.id)
        self.assertIsNone(recalc.claim_next("w2"))  # бошқаси ололмайди

        real = recalc.recompute

        def edit_meanwhile(job):
            self._move("L")  # ҳисоб давомида смета­чи яна ўзгартирди
            return real(job)

        with mock.patch.object(recalc, "recompute", side_effect=edit_meanwhile):
            recalc.run(self.job.id, "w1")
        self.job.refresh_from_db()
        self.assertTrue(self.job.recalc_pending)  # кейинги ўзгариш учун навбатда қолди
        self.assertEqual(self.job.categories, {str(self.row): "L"})  # натижа созлама устидан ёзмади
        self._work()
        self.assertEqual(self._category(), "L")
        self.assertFalse(self.job.recalc_pending)

    def test_stale_lease_is_taken_over(self):
        self._move()
        self.assertEqual(recalc.claim_next("old#0"), self.job.id)
        Job.objects.filter(id=self.job.id).update(recalc_heartbeat=timezone.now() - timedelta(minutes=10))
        self.assertEqual(recalc.claim_next("new#0"), self.job.id)
        recalc.run(self.job.id, "old#0")  # эски worker кеч тугатди — натижа ёзилмайди
        self.job.refresh_from_db()
        self.assertTrue(self.job.recalc_pending)
        recalc.run(self.job.id, "new#0")
        self.assertEqual(self._category(), "M")
