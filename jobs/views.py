import json
from datetime import timedelta

from django.contrib.auth.decorators import login_not_required
from django.contrib.auth.views import redirect_to_login
from django.core.cache import cache
from django.db.models import Prefetch
from django.db.models.fields.json import KeyTransform
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

from openpyxl.utils import column_index_from_string, get_column_letter

from core.batch import KIND_LABELS
from core.pipeline import UNPRICED_STATUSES, extra_columns_of
from core.svodka import DEFAULT_PERCENTS, PERCENT_LABELS, parse_text_number, percent_text, svodka_layout

from panel.audit import log as audit
from panel.models import Event

from . import preview, recalc, xlview
from .archives import is_archive
from .guide import FLOW, GUIDE_PAGES, help_groups
from .folders import (
    build_project_zip, build_tree, create_project, delete_project, ensure_running, is_shown, queue_part,
    ready_stage, resume_project, settle, start_pricing, stop_project,
)
from .models import Batch, Job, Project
from .services import (
    BASE_RES, UploadError, analyze_batch, base_period_of, check_quota, pair_ready, build_combined, confirm_batch, create_batch, create_job, delete_batch, delete_job,
    parse_percents, remember_base_period, remember_percents, remembered_percents, unpriced_res,
)
from .templatetags.narx import file_title, pct as pct_text, pct_input, readable, som

STAT_CARDS = {
    "compat": [
        ("EXACT", "Аниқ топилди", "emerald"),
        ("FUZZY", "Ўхшаш (fuzzy)", "amber"),
        ("AMBIGUOUS", "ШАБҲАЛИ", "orange"),
        ("NOT_FOUND", "ТОПИЛМАДИ", "rose"),
        ("ZERO_RULE", "Нол қоидаси", "slate"),
        ("NON_NUMERIC", "G да рақам эмас", "rose"),
    ],
    "smart": [
        ("EXACT", "Аниқ топилди", "emerald"),
        ("NOT_PRICED", "Нархланмади — қўлда", "rose"),
        ("FUZZY", "Ўхшаш (аниқ эмас)", "rose"),
        ("AMBIGUOUS", "ШАБҲАЛИ", "rose"),
        ("UNIT_MISMATCH", "Бирлик бошқа", "rose"),
        ("NO_PRICE", "Нархсиз ресурс", "rose"),
        ("NOT_FOUND", "ТОПИЛМАДИ", "rose"),
        ("BASE", "Ресурс базасидан", "violet"),
        ("MANUAL", "Қўлда нархланди", "sky"),
        ("WORK_HEADER", "Иш сарлавҳаси → 0", "slate"),
        ("MACHINIST", "Машинист → 0", "slate"),
    ],
}
MODE_LABELS = {"smart": "Smart (тавсия)", "compat": "Эски дастур каби (compat)"}
CATEGORY_LABELS = {"I": "Иш ҳақи", "J": "Машина", "K": "Материал", "L": "Кабель", "M": "Ускуна"}
LAST_EXTRA_COLUMN = "Z"


_KIND_OF = {Project: "folder", Batch: "object", Job: "smeta"}


def _get_work(request, model, **kw):
    """Ўз иши. Суперфойдаланувчи текширув учун бошқанинг ишини ҳам очади — фақат GET (кўриш ва юклаб олиш);
    ўзгартирадиган (POST) сўровлар фақат эгасига. Бошқанинг иши очилса — журналга ёзилади."""
    if request.method in ("GET", "HEAD") and request.user.is_superuser:
        obj = get_object_or_404(model, **kw)
        if obj.owner_id != request.user.pk:
            request.viewing_owner = obj.owner
            # саҳифа ўзи янгиланиб туради — журналга 10 дақиқада бир марта
            if cache.add(f"panel:viewed:{request.user.pk}:{obj.pk}", 1, 600):
                audit(request, Event.Kind.ADMIN_VIEW, target=obj.owner, work=obj, work_kind=_KIND_OF[model],
                      detail=str(getattr(obj, "name", "") or getattr(obj, "local_name", "")))
        return obj
    return get_object_or_404(model, owner=request.user, **kw)


def _audit_download(request, obj, kind, name):
    if getattr(request, "viewing_owner", None):
        audit(request, Event.Kind.ADMIN_DOWNLOAD, target=obj.owner, work=obj, work_kind=kind, detail=f"натижа: {name}")


def column_labels(job) -> dict:
    """Жадвал устунлари ва номлари: I..M (смета­чи қайта номлаган бўлиши мумкин) + қўшилган N, O, …"""
    names = job.column_names or {}
    labels = {c: (names.get(c) or label) for c, label in CATEGORY_LABELS.items()}
    for col, name in extra_columns_of(names):
        labels[col] = name or f"Устун {col}"
    return labels


@login_not_required
@require_http_methods(["GET", "POST"])
def index(request):
    """Бош саҳифа: биринчи келган (кирмаган) одамга — «Тизим ҳақида» (landing), кирганга — юклаш."""
    if not request.user.is_authenticated:
        if request.method == "POST":
            return redirect_to_login(request.get_full_path())
        return landing(request)
    error = None
    if request.method == "POST":
        try:
            check_quota(request.user)  # файллар ўқилишидан олдин
            files = request.FILES.getlist("files")
            local = request.FILES.get("local")
            if files and any(is_archive(f.name) for f in files):
                # ZIP/RAR — ичида папкалар бўлади: папка каби очилади (ҳисоблаш смета­чи босгандагина)
                project = create_project(request.user, files, [f.name for f in files])
                audit(request, Event.Kind.UPLOAD, work=project, work_kind="folder", detail=project.name)
                return redirect("project", project_id=project.id)
            if files:
                # Universal yuklash: bitta obyekt fayli yoki bir nechta fayl — tizim o'zi ajratadi
                batch = create_batch(request.user, files)
                audit(request, Event.Kind.UPLOAD, work=batch, work_kind="object", detail=batch.name)
                return redirect("batch", batch_id=batch.id)
            if local:
                job = create_job(request.user, local, request.FILES.get("res"))  # eski ikki maydonli usul
                audit(request, Event.Kind.UPLOAD, work=job, work_kind="smeta", detail=job.local_name)
                return redirect("job", job_id=job.id)
            error = "Файлни танланг."
        except UploadError as e:
            error = str(e)
    recent = work_items(request.user, limit=6)
    return render(request, "jobs/index.html", {"error": error, "recent": recent, "resume": recent[0] if recent else None,
                                               "nav": "upload", **_base_ctx(base_period_of(request.user))})


def _base_ctx(period) -> dict:
    """«Нарх манбаи» танлови учун (ресурс базаси ёқилган бўлса): жорий давр ва рўйхатлар."""
    from resbase.models import QUARTERS, BaseConfig
    from resbase.regions import REGIONS
    from resbase.services import period_label
    if not BaseConfig.get().enabled:
        return {"base_on": False}
    y = timezone.now().year
    return {"base_on": True, "base_period": period, "base_label": period_label(period), "regions": REGIONS,
            "years": list(range(y + 1, 2017, -1)), "quarters": QUARTERS}


@require_POST
def base_period(request):
    """Юклаш саҳифасида смета­чи танлаган вилоят ва давр — кейинги юклашлар шу билан нархланади."""
    from resbase.services import parse_period, period_label
    period = parse_period(request.POST)
    remember_base_period(request.user, period)
    return JsonResponse({"ok": True, "label": period_label(period)})


# Рўйхатлар учун сметанинг фақат жами суммаси ва ҳолати керак: тўлиқ ``result`` (минглаб қатор) юкланмайди
_SUMMARY = {
    "r_total": KeyTransform("s_nds", KeyTransform("totals", "result")),
    "r_not_priced": KeyTransform("NOT_PRICED", KeyTransform("stats", "result")),
    "r_checks": KeyTransform("checks", "result"),
}


def light_jobs(qs):
    return qs.defer("result").annotate(**_SUMMARY)


def _summary(j) -> dict:
    """``result``дан керакли қисм — ``light_jobs`` билан олинган бўлса, тўлиқ натижани юкламасдан."""
    if hasattr(j, "r_total"):
        return {"totals": {"s_nds": j.r_total}, "stats": {"NOT_PRICED": j.r_not_priced or 0}, "checks": j.r_checks}
    return j.result


def _job_sums(jobs):
    total = sum(t for j in jobs for t in [_summary(j).get("totals", {}).get("s_nds")] if isinstance(t, (int, float)))
    return total, sum(_summary(j).get("stats", {}).get("NOT_PRICED", 0) or 0 for j in jobs if j.status == Job.Status.DONE)


# Ишнинг умумий ҳолати (тарих ва дашборд): номи, ранги, белгиси
WORK_STATES = {
    "working": ("Ишланяпти", "sky", "loader"),
    "awaiting": ("Ҳисоблашга тайёр", "violet", "clock"),
    "stopped": ("Тўхтатилди", "slate", "stop"),
    "review": ("Текширинг", "amber", "alert"),
    "failed": ("Хато", "rose", "x-circle"),
    "warn": ("Нархланмаган бор", "rose", "alert-circle"),
    "done": ("Тайёр", "emerald", "check"),
    "empty": ("Смета топилмади", "slate", "minus"),
}
# Нега смета топилмади — ўтказиб юборилган листлар тури (core.batch.KIND_*) бўйича
SKIP_REASONS = {
    "priced": "локал смета аллақачон нархланган",
    "act": "дефект акти — ишлар рўйхати",
    "summary": "сводка / умумий ҳисоб",
    "catalog": "нарх каталоги (миқдор йўқ)",
    "empty": "бўш лист",
}


def _skip_reason(sheets) -> str:
    """Листлар орасида локал смета йўқлигининг сабаби: «локал смета аллақачон нархланган · сводка»."""
    kinds = {s.get("kind") for s in sheets if s.get("role") != "local"}
    reasons = [r for k, r in SKIP_REASONS.items() if k in kinds]
    return " · ".join(reasons) or "«Наименование» ва «Ед. изм.» жадвали топилмади"
# Смета (папка қисми) босқичлари — дашборддаги «қувур»
PIPELINE = [
    ("opening", "Очиляпти", "sky", "Навбатда ёки файллари очиляпти"),
    ("pricing", "Ҳисобланяпти", "indigo", "Ҳозир фонда нархланяпти"),
    ("pending", "Ҳисобланмаган", "violet", "Тайёр ёки тўхтатилган — «Ҳисоблашни бошлаш»ни кутяпти"),
    ("review", "Текширинг", "amber", "Ресурс танлаш ёки хатони кўриш керак"),
    ("warn", "Нархланмаган қатор бор", "rose", "Ҳисобланди, лекин айрим қаторлар қўлда нархланиши керак"),
    ("done", "Ҳисобланди", "emerald", "Тўлиқ нархланди"),
]
_STAGE_PIPE = {"queued": "opening", "converting": "opening", "pricing": "pricing", "ready": "pending",
               "stopped": "pending", "review": "review", "failed": "review"}


def _pipe_of_jobs(jobs, fallback):
    if any(_summary(j).get("checks") for j in jobs if j.status == Job.Status.DONE):
        return "review"
    if any(j.status == Job.Status.DONE for j in jobs):
        return "warn" if any(_summary(j).get("stats", {}).get("NOT_PRICED") for j in jobs) else "done"
    return fallback


def work_items(user, limit=None):
    """Ишлар (папка, объект, алоҳида смета) — энг янгиси биринчи; ҳолати ва суммаси билан.
    ``pipes`` — ҳар бир смета (папка қисми) қайси босқичда: дашборд статистикаси шундан йиғилади."""
    items = []
    for pr in Project.objects.filter(owner=user).prefetch_related(Prefetch("parts__jobs", light_jobs(Job.objects)))[:limit]:
        parts = [p for p in pr.parts.all() if p.stage != Batch.Stage.EMPTY]
        jobs = [j for p in parts for j in p.jobs.all()]
        total, not_priced = _job_sums(jobs)
        stages = [p.stage for p in parts]
        finished = sum(1 for s in stages if s not in ("queued", "converting", "pricing"))
        stopped, awaiting = stages.count(Batch.Stage.STOPPED), stages.count(Batch.Stage.READY)
        review = sum(1 for s in stages if s in ("review", "failed"))
        reason = ""
        if pr.status == Project.Status.PROCESSING:
            state = "working"
        elif not parts:
            state, reason = "empty", _skip_reason([s for p in pr.parts.all() for s in p.sheets])
        else:
            state = ("stopped" if stopped else "awaiting" if awaiting else "review" if review
                     else "warn" if not_priced else "done")
        items.append({"kind": "folder", "name": pr.name, "url": reverse("project", args=[pr.id]), "created": pr.created,
                      "delete_url": reverse("work_delete", args=["folder", pr.id]),
                      "stop_url": reverse("project_stop", args=[pr.id]),
                      "meta": f"{len(parts)} та смета", "reason": reason, "working": state == "working",
                      "progress": f"{finished}/{len(parts)}", "percent": round(100 * finished / len(parts)) if parts else 100,
                      "total": total, "not_priced": not_priced, "review": review, "stopped": stopped, "awaiting": awaiting,
                      "state": state, "pipes": [_STAGE_PIPE.get(p.stage) or _pipe_of_jobs(list(p.jobs.all()), "done")
                                                for p in parts]})
    for b in Batch.objects.filter(owner=user, project__isnull=True).prefetch_related(
            Prefetch("jobs", light_jobs(Job.objects)))[:limit]:
        jobs = list(b.jobs.all())
        total, not_priced = _job_sums(jobs)
        state = _pipe_of_jobs(jobs, "review") if b.status == Batch.Status.DONE else "review" if b.pairs else "empty"
        items.append({"kind": "object", "name": b.name, "url": reverse("batch", args=[b.id]), "created": b.created,
                      "delete_url": reverse("work_delete", args=["object", b.id]),
                      "meta": f"{len(b.pairs)} та жуфтлик", "working": False,
                      "reason": _skip_reason(b.sheets) if state == "empty" else "", "total": total, "not_priced": not_priced,
                      "review": 0 if b.status == Batch.Status.DONE else 1, "state": state, "pipes": [state]})
    for j in light_jobs(Job.objects.filter(owner=user, batch__isnull=True))[:limit]:
        total, not_priced = _job_sums([j]) if j.status == Job.Status.DONE else (0, 0)
        state = ("failed" if j.status == Job.Status.FAILED else "review" if j.status != Job.Status.DONE
                 else "warn" if not_priced else "done")
        items.append({"kind": "smeta", "name": j.local_name, "url": reverse("job", args=[j.id]), "created": j.created,
                      "delete_url": reverse("work_delete", args=["smeta", j.id]),
                      "meta": j.local_sheet, "reason": "", "working": False, "total": total, "not_priced": not_priced,
                      "review": 1 if j.status != Job.Status.DONE else 0, "state": state,
                      "pipes": [state if state in ("warn", "done") else "review"]})
    items.sort(key=lambda x: x["created"], reverse=True)
    for it in items:
        it["state_label"], it["state_color"], it["state_icon"] = WORK_STATES[it["state"]]
    return items[:limit] if limit else items


def _tn_summary(cards) -> dict:
    """Техназор объектлари: умумий шартнома, бажарилган ва қолдиқ (пудратчи, ҚҚС билан), энг янгилари."""
    contract = sum(c["contract"] for c in cards)
    done = sum(c["done"] for c in cards)
    return {"cards": cards[:5], "count": len(cards), "contract": contract, "done": done, "left": contract - done,
            "pct": 100 * done / contract if contract else 0.0}


# «Эътибор керак»: ҳолат -> смета­чи нима қилиши керак (иш маълумотидан)
def _todo(it) -> str:
    st = it["state"]
    if st == "warn":
        return f"{it['not_priced']} та қаторга нарх ёзинг"
    if st == "awaiting":
        return f"{it['awaiting']} та қисмни ҳисоблашни бошланг" if it["kind"] == "folder" else "Ҳисоблашни бошланг"
    if st == "stopped":
        return "Тўхтатилган — давом эттиринг"
    if st == "failed":
        return "Очиб бўлмади — сабабини кўринг"
    if it["kind"] == "folder" and it.get("review"):
        return f"{it['review']} та қисмни текширинг"
    return "Ресурсни танланг ёки хатони кўринг"


ATTENTION = ("warn", "awaiting", "stopped", "review", "failed")


def dashboard(request):
    """Дашборд: нима тайёр ва нима смета­чини кутяпти (нархлаш ҳамда техназор), сметалар босқичлари, барча ишлар."""
    from texnazor.services import object_cards
    items = work_items(request.user)
    attention = [dict(it, todo=_todo(it)) for it in items if it["state"] in ATTENTION]
    pipes = [p for it in items for p in it["pipes"]]
    pipeline = [{"key": k, "label": label, "color": color, "hint": hint, "count": pipes.count(k),
                 "pct": round(100 * pipes.count(k) / len(pipes), 1) if pipes else 0}
                for k, label, color, hint in PIPELINE]
    states = [{"key": k, "label": label, "color": color, "icon": ic, "count": sum(1 for it in items if it["state"] == k)}
              for k, (label, color, ic) in WORK_STATES.items()]
    # охирги 14 кун: ҳар куни нечта иш юкланди
    today = timezone.localdate()
    per_day = {today - timedelta(days=i): 0 for i in range(13, -1, -1)}
    for it in items:
        d = timezone.localtime(it["created"]).date()
        if d in per_day:
            per_day[d] += 1
    peak = max(per_day.values()) or 1
    return render(request, "jobs/dashboard.html", {
        "nav": "dashboard",
        "items": items,
        "latest": items[:2],
        "last_total": next((it for it in items if it["total"]), None),
        "pipeline": pipeline,
        "states": states,
        "kinds": {k: sum(1 for it in items if it["kind"] == k) for k in ("folder", "object", "smeta")},
        "count_smetas": len(pipes),
        "count_priced": pipes.count("done") + pipes.count("warn"),
        "count_active": pipes.count("opening") + pipes.count("pricing"),
        "count_pending": pipes.count("pending") + pipes.count("review"),
        "count_done": pipes.count("done"),
        "count_warn": pipes.count("warn"),
        "done_pct": round(100 * pipes.count("done") / len(pipes)) if pipes else 0,
        "attention": attention[:6],
        "attention_total": len(attention),
        "tn": _tn_summary(object_cards(request.user)),
        "priced_pct": round(100 * (pipes.count("done") + pipes.count("warn")) / len(pipes)) if pipes else 0,
        "any_working": any(it["working"] for it in items),
        "days": [{"date": d, "count": n, "h": round(100 * n / peak)} for d, n in per_day.items()],
        "days_total": sum(per_day.values()),
    })


@require_POST
def project_stop(request, project_id):
    """«Тўхтатиш»: фондаги иш кейинги файлдан олдин тўхтайди (очилганлари сақланади)."""
    project = _get_work(request, Project, id=project_id)
    stop_project(project)
    audit(request, Event.Kind.STOP, work=project, work_kind="folder", detail=project.name)
    if request.headers.get("x-requested-with") == "fetch":
        return JsonResponse({"ok": True})
    return redirect("project", project_id=project.id)


@require_POST
def project_resume(request, project_id):
    """«Давом эттириш»: тўхтатилган қисмлар қайта очилади."""
    project = _get_work(request, Project, id=project_id)
    resume_project(project)
    return redirect("project", project_id=project.id)


@require_POST
def work_delete(request, kind, work_id):
    """Тарихдан ўчириш: папка, объект ёки алоҳида смета — базадан ҳам, дискдан ҳам (қайтариб бўлмайди)."""
    try:
        if kind == "folder":
            work = _get_work(request, Project, id=work_id)
            name = work.name
            delete_project(work, force=True)  # ишлаётган бўлса — аввал тўхтатилади
        elif kind == "object":
            work = _get_work(request, Batch, id=work_id, project__isnull=True)
            name = work.name
            delete_batch(work)
        elif kind == "smeta":
            work = _get_work(request, Job, id=work_id, batch__isnull=True)
            name = work.local_name
            delete_job(work)
        else:
            raise Http404("Нотўғри тур")
    except UploadError as e:
        return JsonResponse({"ok": False, "error": str(e)}, status=409)
    audit(request, Event.Kind.DELETE, work_kind=kind, detail=f"{name} ({work_id})")
    if request.headers.get("x-requested-with") == "fetch":
        return JsonResponse({"ok": True})
    return redirect("index")


@require_POST
def work_rename(request, kind, work_id):
    """Папка ёки объектни қайта номлаш (фақат эгаси)."""
    model = {"folder": Project, "object": Batch}.get(kind)
    if model is None:
        raise Http404("Нотўғри тур")
    work = _get_work(request, model, id=work_id)
    name = " ".join(request.POST.get("name", "").split())[:255]
    if name and name != work.name:
        old, work.name = work.name, name
        work.save(update_fields=["name"])
        audit(request, Event.Kind.RENAME, work=work, work_kind=kind, detail=f"{old} → {name}")
    return redirect("project" if kind == "folder" else "batch", work.id)


# ===================== sahifa yo'li (breadcrumb) =====================
def crumbs_for(project=None, batch=None, job=None):
    """Папка › ички папка › смета: ҳар бир саҳифадан орқага қайтиш учун."""
    batch = batch or (job.batch if job and job.batch_id else None)
    project = project or (batch.project if batch and batch.project_id else None)
    out = []
    if project:
        out.append({"label": project.name, "url": reverse("project", args=[project.id]), "icon": "folder", "color": "text-amber-500"})
    if batch:
        label = batch.folder or batch.name if project else batch.name
        out.append({"label": label, "url": reverse("batch", args=[batch.id]),
                    "icon": "folder-open" if project else "layers", "color": "text-amber-500" if project else "text-slate-400"})
    if job:
        name = job.local_name
        if batch and job.batch_file is not None and job.batch_file < len(batch.files):
            name = batch.files[job.batch_file]["name"]  # локал файл номи (лист номи «-2026_ЛРВ» тушунарсиз)
        out.append({"label": name, "url": "", "icon": "sheet", "color": "text-emerald-600"})
    return out


# ===================== butun papka =====================
PART_STATES = {
    "queued": ("Навбатда", "slate"),
    "converting": ("Файллар очиляпти", "sky"),
    "ready": ("Ҳисоблашга тайёр", "violet"),
    "stopped": ("Тўхтатилди", "slate"),
    "pricing": ("Нархланяпти", "sky"),
    "done": ("Тайёр", "emerald"),
    "warn": ("Нархланмаган қатор бор", "rose"),
    "review": ("Текширинг", "amber"),
    "empty": ("Локал топилмади", "slate"),
    "failed": ("Хато", "rose"),
}


@require_POST
def project_create(request):
    """Папка юклаш (JS): ``files`` + ``paths`` (ҳар бирининг папка ичидаги йўли) + ``skipped``."""
    try:
        check_quota(request.user)  # файллар ўқилишидан олдин
        project = create_project(request.user, request.FILES.getlist("files"), request.POST.getlist("paths"),
                                 request.POST.getlist("skipped"))
    except UploadError as e:
        return JsonResponse({"ok": False, "error": str(e)}, status=400)
    audit(request, Event.Kind.UPLOAD, work=project, work_kind="folder", detail=project.name)
    return JsonResponse({"ok": True, "url": reverse("project", args=[project.id])})


def _part_nodes(node, out):
    if node["part"]:
        out.append(node)
    for c in node["folders"]:
        _part_nodes(c, out)
    return out


def project_detail(request, project_id):
    project = _get_work(request, Project, id=project_id)
    ensure_running(project)
    parts = list(project.parts.prefetch_related("jobs").order_by("folder"))
    tree = build_tree(project, parts)
    # нархланадиган лист йўқ қисмлар (масалан «Умумий қиймати») — жадвал охирида
    nodes = sorted(_part_nodes(tree, []), key=lambda n: n["part"]["state"] == "empty")
    states = [n["part"]["state"] for n in nodes]
    finished = sum(1 for s in states if s not in ("queued", "converting", "pricing"))
    current = next((n for n in nodes if n["part"]["state"] in ("converting", "pricing")), None)
    ctx = {
        "crumbs": crumbs_for(project=project),
        "project": project,
        "tree": tree,
        "nodes": nodes,
        "states": PART_STATES,
        "working": project.status == Project.Status.PROCESSING,
        "awaiting": project.status == Project.Status.READY,
        "stopping": project.stop_requested,
        "count_stopped": states.count("stopped"),
        "count_ready": states.count("ready"),
        "count_opening": states.count("queued") + states.count("converting"),
        "count_pricing": states.count("pricing"),
        "count_pairs_ready": sum(len(n["part"]["planned"]) for n in nodes if n["part"]["state"] == "ready"),
        "current": current,
        "finished": finished,
        "progress": round(100 * finished / len(nodes)) if nodes else 100,
        "count_done": states.count("done"),
        "count_warn": states.count("warn"),
        "count_review": states.count("review") + states.count("failed"),
        "count_skipped": sum(1 for f in project.files if f["skip"] and is_shown(f["path"])),
        "grand_total": sum(n["part"]["total"] for n in nodes),
        "not_priced_total": sum(n["part"]["not_priced"] for n in nodes),
        "can_download": any(s in ("done", "warn", "review") and n["part"]["jobs"] for s, n in zip(states, nodes)),
        "pct_form": _pct_form(project.owner_id),
    }
    if request.GET.get("partial"):
        return render(request, "jobs/_project_body.html", ctx)
    return render(request, "jobs/project.html", ctx)


def project_viewer(request, project_id, batch_id):
    """Ko'rish oynasi (Excel kabi) uchun bitta qism: listlari va juftliklardagi nomzod resurslar.
    Sahifaga oldindan joylanmaydi — papkada 60 qism bo'lsa, sahifa har yangilanishda yuzlab KB bo'lib ketardi."""
    project = _get_work(request, Project, id=project_id)
    batch = get_object_or_404(Batch, id=batch_id, project=project)
    return JsonResponse({"url": reverse("batch_sheet", args=[batch.id]),
                         "sheets": _sheets_js(batch), "cands": _cands_js(batch)})


def _pct_form(owner_id) -> list:
    """Папкани ҳисоблашдан олдинги сводка фоизлари: охирги танлов (бўлмаса — стандарт)."""
    saved = {**DEFAULT_PERCENTS, **remembered_percents(owner_id)}
    return [{"key": k, "text": PERCENT_LABELS.get(k, k), "pct": saved[k], "default": d} for k, d in DEFAULT_PERCENTS.items()]


@require_POST
def project_price(request, project_id):
    """«Ҳисоблашни бошлаш»: смета­чи файлларни текширди — тайёр қисмлар фонда нархланади.
    Формадаги сводка фоизлари эслаб қолинади: папкадаги ҳамма смета шу фоизлар билан ҳисобланади."""
    project = _get_work(request, Project, id=project_id)
    if any(k.startswith("pct_") for k in request.POST):
        saved = remembered_percents(project.owner_id)
        # бўш катак — 0 (смета саҳифасидагидек)
        data = {f"pct_{k}": (request.POST.get(f"pct_{k}", "").strip() or "0") for k in DEFAULT_PERCENTS
                if f"pct_{k}" in request.POST}
        remember_percents(project.owner_id, {**saved, **parse_percents(data, saved)})
    start_pricing(project)
    return redirect("project", project_id=project.id)


def project_download(request, project_id):
    project = _get_work(request, Project, id=project_id)
    try:
        path = build_project_zip(project)
    except UploadError as e:
        raise Http404(str(e))
    _audit_download(request, project, "folder", path.name)
    return FileResponse(open(path, "rb"), as_attachment=True, filename=path.name)


# ===================== tizim haqida (landing) =====================
LANDING_FEATURES = [
    ("layers", "Локал ва RES'ни ўзи жуфтлайди",
     "Файлдаги ҳар бир листни мазмунидан танийди: локал смета, нарх рўйхати ёки свод. Ҳар бир локалга энг мос RES'ни мослик фоизи билан таклиф қилади."),
    ("search", "Ном ва бирлик бўйича мослаштириш",
     "Ресурс номи ва ўлчов бирлиги RES билан солиштирилади: аниқ мос бўлса — нарх ёзилади, ўхшаш ёки шубҳали бўлса — таклиф берилади, смета­чи тасдиқлайди."),
    ("calculator", "Сумма ва свод ҳисобланади",
     "Жами миқдор × нарх, иш ҳақи, машина, материаллар, транспорт, прочие, суғурта ва НДС — фоизларни ўзгартириш мумкин."),
    ("alert-circle", "Топилмаганлар қизил",
     "RES'да йўқ ресурс «ТОПИЛМАДИ» деб қизил белгиланади ва энг яқин ўхшаши таклиф қилинади — ҳеч нарса яширинмайди."),
    ("folder-open", "Бутун папкани юклаш",
     "ZIP, RAR ёки папка — ичидаги ҳамма локал ва RES топилиб, қисмлар бўйича нархланади. Натижа битта ZIP ёки битта Excel файлда."),
    ("building", "Техназор: ойма-ой бажарилган иш",
     "Офёртани юкланг — блок ва бўлимлар бўйича ҳар ойда бажарилган ҳажмни ёзасиз. Ресурслар меъёр бўйича ўзи ҳисобланади, қолдиқ ва Ф-3 фарқи кўринади."),
]
LANDING_STEPS = [
    ("Юклаш", "Локал смета ва RES файлини (.xls, .xlsx) ёки бутун объект папкасини юкланг."),
    ("Текшириш", "Тизим листларни танийди ва жуфтлайди. Нотўғри бўлса — тури ёки устунини бир босишда тузатинг."),
    ("Нархлаш", "Ҳар бир ресурс RES билан солиштирилади, нарх ва сумма ёзилади, свод ҳисобланади."),
    ("Юклаб олиш", "Тайёр Excel: нархлар, ранг билан белгиланган ҳолатлар, свод ва LOG листи."),
]
LANDING_SECURITY = [
    ("Интернетга юборилмайди", "Файллар фақат тизим ўрнатилган серверда ишланади ва сақланади. Ташқи хизматларга, сунъий интеллектга ёки бошқа серверга юборилмайди."),
    ("Асл файл ўзгармайди", "Юкланган смета нусхаси устида ишланади. Натижа алоҳида «…_narxlandi.xlsx» файлда — асл ҳужжат ҳар доим сақланиб қолади."),
    ("Маълумот бузилмайди", "Тизим устунларни ишонч билан жойлай олмаса — нархламайди ва сабабини айтади. Номаълум маълумот устидан ҳеч қачон ёзилмайди."),
    ("Ҳар бир қадам кўринади", "Қайси ресурсга қайси нарх олингани, ўхшашлик фоизи ва смета­чининг қўлда ўзгаришлари LOG листида ёзилади."),
    ("Тасодифий ҳавола", "Ҳар бир иш ва объект тахмин қилиб бўлмайдиган узун ҳавола (UUID) орқали очилади."),
    ("Файл чекланган", "Фақат Excel файллари (.xls, .xlsx, .xlsm) қабул қилинади, битта файл 30 МБ гача. Бошқа турдаги файллар рад этилади."),
]


@login_not_required
def landing(request):
    """Тизим ҳақида: нима қилади, қандай ишлайди, тезлик, хавфсизлик."""
    return render(request, "jobs/landing.html",
                  {"features": LANDING_FEATURES, "steps": LANDING_STEPS, "security": LANDING_SECURITY})


# ===================== obyekt (bir/ko'p fayl) =====================
ROLE_LABELS = {"local": "Локал", "resource": "Ресурс", "unknown": "Ўтказиш"}
# Устунлар вазифаси (smetachi qo'lda ko'rsatishi mumkin): калит, номи, қайси лист турига тегишли, нима эканлиги
COLUMN_ROLES = [
    ("name", "Ном", ("local", "resource"),
     "Иш ёки ресурс номи (матн). Масалан: «Щебень фракции 20-40 мм». Шу ном бўйича RES'дан нарх қидирилади."),
    ("unit", "Бирлик", ("local", "resource"),
     "Ўлчов бирлиги: м3, т, км, чел.-ч, маш.-ч. Ном билан бирга солиштирилади — бирлиги бошқа бўлса, нарх олинмайди."),
    ("qty_unit", "Миқдор (бирликка)", ("local",),
     "1 бирлик иш учун сарф — норма. Масалан: 1000 м3 грунт қазишга 21,4 маш.-ч экскаватор. Бу устун бўлмаса — «йўқ»."),
    ("qty_total", "Миқдор (жами)", ("local",),
     "Бутун иш ҳажмига сарф = норма × иш ҳажми (21,4 × 0,84 = 17,98 маш.-ч). Сумма шундан ҳисобланади: жами × нарх."),
    ("price", "Нарх", ("resource",),
     "RES'даги 1 бирлик ресурснинг нархи (сўм). Мос келган ресурснинг шу нархи локалга ёзилади."),
]
ROW_HINTS = {
    "start": "Биринчи иш ёки ресурс ёзилган қатор — сарлавҳа ва «1 2 3 4…» рақамлаш қаторидан кейингиси.",
    "end": "Охирги ресурс қатори. Ундан пасти («Итого», имзолар) нархланмайди.",
}


def _column_view(s, manual=None):
    """«Листлар тури» жадвали учун: тизим (ёки смета­чи) қайси устунни нима деб олди ва таҳрир майдонлари."""
    found = s.get("columns") or {}
    letter = lambda c: get_column_letter(int(c)) if c else ""
    parts = []
    for key, label, roles, hint in COLUMN_ROLES:
        if s["role"] not in roles or (key == "qty_unit" and not found.get(key)):
            continue
        parts.append({"key": key, "label": label, "hint": hint, "value": letter(found.get(key))})
    options = [(letter(c), t) for c, t in sorted((s.get("headers") or {}).items(), key=lambda x: int(x[0]))]
    fields = [{"key": key, "label": label, "roles": " ".join(roles), "hint": hint, "value": letter(found.get(key))}
              for key, label, roles, hint in COLUMN_ROLES]
    # қаторлар катагида — фақат смета­чи ёзгани (тизим топгани placeholder: акс ҳолда юборилганда «қўлда» бўлиб қотиб қолади)
    manual = manual or {}
    return {"parts": parts, "start": found.get("start"), "end": found.get("end"), "options": options, "fields": fields,
            "manual_start": manual.get("start"), "manual_end": manual.get("end"), "row_hints": ROW_HINTS,
            "missing": any(not p["value"] for p in parts)}


def _sheet_verdict(s, cv, by_key) -> dict:
    """Лист бир қарашда: ранг (tone), қисқа хулоса ва — керак бўлса — нима қилиш кераклиги."""
    if s.get("duplicate_of"):
        twin = by_key.get(s["duplicate_of"], {})
        return {"tone": "off", "text": "Нусха", "hint": f'«{twin.get("sheet", "")}» билан бир хил — ўтказилди'}
    if s["role"] == "unknown":
        kind = KIND_LABELS.get(s.get("kind", ""), "")
        hint = "Қайта нархлаш керак бўлса — турини «Локал» қилинг" if s.get("kind") == "priced" else (
            "" if kind else "Локал ёки ресурс бўлса — турини ўзгартиринг")
        return {"tone": "off", "text": f"Ўтказилди — {kind}" if kind else "Ўтказилди", "hint": hint}
    if cv["missing"]:
        return {"tone": "warn", "text": "Устун топилмади", "hint": "«тузатиш»ни босиб, қизил устунни кўрсатинг"}
    if not s.get("rows"):
        return {"tone": "warn", "text": "Маълумот топилмади" if s["role"] == "local" else "Нарх топилмади",
                "hint": "«тузатиш»да қаторлар оралиғи ва устунларни текширинг"}
    if s["role"] == "local":
        return {"tone": "ok", "text": "Нархланади", "hint": f'{s.get("rows") or 0} та қатор'}
    return {"tone": "res", "text": "Нарх манбаи", "hint": f'{s.get("rows") or 0} та нарх'}


def _sheet_groups(batch) -> list:
    """«Листлар тури»: файл бўйича гуруҳ, ҳар бир лист — (форма рақами, лист, устунлар, хулоса)."""
    by_key = {s["key"]: s for s in batch.sheets}
    groups = []
    for fi, f in enumerate(batch.files):
        rows = []
        for i, s in enumerate(batch.sheets):
            if s["file"] == fi:
                cv = _column_view(s, (batch.columns or {}).get(s["key"]))
                rows.append((i, s, cv, _sheet_verdict(s, cv, by_key)))
        if rows:
            groups.append({"name": f["name"], "rows": rows})
    return groups


def _parse_columns(post, i) -> dict:
    """POST'дан битта листнинг қўлда устунлари: {"name": 3, …, "start": 22, "end": 850}; бўш — автомат."""
    if post.get(f"cols_{i}_on") != "1":
        return {}
    out = {}
    for key, *_rest in COLUMN_ROLES:
        v = (post.get(f"cols_{i}_{key}") or "").strip().upper()
        if v:
            try:
                out[key] = column_index_from_string(v)
            except ValueError:
                pass
    for key in ("start", "end"):
        v = (post.get(f"cols_{i}_{key}") or "").strip()
        if v.isdigit() and int(v) > 0:
            out[key] = int(v)
    if out.get("start") and out.get("end") and out["end"] < out["start"]:
        out.pop("end")
    return out


def _batch_tags(batch):
    """Боғланишни кўрсатиш учун белгилар: ҳар бир локал «Л1, Л2…» (жуфтлик тартибида), ресурс «Р1, Р2…»."""
    loc_tag = {p["local"]: f"Л{n + 1}" for n, p in enumerate(batch.pairs)}
    resources = [s for s in batch.sheets if s["role"] == "resource" and not s.get("duplicate_of")]
    res_tag = {s["key"]: f"Р{i + 1}" for i, s in enumerate(resources)}
    return loc_tag, res_tag, resources


def _pair_options(p, resources, res_tag, sheets):
    """Танлаш рўйхати: ҳар бир ресурс листи мослиги билан, энг мос биринчи (эски объектларда — рўйхат тартибида)."""
    loc = sheets[p["local"]]
    by_key = {o["res"]: o for o in p.get("options") or []}
    out = []
    for s in resources:
        o = by_key.get(s["key"], {})
        out.append({"key": s["key"], "tag": res_tag[s["key"]], "sheet": s["sheet"], "file_name": s["file_name"],
                    "title": sheet_title(s), "rows": s.get("rows", 0), "rows": s.get("rows", 0), "number": s.get("number", ""), "hidden": s.get("hidden"),
                    "overlap": o.get("overlap"), "score": o.get("score", 0), "same_file": s["file"] == loc["file"]})
    # локалнинг ўз файлидаги ресурс листлари — мослиги паст бўлса ҳам биринчи (одатда тўғриси шу), кейин мослик бўйича
    out.sort(key=lambda o: (not o["same_file"], -(o["score"] or 0)))
    best = max(out, key=lambda o: o["overlap"] or 0, default=None)
    if best and best["overlap"]:
        best["best"] = True
    return out


def sheet_title(s) -> str:
    """Листнинг ўқиладиган номи: «LRV», «Лист5» ўрнига қисм номи ёки файл номи."""
    return readable(s.get("part")) or file_title(s["file_name"])


def _cands_js(batch):
    """Модал учун: ҳар бир жуфтликнинг номзод ресурслари (ўз файлидагилари биринчи, кейин мослик бўйича)."""
    sheets = {s["key"]: s for s in batch.sheets}
    _loc_tag, res_tag, resources = _batch_tags(batch)
    return [[{"k": o["key"], "ov": o["overlap"], "same": o["same_file"]} for o in _pair_options(p, resources, res_tag, sheets)]
            for p in batch.pairs]


def _batch_tree(batch):
    """Стройка → Объект → жуфтликлар (кўрсатиш учун)."""
    sheets = {s["key"]: s for s in batch.sheets}
    loc_tag, res_tag, resources = _batch_tags(batch)
    tree = {}
    for n, p in enumerate(batch.pairs):
        loc = sheets[p["local"]]
        res = sheets.get(p["res"])
        st = tree.setdefault(loc.get("stroyka") or "—", {})
        st.setdefault(loc.get("obj") or "—", []).append(dict(
            p, n=n, loc=loc, res_sheet=res, tag=loc_tag[p["local"]], res_tag=res_tag.get(p["res"], ""),
            title=readable(p.get("label")) or sheet_title(loc),
            same_file=bool(res) and res["file"] == loc["file"],
            choices=_pair_options(p, resources, res_tag, sheets)))
    return [{"stroyka": s, "objects": [{"obj": o, "pairs": ps} for o, ps in objs.items()]} for s, objs in tree.items()]


PREVIEW_ROWS = 80


def batch_sheet(request, batch_id):
    """Объект листи Excel кўринишида (бошидан ``n`` қатор): ``k`` — лист калити «файл:лист»."""
    batch = _get_work(request, Batch, id=batch_id)
    key = request.GET.get("k", "")
    info = batch.sheet(key)
    if info is None:
        raise Http404("Лист топилмади")
    f = batch.files[info["file"]]
    path = batch.abs(f["path"])
    if not f.get("path") or not path.is_file():
        raise Http404("Файл ҳали очилмаган")
    try:
        n = max(20, min(int(request.GET.get("n") or PREVIEW_ROWS), xlview.MAX_ROWS))
    except ValueError:
        n = PREVIEW_ROWS
    try:
        data = xlview.cached(batch.dir / "preview", f"{key}|{n}", path,
                             lambda: xlview.render(path, info["sheet"], max_rows=n))
    except Exception as e:
        return JsonResponse({"error": f"{type(e).__name__}: {e}"}, status=500)
    return JsonResponse({**data, "file": f["name"], "role": info["role"], "limit": n})


def _sheets_js(batch):
    """Модал учун: ҳар бир лист (файл тартибида) — тури ва белгиси (Л1/Р2), файл ичида листдан листга ўтиш учун."""
    loc_tag, res_tag, _resources = _batch_tags(batch)
    return [{"k": s["key"], "file": s["file"], "file_name": s["file_name"], "sheet": s["sheet"], "role": s["role"],
             "title": sheet_title(s),
             "tag": loc_tag.get(s["key"]) or res_tag.get(s["key"], ""), "dup": bool(s.get("duplicate_of")),
             "kind": KIND_LABELS.get(s.get("kind", ""), ""), "hidden": bool(s.get("hidden"))} for s in batch.sheets]


def batch_detail(request, batch_id):
    batch = _get_work(request, Batch, id=batch_id)
    sheets = batch.sheets
    used_res = {p["res"] for p in batch.pairs if p.get("res")}
    jobs = list(batch.jobs.order_by("batch_order"))
    by_key = {s["key"]: s for s in sheets}
    # ишлатилмаган ресурсни бириктириш учун: ҳар бир локал (белгиси, номи, файли)
    loc_titles = [{"n": n, "tag": f"Л{n + 1}", "file": by_key[p["local"]]["file"], "has_res": bool(p.get("res")),
                   "title": readable(p.get("label")) or sheet_title(by_key[p["local"]])} for n, p in enumerate(batch.pairs)]
    ctx = {
        "batch": batch,
        "crumbs": crumbs_for(batch=batch),
        "tree": _batch_tree(batch),
        "resources": [s for s in sheets if s["role"] == "resource" and not s.get("duplicate_of")],
        "unused": [dict(s, title=sheet_title(s), same_files=[t for t in loc_titles if t["file"] == s["file"]])
                   for s in sheets if s["role"] == "resource" and not s.get("duplicate_of") and s["key"] not in used_res],
        "locals": loc_titles,
        "duplicates": [s for s in sheets if s.get("duplicate_of")],
        "unknown": [s for s in sheets if s["role"] == "unknown"],
        "sheets_by_file": [(i, f, [s for s in sheets if s["file"] == i]) for i, f in enumerate(batch.files)],
        "sheet_groups": _sheet_groups(batch),
        "row_hints": ROW_HINTS,
        "role_counts": {r: sum(1 for s in sheets if s["role"] == r and not s.get("duplicate_of")) for r in ROLE_LABELS},
        "kind_labels": KIND_LABELS,
        "sheets_js": _sheets_js(batch),
        "role_labels": ROLE_LABELS.items(),
        "ready": sum(1 for p in batch.pairs if pair_ready(p)),
        **_base_ctx(base_period_of(batch.owner_id)), "base_res": BASE_RES,
        "jobs": jobs,
        "grand_total": sum(j.result.get("totals", {}).get("s_nds") or 0 for j in jobs
                           if isinstance(j.result.get("totals", {}).get("s_nds"), (int, float))),
        "not_priced_total": sum(j.result.get("stats", {}).get("NOT_PRICED", 0) for j in jobs),
        # танланган RES'да нарх йўқ — нархлашдан олдин огоҳлантириш
        "unpriced": [readable(p.get("label")) or p["local"] for p in batch.pairs
                     if p.get("include") and p.get("res") in unpriced_res(batch)],
    }
    return render(request, "jobs/batch.html", ctx)


@require_POST
def batch_replan(request, batch_id):
    batch = _get_work(request, Batch, id=batch_id)
    roles, columns = {}, {}
    for i, s in enumerate(batch.sheets):
        role = request.POST.get(f"role_{i}")
        if role in ROLE_LABELS and role != s["role"]:
            roles[s["key"]] = role
        elif s["key"] in (batch.roles or {}) and role == s["role"]:
            roles[s["key"]] = role  # аввал қўлда белгиланган — сақланади
        cols = _parse_columns(request.POST, i)
        if cols:
            columns[s["key"]] = cols
    analyze_batch(batch, roles, columns)
    batch.status = Batch.Status.REVIEW
    batch.save()
    if batch.project_id and not batch.jobs.exists():
        batch.stage, batch.message = ready_stage(batch)
        batch.save(update_fields=["stage", "message"])
        settle(batch.project)
    return redirect("batch", batch_id=batch.id)


@require_POST
def batch_confirm(request, batch_id):
    batch = _get_work(request, Batch, id=batch_id)
    decisions = {}
    for n in range(len(batch.pairs)):
        decisions[n] = {
            "include": request.POST.get(f"include_{n}") == "1",
            "res": request.POST.get(f"res_{n}", ""),
            "label": request.POST.get(f"label_{n}", ""),
        }
    if batch.project_id:
        queue_part(batch, decisions)  # папка қисми: фонда нархланади, жараён папка саҳифасида жонли
        return redirect("project", project_id=batch.project_id)
    else:
        confirm_batch(batch, decisions)
    return redirect("batch", batch_id=batch.id)


def batch_download(request, batch_id):
    batch = _get_work(request, Batch, id=batch_id)
    try:
        path = build_combined(batch)
    except UploadError as e:
        raise Http404(str(e))
    _audit_download(request, batch, "object", path.name)
    return FileResponse(open(path, "rb"), as_attachment=True, filename=path.name)


def job_detail(request, job_id):
    job = _get_work(request, Job, id=job_id)
    ctx = {"job": job, "modes": MODE_LABELS.items(), "crumbs": crumbs_for(job=job),
           "siblings": list(job.batch.jobs.order_by("batch_order")) if job.batch_id else [],
           **_base_ctx(job.base_period)}
    if job.status == Job.Status.DONE:
        res = job.result
        totals = res["totals"]
        mode = res.get("mode", "compat")
        labels = column_labels(job) if mode == "smart" else dict(CATEGORY_LABELS)
        res_groups = _res_groups(res, labels) if mode == "smart" else []
        view = request.GET.get("view") or ("res" if res_groups else "local")
        if view not in XL_FILES and not res_groups:
            view = "local"
        ctx.update(
            view=view,
            grid_view="local" if view == "local" or not res_groups else ("res" if view in XL_FILES else view),
            xl_file=XL_FILES.get(view),
            xl_files=[(v, label) for v, (_f, label) in XL_FILES.items() if v != "xl_res" or job.res_path],
            res_groups=res_groups,
            result_mode=mode,
            # «N та қатор нархланмади» — ён панелда ва сводкада бор, бу ерда такрорламаймиз
            warnings=[w for w in res.get("warnings", []) if "нархланмади" not in w and not w.startswith("Текширувдан")],
            checks=res.get("checks", []),
            check_groups=check_groups(res.get("checks", []), res["rows"]),
            columns=res.get("columns", {}),
            # smart: фақат бор нарсалар (0 бўлган карталар чалғитмасин); compat — эски дастур каби ҳаммаси
            stats=[(key, label, res["stats"].get(key, 0), color) for key, label, color in STAT_CARDS[mode]
                   if mode != "smart" or res["stats"].get(key, 0) or key in ("EXACT", "NOT_PRICED")],
            svodka=_svodka_rows(res, mode),
            grid_columns=list(labels.items()),
            grid_span=len(labels) + 6,
            extra_letters=[c for c in labels if c not in CATEGORY_LABELS],
            renamed=[c for c in CATEGORY_LABELS if (job.column_names or {}).get(c)],
            can_add_column=mode == "smart" and _next_column(labels) is not None,
            percents_changed=mode == "smart" and {**DEFAULT_PERCENTS, **(job.percents or {})} != DEFAULT_PERCENTS,
            pct_memory=_pct_memory(job) if mode == "smart" else None,
            s_nds=totals.get("s_nds"),
            has_errors=any(isinstance(v, str) for v in totals.values()),
            rows=res["rows"],
            problems=res["problems"],
            category_counts=[(c, label, sum(1 for r in res["rows"] if _is_resource(r) and r.get("category", "K") == c))
                             for c, label in labels.items()],
            category_choices=labels.items(),
            # L/M таклифлари 2-босқич карталарида; бу ерда қолганлари (масалан I/J/K)
            category_suggestions=[s for s in _category_suggestions(res["rows"]) if mode != "smart" or s["cat"] not in LM_COLUMNS],
            override_count=sum(1 for r in res["rows"] if r.get("override")),
            bez_nds=totals.get("bez_nds"),
            not_priced=res["stats"].get("NOT_PRICED", res["stats"].get("NON_NUMERIC", 0)),
            sugg_price_count=sum(1 for r in res["rows"] if r.get("suggested_price") is not None and not r.get("manual")),
            manual_count=sum(1 for r in res["rows"] if r.get("manual")),
            sugg_cat_count=sum(1 for r in res["rows"] if r.get("suggested_category")),
            nf=_not_found_stage(res["rows"]) if mode == "smart" else None,
            lm=_lm_stage(res["rows"], res_groups, labels) if mode == "smart" else None,
        )
        ctx["sugg_marks"] = _sugg_marks(res["rows"]) if mode == "smart" else []
        ctx["price_marks"] = _price_marks(res["rows"]) if mode == "smart" else []
        ctx["stages"] = _stages(job, ctx["nf"], ctx["lm"]) if mode == "smart" else []
        # саҳифа қайси босқичда очилади: биринчи тугамагани (compat — жадвал)
        ctx["default_tab"] = next((st["n"] for st in ctx["stages"] if st["state"] == "current"), 3 if ctx["stages"] else 2)
    return render(request, "jobs/job.html", ctx)


def _sugg_marks(rows):
    """Устун таклифлари — Excel кўринишида белгилаш ва бирма-бир ўтиш учун: қатор, қаерга (L, M, …),
    ҳозир қаерда, сабаби. ``_lm_stage`` дан кейин чақирилади (``lm_target`` шу ерда қўйилади)."""
    out = []
    for r in rows:
        to = r.get("lm_target") or r.get("suggested_category")
        if to and to != r.get("category") and _is_resource(r):
            out.append({"row": r["row"], "to": to, "from": r.get("category") or "", "name": r.get("name") or "",
                        "why": r.get("suggest_reason") or ("ресурс файлида шу бўлимда" if r.get("lm_target") else ""),
                        "h": r["h"] if isinstance(r.get("h"), (int, float)) and not isinstance(r["h"], bool) else None})
    return out


def _price_marks(rows):
    """Excel кўринишида G катагидан нарх ёзиш учун: нархланмаган (қизил) ва қўлда нархланган қаторлар —
    номи, бирлиги, миқдори, ресурсдаги энг яқин таклиф. ``_not_found_stage`` дан кейин (``near`` шу ерда тўлади)."""
    out = []
    for r in rows:
        unpriced = r.get("status") in UNPRICED_STATUSES and not r.get("cleared") and _has_unit(r)
        if not (unpriced or r.get("manual")):
            continue
        near = r.get("near") or {}
        out.append({"row": r["row"], "name": r.get("name") or "", "unit": r.get("unit") or "",
                    "qty": r.get("qty") if isinstance(r.get("qty"), (int, float, str)) else None,
                    "g": r.get("g") if isinstance(r.get("g"), (int, float, str)) else None,
                    "manual": bool(r.get("manual")), "sugg": r.get("suggested_price"),
                    "near": {"name": near.get("name"), "unit": near.get("unit"), "price": near.get("price"),
                             "score": near.get("score"), "same_unit": near.get("same_unit", True)} if near else None})
    return out


def _is_resource(row) -> bool:
    """Жадвалда танлаш/ўтказиш мумкин бўлган қатор: номи бор."""
    return bool(row.get("name"))


ZERO_STATUSES = ("WORK_HEADER", "MACHINIST")


def _next_column(labels):
    last = max(column_index_from_string(c) for c in labels)
    nxt = get_column_letter(last + 1)
    return nxt if column_index_from_string(nxt) <= column_index_from_string(LAST_EXTRA_COLUMN) else None


def _res_groups(res, labels=None):
    """Smart: қаторлар ресурс файлидаги бўлимлар тартибида («ТРУДОВЫЕ РЕСУРСЫ», «…МАШИНЫ…»,
    «МАТЕРИАЛЬНЫЕ…», «ОБОРУДОВАНИЕ»); бўлим ичида — ресурслар RES даги тартибда."""
    sections = res.get("res_sections") or []
    if not sections:
        return []
    labels = labels or CATEGORY_LABELS
    groups = [{"title": s["title"], "column": s["column"], "label": labels[s["column"]], "rows": []}
              for s in sections]
    missing = {"title": "Ресурс файлида топилмаганлар", "column": None, "label": "", "rows": [],
               "hint": "аниқ мос ресурс йўқ — нархни қўлда киритинг"}
    other = {"title": "RES бўлимидан ташқари", "column": None, "label": "", "rows": [], "hint": ""}
    zero = {"title": "Иш сарлавҳалари ва машинистлар", "column": None, "label": "", "rows": [],
            "hint": "суммага кирмайди (0)"}
    for r in res["rows"]:
        if not _is_resource(r):
            continue
        idx = r.get("res_section")
        if r.get("status") in ZERO_STATUSES:
            zero["rows"].append(r)
        elif idx is None:
            missing["rows"].append(r)
        elif idx < 0:
            other["rows"].append(r)
        else:
            groups[idx]["rows"].append(r)
    for g in groups:
        g["rows"].sort(key=lambda r: (r.get("res_row") or 0, r["row"]))
    out = []
    for i, g in enumerate(groups + [other, missing, zero]):
        if not g["rows"] and g["column"] is None:
            continue
        g["id"] = i
        g["count"] = len(g["rows"])
        g["sum"] = sum(r["h"] for r in g["rows"] if isinstance(r.get("h"), (int, float)) and not isinstance(r["h"], bool))
        # ресурс бўлими бошқа устунни кўрсатяпти: «M га ўтказамизми?»
        g["to_move"] = [r["row"] for r in g["rows"] if g["column"] and r.get("category") != g["column"]
                        and not r.get("override")]
        out.append(g)
    return out


def _category_suggestions(rows):
    groups = {}
    for r in rows:
        cat = r.get("suggested_category")
        if cat:
            g = groups.setdefault(cat, {"cat": cat, "label": CATEGORY_LABELS[cat], "count": 0, "reasons": []})
            g["count"] += 1
            if r["suggest_reason"] not in g["reasons"] and len(g["reasons"]) < 3:
                g["reasons"].append(r["suggest_reason"])
    return [groups[c] for c in CATEGORY_LABELS if c in groups]


@require_POST
def job_categories(request, job_id):
    """JSON: ``{"rows": [183, 281], "category": "M"}``; ``category: ""`` — asl holiga qaytarish."""
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or job.mode != "smart":
        return JsonResponse({"ok": False, "error": "Устунни фақат Smart режимдаги натижада ўзгартириш мумкин"}, status=400)
    try:
        data = json.loads(request.body or b"{}")
        rows = [int(r) for r in data["rows"]]
        category = data.get("category") or ""
    except (ValueError, TypeError, KeyError):
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    if category and category not in column_labels(job):
        return JsonResponse({"ok": False, "error": f"нотўғри устун: {category}"}, status=400)
    known = {r["row"]: r for r in job.result.get("rows", []) if _is_resource(r)}
    unknown = [r for r in rows if r not in known]
    if not rows or unknown:
        return JsonResponse({"ok": False, "error": f"бундай қатор йўқ: {unknown[:5]}"}, status=400)

    categories = dict(job.categories or {})
    for r in rows:
        if not category or category == known[r]["default_category"]:
            categories.pop(str(r), None)  # дастур аниқлагани билан бир хил — қўлда белги шарт эмас
        else:
            categories[str(r)] = category
    recalc.request(job, categories=categories)
    return JsonResponse({"ok": True, "queued": True, "changed": len(rows), "overrides": len(job.categories)})


def _as_price(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return parse_text_number(str(v)) if v is not None else None


@require_POST
def job_prices(request, job_id):
    """JSON: ``{"rows": [...], "action": "accept" | "set" | "clear", "price": "4 600 000"}``.
    accept — taklif narxini qabul qilish; set — bitta narxni hammasiga; clear — qo'lda narxni bekor qilish."""
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or job.mode != "smart":
        return JsonResponse({"ok": False, "error": "Нархни фақат Smart режимдаги натижада ўзгартириш мумкин"}, status=400)
    try:
        data = json.loads(request.body or b"{}")
        rows = [int(r) for r in data["rows"]]
        action = data["action"]
    except (ValueError, TypeError, KeyError):
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    known = {r["row"]: r for r in job.result.get("rows", []) if _is_resource(r)}
    unknown = [r for r in rows if r not in known]
    if not rows or unknown or action not in ("accept", "set", "clear"):
        return JsonResponse({"ok": False, "error": f"нотўғри сўров: {unknown[:5] or action}"}, status=400)

    prices = dict(job.prices or {})
    skipped = []
    if action == "set":
        price = _as_price(data.get("price"))
        if price is None or price < 0:
            return JsonResponse({"ok": False, "error": "нарх рақам бўлиши керак (масалан 4 600 000)"}, status=400)
        for r in rows:
            prices[str(r)] = price
    elif action == "accept":
        for r in rows:
            price = _as_price(known[r].get("suggested_price"))
            if price is None:
                skipped.append(r)
            else:
                prices[str(r)] = price
    else:
        for r in rows:
            prices.pop(str(r), None)
    changed = len(rows) - len(skipped)
    if changed:
        recalc.request(job, prices=prices)
    return JsonResponse({"ok": True, "queued": bool(changed), "changed": changed, "skipped": len(skipped)})


def _has_unit(row) -> bool:
    """Бирлик бор. Номи билан бир хил «бирлик» — C:D бирлаштирилган катак (бирлик эмас)."""
    unit = str(row.get("unit") or "").strip()
    return unit not in ("", "None") and unit != str(row.get("name") or "").strip()


def _not_found_stage(rows):
    """1-босқич: «ТОПИЛМАДИ» ёзилган қаторлар — бирлиги бор (ресурс, нархи йўқ) ва бирлиги йўқ
    (бўлим номи ва ш.к.) га ажратилган; аввал тозаланганлари алоҳида (қайтариш учун).
    Ҳар бир қаторга смета ичидаги ўрни қўшилади: қайси бўлим остида (``ctx_section``) ва қайси иш
    таркибида (``ctx_work`` — «12 Е0601-26-4 УСТРОЙСТВО…»), смета­чи Excel билан осон солиштирсин."""
    section, work = "", ""
    pending = []
    for r in rows:  # смета тартибида
        if not r.get("name"):
            continue
        r["is_heading"] = not _has_unit(r) and not str(r.get("num") or "").strip()
        if r["is_heading"]:
            r["ctx_section"], r["ctx_work"] = section, ""
            section, work = str(r["name"]).strip(), ""
        else:
            if r.get("status") == "WORK_HEADER":
                work = " ".join(str(x).strip() for x in (r.get("num"), r.get("code"), r.get("name")) if x)
            elif "." not in str(r.get("num") or "") and r.get("num"):
                work = ""  # мустақил ресурс («70 СЦЕНА …») — аввалги ишга тегишли эмас
            r["ctx_section"], r["ctx_work"] = section, work
        r["nf_kind"] = "unit" if _has_unit(r) else "nounit"
        if r.get("status") in UNPRICED_STATUSES and not r.get("cleared"):
            if not r.get("near") and r.get("suggested_price") is not None and r.get("picked_name"):
                # эски натижа (яқин ресурс сақланмаган): LOG'даги таклиф
                score = r.get("score")
                r["near"] = {"name": r["picked_name"], "unit": "", "price": r["suggested_price"],
                             "score": score if isinstance(score, (int, float)) else 0, "same_unit": True}
            if r.get("near"):
                # H формуласи каби: миқдор ЛОКАЛДАН (F, бўш бўлса E), нарх RES дан
                r["qty_col"] = "E" if r.get("qty_total") in (None, "") else "F"
                q, p = _as_price(r.get("qty")), _as_price(r["near"].get("price"))
                r["near_sum"] = q * p if q is not None and p is not None else None
            pending.append(r)
    # бўлим номи учун — остида нима бор: кейинги бўлимгача қаторлар сони ва биринчи иш
    named = [r for r in rows if r.get("name")]
    for i, r in enumerate(named):
        if not r.get("is_heading"):
            continue
        below = []
        for x in named[i + 1:]:
            if x.get("is_heading"):
                break
            below.append(x)
        first = next((x for x in below if x.get("status") == "WORK_HEADER"), below[0] if below else None)
        r["ctx_count"] = len(below)
        r["ctx_first"] = " ".join(str(x).strip() for x in (first.get("num"), first.get("name")) if x) if first else ""
    return {
        "pending": len(pending),
        "rows": pending,
        "with_unit": [r for r in pending if _has_unit(r)],
        "no_unit": [r for r in pending if not _has_unit(r)],
        # гуруҳлар (смета­чи тилида): A — бирлиги ҳам, таклифи ҳам йўқ; B — бирлиги бор, нарх йўқ (таклифли/таклифсиз)
        "group_a": [r for r in pending if not _has_unit(r) and not r.get("near")],
        "group_b_near": [r for r in pending if _has_unit(r) and r.get("near")],
        "group_b_none": [r for r in pending if _has_unit(r) and not r.get("near")],
        "cleared": [r for r in rows if r.get("cleared")],
    }


@require_POST
def job_not_found(request, job_id):
    """JSON: ``{"rows": [...], "action": "clear" | "mark" | "restore"}``. clear — «ТОПИЛМАДИ» сўзини олиб ташлаш:
    бирлиги бор қаторда G бўш, қизил қолади; бирлиги йўқда G ва H бўш. mark — сўз олинади, G сариқ
    («кўриб чиқилди», нарх кейин). restore — қайтариш."""
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or job.mode != "smart":
        return JsonResponse({"ok": False, "error": "Фақат Smart режимдаги натижада"}, status=400)
    try:
        data = json.loads(request.body or b"{}")
        rows = {int(r) for r in data["rows"]}
        action = data["action"]
    except (ValueError, TypeError, KeyError):
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    by_row = {r["row"]: r for r in job.result.get("rows", [])}
    allowed = {r for r, x in by_row.items() if x.get("status") in UNPRICED_STATUSES or x.get("cleared")}
    if not rows or action not in ("clear", "mark", "restore") or not rows <= allowed:
        return JsonResponse({"ok": False, "error": f"нотўғри сўров: {sorted(rows - allowed)[:5] or action}"}, status=400)
    cleared, marked = set(job.cleared or []), set(job.marked or [])
    if action == "clear":  # сўз олинади: бирлиги борида G қизил, йўғида G ва H бўш
        cleared, marked = cleared | rows, marked - rows
    elif action == "mark":  # сўз олинади, G сариқ — «кўриб чиқилди, нарх кейин»
        cleared, marked = cleared - rows, marked | rows
    else:
        cleared, marked = cleared - rows, marked - rows
    recalc.request(job, cleared=sorted(cleared), marked=sorted(marked))
    return JsonResponse({"ok": True, "queued": True, "changed": len(rows)})


LM_COLUMNS = ("L", "M")


def _lm_stage(rows, res_groups, labels):
    """2-босқич: L (кабель) ва M (ускуна) га ўтказиш керак бўлиши мумкин қаторлар — таклиф қилинганлар
    ва RES'да шу бўлимда турганлари. Ҳар бир қаторга ``lm_target`` белгиси қўйилади (жадвалда танлаш учун)."""
    by_row = {r["row"]: r for r in rows}
    targets = {}
    for r in rows:
        if r.get("suggested_category") in LM_COLUMNS:
            targets[r["row"]] = r["suggested_category"]
    for g in res_groups:
        if g["column"] in LM_COLUMNS:
            for n in g["to_move"]:
                targets.setdefault(n, g["column"])
    cards = []
    for c in LM_COLUMNS:
        # бўлим номлари (бирлиги йўқ), иш сарлавҳалари ва машинистлар ўтказилмайди
        pending = [by_row[n] for n in sorted(targets) if targets[n] == c and by_row[n].get("category") != c
                   and _has_unit(by_row[n]) and by_row[n].get("status") not in ZERO_STATUSES + ("CLEARED",)]
        for r in pending:
            r["lm_target"] = c
        cards.append({"col": c, "label": labels.get(c, c), "rows": pending,
                      "done": sum(1 for r in rows if _is_resource(r) and r.get("category") == c),
                      "sum": sum(r["h"] for r in pending if isinstance(r.get("h"), (int, float)) and not isinstance(r["h"], bool))})
    return {"cards": cards, "pending": sum(len(c["rows"]) for c in cards)}


@require_POST
def job_table_checked(request, job_id):
    """2-босқични якунлаш (``done=1``) ёки қайта очиш (``done=0``)."""
    job = _get_work(request, Job, id=job_id)
    job.table_checked = request.POST.get("done") == "1"
    job.save(update_fields=["table_checked"])
    return redirect(reverse("job", args=[job.id]) + ("#t-3" if job.table_checked else "#t-2"))


def _stages(job, nf, lm):
    """Смета­чи йўли: 1 — топилмаганлар, 2 — жадвал (кабель/ускунани ажратиш), 3 — сводка ва фоизлар (кейин юклаб олиш)."""
    s1 = not nf["pending"]
    s2 = s1 and job.table_checked
    return [
        {"n": 1, "anchor": "stage-nf", "title": "Топилмаганлар", "state": "done" if s1 else "current",
         "note": "тайёр" if s1 else f"{nf['pending']} та нархланмаган"},
        {"n": 2, "anchor": "stage-table", "title": "Жадвал",
         "state": "done" if s2 else ("current" if s1 else "locked"),
         "note": "текширилди" if s2 else (f"{lm['pending']} та кабель/ускуна таклифи" if lm["pending"]
                                          else "жадвални кўриш")},
        {"n": 3, "anchor": "svodka", "title": "Сводка ва фоизлар", "state": "current" if s2 else "locked",
         "note": "фоизлар, юклаб олиш"},
    ]


def _pct_memory(job):
    """Эслаб қолинган фоизлар: стандартдан фарқлилари ва шу смета улардан фарқ қиладими."""
    saved = remembered_percents(job.owner_id)
    labels = dict((k, t) for k, t, p in svodka_layout([]) if p is not None)
    diff = [{"text": labels.get(k, k), "pct": pct_text(v), "default": pct_text(DEFAULT_PERCENTS[k])}
            for k, v in saved.items() if k in DEFAULT_PERCENTS and v != DEFAULT_PERCENTS[k]]
    own = {**DEFAULT_PERCENTS, **(job.percents or {})}
    mine = {**DEFAULT_PERCENTS, **saved}
    return {"diff": diff, "differs": own != mine}


def _svodka_rows(res, mode):
    used = res.get("percents") or DEFAULT_PERCENTS
    rows = []
    extra = [tuple(c) for c in res.get("extra_columns", [])]
    for i, (key, text, pct) in enumerate(svodka_layout(extra)):
        p = used.get(key, pct) if pct is not None else None
        if key == "s_nds" and mode == "smart":
            text = f"Всего с НДС      {percent_text(p)} %"
        rows.append({"row": res["svodka_start"] + i, "key": key, "text": text, "pct": p,
                     "default": pct, "value": res["totals"].get(key)})
    return rows


@require_POST
def job_columns(request, job_id):
    """JSON: ``{"action": "add", "name": "Монтаж"}`` — кейинги бўш устун (N, O, …);
    ``{"action": "rename", "col": "K", "name": "…"}`` (бўш ном — асл номига қайтади);
    ``{"action": "remove", "col": "N"}`` — қўшилган устун, ундаги қаторлар асл устунига қайтади."""
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or job.mode != "smart":
        return JsonResponse({"ok": False, "error": "Устунлар фақат Smart режимда"}, status=400)
    try:
        data = json.loads(request.body or b"{}")
        action = data["action"]
        col = str(data.get("col") or "").strip().upper()
        name = " ".join(str(data.get("name") or "").split())[:60]
    except (ValueError, TypeError, KeyError):
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    labels = column_labels(job)
    names = dict(job.column_names or {})
    categories = None
    if action == "add":
        col = _next_column(labels)
        if not col:
            return JsonResponse({"ok": False, "error": f"{LAST_EXTRA_COLUMN} устунидан кейин қўшиб бўлмайди"}, status=400)
        names[col] = name or f"Устун {col}"
    elif action == "rename" and col in labels:
        if col in CATEGORY_LABELS and not name:
            names.pop(col, None)
        else:
            names[col] = name or f"Устун {col}"
    elif action == "remove" and col in labels and col not in CATEGORY_LABELS:
        names.pop(col)
        categories = {r: c for r, c in (job.categories or {}).items() if c != col}
    else:
        return JsonResponse({"ok": False, "error": f"нотўғри амал ёки устун: {action} {col}"}, status=400)
    job.column_names = names
    recalc.request(job, categories=categories)
    return JsonResponse({"ok": True, "queued": True, "col": col})


@require_POST
def job_base(request, job_id):
    """Сметанинг ресурс базаси даври ўзгарди — қайта нархланади (смета­чи танлови ҳам эслаб қолинади)."""
    from resbase.services import parse_period
    job = _get_work(request, Job, id=job_id)
    job.base_period = parse_period(request.POST)
    remember_base_period(job.owner_id, job.base_period)
    recalc.request(job)
    return redirect("job", job_id=job.id)


@require_POST
def job_run(request, job_id):
    job = _get_work(request, Job, id=job_id)
    if "reset_percents" in request.POST:
        percents = dict(DEFAULT_PERCENTS)  # аниқ «стандарт» (бўш = эслаб қолинган фоизлар)
    elif any(k.startswith("pct_") for k in request.POST):
        percents = parse_percents(request.POST, job.percents)
    else:
        percents = None  # yuqoridagi forma: foizlar o'zgarmaydi
    recalc.request(job, request.POST.get("local_sheet"), request.POST.get("res_sheet"),
                   request.POST.get("fuzzy") or job.fuzzy, request.POST.get("mode"), percents)
    if percents is not None:
        remember_percents(job.owner_id, job.percents)
    return redirect("job", job_id=job.id)


@require_POST
def job_percents(request, job_id):
    """JSON: ``{"pct": {"transport_mat": "5", ...}}`` | ``{"reset": true}`` | ``{"remembered": true}``.
    Фоиз ўзгариши билан қайта ҳисоблайди (тугма шарт эмас) ва уни эслаб қолади — кейинги сметалар
    шу фоизлар билан бошланади. Жавобда янги сводка (саҳифа қайта юкланмайди)."""
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or job.mode != "smart":
        return JsonResponse({"ok": False, "error": "Фоизлар фақат Smart режимда"}, status=400)
    try:
        data = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    if data.get("reset"):
        percents = dict(DEFAULT_PERCENTS)
    elif data.get("remembered"):
        percents = remembered_percents(job.owner_id)
    elif isinstance(data.get("pct"), dict):
        # бўш катак — 0 (смета­чи ўчириб қўйди): «стандарт» эмас, акс ҳолда саҳифада бўш, Excel'да 2% бўлиб қолади
        percents = parse_percents({f"pct_{k}": str(v).strip() or "0" for k, v in data["pct"].items()}, job.percents)
    else:
        return JsonResponse({"ok": False, "error": "нотўғри сўров"}, status=400)
    recalc.request(job, percents=percents)
    remember_percents(job.owner_id, job.percents)
    return JsonResponse({"ok": True, "queued": True})


def _svodka_payload(job):
    """Сводка (фоизлар ва суммалар) — фоиз ўзгаргач саҳифа қайта юкланмай янгиланади."""
    totals = job.result["totals"]
    return {
        "rows": [{"key": s["key"], "value": som(s["value"]), "pct": pct_input(s["pct"]) if s["pct"] is not None else None,
                  "changed": s["pct"] is not None and s["pct"] != s["default"]} for s in _svodka_rows(job.result, "smart")],
        "s_nds": som(totals.get("s_nds")), "bez_nds": som(totals.get("bez_nds")),
        "changed": {**DEFAULT_PERCENTS, **(job.percents or {})} != DEFAULT_PERCENTS,
    }


def job_state(request, job_id):
    """Қайта ҳисоб ҳолати (саҳифа шуни сўраб туради): ``pending`` — навбатда ёки ҳисобланяпти; тугагач —
    натижа ҳолати ва smart режимда янги сводка."""
    job = _get_work(request, Job, id=job_id)
    data = {"pending": job.recalc_pending, "status": job.status, "error": job.error if job.status == Job.Status.FAILED else "",
            "running": bool(job.recalc_worker), "duration": job.duration,
            "elapsed": round((timezone.now() - job.recalc_requested).total_seconds()) if job.recalc_requested else None}
    if not job.recalc_pending and job.status == Job.Status.DONE and job.result.get("mode") == "smart":
        data["svodka"] = _svodka_payload(job)
    return JsonResponse(data)


# Нархлашдан кейинги текширув: тури -> (сарлавҳа, нима бўлди, нима қилиш керак) — смета­чига оддий тилда
CHECK_KINDS = {
    "no_qty": ("Миқдори йўқ — сумма 0 бўлиб қолди",
               "Бу қаторларда нарх топилди, лекин миқдор («Кол-во») катаги бўш ёки унда сон эмас, матн турибди. "
               "Сумма = миқдор × нарх бўлгани учун 0 чиқди ва смета жамига кирмади.",
               "Локал файлда шу қаторнинг миқдорини текширинг (одатда F — «по проектным данным»). Миқдор бошқа устунда "
               "бўлса — объект саҳифасида «Листлар тури» → «тузатиш» орқали устунларни кўрсатинг ва қайта нархланг. "
               "Қатор изоҳ ёки ортиқча бўлса — эътибор бермасангиз ҳам бўлади."),
    "amount": ("Сумма нотўғри ҳисобланди",
               "Сумма (H) миқдор × нархга тенг чиқмади.",
               "Одатда миқдор устуни нотўғри аниқланганда бўлади: объект саҳифасида устунларни текширинг ва қайта нархланг."),
    "column": ("Сумма I–M устунига тушмади",
               "Сумма иш ҳақи, машина, материал… устунларидан бирига ўтмади — жамида кўринмайди.",
               "Қаторни белгиланг ва пастки панелдан тўғри устунни танланг."),
    "overwrite": ("Смета маълумоти устидан ёзилди",
                  "Нарх ёки сумма локалнинг номи, бирлиги ёки миқдори турган устунга ёзилиб қолди.",
                  "Локал устунлари нотўғри аниқланган: объект саҳифасида устунларни кўрсатинг ва қайта нархланг."),
    "zero_prices": ("RES'даги нархлар 0",
                    "Ресурс файлида топилган нархларнинг деярли ҳаммаси 0 — RES нархланмаган бўлса керак.",
                    "Нархланган RES файлини юкланг ёки RES'нинг нарх устунини текширинг."),
    "zero_total": ("Жами сумма 0",
                   "Нархлар топилди, лекин смета жами 0 чиқди.",
                   "Миқдор ёки сумма устуни нотўғри аниқланган — объект саҳифасида устунларни текширинг."),
    "other": ("Бошқа муаммолар", "Натижа текширувидан ўтмади.", "Қаторни очиб текширинг."),
}
_CHECK_TEXT = [("миқдори йўқ", "no_qty"), ("кутилгани", "amount"), ("устунига тушмади", "column"),
               ("манба устуни", "overwrite"), ("нархи 0", "zero_prices"), ("жами сумма 0", "zero_total")]


def check_groups(checks, rows) -> list:
    """Текширув муаммолари турлари бўйича: сарлавҳа, сабаб, ечим ва қаторлар (номи, нархи) — техник матн ўрнига."""
    by_row = {r["row"]: r for r in rows}
    groups = {}
    for c in checks:
        kind = c.get("kind") or next((k for t, k in _CHECK_TEXT if t in c.get("problem", "")), "other")
        title, why, fix = CHECK_KINDS.get(kind, CHECK_KINDS["other"])
        g = groups.setdefault(kind, {"kind": kind, "title": title, "why": why, "fix": fix, "rows": [], "general": []})
        row = by_row.get(c.get("row"))
        if row:
            g["rows"].append({"row": row["row"], "name": row.get("name") or "", "unit": row.get("unit") or "",
                              "g": row.get("g"), "qty": row.get("qty")})
        else:
            g["general"].append(c.get("problem", ""))
    order = list(CHECK_KINDS)
    return sorted(groups.values(), key=lambda g: order.index(g["kind"]))


# Жадвалнинг «файл кўриниши»: ?view=… -> (файл, тугма номи)
XL_FILES = {
    "xl_out": ("out", "Нархланган смета"),
    "xl_local": ("local", "Юкланган локал"),
    "xl_res": ("res", "Ресурс файли"),
}


def job_sheet(request, job_id):
    """Excel листи HTML жадвал бўлиб (``f`` — out | local | res, ``s`` — лист номи)."""
    job = _get_work(request, Job, id=job_id)
    try:
        data = preview.sheet_data(job, request.GET.get("f"), request.GET.get("s") or None)
    except preview.NotAvailable as e:
        raise Http404(str(e))
    except Exception as e:  # кўриниш бузилса ҳам саҳифа ишласин
        return JsonResponse({"error": f"{type(e).__name__}: {e}"}, status=500)
    if request.GET.get("head"):
        data = preview.head(data)  # аввал боши — дарҳол кўринади, кейин бутун жадвал
    return JsonResponse(data)


def job_download(request, job_id):
    job = _get_work(request, Job, id=job_id)
    if job.status != Job.Status.DONE or not job.output_path:
        raise Http404("Натижа ҳали тайёр эмас")
    path = job.abs(job.output_path)
    if not path.is_file():
        raise Http404("Файл топилмади")
    _audit_download(request, job, "smeta", job.download_name)
    return FileResponse(open(path, "rb"), as_attachment=True, filename=job.download_name)


# ===================== йўриқнома =====================
def help_page(request):
    """Йўриқнома: иш тартиби қадамма-қадам ва ҳар бир ҳолат / огоҳлантириш нимани билдиради."""
    return render(request, "jobs/help.html", {"nav": "help", "groups": help_groups(), "flow": FLOW})


@require_POST
def guide_seen(request):
    """Саҳифа қўлланмаси кўрилди (ёки ёпилди) — кейин ўзи очилмайди. ``reset`` — ҳаммасини қайта кўрсатиш."""
    user = request.user
    if request.POST.get("reset"):
        user.guides_seen = []
    else:
        page = request.POST.get("page", "")
        if page not in GUIDE_PAGES:
            raise Http404
        if page not in user.guides_seen:
            user.guides_seen = [*user.guides_seen, page]
    user.save(update_fields=["guides_seen"])
    if request.POST.get("reset"):
        return redirect("index")
    return JsonResponse({"ok": True})
