import json
import shutil
import tempfile
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from core.pipeline import UNPRICED_STATUSES
from core.svodka import DEFAULT_PERCENTS
from core.tests.conftest import build_local, build_res

from accounts.models import User

from .models import Batch, Job, Project

TMP_MEDIA = tempfile.mkdtemp(prefix="narxlash_media_")


class LoggedInTestCase(TestCase):
    """Барча саҳифалар кириш талаб қилади — ҳар бир тест смета­чи сифатида ишлайди."""

    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user("smetachi@test.uz", "Kuchli-parol-2026", first_name="Ali", last_name="Valiyev")
        self.client.force_login(self.user)


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class FlowTests(LoggedInTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.src = Path(tempfile.mkdtemp(prefix="narxlash_src_"))
        cls.local = build_local(cls.src / "Локал 3.xlsx")
        cls.res = build_res(cls.src / "Ресурс 3.xlsx")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.src, ignore_errors=True)
        shutil.rmtree(TMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def _upload(self, **files):
        data = {k: SimpleUploadedFile(p.name, p.read_bytes()) for k, p in files.items()}
        return self.client.post(reverse("index"), data)

    def test_index_page(self):
        r = self.client.get(reverse("index"))
        self.assertContains(r, 'name="files"')
        self.assertContains(r, "1 та катта файл")

    # ---------------- объект: бир/кўп файл ----------------
    def _upload_files(self, *paths):
        files = [SimpleUploadedFile(Path(p).name, Path(p).read_bytes()) for p in paths]
        return self.client.post(reverse("index"), {"files": files})

    def test_object_named_after_local_and_renamed(self):
        from core.tests.test_batch import OBJ_LOCAL, OBJ_RES
        self._upload_files(OBJ_RES, OBJ_LOCAL)  # ресурс файли биринчи бўлса ҳам — ном локалдан
        batch = Batch.objects.get()
        self.assertEqual(batch.name, "obekt_lokal_7list")
        page = self.client.get(reverse("batch", args=[batch.id]))
        self.assertContains(page, reverse("work_rename", args=["object", batch.id]))
        r = self.client.post(reverse("work_rename", args=["object", batch.id]), {"name": "  Мактаб   №5  "})
        self.assertRedirects(r, reverse("batch", args=[batch.id]), fetch_redirect_response=False)
        batch.refresh_from_db()
        self.assertEqual(batch.name, "Мактаб №5")
        other = User.objects.create_user("boshqa@test.uz", "Kuchli-parol-2026")
        self.client.force_login(other)
        r = self.client.post(reverse("work_rename", args=["object", batch.id]), {"name": "X"})
        self.assertEqual(r.status_code, 404)  # бошқанинг ишини ўзгартириб бўлмайди

    def test_object_two_files_review_confirm_and_combined_download(self):
        from core.tests.test_batch import OBJ_LOCAL, OBJ_RES
        r = self._upload_files(OBJ_LOCAL, OBJ_RES)
        batch = Batch.objects.get()
        self.assertRedirects(r, reverse("batch", args=[batch.id]))
        self.assertEqual(len(batch.pairs), 4)
        page = self.client.get(reverse("batch", args=[batch.id]))
        self.assertContains(page, "Стройка:")
        self.assertContains(page, "МОСТ НА ПК 75+34")
        self.assertContains(page, "Нусхалар — бир марта олинди (6)")
        self.assertContains(page, "93% мос")              # ҳар бир вариантнинг мослиги
        self.assertContains(page, "★ ")                    # энг мос вариант белгиланган
        self.assertContains(page, "Л1")                    # боғланиш белгилари: Л1 → Р…

        # smetachi: birinchisini o'chiradi, oxirgisining nomini o'zgartiradi
        data = {}
        for n, p in enumerate(batch.pairs):
            data.update({f"res_{n}": p["res"], f"label_{n}": p["label"]})
            if n:
                data[f"include_{n}"] = "1"
        data["label_3"] = "Мост ПК 75+34"
        self.client.post(reverse("batch_confirm", args=[batch.id]), data)
        batch.refresh_from_db()
        jobs = list(batch.jobs.order_by("batch_order"))
        self.assertEqual(batch.status, Batch.Status.DONE)
        self.assertEqual([j.status for j in jobs], ["done"] * 3)
        self.assertEqual(jobs[-1].batch_label, "Мост ПК 75+34")
        self.assertEqual((jobs[-1].local_sheet, jobs[-1].res_sheet), ("202611450_БВ", "202611450_БР"))
        page = self.client.get(reverse("batch", args=[batch.id]))
        self.assertContains(page, "Ҳаммасини битта файлда")
        job_page = self.client.get(reverse("job", args=[jobs[0].id]))
        self.assertContains(job_page, "Объект қисмлари")

        dl = self.client.get(reverse("batch_download", args=[batch.id]))
        self.assertEqual(dl.status_code, 200)
        out = self.src / "obj.xlsx"
        out.write_bytes(b"".join(dl.streaming_content))
        names = load_workbook(out).sheetnames
        self.assertEqual([n for n in names if n.startswith("LOG")],
                         ["LOG 202611410_БВ", "LOG 202611420_БВ", "LOG 202611450_БВ"])

    def test_object_one_file_and_manual_role(self):
        from core.tests.test_batch import OBJ_LOCAL
        self._upload_files(OBJ_LOCAL)
        batch = Batch.objects.get()
        most = next(p for p in batch.pairs if p["local"] == "0:202611450_БВ")
        self.assertIsNone(most["res"])                       # bu faylda 450_БР yo'q
        page = self.client.get(reverse("batch", args=[batch.id]))
        self.assertContains(page, 'data-k="0:202611450_БВ"')   # листни Excel кўринишида очиш
        self.assertContains(page, "алмаштириш")
        prev = self.client.get(reverse("batch_sheet", args=[batch.id]), {"k": "0:202611450_БВ", "n": "30"}).json()
        self.assertEqual((prev["sheet"], prev["rows"], prev["truncated"]), ("202611450_БВ", 30, True))
        self.assertEqual(self.client.get(reverse("batch_sheet", args=[batch.id]), {"k": "0:йўқ"}).status_code, 404)
        # smetachi 400_БВ ni «o'tkazish» qiladi
        idx = next(i for i, s in enumerate(batch.sheets) if s["sheet"] == "202611400_БВ")
        roles = {f"role_{i}": s["role"] for i, s in enumerate(batch.sheets)}
        roles[f"role_{idx}"] = "unknown"
        self.client.post(reverse("batch_replan", args=[batch.id]), roles)
        batch.refresh_from_db()
        self.assertEqual(len(batch.pairs), 3)
        self.assertEqual(batch.roles, {"0:202611400_БВ": "unknown"})

        # smetachi 410_БВ ustunlarini qo'lda ko'rsatadi: saqlanadi, sahifada «қўлда», narxlashga o'tadi
        idx = next(i for i, s in enumerate(batch.sheets) if s["sheet"] == "202611410_БВ")
        loc = batch.sheets[idx]["columns"]
        self.assertTrue(loc["name"] and loc["unit"] and loc["start"])
        form = {f"role_{i}": s["role"] for i, s in enumerate(batch.sheets)}
        form.update({f"cols_{idx}_on": "1", f"cols_{idx}_name": "c", f"cols_{idx}_unit": "D",
                     f"cols_{idx}_start": str(loc["start"]), f"cols_{idx}_end": "1", f"cols_{idx}_price": "ZZZZ9"})
        self.client.post(reverse("batch_replan", args=[batch.id]), form)
        batch.refresh_from_db()
        self.assertEqual(batch.columns, {"0:202611410_БВ": {"name": 3, "unit": 4, "start": loc["start"]}})  # нотўғрилари ташланди
        self.assertTrue(batch.sheet("0:202611410_БВ")["manual"])
        self.assertContains(self.client.get(reverse("batch", args=[batch.id])), "қўлда")
        pair = next(n for n, p in enumerate(batch.pairs) if p["local"] == "0:202611410_БВ")
        self.assertTrue(batch.pairs[pair]["res"])
        self.client.post(reverse("batch_confirm", args=[batch.id]), {f"include_{pair}": "1", f"res_{pair}": batch.pairs[pair]["res"]})
        job = batch.jobs.get(local_sheet="202611410_БВ")
        self.assertEqual(job.local_columns, {"name": 3, "unit": 4, "start": loc["start"]})
        self.assertEqual(job.status, "done")

    def test_upload_errors(self):
        r = self.client.post(reverse("index"), {"files": [SimpleUploadedFile("a.csv", b"x")]})
        self.assertContains(r, "фақат .xls")
        self.assertFalse(Batch.objects.exists())

    def test_full_flow(self):
        r = self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.assertRedirects(r, reverse("job", args=[job.id]))
        self.assertEqual((job.local_sheet, job.res_sheet), ("202611450_БВ", "202611450_БР"))

        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, "Нархлаш")

        r = self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95"})
        self.assertRedirects(r, reverse("job", args=[job.id]))
        job.refresh_from_db()
        self.assertEqual(job.status, Job.Status.DONE, job.error)
        self.assertEqual(job.mode, "smart")
        self.assertEqual(job.result["stats"]["FUZZY"], 1)
        self.assertEqual(job.result["stats"]["MACHINIST"], 1)

        page = self.client.get(reverse("job", args=[job.id]))
        self.assertNotContains(page, "#ЗНАЧ!")
        self.assertContains(page, "29 818,17")       # ишчилар нархи
        self.assertContains(page, "машинист")

        # compat: эски дастур каби — «--» туфайли Excel'дагидек #ЗНАЧ!
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "compat"})
        job.refresh_from_db()
        self.assertEqual(job.mode, "compat")
        self.assertContains(self.client.get(reverse("job", args=[job.id])), "#ЗНАЧ!")

        dl = self.client.get(reverse("job_download", args=[job.id]))
        self.assertEqual(dl.status_code, 200)
        self.assertIn("narxlandi.xlsx", dl["Content-Disposition"])
        out = self.src / "dl.xlsx"
        out.write_bytes(b"".join(dl.streaming_content))
        self.assertEqual(load_workbook(out).sheetnames, ["LOG", "202611450_БВ"])

    def test_edit_percents(self):
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        url = reverse("job_run", args=[job.id])
        base = {"local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"}
        self.client.post(url, base)
        job.refresh_from_db()
        default_total = job.result["totals"]["s_nds"]
        self.assertEqual(job.result["percents"]["prochie"], 0.19)

        self.client.post(url, dict(base, pct_prochie="17,27", pct_s_nds="20", pct_strahovanie="0,5"))
        job.refresh_from_db()
        self.assertEqual(job.percents, {"prochie": 0.1727, "strahovanie": 0.005, "s_nds": 0.2})
        # noto'g'ri yozilgan foiz avvalgi qiymatda qoladi
        self.client.post(url, dict(base, pct_prochie="17,27", pct_s_nds="20", pct_strahovanie="abc"))
        job.refresh_from_db()
        self.assertEqual(job.percents["strahovanie"], 0.005)
        t = job.result["totals"]
        self.assertAlmostEqual(t["prochie"], (t["pryamye"] - t["transport_obor"] - t["oborudovanie"]) * 0.1727)
        self.assertAlmostEqual(t["s_nds"], t["bez_nds"] * 1.2)
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, 'value="17,27"')
        self.assertContains(page, "Всего с НДС      20 %")

        # yuqoridagi forma (pct_ maydonlarisiz) foizlarni saqlab qoladi
        self.client.post(url, dict(base, fuzzy="90"))
        job.refresh_from_db()
        self.assertEqual(job.result["percents"]["prochie"], 0.1727)

        self.client.post(url, dict(base, reset_percents="1"))
        job.refresh_from_db()
        self.assertEqual(job.percents, dict(DEFAULT_PERCENTS))
        self.assertAlmostEqual(job.result["totals"]["s_nds"], default_total)

    def test_percents_autosave_and_remembered_for_next_jobs(self):
        """Фоиз ўзгарса — тугмасиз ҳисобланади (JSON) ва эслаб қолинади; кейинги смета шу фоиз билан бошланади."""
        from .models import Setting
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        url = reverse("job_percents", args=[job.id])
        post = lambda payload: self.client.post(url, json.dumps(payload), content_type="application/json").json()

        d = post({"pct": {"transport_mat": "5", "prochie": "19"}})
        self.assertTrue(d["ok"])
        job.refresh_from_db()
        self.assertEqual(job.percents["transport_mat"], 0.05)
        row = next(r for r in d["rows"] if r["key"] == "transport_mat")
        self.assertEqual((row["pct"], row["changed"]), ("5", True))
        self.assertEqual(Setting.objects.get(key="percents").value["transport_mat"], 0.05)

        self._upload(local=self.local, res=self.res)  # янги смета — эслаб қолинган фоиз билан
        new = Job.objects.exclude(id=job.id).get()
        self.assertEqual(new.percents["transport_mat"], 0.05)
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, "Эслаб қолинган (стандартдан фарқли)")

        d = post({"reset": True})
        job.refresh_from_db()
        self.assertEqual((job.percents, d["changed"]), (dict(DEFAULT_PERCENTS), False))
        self.assertEqual(Setting.objects.get(key="percents").value, dict(DEFAULT_PERCENTS))  # ўчмайди — стандарт эслаб қолинади

        # ўз фоизи йўқ эски смета ҳам қайта ҳисобланганда эслаб қолинган фоизни олади
        post({"pct": {"transport_mat": "4"}})
        old = Job.objects.exclude(id=job.id).get()
        old.percents = {}
        old.save()
        self.client.post(reverse("job_run", args=[old.id]), {"local_sheet": old.local_sheet, "res_sheet": old.res_sheet,
                                                             "fuzzy": "95", "mode": "smart"})
        old.refresh_from_db()
        self.assertEqual(old.result["percents"]["transport_mat"], 0.04)
        self.assertEqual(self.client.post(url, "{}", content_type="application/json").status_code, 400)

    def test_move_rows_between_categories(self):
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        job.refresh_from_db()
        rows = {r["row"]: r for r in job.result["rows"]}
        water = next(r for r, x in rows.items() if x["name"] == "ВОДА")
        cloth = next(r for r, x in rows.items() if x["name"] == "ТКАНЬ МЕШОЧНАЯ")
        url = reverse("job_categories", args=[job.id])
        post = lambda payload: self.client.post(url, json.dumps(payload), content_type="application/json")

        r = post({"rows": [water, cloth], "category": "M"})
        self.assertEqual(r.json(), {"ok": True, "changed": 2, "overrides": 2})
        job.refresh_from_db()
        self.assertEqual(job.categories, {str(water): "M", str(cloth): "M"})
        rows = {x["row"]: x for x in job.result["rows"]}
        self.assertTrue(rows[water]["override"] and rows[water]["category"] == "M")
        self.assertAlmostEqual(job.result["totals"]["oborudovanie"], rows[water]["h"] + rows[cloth]["h"])
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, 'title="Қўлда ўзгартирилган (асли K)"')  # M ustunidagi ✎ belgisi
        self.assertContains(page, 'data-filter="OVERRIDE"')

        # o'zining asl ustuniga qaytarish — qo'lda belgi olib tashlanadi
        post({"rows": [cloth], "category": "K"})
        job.refresh_from_db()
        self.assertEqual(job.categories, {str(water): "M"})
        post({"rows": [water], "category": ""})
        job.refresh_from_db()
        self.assertEqual(job.categories, {})
        self.assertEqual(job.result["totals"]["oborudovanie"], 0)

        self.assertEqual(post({"rows": [water], "category": "Z"}).status_code, 400)
        self.assertEqual(post({"rows": [99999], "category": "M"}).status_code, 400)
        self.assertEqual(post({"rows": "x"}).status_code, 400)

        # lokal list almashsa qo'lda kategoriyalar tozalanadi
        post({"rows": [water], "category": "L"})
        job.refresh_from_db()
        from .services import run_job
        run_job(job, "boshqa", None, 95)
        self.assertEqual(job.categories, {})

    def test_prices_accept_set_clear(self):
        from core.tests.test_smart import TESTS
        d = TESTS / "excel_calc/bolnitsa7_kr_blok2"
        self._upload(local=d / "input_local.xlsx", res=d / "input_res.xlsx")
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        job.refresh_from_db()
        red = [r["row"] for r in job.result["rows"] if r.get("suggested_price") is not None]
        exact = next(r["row"] for r in job.result["rows"] if r["status"] == "EXACT")
        self.assertEqual(len(red), 12)
        url = reverse("job_prices", args=[job.id])
        post = lambda payload: self.client.post(url, json.dumps(payload), content_type="application/json")

        # 10 tasini taklif bo'yicha, bittasini qo'lda, bitta EXACT qatorda taklif yo'q — o'tkazib yuboriladi
        r = post({"rows": red[:10] + [exact], "action": "accept"})
        self.assertEqual(r.json(), {"ok": True, "changed": 10, "skipped": 1})
        r = post({"rows": [red[10]], "action": "set", "price": "4 500 000"})
        self.assertTrue(r.json()["ok"])
        job.refresh_from_db()
        self.assertEqual(job.result["stats"]["MANUAL"], 11)
        self.assertEqual(job.result["stats"]["NOT_PRICED"], 1)
        self.assertEqual(job.prices[str(red[10])], 4_500_000)
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, ">қўлда</span>")
        self.assertContains(page, "1 та қатор нархланмади")

        post({"rows": [red[11]], "action": "accept"})
        job.refresh_from_db()
        self.assertContains(self.client.get(reverse("job", args=[job.id])), "ҳаммаси нархланган")

        post({"rows": red, "action": "clear"})
        job.refresh_from_db()
        self.assertEqual(job.prices, {})
        self.assertEqual(job.result["stats"]["NOT_PRICED"], 12)
        self.assertEqual(post({"rows": [red[0]], "action": "set", "price": "abc"}).status_code, 400)
        self.assertEqual(post({"rows": [red[0]], "action": "boshqa"}).status_code, 400)

    def test_not_found_stage_clear_and_restore(self):
        """1-босқич: «ТОПИЛМАДИ» қаторлари саҳифада бирлиги бор/йўқ бўйича; олиб ташлаш ва қайтариш."""
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        job.refresh_from_db()
        nf = [r["row"] for r in job.result["rows"] if r["status"] in UNPRICED_STATUSES]  # ТОПИЛМАДИ, ТЕКШИРИНГ …
        self.assertTrue(nf)
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, "1-босқич")
        self.assertIn('DEFAULT_TAB = "1"', page.content.decode())  # саҳифа 1-босқичда очилади

        url = reverse("job_not_found", args=[job.id])
        post = lambda payload: self.client.post(url, json.dumps(payload), content_type="application/json")
        self.assertEqual(post({"rows": nf, "action": "clear"}).json(), {"ok": True, "changed": len(nf)})
        job.refresh_from_db()
        self.assertEqual(job.cleared, nf)
        self.assertEqual(job.result["cleared"], nf)
        ws = load_workbook(job.abs(job.output_path))[job.local_sheet]
        self.assertIsNone(ws[f"G{nf[0]}"].value)
        self.assertContains(self.client.get(reverse("job", args=[job.id])), "ёзуви қолмади")

        post({"rows": nf[:1], "action": "restore"})
        job.refresh_from_db()
        self.assertEqual(job.cleared, nf[1:])
        exact = next(r["row"] for r in job.result["rows"] if r["status"] == "EXACT")
        self.assertEqual(post({"rows": [exact], "action": "clear"}).status_code, 400)
        self.assertEqual(post({"rows": nf, "action": "boshqa"}).status_code, 400)

    def test_table_checked_stage(self):
        """2-босқич якунланмагунча сводка ва қолганлари йиғилган; якунлангач очилади, қайта очиш мумкин."""
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        job.refresh_from_db()
        nf = [r["row"] for r in job.result["rows"] if r["status"] in UNPRICED_STATUSES]  # ТОПИЛМАДИ, ТЕКШИРИНГ …
        self.assertIn('DEFAULT_TAB = "1"', self.client.get(reverse("job", args=[job.id])).content.decode())
        self.client.post(reverse("job_not_found", args=[job.id]), json.dumps({"rows": nf, "action": "clear"}),
                         content_type="application/json")
        page = self.client.get(reverse("job", args=[job.id])).content.decode()
        self.assertIn("2-босқич", page)
        self.assertIn('DEFAULT_TAB = "2"', page)  # 1-босқич тугади — саҳифа 2-босқичда очилади
        r = self.client.post(reverse("job_table_checked", args=[job.id]), {"done": "1"})
        self.assertRedirects(r, reverse("job", args=[job.id]) + "#t-3", fetch_redirect_response=False)
        job.refresh_from_db()
        self.assertTrue(job.table_checked)
        page = self.client.get(reverse("job", args=[job.id])).content.decode()
        self.assertIn('DEFAULT_TAB = "3"', page)
        self.client.post(reverse("job_table_checked", args=[job.id]), {"done": "0"})
        job.refresh_from_db()
        self.assertFalse(job.table_checked)

    def test_near_suggestion_saved_and_accepted(self):
        """Топилмаганга энг яқин ресурс натижада сақланади (JSON dict) ва 1-босқичдан нархи айнан қўйилади."""
        wb = load_workbook(self.local)
        ws = wb.active
        row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value == "НЕСУЩЕСТВУЮЩИЙ МАТЕРИАЛ")
        ws.cell(row, 3).value, ws.cell(row, 4).value = "ТКАНЬ МЕШОЧНАЯ ДВОЙНАЯ", "10М2"
        local = self.src / "near.xlsx"
        wb.save(local)
        self._upload(local=local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        job.refresh_from_db()
        x = next(r for r in job.result["rows"] if r["row"] == row)
        self.assertEqual((x["near"]["name"], x["near"]["price"]), ("ТКАНЬ МЕШОЧНАЯ", 105000))
        page = self.client.get(reverse("job", args=[job.id]))
        self.assertContains(page, "нархни қўйиш")
        self.assertContains(page, "локал · F")  # миқдор локал сметадан
        # «Кўриш»: локал / натижа — шу қатор, ресурс — таклиф қилинган ресурснинг RES қатори
        res_row = x["near"]["res_row"]
        self.assertIsInstance(res_row, int)
        self.assertContains(page, f'data-view="xl_res" data-row="{res_row}"')
        self.assertContains(page, f'data-view="xl_out" data-row="{row}"')
        self.assertIsNotNone(next(r for r in page.context["nf"]["rows"] if r["row"] == row)["near_sum"])
        r = self.client.post(reverse("job_prices", args=[job.id]), json.dumps({"rows": [row], "action": "accept"}),
                             content_type="application/json")
        self.assertEqual(r.json()["changed"], 1)
        job.refresh_from_db()
        self.assertEqual(job.prices[str(row)], 105000)
        self.assertEqual(next(r for r in job.result["rows"] if r["row"] == row)["status"], "MANUAL")

    def test_file_view_like_excel(self):
        """2-босқич «Файл» кўриниши: нархланган смета, юкланган локал ва ресурс — листлари билан."""
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {
            "local_sheet": job.local_sheet, "res_sheet": job.res_sheet, "fuzzy": "95", "mode": "smart"})
        page = self.client.get(reverse("job", args=[job.id]) + "?view=xl_out")
        self.assertContains(page, "Нархланган смета")
        self.assertContains(page, 'id="xl-view"')

        url = reverse("job_sheet", args=[job.id])
        out = self.client.get(url, {"f": "out"}).json()
        self.assertEqual(out["sheet"], "202611450_БВ")
        self.assertIn("LOG", out["sheets"])
        self.assertIn("29 818", out["html"])      # ишчилар нархи, Excel форматида
        self.assertIn("#FFC7CE", out["html"])     # нархланмаган G — қизил
        self.assertEqual(self.client.get(url, {"f": "out", "s": "LOG"}).json()["sheet"], "LOG")
        self.assertEqual(self.client.get(url, {"f": "res"}).json()["sheet"], "202611450_БР")
        self.assertEqual(self.client.get(url, {"f": "local"}).json()["sheet"], "202611450_БВ")
        self.assertEqual(self.client.get(url, {"f": "x"}).status_code, 404)

    def test_single_file_for_both(self):
        self._upload(local=self.res)
        job = Job.objects.get()
        self.assertEqual(job.local_path, job.res_path)

    def test_bad_extension(self):
        r = self.client.post(reverse("index"), {"local": SimpleUploadedFile("a.csv", b"x")})
        self.assertContains(r, "фақат .xls")
        self.assertFalse(Job.objects.exists())

    def test_wrong_sheet_shows_error(self):
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.client.post(reverse("job_run", args=[job.id]), {"local_sheet": "йўқ", "res_sheet": "", "fuzzy": "95"})
        job.refresh_from_db()
        self.assertEqual(job.status, Job.Status.FAILED)
        self.assertContains(self.client.get(reverse("job", args=[job.id])), "Нархлашда хато")

    def test_download_before_run_404(self):
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get()
        self.assertEqual(self.client.get(reverse("job_download", args=[job.id])).status_code, 404)

    def test_delete_from_history(self):
        """Тарихдан ўчириш: объект ва алоҳида смета — базадан ҳам, дискдан ҳам."""
        from core.tests.test_batch import OBJ_LOCAL
        self._upload_files(OBJ_LOCAL)
        batch = Batch.objects.get()
        self.client.post(reverse("batch_confirm", args=[batch.id]), {})
        job_dirs = [j.dir for j in batch.jobs.all()]
        self._upload(local=self.local, res=self.res)
        job = Job.objects.get(batch__isnull=True)
        page = self.client.get(reverse("index"))
        self.assertContains(page, reverse("work_delete", args=["object", batch.id]))
        self.assertContains(page, reverse("work_delete", args=["smeta", job.id]))

        r = self.client.post(reverse("work_delete", args=["object", batch.id]), HTTP_X_REQUESTED_WITH="fetch")
        self.assertEqual(r.json(), {"ok": True})
        self.assertFalse(Batch.objects.exists())
        self.assertFalse(batch.dir.exists())
        self.assertFalse(any(d.exists() for d in job_dirs))
        self.assertEqual(Job.objects.count(), 1)  # алоҳида смета қолди

        dash = self.client.get(reverse("dashboard"))
        self.assertContains(dash, "Барча ишлар")
        self.assertContains(dash, reverse("work_delete", args=["smeta", job.id]))
        self.assertEqual(sum(p["count"] for p in dash.context["pipeline"]), dash.context["count_smetas"])

        r = self.client.post(reverse("work_delete", args=["smeta", job.id]))
        self.assertRedirects(r, reverse("index"))
        self.assertContains(self.client.get(reverse("dashboard")), "Ҳали ҳеч нарса юкланмаган")
        self.assertFalse(Job.objects.exists())
        self.assertFalse(job.dir.exists())
        self.assertEqual(self.client.get(reverse("work_delete", args=["smeta", job.id])).status_code, 405)  # фақат POST


class ResGroupsTests(TestCase):
    def test_rows_grouped_in_res_order_with_move_offer(self):
        from .views import _res_groups
        res = {
            "res_sections": [{"kind": "workers", "column": "I", "title": "ТРУДОВЫЕ РЕСУРСЫ"},
                             {"kind": "equipment", "column": "M", "title": "ОБОРУДОВАНИЕ"}],
            "rows": [
                {"row": 5, "name": "НАСОС Б", "category": "K", "res_section": 1, "res_row": 30, "h": 200.0, "status": "EXACT"},
                {"row": 6, "name": "НАСОС А", "category": "K", "res_section": 1, "res_row": 20, "h": 100.0, "status": "EXACT"},
                {"row": 7, "name": "ТРУД", "category": "I", "res_section": 0, "res_row": 3, "h": 50.0, "status": "EXACT"},
                {"row": 9, "name": "ЙЎҚ", "category": "K", "res_section": None, "h": 0, "status": "NOT_FOUND"},
                {"row": 10, "name": "ИШ", "category": "K", "res_section": None, "h": 0, "status": "WORK_HEADER"},
            ],
        }
        groups = _res_groups(res)
        self.assertEqual([g["title"] for g in groups],
                         ["ТРУДОВЫЕ РЕСУРСЫ", "ОБОРУДОВАНИЕ", "Ресурс файлида топилмаганлар", "Иш сарлавҳалари ва машинистлар"])
        eq = groups[1]
        self.assertEqual([r["row"] for r in eq["rows"]], [6, 5])  # RES даги тартибда
        self.assertEqual(eq["to_move"], [6, 5])
        self.assertEqual(eq["sum"], 300.0)
        self.assertEqual(groups[0]["to_move"], [])

    def test_not_found_context(self):
        """1-босқич: ҳар бир «ТОПИЛМАДИ» қаторига бўлим ва иш (норматив) номи; бўлим номи — бирлиги йўқ тур."""
        from .views import _not_found_stage
        rows = [
            {"row": 16, "name": "ЗЕМЛЯНЫЕ РАБОТЫ", "unit": None, "num": None, "status": "NOT_FOUND"},
            {"row": 17, "name": "РАЗРАБОТКА ГРУНТА", "unit": "1000М3", "num": "1", "code": "Е0101", "status": "WORK_HEADER"},
            {"row": 18, "name": "ЗАТРАТЫ ТРУДА", "unit": "ЧЕЛ.-Ч", "num": "1.1", "code": "1", "status": "EXACT"},
            {"row": 19, "name": "ШАЙБА", "unit": "ШТ", "num": "1.2", "code": "333", "status": "NOT_FOUND"},
            {"row": 20, "name": "БОЛТ", "unit": "ШТ", "num": "2", "code": "СЦЕНА", "status": "NOT_FOUND"},
        ]
        nf = _not_found_stage(rows)
        by = {r["row"]: r for r in nf["rows"]}
        self.assertEqual(sorted(by), [16, 19, 20])
        self.assertEqual((by[16]["nf_kind"], by[16]["ctx_section"]), ("nounit", ""))
        self.assertEqual((by[16]["ctx_count"], by[16]["ctx_first"]), (4, "1 РАЗРАБОТКА ГРУНТА"))  # остида нима бор
        self.assertEqual((by[19]["ctx_section"], by[19]["ctx_work"]), ("ЗЕМЛЯНЫЕ РАБОТЫ", "1 Е0101 РАЗРАБОТКА ГРУНТА"))
        self.assertEqual((by[20]["nf_kind"], by[20]["ctx_work"]), ("unit", ""))  # мустақил ресурс — ишга тегишли эмас

    def test_lm_stage_cards(self):
        """2-босқич: M картасида RES «ОБОРУДОВАНИЕ» даги ва таклиф қилинган қаторлар, L да кабель таклифи;
        аллақачон шу устундагилар рўйхатда йўқ."""
        from .views import _lm_stage, _res_groups
        res = {
            "res_sections": [{"kind": "equipment", "column": "M", "title": "ОБОРУДОВАНИЕ"}],
            "rows": [
                {"row": 5, "name": "НАСОС", "unit": "ШТ", "category": "K", "res_section": 0, "res_row": 3, "h": 200.0, "status": "EXACT"},
                {"row": 4, "name": "ОБОРУДОВАНИЕ", "category": "K", "suggested_category": "M", "h": 0, "status": "NOT_FOUND"},
                {"row": 6, "name": "ЩИТ", "category": "M", "res_section": 0, "res_row": 4, "h": 50.0, "status": "EXACT"},
                {"row": 7, "name": "КАБЕЛЬ ВВГ", "unit": "М", "category": "K", "suggested_category": "L", "res_section": None,
                 "h": 10.0, "status": "EXACT"},
            ],
        }
        lm = _lm_stage(res["rows"], _res_groups(res), {"L": "Кабель", "M": "Ускуна"})
        cards = {c["col"]: c for c in lm["cards"]}
        self.assertEqual([r["row"] for r in cards["M"]["rows"]], [5])
        self.assertEqual([r["row"] for r in cards["L"]["rows"]], [7])
        self.assertEqual((cards["M"]["done"], cards["M"]["sum"], lm["pending"]), (1, 200.0, 2))
        self.assertEqual(res["rows"][0]["lm_target"], "M")
        self.assertNotIn("lm_target", res["rows"][1])  # бўлим номи (бирлиги йўқ) таклифга кирмайди


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class FolderTests(LoggedInTestCase):
    """Бутун папка: ҳар бир ички папка алоҳида қисм, ўзи жуфтланади ва нархланади; ZIP да тузилиш сақланади."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        golden = Path(__file__).resolve().parent.parent / "core/tests/golden/202611450"
        cls.local = (golden / "input_local.xlsx").read_bytes()
        cls.res = (golden / "input_res.xlsx").read_bytes()

    def _upload(self):
        from unittest import mock

        from . import folders

        files = {
            "МТТ/1. Блок-А/Подзем/A_lokal.xlsx": self.local,
            "МТТ/1. Блок-А/Подзем/A_res.xlsx": self.res,
            "МТТ/2. ОВ/ov_lokal.xlsx": self.local,
            "МТТ/2. ОВ/ov_res.xlsx": self.res,
            "МТТ/2. ОВ/ov_lokal_narxlandi.xlsx": self.local,  # олдинги натижа — олинмайди
        }
        data = {"files": [SimpleUploadedFile(p.split("/")[-1], b) for p, b in files.items()],
                "paths": list(files), "skipped": ["МТТ/Yangi.zip", "МТТ/2. ОВ/~$ov_lokal.xlsx"]}
        # фондаги иш тестда шу заҳоти бажарилсин
        with mock.patch.object(folders, "start", lambda project: folders._run(project.id)):
            r = self.client.post(reverse("project_create"), data)
        self.assertEqual(r.status_code, 200, r.content)
        return Project.objects.get(), r.json()

    def test_folder_upload_prices_each_subfolder(self):
        from unittest import mock

        from . import folders

        project, data = self._upload()
        self.assertEqual(project.name, "МТТ")
        # юклангач фақат очилади ва жуфтланади — ҳисоблаш смета­чи босгандагина
        self.assertEqual(project.status, Project.Status.READY)
        parts = {p.folder: p for p in project.parts.all()}
        self.assertEqual(set(parts), {"1. Блок-А/Подзем", "2. ОВ"})
        for p in parts.values():
            self.assertEqual(p.stage, Batch.Stage.READY, p.message)
            self.assertEqual(p.jobs.count(), 0)
        page = self.client.get(data["url"])
        self.assertContains(page, "Ҳисоблашни бошлаш")
        self.assertContains(page, "A_lokal.xlsx")         # режа: қайси локал → қайси ресурс
        self.assertContains(page, "A_res.xlsx")
        self.assertContains(page, 'class="pv-open')              # локал/ресурс — Excel кўриниши
        self.assertContains(page, 'data-viewer-url="')           # листлар ва номзодлар — босилганда сўралади
        part = next(iter(parts.values()))
        viewer = self.client.get(reverse("project_viewer", args=[project.id, part.id])).json()
        self.assertEqual(viewer["url"], reverse("batch_sheet", args=[part.id]))
        self.assertTrue(viewer["sheets"])

        with mock.patch.object(folders, "start", lambda project: folders._run(project.id)):
            r = self.client.post(reverse("project_price", args=[project.id]))
        self.assertRedirects(r, data["url"])
        project.refresh_from_db()
        self.assertEqual(project.status, Project.Status.DONE)
        for p in project.parts.all():
            self.assertEqual(p.stage, Batch.Stage.DONE, p.message)
            self.assertEqual(p.jobs.count(), 1)
        skipped = {f["path"]: f["skip"] for f in project.files if f["skip"]}
        self.assertEqual(skipped["2. ОВ/ov_lokal_narxlandi.xlsx"], "олдинги натижа")
        # Excel очиқлигидаги «~$» ва Excel бўлмаган файллар умуман кўрсатилмайди (чалғитмасин)
        self.assertFalse(any("~$" in f["path"] or not f["path"].endswith((".xls", ".xlsx", ".xlsm"))
                             for f in project.files))

        page = self.client.get(data["url"])
        self.assertContains(page, "Подзем")
        self.assertContains(page, "ov_lokal_narxlandi.xlsx")  # натижа файли дарахтда
        self.assertContains(page, 'data-working="0"')
        self.assertContains(self.client.get(data["url"] + "?partial=1"), "project-state")

        import io
        import zipfile
        z = zipfile.ZipFile(io.BytesIO(b"".join(self.client.get(reverse("project_download", args=[project.id])).streaming_content)))
        self.assertEqual(sorted(z.namelist()), ["1. Блок-А/Подзем/A_lokal_narxlandi.xlsx", "2. ОВ/ov_lokal_narxlandi.xlsx"])

    def test_folder_without_excel_rejected(self):
        r = self.client.post(reverse("project_create"), {"skipped": ["x/a.zip"]})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(Project.objects.exists())

    def _zip(self):
        """«Объект.zip»: ичида PDF, ички ZIP (ўзида папка) ва «../» ли зарарли йўл."""
        import io
        import zipfile
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, "w") as z:
            z.writestr("смета/lokal.xlsx", self.local)
            z.writestr("смета/res.xlsx", self.res)
            z.writestr("смета/лойиҳа.pdf", b"%PDF")
        outer = io.BytesIO()
        with zipfile.ZipFile(outer, "w") as z:
            z.writestr("Объект/1. Блок-А/a_lokal.xlsx", self.local)
            z.writestr("Объект/1. Блок-А/a_res.xlsx", self.res)
            z.writestr("Объект/2. ОВ.zip", inner.getvalue())
            z.writestr("Объект/экспертиза.pdf", b"%PDF")
            z.writestr("Объект/3. Свод/Умумий ?иймати .xlsx", self.local)  # «?» — Windows'да тақиқланган
            z.writestr("Объект/~$a_lokal.xlsx", b"x")
            z.writestr("Объект/../../evil.xlsx", self.local)
        return outer.getvalue()

    def test_zip_archive_opened_nested_only_excel(self):
        from unittest import mock

        from . import folders

        data = {"files": [SimpleUploadedFile("Объект.zip", self._zip())], "paths": ["Объект.zip"]}
        with mock.patch.object(folders, "start", lambda project: folders._run(project.id)):
            r = self.client.post(reverse("project_create"), data)
        self.assertEqual(r.status_code, 200, r.content)
        project = Project.objects.get()
        self.assertEqual(project.name, "Объект")
        self.assertEqual(sorted(f["path"] for f in project.files),
                         ["1. Блок-А/a_lokal.xlsx", "1. Блок-А/a_res.xlsx", "2. ОВ/смета/lokal.xlsx",
                          "2. ОВ/смета/res.xlsx", "3. Свод/Умумий _иймати .xlsx", "evil.xlsx"])  # PDF, «~$» йўқ; «../» архив ичида қолди
        self.assertFalse((project.dir.parent / "evil.xlsx").exists())
        parts = {p.folder: p for p in project.parts.all()}
        self.assertEqual(parts["1. Блок-А"].stage, Batch.Stage.READY, parts["1. Блок-А"].message)
        self.assertEqual(parts["2. ОВ/смета"].stage, Batch.Stage.READY, parts["2. ОВ/смета"].message)
        self.assertFalse((project.dir / "unpack").exists())

    def test_zip_from_plain_files_form_goes_to_folder_flow(self):
        from unittest import mock

        from . import folders

        with mock.patch.object(folders, "start", lambda project: None):
            r = self.client.post(reverse("index"), {"files": [SimpleUploadedFile("Объект.zip", self._zip())]})
        project = Project.objects.get()
        self.assertRedirects(r, reverse("project", args=[project.id]), fetch_redirect_response=False)
        self.assertEqual(project.parts.count(), 4)

    def test_broken_archive_error(self):
        r = self.client.post(reverse("project_create"), {"files": [SimpleUploadedFile("bad.zip", b"not a zip")],
                                                         "paths": ["bad.zip"]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("ZIP", r.json()["error"])
        self.assertFalse(Project.objects.exists())

    def test_delete_folder_stops_running_work_first(self):
        from django.utils import timezone

        project, data = self._upload()
        part_dirs = [p.dir for p in project.parts.all()]
        Project.objects.filter(id=project.id).update(status=Project.Status.PROCESSING, heartbeat=timezone.now())
        self.assertContains(self.client.get(reverse("index")), reverse("project_stop", args=[project.id]))
        # ишлаётган папка ҳам ўчирилади: аввал тўхтатилади
        r = self.client.post(reverse("work_delete", args=["folder", project.id]), HTTP_X_REQUESTED_WITH="fetch")
        self.assertEqual(r.json(), {"ok": True})
        self.assertFalse(Project.objects.exists())
        self.assertFalse(Batch.objects.exists())
        self.assertFalse(project.dir.exists())
        self.assertFalse(any(d.exists() for d in part_dirs))

    def test_stop_and_resume(self):
        from unittest import mock

        from . import folders

        project, data = self._upload()
        # фонда иш тўхтаб қолган ҳолат: қисмлар навбатда
        project.parts.update(stage=Batch.Stage.QUEUED)
        Project.objects.filter(id=project.id).update(status=Project.Status.PROCESSING)
        r = self.client.post(reverse("project_stop", args=[project.id]))
        self.assertRedirects(r, data["url"])
        project.refresh_from_db()
        self.assertEqual(project.status, Project.Status.READY)
        self.assertFalse(project.stop_requested)
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.STOPPED})
        self.assertContains(self.client.get(data["url"]), "Давом эттириш")

        with mock.patch.object(folders, "start", lambda project: folders._run(project.id)):
            self.client.post(reverse("project_resume", args=[project.id]))
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.READY})

    def test_stop_flag_interrupts_background_run(self):
        """Фондаги иш ҳар бир файлдан олдин «Тўхтатиш»ни текширади."""
        from . import folders

        project, data = self._upload()
        project.parts.update(stage=Batch.Stage.QUEUED)
        Project.objects.filter(id=project.id).update(status=Project.Status.PROCESSING, stop_requested=True)
        folders._run(project.id)
        project.refresh_from_db()
        self.assertEqual({p.stage for p in project.parts.all()}, {Batch.Stage.STOPPED})
        self.assertFalse(project.stop_requested)
        self.assertEqual(project.status, Project.Status.READY)

class XlViewFormatTests(TestCase):
    def test_number_formats_like_excel(self):
        from .xlview import display
        cases = [
            (1234567.891, "#,##0", "1 234 568"),
            (-1234.5, '#,##0;-#,##0;"-"', "-1 235"),
            (0, '#,##0;-#,##0;"-"', "-"),
            (0.12, "0.00%", "12,00%"),
            (39.7, "General", "39,7"),
            (3.0, "0.0#", "3,0"),
            (2.5, '#,##0.00" сум"', "2,50 сум"),
            ("#VALUE!", "General", "#ЗНАЧ!"),
            (None, "General", ""),
        ]
        for v, fmt, want in cases:
            self.assertEqual(display(v, fmt), want, (v, fmt))


class JobPageExplainTests(TestCase):
    """Смета саҳифаси: текширув муаммолари ва 0 нархлар оддий тилда тушунтирилади."""

    def test_check_groups(self):
        from .views import check_groups
        rows = [{"row": 10, "name": "ЦЕМЕНТ", "unit": "Т", "g": 26800, "qty": None},
                {"row": 11, "name": "ВОДА", "unit": "М3", "g": 5000, "qty": 2}]
        checks = [{"row": 10, "kind": "no_qty", "problem": "…"},
                  {"row": 11, "problem": "H = 1, кутилгани 10 000,00 (миқдор 2 × нарх 5000)"},  # эски натижа: тури йўқ
                  {"row": None, "kind": "zero_total", "problem": "жами 0"}]
        groups = check_groups(checks, rows)
        self.assertEqual([g["kind"] for g in groups], ["no_qty", "amount", "zero_total"])
        self.assertEqual(groups[0]["rows"][0]["name"], "ЦЕМЕНТ")
        self.assertIn("Миқдори йўқ", groups[0]["title"])
        self.assertEqual(groups[2]["general"], ["жами 0"])

    def test_zero_reason(self):
        from .templatetags.narx import zero_reason
        self.assertEqual(zero_reason({"status": "EXACT", "g": 0, "name": "ЗАТРАТЫ ТРУДА МАШИНИСТОВ"})["tag"], "машинист · 0")
        self.assertEqual(zero_reason({"status": "MACHINIST", "g": 0, "name": "x"})["tag"], "машинист · 0")
        self.assertEqual(zero_reason({"status": "EXACT", "g": 0, "name": "ГВОЗДИ"})["tag"], "RES'да 0")
        self.assertEqual(zero_reason({"status": "WORK_HEADER", "g": 0, "name": "x"})["tag"], "иш сарлавҳаси")
        self.assertEqual(zero_reason({"status": "EXACT", "g": 100, "name": "ГВОЗДИ"}), {})
