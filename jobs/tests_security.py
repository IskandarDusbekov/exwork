"""Begona (zararli bo'lishi mumkin) Excel va arxiv fayllar: XSS, «arxiv bombasi», sirlarning tashqi dasturlarga
o'tmasligi. Har bir test — xavfsizlik tekshiruvida amalda tasdiqlangan holat."""
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import xlview

PAYLOAD = "X'}</style><img src=x onerror=alert(document.cookie)><style>"


def test_font_name_cannot_inject_html(tmp_path):
    """Shrift nomi <style> ichiga yoziladi — undan chiqib HTML/skript qo'shib bo'lmasin."""
    p = tmp_path / "x.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "матн"
    ws["A1"].font = Font(name=PAYLOAD)
    ws["B1"] = "Times"
    ws["B1"].font = Font(name="Times New Roman")
    wb.save(p)
    html = xlview.render(p)["html"]
    assert "<img" not in html and "onerror=" not in html
    assert html.count("</style>") == 1 and html.count("<style>") == 1  # faqat bizning <style> bloki
    style = html[:html.index("</style>")]
    assert style.count("'") % 2 == 0 and "<" not in style[len("<style>"):]  # qo'shtirnoqdan chiqib bo'lmaydi
    assert "font-family:'Times New Roman'" in html  # oddiy shrift nomi o'zgarmaydi


def test_color_must_be_hex():
    styles = xlview._Styles.__new__(xlview._Styles)
    styles.theme = []

    class C:
        type, rgb, tint = "rgb", "00;}</style><x>", 0
    assert styles.color(C) is None
    C.rgb = "FF00AA11"
    assert styles.color(C) == "#00AA11"


def test_fill_colour_still_rendered(tmp_path):
    p = tmp_path / "f.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.active["A1"].fill = PatternFill("solid", fgColor="FFFFCC00")
    wb.save(p)
    assert "background:#FFCC00" in xlview.render(p)["html"]


# ---------- «arxiv bombasi» ----------
def _zip(path, members):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, size in members:
            with z.open(name, "w") as m:
                for _ in range(size // (1 << 20)):
                    m.write(b"\0" * (1 << 20))
    return path


def test_oversized_excel_in_archive_is_skipped(tmp_path, monkeypatch):
    from . import archives
    monkeypatch.setattr(archives, "MAX_EXCEL_BYTES", 2 << 20)
    z = _zip(tmp_path / "a.zip", [("katta.xlsx", 5 << 20), ("kichik.xlsx", 1 << 20)])
    out = archives.extract_excel(z, tmp_path / "s")
    assert [str(r) for r, _ in out] == ["a/kichik.xlsx"]  # kattasi tashlandi, qolgani olindi
    assert sum(p.stat().st_size for p in (tmp_path / "s").iterdir()) == 1 << 20  # diskda qoldiq yo'q


def test_total_unpacked_budget_stops_upload(tmp_path, monkeypatch):
    from . import archives
    monkeypatch.setattr(archives, "MAX_TOTAL_BYTES", 3 << 20)
    inner = _zip(tmp_path / "ichki.zip", [(f"{i}.xlsx", 1 << 20) for i in range(5)])
    outer = tmp_path / "tashqi.zip"
    with zipfile.ZipFile(outer, "w") as z:
        z.write(inner, "ichki.zip")
    with pytest.raises(archives.TooBig):
        archives.extract_excel(outer, tmp_path / "s")


def test_declared_size_is_not_trusted(tmp_path, monkeypatch):
    """Sarlavhada hajm kichik yozilgan bo'lsa ham — o'qilgani chegaradan oshsa to'xtaydi."""
    from . import archives
    calls = []

    def lying_members(path):
        def write(target, limit):
            calls.append(limit)
            return archives._copy_capped(__import__("io").BytesIO(b"\0" * (3 << 20)), target, limit)
        yield "x.xlsx", 10, write
    monkeypatch.setattr(archives, "members", lying_members)
    monkeypatch.setattr(archives, "MAX_EXCEL_BYTES", 1 << 20)
    assert archives.extract_excel(tmp_path / "a.zip", tmp_path / "s") == []
    assert calls == [1 << 20] and not any((tmp_path / "s").iterdir())


def test_xlsx_bomb_is_rejected_before_opening(tmp_path, monkeypatch):
    from core import convert
    monkeypatch.setattr(convert, "MAX_XLSX_UNPACKED", 2 << 20)
    bomb = _zip(tmp_path / "bomba.xlsx", [("xl/worksheets/sheet1.xml", 5 << 20)])
    with pytest.raises(convert.ConversionError, match="жуда катта"):
        convert.ensure_xlsx(bomb, tmp_path)
    ok = tmp_path / "ok.xlsx"
    Workbook().save(ok)
    assert convert.ensure_xlsx(ok, tmp_path) == ok


def test_xxe_entity_is_not_resolved(tmp_path):
    """XML ichidagi tashqi «entity» (XXE) — mahalliy faylni o'qib, katakka qo'yib bo'lmasin."""
    import defusedxml
    from openpyxl import load_workbook
    secret = tmp_path / "maxfiy.txt"
    secret.write_text("SIR-12345", encoding="utf-8")
    ok = tmp_path / "ok.xlsx"
    wb = Workbook()
    wb.active["A1"] = "salom"
    wb.save(ok)
    bad = tmp_path / "xxe.xlsx"
    with zipfile.ZipFile(ok) as zi, zipfile.ZipFile(bad, "w") as zo:
        for item in zi.infolist():
            data = zi.read(item)
            if item.filename == "xl/sharedStrings.xml":
                data = data.replace(b"<sst", f'<!DOCTYPE sst [<!ENTITY e SYSTEM "{secret.as_uri()}">]><sst'.encode(), 1)
                data = data.replace(b"salom", b"&e;")
            zo.writestr(item, data)
    try:
        wb = load_workbook(bad)
    except (defusedxml.DefusedXmlException, ValueError, SyntaxError):
        return  # rad etildi — yaxshi
    assert "SIR-12345" not in str(wb.active["A1"].value)


# ---------- sirlar tashqi dasturga o'tmaydi ----------
def test_tool_env_has_no_secrets(monkeypatch):
    from core.convert import tool_env
    for k in ("DJANGO_SECRET_KEY", "DATABASE_URL", "EMAIL_HOST_PASSWORD", "REDIS_URL"):
        monkeypatch.setenv(k, "maxfiy")
    monkeypatch.setenv("PATH", "/usr/bin")
    env = tool_env()
    assert "maxfiy" not in env.values() and env.get("PATH") == "/usr/bin"


# ---------- ro'yxatdan o'tish va kvotalar ----------
from django.core.files.uploadedfile import SimpleUploadedFile  # noqa: E402
from django.test import TestCase, override_settings  # noqa: E402
from django.urls import reverse  # noqa: E402

from accounts.models import User  # noqa: E402
from panel.models import Event  # noqa: E402


@override_settings(MEDIA_ROOT=__import__("tempfile").mkdtemp(prefix="narxlash_sec_"))
class SignupAndQuotaTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("ali@test.uz", "Kuchli-parol-2026", first_name="Ali", last_name="V")

    @override_settings(NARXLASH_SIGNUP="closed")
    def test_closed_signup_rejects_registration(self):
        r = self.client.post(reverse("register"), {"first_name": "A", "last_name": "B", "email": "yangi@test.uz",
                                                   "password1": "Kuchli-parol-2026", "password2": "Kuchli-parol-2026"})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(User.objects.filter(email="yangi@test.uz").exists())
        self.assertNotContains(self.client.get(reverse("login")), reverse("register"))

    def test_open_signup_shows_link(self):
        self.assertContains(self.client.get(reverse("login")), reverse("register"))

    @override_settings(NARXLASH_UPLOADS_PER_HOUR=2)
    def test_upload_rate_limit(self):
        self.client.force_login(self.user)
        for _ in range(2):
            Event.objects.create(kind=Event.Kind.UPLOAD, actor=self.user)
        f = SimpleUploadedFile("a.xlsx", b"x")
        r = self.client.post(reverse("project_create"), {"files": [f], "paths": ["P/a.xlsx"]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("соатда", r.json()["error"])

    @override_settings(NARXLASH_USER_QUOTA_MB=1)
    def test_disk_quota(self):
        from jobs.models import Project
        from jobs.services import UploadError, check_quota
        p = Project.objects.create(owner=self.user, name="P")
        p.dir.mkdir(parents=True, exist_ok=True)
        (p.dir / "katta.bin").write_bytes(b"\0" * (2 << 20))
        try:
            with self.assertRaisesRegex(UploadError, "жойингиз тўлди"):
                check_quota(self.user)
            self.user.is_staff = True
            check_quota(self.user)  # xodim — cheklovsiz
        finally:
            import shutil
            shutil.rmtree(p.dir, ignore_errors=True)


def test_preview_head_cuts_rows_and_closes_table():
    from .preview import head
    rows = "".join(f"<tr><th>{i}</th><td>x</td></tr>" for i in range(1, 1001))
    data = {"rows": 1000, "html": "<style></style><table><thead><tr><th></th></tr></thead><tbody>" + rows + "</tbody></table>"}
    h = head(data, 400)
    assert h["partial"] and h["shown"] == 400
    assert h["html"].count("<tr>") == 401 and h["html"].endswith("</tr></tbody></table>")
    assert head({"rows": 10, "html": "x"}, 400) == {"rows": 10, "html": "x"}
