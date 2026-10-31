"""Ofyorta (shartnoma smetasi) o'quvchisi: Excel -> bo'limlar, ishlar, resurslar va svodka foizlari.

Ofyortada 4 xil qator bor (A ustuni — tartib raqami):

* **bo'lim** — A bo'sh, faqat nom («СТОЙКА БАСКЕТБОЛЬНАЯ В СБОРЕ», «на АС Блок-1»);
* **ish** — A butun son (40), ostida «40.1, 40.2 …» resurslari bor; E — loyiha bo'yicha hajm (0,639 т);
* **resurs** — A «40.1»: ishning resursi; E — 1 birlik ishga me'yor, F = E × ish hajmi, I = G × F;
* **mustaqil resurs** — A butun son, ostida resurs yo'q (alohida material/uskuna); I = G × E.

Summa J..O ustunlarning bittasiga tushadi (turi): J ish haqi, K mashina, L material, M pudratchi uskunasi,
N buyurtmachi uskunasi, O kabel. Pastdagi svodka — shu ustunlar yig'indisidan foizlar bilan (``calc.chain``).

Texnazor faqat **ish** va **mustaqil resurs** hajmini belgilaydi; resurslar ishga ergashadi — shuning uchun
har bir ish uchun «1 birlik ishning narxi turlar bo'yicha» (``unit_cost``) oldindan hisoblanadi.
"""
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional

from openpyxl import load_workbook

CATEGORIES = ("J", "K", "L", "M", "N", "O")
CATEGORY_LABELS = {"J": "Иш ҳақи", "K": "Машина", "L": "Материал", "M": "Ускуна (пудратчи)",
                   "N": "Ускуна (буюртмачи)", "O": "Кабель"}
# svodka foizlari: faylda topilmasa — shu standart (40-maktab ofyortasidagidek)
DEFAULT_RATES = {"tr_material": 0.03, "tr_cable": 0.015, "overhead": 0.17, "tr_eq_customer": 0.02,
                 "tr_eq_contractor": 0.02, "insurance": 0.0032, "vat": 0.12}

# sarlavha matni -> ustun ma'nosi (kichik harfda, birinchi mos keladigan)
_HEADERS = [
    ("code", ("шифр",)), ("name", ("наименование",)), ("unit", ("единица",)),
    ("qty", ("на единицу",)), ("total", ("по проектным",)), ("price", ("цена подр",)),
    ("price_customer", ("цена заказ",)), ("amount", ("сумма",)),
    ("J", ("чел-ч",)), ("K", ("маш-ч",)), ("L", ("матриал", "материал")),
    ("M", ("жихоз подр", "оборудование подр")), ("N", ("жихоз заказ", "оборудование заказ")), ("O", ("кабель",)),
]
_FALLBACK = {"num": 1, "code": 2, "name": 3, "unit": 4, "qty": 5, "total": 6, "price": 7, "price_customer": 8,
             "amount": 9, "J": 10, "K": 11, "L": 12, "M": 13, "N": 14, "O": 15}
_NUM_TOP = re.compile(r"^\d+$")
NO_SECTION = "Бўлим номи кўрсатилмаган"
_NUM_SUB = re.compile(r"^(\d+)\.\d+$")


class OfyortaError(Exception):
    pass


@dataclass
class Row:
    excel_row: int
    kind: str                      # section | work | res | free | extra
    name: str
    num: str = ""
    code: str = ""
    unit: str = ""
    qty: float = 0.0               # ish/mustaqil: hajm; resurs: me'yor (1 birlik ishga)
    total_qty: float = 0.0         # resurs: jami (F); ish/mustaqil: hajm
    price: float = 0.0             # G — pudratchi narxi
    price_customer: float = 0.0    # H
    amount: float = 0.0            # shartnoma summasi (biz hisoblagan)
    file_amount: Optional[float] = None  # faylning I ustuni (tekshiruv uchun)
    category: str = ""             # J..O
    parent: Optional[int] = None   # resurs -> ishning indeksi (rows ichida)
    section: Optional[int] = None  # bo'lim indeksi (sections ichida)
    unit_cost: Dict[str, float] = field(default_factory=dict)  # ish/mustaqil: 1 birlik narxi turlar bo'yicha


@dataclass
class Ofyorta:
    sheet: str
    title: str = ""
    object_type: str = ""
    sections: List[str] = field(default_factory=list)
    blocks: set = field(default_factory=set)  # blok sarlavhasi bo'lgan bo'limlar (indeksi) — qizil yoki «НА … / БЛОК-1»
    rows: List[Row] = field(default_factory=list)
    rates: Dict[str, float] = field(default_factory=lambda: dict(DEFAULT_RATES))
    file_totals: Dict[str, float] = field(default_factory=dict)  # svodkadagi summalar (taqqoslash uchun)
    warnings: List[dict] = field(default_factory=list)  # {"row": Excel qatori, "title", "text"} — oddiy tilda
    amount_diff: float = 0.0  # faylning I ustuni − biz hisoblagan (to'g'ridan-to'g'ri, ofyortadagi xatolar)

    @property
    def items(self) -> List[Row]:
        """Texnazor belgilaydigan qatorlar: ish, mustaqil resurs, qo'shimcha."""
        return [r for r in self.rows if r.kind in ("work", "free", "extra")]


def _num(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace("\xa0", "").replace(" ", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def pick_sheet(wb) -> str:
    for name in wb.sheetnames:
        low = name.lower().replace("ё", "е")
        if "офер" in low or "oferta" in low or "ofyorta" in low or "оферт" in low:
            return name
    return wb.sheetnames[0]


def _find_header(rows: List[tuple]) -> Optional[int]:
    for i, r in enumerate(rows[:60]):
        if any("наименование" in _text(v).lower() for v in r):
            return i
    return None


def _columns(rows: List[tuple], h: int) -> Dict[str, int]:
    """Sarlavhadan (2–3 qator) ustunlar; topilmagani — standart joyi."""
    cols = {}
    for r in rows[h:h + 3]:
        for j, v in enumerate(r):
            t = _text(v).lower()
            if not t:
                continue
            for key, needles in _HEADERS:
                if key not in cols and any(n in t for n in needles):
                    cols[key] = j
                    break
    out = {k: v - 1 for k, v in _FALLBACK.items()}
    out.update(cols)
    out["num"] = 0
    return out


_SVODKA_START = re.compile(r"ISH\s*HAQQI|ИШ\s*ҲАҚИ|ЗАРАБОТНАЯ\s+ПЛАТА", re.I)
_PCT = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")


def parse(path: Path, sheet: Optional[str] = None, progress: Callable[[int, int], None] = None) -> Ofyorta:
    """``progress(o'qilgan, jami)`` — har 500 qatorda (jarayon oynasi uchun)."""
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        name = sheet if sheet and sheet in wb.sheetnames else pick_sheet(wb)
        ws = wb[name]
        rows = []
        total = ws.max_row or 0
        for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            rows.append(r)
            if progress and i % 500 == 0:
                progress(i, total)
    finally:
        wb.close()
    of = Ofyorta(sheet=name)
    h = _find_header(rows)
    if h is None:
        raise OfyortaError("Офёрта жадвали топилмади: «Наименование работ…» сарлавҳали устун йўқ")
    cols = _columns(rows, h)
    # obyekt nomi va turi — sarlavhadan yuqoridagi matnlar («obyekt nomi» / «obyekt turi» ostida)
    texts = [_text(v) for r in rows[:h] for v in r if _text(v)]
    for i, t in enumerate(texts):
        low = t.lower()
        if low in ("obyekt nomi", "объект номи") and i > 0:
            of.title = texts[i - 1]
        if low in ("obyekt turi", "объект тури") and i > 0:
            of.object_type = texts[i - 1]
    if not of.title:
        of.title = next((t for t in texts if len(t) > 10 and t.upper() == t and "OFYORTA" not in t.upper()), "")

    get = lambda r, key: r[cols[key]] if cols[key] < len(r) else None  # noqa: E731
    start = h + 1
    # sarlavha ostidagi «1 2 3 …» raqamlar qatori va ikkinchi sarlavha qatori o'tkaziladi
    while start < len(rows) and start < h + 4:
        r = rows[start]
        if _text(get(r, "name")) in ("", "3") or _text(get(r, "name")).isdigit():
            start += 1
        else:
            break
    # svodka: «JAMI:» qatori yoki «… ISH HAQQI» (svodka qatorlarida ham A da raqam bo'ladi: 1, 2, 3 …)
    end = len(rows)
    for i in range(start, len(rows)):
        nm = _text(get(rows[i], "name"))
        if nm.upper().rstrip(": ") == "JAMI" or _SVODKA_START.search(nm):
            end = i
            break

    def next_num(i: int) -> str:
        for j in range(i + 1, end):
            n, t = _text(get(rows[j], "num")), _text(get(rows[j], "name"))
            if n or t:
                return n
        return ""

    def is_numbered_header(r, num: str, i: int) -> bool:
        """«1 | ДЕМОНТАЖ» — A da raqam bor, lekin birligi, miqdori, narxi, summasi yo'q va ostida «1.1» resurs yo'q:
        bu ish emas, bo'lim sarlavhasi (ba'zi ofyortalarda sarlavha ham raqamlanadi)."""
        if not _NUM_TOP.match(num) or _text(get(r, "unit")):
            return False
        if any(_num(get(r, k)) for k in ("qty", "total", "price", "price_customer", "amount")):
            return False
        return not next_num(i).startswith(num + ".")

    section = None
    last_work = None
    for i in range(start, end):
        r = rows[i]
        num, nm = _text(get(r, "num")), _text(get(r, "name"))
        if not num and not nm:
            continue
        if not num or not (_NUM_TOP.match(num) or _NUM_SUB.match(num)) or (nm and is_numbered_header(r, num, i)):
            if nm:
                of.sections.append(nm)
                section = len(of.sections) - 1
                of.rows.append(Row(excel_row=i + 1, kind="section", name=nm, section=section))
                last_work = None
            continue
        if section is None:  # birinchi sarlavhadan oldingi (yoki sarlavhasiz ofyortadagi) ishlar — yo'qolib qolmasin
            of.sections.append(NO_SECTION)
            section = len(of.sections) - 1
            of.rows.append(Row(excel_row=i + 1, kind="section", name=NO_SECTION, section=section))
            of.warnings.append({"row": i + 1, "title": "Бўлим сарлавҳаси йўқ ишлар бор",
                                "text": f"{i + 1}-қатордан бошлаб ишлар тепасида бўлим номи (А устуни бўш, C да ном) йўқ. "
                                        f"Тизим уларни «{NO_SECTION}» бўлимига олди — ишлар ва пул ҳисоби тўлиқ."})
        qty = _num(get(r, "qty")) or 0.0
        g, hc = _num(get(r, "price")) or 0.0, _num(get(r, "price_customer")) or 0.0
        # buyurtmachi uskunasida narx H da (G bo'sh): I = H × E
        row = Row(excel_row=i + 1, kind="res" if _NUM_SUB.match(num) else "top", name=nm, num=num,
                  code=_text(get(r, "code")), unit=_text(get(r, "unit")), qty=qty,
                  price=g or hc, price_customer=hc, file_amount=_num(get(r, "amount")), section=section)
        cat = ""
        for c in CATEGORIES:
            v = _num(get(r, c))
            if v and row.file_amount and abs(v - row.file_amount) <= max(0.01, abs(row.file_amount) * 1e-9):
                cat = c
                break
        row.category = cat or _guess_category(row.unit)
        if row.kind == "res":
            parent = last_work if last_work is not None and _NUM_SUB.match(num).group(1) == of.rows[last_work].num else None
            if parent is None:
                of.warnings.append({"row": i + 1, "title": f"«{num}» ресурсининг иши топилмади",
                                    "text": f"Рақами «{num}», лекин тепасида {num.split('.')[0]}-иш йўқ. Тизим уни алоҳида "
                                            "материал деб олди: ҳажмини ўзингиз белгилайсиз."})
                row.kind = "top"
            else:
                row.parent = parent
                f = _num(get(r, "total"))
                row.total_qty = f if f is not None else qty * of.rows[parent].qty
                row.amount = row.price * row.total_qty
                of.rows[parent].kind = "work"
        if row.kind == "top":
            row.total_qty = qty
            row.amount = row.price * qty
            last_work = len(of.rows)
        of.rows.append(row)
    for row in of.rows:
        if row.kind == "top":
            row.kind = "free"
    _unit_costs(of)
    _check_amounts(of)
    _svodka(of, rows[end:], cols, end)
    _mark_blocks(of, path, name, cols["name"])
    return of


BLOCK_RE = re.compile(r"^\s*НА\s+|\bБЛОК[\s№-]*\d", re.I)  # «НА АС БЛОК-1», «НА ЭО», «АС БЛОК-2»


def _red(cell) -> bool:
    c = cell.font.color if cell.font is not None else None
    return bool(c is not None and c.type == "rgb" and isinstance(c.rgb, str) and c.rgb.upper().endswith("FF0000"))


def _mark_blocks(of: Ofyorta, path: Path, sheet: str, name_col: int):
    """Qaysi sarlavha blok: nomi qizil (uning ofyortalarida shunday) yoki nomida «НА …» / «БЛОК-1». Faqat nom
    ustuni o'qiladi (tez). Tizimda yasalgan ofyortada bu taxmin emas — tuzilish alohida saqlanadi (``MyOfyorta``)."""
    heads = {r.excel_row: r.section for r in of.rows if r.kind == "section"}
    red = set()
    try:
        wb = load_workbook(path, read_only=True, keep_links=False)
        try:
            for (cell,) in wb[sheet].iter_rows(min_col=name_col + 1, max_col=name_col + 1):
                if getattr(cell, "row", None) in heads and _red(cell):
                    red.add(heads[cell.row])
        finally:
            wb.close()
    except Exception:  # noqa: BLE001 — rang o'qilmasa ham nomi bo'yicha aniqlanadi
        pass
    of.blocks = {i for i, t in enumerate(of.sections) if i in red or BLOCK_RE.search(t)}


def _check_amounts(of: Ofyorta):
    """Faylning I ustuni biz hisoblagan summadan farq qilsa — ofyortadagi xato (masalan, hajmi 0 bo'lgan ishning
    resurslarida I = G × E yozilgan, G × F o'rniga). Ish bo'yicha guruhlab ogohlantiriladi; hisob — to'g'ri formula bilan."""
    by_work: Dict[int, list] = {}
    for i, row in enumerate(of.rows):
        if row.kind not in ("res", "free") or row.file_amount is None:
            continue
        if abs(row.file_amount - row.amount) > max(1.0, abs(row.amount) * 1e-6):
            by_work.setdefault(row.parent if row.kind == "res" else i, []).append(row)
    for key, rows in by_work.items():
        head = of.rows[key]
        diff = sum(r.file_amount - r.amount for r in rows)
        name = f"{head.num}-иш «{head.name[:70].strip()}»" if head.kind == "work" else f"«{head.name[:70].strip()}»"
        if head.kind == "work" and not head.qty:
            text = (f"Шартномада бу ишнинг ҳажми 0, лекин Excel'да унинг {len(rows)} та ресурсига жами {money(diff)} сўм пул "
                    "ёзилган. Сабаби — формула хатоси: ресурс суммаси «нарх × жами ҳажм» ўрнига «нарх × меъёр» бўлиб "
                    "ҳисобланган. Ҳажми 0 бўлган иш учун пул бўлмаслиги керак, шунинг учун тизим бу суммани ҳисобга олмайди. "
                    "Шу сабаб тизимдаги акт Excel'дагидан шунча (устамалари билан бироз кўпроқ) кам чиқади. "
                    "Пудратчи ёки смета тузувчи билан аниқлаштиринг.")
        else:
            text = (f"Excel'да {len(rows)} та қатор суммаси «нарх × миқдор»га тўғри келмайди, фарқ {money(diff)} сўм. "
                    "Тизим тўғри формула билан ҳисоблайди. Excel'даги формулани текширинг.")
        of.warnings.append({"row": head.excel_row, "title": f"{name}: Excel'да ортиқча {money(diff)} сўм", "text": text})
    of.amount_diff = sum(r.file_amount - r.amount for rows in by_work.values() for r in rows)


def money(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def _guess_category(unit: str) -> str:
    u = unit.upper()
    if "ЧЕЛ" in u:
        return "J"
    if "МАШ" in u:
        return "K"
    return "L"


def _unit_costs(of: Ofyorta):
    """Ish: 1 birlik ishning narxi turlar bo'yicha (resurslar yig'indisi / hajm). Mustaqil: narxning o'zi."""
    for i, row in enumerate(of.rows):
        if row.kind == "free":
            row.unit_cost = {row.category: row.price} if row.price else {}
        elif row.kind == "work":
            row.unit_cost = {}
    for row in of.rows:
        if row.kind == "res" and row.amount:
            work = of.rows[row.parent]
            if work.qty:
                work.unit_cost[row.category] = work.unit_cost.get(row.category, 0.0) + row.amount / work.qty
            else:
                of.warnings.append({"row": work.excel_row, "title": f"{work.num}-ишнинг ҳажми 0",
                                    "text": f"Ресурсларида {money(row.amount)} сўм бор, лекин иш ҳажми 0 — бу пул ҳисобга олинмади."})
    for row in of.rows:
        if row.kind == "work":
            row.amount = sum(v * row.qty for v in row.unit_cost.values())


def _svodka(of: Ofyorta, rows: List[tuple], cols: Dict[str, int], offset: int):
    """Pastdagi svodka: foizlar (transport, boshqa xarajatlar, sug'urta, QQS), qo'shimcha qatorlar va jami summalar."""
    amount_col = cols["amount"]
    pending = None  # oxirgi ko'rilgan bazaviy qator — keyingi «transport …%» unga tegishli
    after_insurance = False
    for k, r in enumerate(rows):
        label = _text(r[cols["name"]]) if cols["name"] < len(r) else ""
        if not label:
            continue
        up = label.upper().replace("`", "'").replace("’", "'")
        value = _num(r[amount_col]) if amount_col < len(r) else None
        m = _PCT.search(label)
        pct = float(m.group(1).replace(",", ".")) / 100 if m else None
        if "ISH HAQQI" in up:
            pending = "J"
        elif "MASHINA" in up:
            pending = "K"
        elif "MATERIAL" in up and "TRANSPORT" not in up:
            pending = "material"
        elif "KABEL" in up and "TRANSPORT" not in up:
            pending = "cable"
        elif "USKUNALAR" in up and "BUYURTMACHI" in up and "TRANSPORT" not in up:
            pending = "eq_customer"
        elif "USKUNALAR" in up and "PUDRATCHI" in up and "TRANSPORT" not in up:
            pending = "eq_contractor"
        elif "TRANSPORT" in up and pct is not None:
            key = {"material": "tr_material", "cable": "tr_cable", "eq_customer": "tr_eq_customer",
                   "eq_contractor": "tr_eq_contractor"}.get(pending)
            if key:
                of.rates[key] = pct
        elif "BOSHQA XARAJAT" in up and pct is not None:
            of.rates["overhead"] = pct
        elif ("SUG'URTA" in up or "СУҒУРТА" in up) and pct is not None and not after_insurance:
            of.rates["insurance"] = pct
            after_insurance = True
        elif "QQS" in up and pct is not None and "BILAN" not in up and "SIZ" not in up and not of.file_totals.get("_vat_seen"):
            of.rates["vat"] = pct
            of.file_totals["_vat_seen"] = 1
        elif after_insurance and "JAMI" not in up and "QQS" not in up and "BUYURTMACHI" not in up \
                and "PUDRATCHI" not in up and "SUG'URTA" not in up and value:
            # sug'urtadan keyin qo'shilgan qator (masalan «Художественный элемент», 1024,7 м²) — alohida ish:
            # sug'urta va transport olinmaydi, faqat QQS (svodkadagidek)
            qty = _num(r[cols["qty"]]) if cols["qty"] < len(r) else None
            unit = _text(r[cols["unit"]]) if cols["unit"] < len(r) else ""
            qty = qty if qty else 1.0
            of.rows.append(Row(excel_row=offset + k + 1, kind="extra", name=label, unit=unit, qty=qty, total_qty=qty,
                               price=value / qty, amount=value, file_amount=value, unit_cost={"X": value / qty}))
        if value is not None:
            if "OB'EKT QIYMATI" in up and "BILAN" in up:
                of.file_totals["total_vat"] = value
            elif "PUDRATCHI XARAJATLARI" in up and "BILAN" in up:
                of.file_totals["contractor_vat"] = value
            elif "BUYURTMACHI JIHOZLAR" in up and "BILAN" in up:
                of.file_totals["customer_vat"] = value
    of.file_totals.pop("_vat_seen", None)
