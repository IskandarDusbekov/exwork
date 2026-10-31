"""Texnazor: ofyortani o'qish, svodka zanjiri, oyma-oy hajmlar, qoldiq, ruxsatlar."""
import json
from pathlib import Path
import shutil
import tempfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from accounts.models import User

from . import services, tasks
from .calc import chain
from .models import Correction, Done, Item, Period, SiteObject
from .ofyorta import parse

TMP_MEDIA = tempfile.mkdtemp(prefix="narxlash_tn_")


def build_ofyorta(path, broken_zero_work=False):
    """Haqiqiy ofyorta tuzilishi: sarlavha, bo'limlar, ish + resurslar, mustaqil resurs, buyurtmachi uskunasi, svodka."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ОФЁРТА"
    ws["C2"], ws["C3"] = "CHILONZOR TUMAN 5-MAKTAB TA'MIRI", "obyekt nomi"
    ws["C4"], ws["C5"] = "MUKAMMAL TA'MIRLASH", "obyekt turi"
    head = ["№ п/п", "Шифр номера нормативов", "Наименование работ и затрат", "Единица измерения", "Количество"]
    for j, v in enumerate(head, start=1):
        ws.cell(row=8, column=j, value=v)
    for j, v in enumerate(["на единицу измерения", "по проектным данным", "Цена подрятчик", "Цена заказчик", "Сумма",
                           "ЧЕЛ-Ч.", "МАШ-Ч.", "МАТРИАЛ", "ЖИХОЗ ПОДРЯТЧИК", "ЖИХОЗ ЗАКАЗЧИК", "КАБЕЛЬ"], start=5):
        ws.cell(row=9, column=j, value=v)
    for j in range(1, 16):
        ws.cell(row=10, column=j, value=j)
    r = 11

    def put(*vals):
        nonlocal r
        for j, v in enumerate(vals, start=1):
            if v is not None:
                ws.cell(row=r, column=j, value=v)
        r += 1
        return r - 1

    put(None, None, "на АС Блок-1")
    put(None, None, "ФУНДАМЕНТ")
    w = put(1, "Е06-01-001-1", "БЕТОН ФУНДАМЕНТОВ", "1 М3", 10)
    # resurs: me'yor (E), F = E × ish hajmi, I = G × F, summa — J/K/L ustunlardan biriga
    for num, code, name, unit, norm, price, col in (("1.1", "1", "ЗАТРАТЫ ТРУДА РАБОЧИХ", "ЧЕЛ-Ч.", 5, 1000, 10),
                                                   ("1.2", "715", "КРАН", "МАШ-Ч.", 0.5, 20000, 11),
                                                   ("1.3", "22006", "БЕТОН", "М3", 1.02, 500000, 12)):
        row = put(num, code, name, unit, norm, norm * 10, price, None, price * norm * 10)
        ws.cell(row=row, column=col, value=price * norm * 10)
    put(2, "30032", "ШВЕЛЛЕР №16", "ТН", 2, None, 9600000, None, 19200000, None, None, 19200000)
    put(None, None, "УСКУНАЛАР")
    put(3, None, "ТЕПЛОВОЙ НАСОС", "КОМПЛЕКТ", 4, None, None, 1000000, 4000000, None, None, None, None, 4000000)
    if broken_zero_work:
        z = put(4, "Е10-01-082-2", "УКЛАДКА ПРОГОНОВ", "1 М3", 0)
        put("4.1", "1", "ЗАТРАТЫ ТРУДА", "ЧЕЛ-Ч.", 15, 0, 1000, None, 15000, 15000)  # I = G × E (xato)
    # svodka (A da ham raqam bor)
    rows = [(None, "JAMI:"), (1, "QURUVCHI-ISHCHILAR ISH HAQQI"), (2, "QURILISH MASHINALARI VA MEXANIZMLARI"),
            (3, "QURILISH MATERIALLARI UCHUN JAMI: /PUDRATCHI/"), (4, "TRANSPORT XARAJATLARI 3%"),
            (5, "KABEL MAHSULOTLARI  /PUDRATCHI/"), (6, "TRANSPORT XARAJATLARI SOTUVCHIDAN 1,5%"),
            (7, "JAMI TO'G'RIDAN TO'G'RI XARAJATLAR"), (8, "PUDRATCHI BOSHQA XARAJATLARI 17%"),
            (10, "USKUNALAR  /BUYURTMACHI/"), (11, "TRANSPORT XARAJATI 2%"), (13, "USKUNALAR  /PUDRATCHI/"),
            (14, "TRANSPORT XARAJATLARI 2%"), (17, "OBYEKTNI SUG'URTALASH 0,32%")]
    for a, label in rows:
        put(a, None, label)
    row = put(18, None, "ХУДОЖЕСТВЕННЫЙ ЭЛЕМЕНТ", "М2", 10)
    ws.cell(row=row, column=9, value=500000)
    put(19, None, "JAMI OB'EKT QIYMATI QQS SIZ")
    put(20, None, "QQS 12%")
    wb.save(path)
    return path


# ======== sof hisob (Django'siz) ========
def test_parse_structure_and_unit_costs(tmp_path):
    of = parse(build_ofyorta(tmp_path / "o.xlsx"))
    kinds = [r.kind for r in of.rows]
    assert kinds.count("section") == 3 and kinds.count("work") == 1 and kinds.count("res") == 3
    assert kinds.count("free") == 2 and kinds.count("extra") == 1
    work = next(r for r in of.rows if r.kind == "work")
    # 1 m3 betonning narxi: mehnat 5×1000, kran 0,5×20000, beton 1,02×500000
    assert work.unit_cost == pytest.approx({"J": 5000, "K": 10000, "L": 510000})
    pump = next(r for r in of.rows if r.name == "ТЕПЛОВОЙ НАСОС")
    assert (pump.category, pump.price, pump.amount) == ("N", 1000000, 4000000)  # narx H da
    assert of.title == "CHILONZOR TUMAN 5-MAKTAB TA'MIRI"
    assert of.rates["tr_material"] == 0.03 and of.rates["overhead"] == 0.17 and of.rates["vat"] == 0.12
    extra = next(r for r in of.rows if r.kind == "extra")
    assert (extra.qty, extra.unit, extra.amount) == (10, "М2", 500000)
    assert not of.warnings


def test_zero_volume_work_with_wrong_formula_is_reported(tmp_path):
    of = parse(build_ofyorta(tmp_path / "o.xlsx", broken_zero_work=True))
    w = of.warnings[0]
    assert len(of.warnings) == 1 and "ҳажми 0" in w["text"] and "15 000" in w["title"] and w["row"]


def test_money_words_russian():
    assert services.money_words(18390330726.71) == ("Восемнадцать миллиардов триста девяносто миллионов триста тридцать "
                                                   "тысяч семьсот двадцать семь сум")
    assert services.money_words(2002) == "Две тысячи два сум"
    assert services.money_words(11000) == "Одиннадцать тысяч сум"


def test_chain_matches_svodka_formula():
    rates = {"tr_material": 0.03, "tr_cable": 0.015, "overhead": 0.17, "tr_eq_customer": 0.02,
             "tr_eq_contractor": 0.02, "insurance": 0.0032, "vat": 0.12}
    c = chain({"J": 100, "K": 50, "L": 1000, "O": 200, "N": 400, "M": 300, "X": 10}, rates)
    direct = 100 + 50 + 1000 * 1.03 + 200 * 1.015
    base = direct * 1.17 + 400 * 1.02 + 300 * 1.02
    total = (base * 1.0032 + 10) * 1.12
    assert c["direct"] == pytest.approx(direct) and c["total"] == pytest.approx(total)
    assert c["customer"] == pytest.approx(400 * 1.02 * 1.0032 * 1.12)
    assert c["contractor"] == pytest.approx(total - c["customer"])


# ======== bazaga yozish, oylar, sahifalar ========
@override_settings(MEDIA_ROOT=TMP_MEDIA)
class TexnazorFlowTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.user = User.objects.create_user("tn@test.uz", "Kuchli-parol-2026", first_name="T", last_name="N")
        self.client.force_login(self.user)
        path = build_ofyorta(f"{TMP_MEDIA}/src.xlsx")
        with open(path, "rb") as f:
            data = f.read()
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("texnazor:create"), {"ofyorta": SimpleUploadedFile("5-maktab.xlsx", data),
                                                              "address": "Чилонзор"})
        self.obj = SiteObject.objects.get()
        self.assertRedirects(r, reverse("texnazor:object", args=[self.obj.pk]), fetch_redirect_response=False)
        tasks.run_now(self.obj.pk)  # testda — darhol (thread o'rniga)
        self.obj.refresh_from_db()
        self.work = Item.objects.get(kind="work")
        self.free = Item.objects.get(name="ШВЕЛЛЕР №16")

    def _post(self, name, body, **kw):
        return self.client.post(reverse(f"texnazor:{name}", kwargs={"obj_id": self.obj.pk, **kw}), json.dumps(body),
                                content_type="application/json")

    def test_f3_two_months_cumulative_and_settings(self):
        """Ф-3: шу ой = Ф-2 жами (пудратчи), йиғма — олдинги ойлар + шу ой; рақам, сана, давр «с» — олдингисидан."""
        from io import BytesIO
        from openpyxl import load_workbook
        add = lambda m: self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": m})  # noqa: E731
        add(10)
        octo = Period.objects.get()
        self._post("fill_items", {"period": octo.pk, "items": [self.work.pk], "pct": "40"})
        self.client.post(reverse("texnazor:f3_period", args=[self.obj.pk, octo.pk]), {"date": "2026-10-28", "start": "2026-09-30"})
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, octo.pk]), {"closed": "1"})
        add(11)
        nov = Period.objects.get(month=11)
        self._post("fill_items", {"period": nov.pk, "items": [self.work.pk], "pct": "100"})
        r = self.client.post(reverse("texnazor:f3_settings", args=[self.obj.pk]),
                             {"period": nov.pk, "customer": "Халқ таълими", "contractor": "ООО Қурувчи", "customer_inn": "312 081 448",
                              "contractor_inn": "305418313", "contractor_account": "2020800", "contract_no": "43/Q",
                              "contract_date": "2026-04-14", "contract_id": "2501", "contractor_head": "Ш.Мамараззаков",
                              "texnazor": "А.Шаймарданов", "year_plan": ""})
        self.assertEqual(r.status_code, 302)
        self.obj.refresh_from_db()
        self.assertEqual(self.obj.contract, "№ 43/Q 14.04.2026")
        self.assertEqual(services.f3_missing(self.obj), [])  # ҳаммаси эслаб қолинди
        nov = Period.objects.get(pk=nov.pk)
        d = services.f3_data(nov)
        t = services.period_totals(nov)
        vat = self.obj.rates["vat"]
        self.assertEqual(d["meta"]["num"], 2)
        self.assertEqual(str(d["meta"]["start"]), "2026-10-28")  # олдинги Ф-3 тузилган санаси
        self.assertAlmostEqual(d["this"]["total"], t["this"]["contractor"])
        self.assertAlmostEqual(d["this"]["novat"] * (1 + vat), t["this"]["contractor"])
        self.assertAlmostEqual(d["cum"]["total"], t["before"]["contractor"] + t["this"]["contractor"])
        self.assertAlmostEqual(d["year"]["total"], d["contract"]["total"])  # жорий йил киритилмаган — бутун шартнома
        self.assertGreater(d["before"]["total"], 0)
        page = self.client.get(reverse("texnazor:object", args=[self.obj.pk]), {"p": nov.pk}).content.decode()
        self.assertIn("Ф-3 №2", page)
        self.assertIn(d["words"], page)
        dl = self.client.get(reverse("texnazor:f3", args=[self.obj.pk, nov.pk]))
        self.assertEqual(dl.status_code, 200)
        ws = load_workbook(BytesIO(dl.content), data_only=True).active
        cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        self.assertIn("Заказчик: Халқ таълими", cells)
        self.assertIn("СТИР  312 081 448", cells)
        self.assertIn("Договор от 14 апреля 2026 г. № 43/Q\nИД 2501", cells)
        self.assertIn("Руководитель подрядчика  ______________  Ш.Мамараззаков", cells)
        self.assertNotIn("МФО: ", cells)  # бўш майдон қатори чиқмайди
        self.assertIn(d["words"], cells)
        self.assertIn("А.Шаймарданов", cells)
        tot = cells["ВСЕГО С НДС"].row
        self.assertAlmostEqual(ws[f"P{tot}"].value, d["this"]["total"] / 1000, places=6)
        self.assertAlmostEqual(ws[f"J{tot}"].value, d["cum"]["total"] / 1000, places=6)
        # жорий йилга алоҳида сумма
        self.client.post(reverse("texnazor:f3_settings", args=[self.obj.pk]), {"period": nov.pk, "year_plan": "1 000 000,5"})
        self.assertAlmostEqual(services.f3_data(Period.objects.get(pk=nov.pk))["year"]["total"], 1000000.5)

    def test_actions_logged_and_visible_in_panel(self):
        """Техназордаги амаллар панел журналига тушади, ҳажм ёзилган вақт сақланади; панел — фақат staff."""
        from panel.models import Event
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 10})
        octo = Period.objects.get()
        self.assertTrue(self._post("done", {"period": octo.pk, "item": self.work.pk, "value": "4"}).json()["ok"])
        self.assertIsNotNone(Done.objects.get().updated)
        self.client.get(reverse("texnazor:f2", args=[self.obj.pk, octo.pk]))
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, octo.pk]), {"closed": "1"})
        texts = list(Event.objects.filter(kind=Event.Kind.TEXNAZOR, work_id=self.obj.pk).values_list("detail", flat=True))
        self.assertTrue(Event.objects.filter(kind=Event.Kind.UPLOAD, work_kind="texnazor", work_id=self.obj.pk).exists())
        self.assertTrue(any("Октябрь 2026 ойи қўшилди" in t for t in texts))
        self.assertTrue(any("Ф-2 юклаб олинди" in t for t in texts))
        self.assertTrue(any("ёпилди" in t for t in texts))

        url = reverse("panel:texnazor")
        self.assertEqual(self.client.get(url).status_code, 404)  # oddiy foydalanuvchi panelni ko'rmaydi
        viewer = User.objects.create_user("v@test.uz", "Kuchli-parol-2026", first_name="V", last_name="V", is_staff=True)
        self.client.force_login(viewer)
        r = self.client.get(url)
        self.assertContains(r, "5-MAKTAB")
        self.assertContains(r, "Ф-2 юклаб олинди")
        r = self.client.get(reverse("panel:texnazor_object", args=[self.obj.pk]))
        self.assertContains(r, "Октябрь 2026")
        self.assertContains(r, "БЕТОН ФУНДАМЕНТОВ")  # охирги ёзилган ҳажмлар
        self.assertEqual(self.client.get(reverse("panel:texnazor_file", args=[self.obj.pk, "src"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("panel:user", args=[self.user.pk])).status_code, 200)
        boss = User.objects.create_superuser("b@test.uz", "Kuchli-parol-2026", first_name="B", last_name="B")
        self.client.force_login(boss)
        self.assertEqual(self.client.get(reverse("panel:texnazor_month_file", args=[self.obj.pk, "f3", octo.pk])).status_code, 200)
        self.assertTrue(Event.objects.filter(kind=Event.Kind.ADMIN_DOWNLOAD, actor=boss).exists())

    @override_settings(NARXLASH_UPLOADS_PER_HOUR=1)
    def test_upload_limit_counts_texnazor(self):
        """Texnazor yuklashi ham soatlik cheklovga kiradi (setUp'da bittasi yuklandi)."""
        r = self.client.post(reverse("texnazor:create"), {"ofyorta": SimpleUploadedFile("b.xlsx", b"x")})
        self.assertContains(r, "Бир соатда кўпи билан 1 та")
        self.assertEqual(SiteObject.objects.count(), 1)

    @override_settings(NARXLASH_USER_QUOTA_MB=1)
    def test_disk_quota_counts_texnazor(self):
        self.obj.abs("big.bin").write_bytes(b"0" * (1024 * 1024 + 1))
        r = self.client.post(reverse("texnazor:create"), {"ofyorta": SimpleUploadedFile("b.xlsx", b"x")})
        self.assertContains(r, "Дискдаги жойингиз тўлди")

    def test_my_ofyorta_hidden_by_default(self):
        self.assertNotContains(self.client.get(reverse("texnazor:index")), "Менинг офёрталарим")
        self.assertNotContains(self.client.get(reverse("texnazor:create") + "?my=1"), "Менинг офёрталаримдан")

    def test_month_cycle_october_close_november(self):
        """Октябрь: бир иш 100% (белгилаб), бири қисман 40%, ёпилади; Ноябрда тугагани «тугаган», қисмани қолдиғи билан."""
        add = lambda m: self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": m})  # noqa: E731
        add(10)
        octo = Period.objects.get()
        self.assertTrue(self._post("fill_items", {"period": octo.pk, "items": [self.free.pk], "pct": "100"}).json()["ok"])
        self.assertTrue(self._post("done", {"period": octo.pk, "item": self.work.pk, "value": "40%"}).json()["ok"])
        f2 = services.month_f2(octo)  # Ф-2 жами: 4 м3 × 5 чел-ч, 4 × 0,5 маш-ч
        self.assertAlmostEqual(f2["hours"]["J"], 20)
        self.assertAlmostEqual(f2["hours"]["K"], 2)
        self.assertAlmostEqual(f2["money"]["J"], 20 * 1000)
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, octo.pk]), {"closed": "1"})
        # Ф-2: faqat shu oy ishlari (resurslari bilan), pastida JAMI va svodka formulalari
        dl = self.client.get(reverse("texnazor:f2", args=[self.obj.pk, octo.pk]))
        self.assertEqual(dl.status_code, 200)
        self.assertIn("attachment", dl["Content-Disposition"])
        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(dl.content)).active
        cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
        self.assertEqual(cells["A8"], "2026 YIL OKTYABR OYI")
        rows = {v: k for k, v in cells.items() if isinstance(v, str) and k.startswith("C")}
        w = int(rows["БЕТОН ФУНДАМЕНТОВ"][1:])
        self.assertEqual(cells[f"E{w}"], 4)  # ish qatorida — shu oydagi hajm
        self.assertEqual(cells[f"F{w + 1}"], f"=E{w + 1}*E{w}")  # resurs = me'yor × ish hajmi
        self.assertEqual(cells[f"J{w + 1}"], f"=+I{w + 1}")  # ish haqi ustuniga
        sh = int(rows["ШВЕЛЛЕР №16"][1:])
        self.assertEqual((cells[f"E{sh}"], cells[f"I{sh}"], cells[f"L{sh}"]), (2, f"=+G{sh}*E{sh}", f"=+I{sh}"))
        self.assertNotIn("ТЕПЛОВОЙ НАСОС", rows)  # shu oyda yozilmagan — Ф-2 da yo'q
        jami = int(rows["JAMI:"][1:])
        self.assertEqual(cells[f"J{jami}"], f"=+SUM(J13:J{jami - 1})")
        self.assertIn("PUDRATCHI XARAJATLARI QQS BILAN", rows)
        # «Ф-2» ko'rinishi sahifaning o'zida: faqat shu oy ishlari, resurslari ochiq, oxirida JAMI (soatlar bilan) va svodka
        fv = self.client.get(reverse("texnazor:rows", args=[self.obj.pk]), {"p": octo.pk, "view": "f2"}).json()
        fh = "".join(c["h"] for c in fv["chunks"])
        self.assertEqual(fv["cols"], 9)
        self.assertEqual(sorted(i for c in fv["chunks"] for i in c["items"]), sorted([self.work.pk, self.free.pk]))
        self.assertNotIn("ТЕПЛОВОЙ НАСОС", fh)
        self.assertIn('class="tn-r is-shown', fh)
        self.assertIn("20</td>", fh.split("Иш ҳақи (чел-ч)")[1])  # 4 м3 × 5 чел-ч
        from .templatetags.texnazor_tags import _som as som
        self.assertIn(som(services.period_totals(octo)["this"]["contractor"]), fh.split("Пудратчи харажатлари ҚҚС билан")[1])
        page = self.client.get(reverse("texnazor:object", args=[self.obj.pk]), {"p": octo.pk})
        self.assertContains(page, "Ноябрь 2026 ни очиш")  # ёпилгач — кейинги ой бир тугмада
        add(11)
        nov = Period.objects.get(month=11)
        rows = {r["item"].pk: r for r in services.item_rows(nov, [self.work, self.free])}
        self.assertEqual(rows[self.free.pk]["prev"], "Октябрь 2026")  # тугаган — ёзилмайди, адашмайди
        self.assertEqual(rows[self.work.pk]["prev"], "")
        self.assertAlmostEqual(rows[self.work.pk]["before_pct"], 40)
        self.assertAlmostEqual(rows[self.work.pk]["left"], 6)
        full = self.client.get(reverse("texnazor:rows", args=[self.obj.pk]), {"p": nov.pk}).json()
        html = "".join(c["h"] for c in full["chunks"])
        self.assertIn("Октябрь 2026да тугаган", html)
        self.assertEqual([i for c in full["chunks"] for i in c["skip"]], [self.free.pk])  # белгилаб бўлмайди
        # «Ҳамма ойлар»: har oy alohida ustun — Oktyabrdagi 4 va qoldiq 6 bir qatorda; bo'lim qatorida — so'mda
        allm = self.client.get(reverse("texnazor:rows", args=[self.obj.pk]), {"p": nov.pk, "view": "months"}).json()
        self.assertEqual(allm["cols"], 5 + 4 + 2 * (2 + 1))
        self.assertIn("Октябрь 2026", allm["panel"])
        self.assertIn("Ноябрь 2026", allm["panel"])
        html = "".join(c["h"] for c in allm["chunks"])

        def blocks(row_id):  # «Офёрта» — 4 ustun (на ед., по проект., цена, сумма), qolgan bloklar — miqdor va summa
            row = html.split(row_id)[1].split("</tr>")[0]
            cells = [c.split(">", 1)[1].split("<")[0] for c in row.split('<td class="num')[1:]]
            return [cells[:4]] + [cells[i:i + 2] for i in range(4, len(cells), 2)]
        # ish (1 м3 = 525 000): офёрта, Октябрь, Ноябрь, қолдиқ (Ф-2 dagidek)
        w = blocks(f'id="tn-i-{self.work.pk}"')
        self.assertEqual(w[0], ["10", "", "", "5 250 000"])
        self.assertEqual(w[1:], [["4", "2 100 000"], ["", ""], ["6", "3 150 000"]])
        # resurs (ish haqi: me'yor 5, narx 1 000): офёртада 50 чел-ч, Oktyabrda 5 × 4 = 20 чел-ч, 20 000 so'm
        res = blocks(f'data-w="{self.work.pk}"')
        self.assertEqual((res[0], res[1]), (["5", "50", "1 000", "50 000"], ["20", "20 000"]))
        # oxirida svodka: Oktyabr blokida pudratchi summasi — Oktyabr akti bilan bir xil
        from .templatetags.texnazor_tags import _som
        self.assertIn(_som(services.period_totals(octo)["this"]["contractor"]), html.split("Пудратчи харажатлари ҚҚС билан")[1])
        # «Ҳамма ойлар»da ham yoziladi: shu ko'rinishdagi qator (yozish katagi bilan) va yangilangan svodka qaytadi
        d = self._post("done", {"period": nov.pk, "item": self.work.pk, "value": "50%", "view": "months", "fit": 3}).json()
        self.assertTrue(d["ok"])
        self.assertIn('class="tn-in num" value="5"', d["row"])
        self.assertIn("Пудратчи харажатлари ҚҚС билан", d["svodka"])
        Done.objects.filter(period=nov, item=self.work).delete()
        # 4 oy, ekranga 3 blok sig'sa — eskilari «Олдинги ойлар» blokida (jadval yonga surilmasin); keng ekranda va
        # «Ҳар ойни алоҳида» bosilganda — hammasi alohida
        add(12)
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2027, "month": 1})
        dec = Period.objects.get(month=12)
        url = reverse("texnazor:rows", args=[self.obj.pk])
        four = self.client.get(url, {"p": dec.pk, "view": "months", "fit": 3}).json()
        self.assertIn("Олдинги ойлар", four["panel"])  # Oktyabr + Noyabr jamlangan, Dekabr va Yanvar — alohida
        self.assertEqual(four["cols"], 5 + 4 + 2 * (3 + 1))
        self.assertEqual(self.client.get(url, {"p": dec.pk, "view": "months", "fit": 5}).json()["cols"], 5 + 4 + 2 * (4 + 1))
        self.assertEqual(self.client.get(url, {"p": dec.pk, "view": "months", "all": "1"}).json()["cols"], 5 + 4 + 2 * (4 + 1))
        Period.objects.filter(month__in=(12, 1)).delete()
        # Ноябрда иккаласини «100%» қилса ҳам — тугагани ўзгармайди, қисманнинг қолган 60% и ёзилади
        self.assertTrue(self._post("fill_items", {"period": nov.pk, "items": [self.work.pk, self.free.pk], "pct": "100"}).json()["ok"])
        self.assertAlmostEqual(Done.objects.get(period=nov, item=self.work).qty, 6)
        self.assertFalse(Done.objects.filter(period=nov, item=self.free).exists())
        self.assertAlmostEqual(Done.objects.get(period=octo, item=self.free).qty, 2)

    def test_import(self):
        self.assertEqual(self.obj.status, "ready")
        self.assertEqual(self.obj.name, "CHILONZOR TUMAN 5-MAKTAB TA'MIRI")  # nomi berilmadi — ofyortadan
        self.assertEqual(self.work.resources.count(), 3)
        self.assertAlmostEqual(self.work.cost_l, 510000)
        self.assertGreater(self.obj.contract_totals["contractor"], 0)
        page = self.client.get(reverse("texnazor:object", args=[self.obj.pk]))
        self.assertContains(page, "Биринчи ҳисобот ойини қўшинг")

    def test_months_carry_over_and_remaining(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self.assertTrue(self._post("done", {"period": july.pk, "item": self.work.pk, "value": "4"}).json()["ok"])
        self.assertTrue(self._post("done", {"period": july.pk, "item": self.free.pk, "value": "50%"}).json()["ok"])
        self.assertEqual(Done.objects.get(item=self.free).qty, 1.0)  # 2 ТН × 50%
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 9})
        sept = Period.objects.get(month=9)
        # sentyabrda iyuldagi 4 m3 «avval» bo'lib turadi; qoldiq 6 — 7 yozib bo'lmaydi
        r = self._post("done", {"period": sept.pk, "item": self.work.pk, "value": "7"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Қолдиқ 6 ", r.json()["error"])
        d = self._post("done", {"period": sept.pk, "item": self.work.pk, "value": "6"}).json()
        self.assertTrue(d["ok"])
        t = services.period_totals(sept)
        per_m3 = 5000 + 10000 + 510000
        self.assertAlmostEqual(t["this"]["direct"], 6 * (5000 + 10000 + 510000 * 1.03))
        self.assertAlmostEqual(t["before"]["L"], 4 * 510000 + 1 * 9600000)
        self.assertAlmostEqual(t["cum"]["J"], 10 * 5000)
        self.assertAlmostEqual(t["left"]["J"], 0)
        self.assertGreater(per_m3, 0)
        self.assertAlmostEqual(d["summary"]["cum"] + d["summary"]["left"], d["summary"]["contract"], places=2)
        rows = services.item_rows(sept, [self.work])
        self.assertEqual((rows[0]["before"], rows[0]["this"], rows[0]["left"]), (4.0, 6.0, 0.0))
        # resurslar ishga ergashadi (me'yor × hajm), bo'lim qatori ham yangilanadi
        for x in rows[0]["res"]:
            self.assertAlmostEqual(x["this"], x["r"].qty * 6)
        self.assertAlmostEqual(sum(x["this_amount"] for x in rows[0]["res"]), rows[0]["this_amount"])
        self.assertAlmostEqual(d["section"]["this"], 6 * per_m3)

    def test_closed_month_and_section_fill(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        sec = self.work.section
        self.assertTrue(self._post("fill", {"period": july.pk, "pct": 100}, section_id=sec.pk).json()["ok"])
        self.assertEqual(Done.objects.get(item=self.work).qty, 10)
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, july.pk]), {"closed": "1"})
        r = self._post("done", {"period": july.pk, "item": self.work.pk, "value": "1"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("ёпилган", r.json()["error"])
        # o'tmishdagi oy qo'shib bo'lmaydi
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 5})
        self.assertEqual(Period.objects.count(), 1)

    def test_pages_and_partials(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self.assertContains(self.client.get(reverse("texnazor:index")), "CHILONZOR")
        page = self.client.get(reverse("texnazor:object", args=[self.obj.pk]))
        self.assertContains(page, "Фундамент")  # mundarija; nomlar oddiy yozuvda (readable)
        self.assertContains(page, "Июль 2026")
        full = self.client.get(reverse("texnazor:rows", args=[self.obj.pk]), {"p": july.pk}).json()  # butun ofyorta, bo'laklarda
        rows = "".join(c["h"] for c in full["chunks"])
        self.assertIn("БЕТОН ФУНДАМЕНТОВ", rows)  # jadvalda nomlar ofyortadagidek (Excel'dagi katta harflar)
        self.assertIn("КРАН", rows)  # ishning resurslari shu jadvalda (yopiq), alohida so'rovsiz
        self.assertIn(f'id="tn-s-{self.work.section_id}"', rows)  # bo'lim sarlavhasi — mundarijadan shunga o'tiladi
        self.assertIn(self.work.pk, [i for c in full["chunks"] for i in c["items"]])
        self.assertContains(self.client.get(reverse("texnazor:search", args=[self.obj.pk]), {"p": july.pk, "q": "швеллер"}), "Швеллер")
        res_row = self.work.resources.first().excel_row  # Excel qatori bo'yicha: resurs qatori -> uning ishi
        self.assertContains(self.client.get(reverse("texnazor:search", args=[self.obj.pk]), {"p": july.pk, "q": str(res_row)}),
                            "Бетон фундаментов")

    def test_correction_changes_contract_and_keeps_history(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self._post("done", {"period": july.pk, "item": self.work.pk, "value": "4"})
        before = SiteObject.objects.get().contract_totals["direct"]
        url = lambda it: reverse("texnazor:correct", args=[self.obj.pk, it.pk])  # noqa: E731
        post = lambda it, body: self.client.post(url(it), json.dumps(body), content_type="application/json")  # noqa: E731
        r = post(self.work, {"period": july.pk, "value": "12", "basis": ""})
        self.assertIn("Асос", r.json()["error"])  # asos hujjatisiz bo'lmaydi
        r = post(self.work, {"period": july.pk, "value": "3", "basis": "ҚК №2"})
        self.assertIn("Бажарилган 4", r.json()["error"])  # bajarilganidan kam bo'lmaydi
        self.assertTrue(post(self.work, {"period": july.pk, "value": "12", "basis": "ҚК №2, 12.11.2026"}).json()["ok"])
        self.work.refresh_from_db()
        self.assertEqual(self.work.qty, 12)
        self.assertEqual(Done.objects.get(item=self.work).qty, 4)  # bajarilgan o'zgarmaydi
        per_m3 = 5000 + 10000 + 510000 * 1.03
        self.assertAlmostEqual(SiteObject.objects.get().contract_totals["direct"] - before, 2 * per_m3, places=2)
        c = Correction.objects.get()
        self.assertEqual((c.old_qty, c.new_qty, c.basis, c.author), (10, 12, "ҚК №2, 12.11.2026", self.user))
        page = self.client.get(reverse("texnazor:object", args=[self.obj.pk]))
        self.assertContains(page, "Офёрта тузатишлари")
        self.assertContains(page, "ҚК №2, 12.11.2026")
        # yopilgan oydan kuchga kirmaydi
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, july.pk]), {"closed": "1"})
        self.assertIn("ёпилган", post(self.work, {"period": july.pk, "value": "13", "basis": "ҚК №3"}).json()["error"])

    def test_paid_difference(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self._post("done", {"period": july.pk, "item": self.work.pk, "value": "4"})
        calc_ = services.period_totals(july)["this"]["contractor"]
        self.client.post(reverse("texnazor:paid", args=[self.obj.pk, july.pk]), {"paid": f"{calc_ + 1000.5:,.2f}".replace(",", " ").replace(".", ",")})
        july.refresh_from_db()
        self.assertAlmostEqual(july.paid, calc_ + 1000.5, places=2)
        self.assertAlmostEqual(services.payments(self.obj)["over"], 1000.5, places=2)
        self.assertContains(self.client.get(reverse("texnazor:object", args=[self.obj.pk])), "ортиқча тўланган")

    def test_delete_month_needs_typed_title_and_open_month(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self._post("done", {"period": july.pk, "item": self.work.pk, "value": "4"})
        url = reverse("texnazor:period_delete", args=[self.obj.pk, july.pk])
        self.client.post(url, {"confirm": "ҳа"})
        self.assertTrue(Period.objects.exists())  # 2-qadam: oy nomi yozilmasa o'chmaydi
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, july.pk]), {"closed": "1"})
        self.client.post(url, {"confirm": "Июль 2026"})
        self.assertTrue(Period.objects.exists())  # yopilgan oy o'chmaydi
        self.client.post(reverse("texnazor:period_close", args=[self.obj.pk, july.pk]), {"closed": "0"})
        self.client.post(url, {"confirm": "июль 2026"})
        self.assertFalse(Period.objects.exists())
        self.assertFalse(Done.objects.exists())

    def test_revert_correction_last_first(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        c1 = services.correct_item(self.user, self.work, july, "12", "ҚК №1")
        c2 = services.correct_item(self.user, Item.objects.get(pk=self.work.pk), july, "15", "ҚК №2")
        rv = lambda c, w: self.client.post(reverse("texnazor:revert", args=[self.obj.pk, c.pk]), json.dumps({"confirm": w}),  # noqa: E731
                                           content_type="application/json").json()
        self.assertIn("БЕКОР", rv(c2, "ҳа")["error"])
        self.assertIn("кейинги", rv(c1, "бекор")["error"])  # oxirgisi birinchi
        self.assertTrue(rv(c2, "бекор")["ok"])
        self.assertEqual(Item.objects.get(pk=self.work.pk).qty, 12)
        self.assertTrue(rv(c1, "БЕКОР")["ok"])
        self.assertEqual(Item.objects.get(pk=self.work.pk).qty, 10)
        self.assertEqual(Correction.objects.filter(reverted__isnull=False).count(), 2)  # tarixda qoladi
        self.assertAlmostEqual(SiteObject.objects.get().contract_totals["direct"], self.obj.contract_totals["direct"], places=2)

    def test_delete_object_needs_word(self):
        self.client.post(reverse("texnazor:delete", args=[self.obj.pk]), {"confirm": ""})
        self.assertTrue(SiteObject.objects.exists())
        self.client.post(reverse("texnazor:delete", args=[self.obj.pk]), {"confirm": "ўчириш"})
        self.assertFalse(SiteObject.objects.exists())

    def test_selected_items_get_total_percent(self):
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        july = Period.objects.get()
        self._post("done", {"period": july.pk, "item": self.work.pk, "value": "4"})
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 9})
        sept = Period.objects.get(month=9)
        fill = lambda ids, pct: self._post("fill_items", {"period": sept.pk, "items": ids, "pct": pct})  # noqa: E731
        self.assertIn("белгиланг", fill([], 100).json()["error"])
        d = fill([self.work.pk, self.free.pk], "50").json()
        self.assertTrue(d["ok"])
        self.assertAlmostEqual(Done.objects.get(item=self.work, period=sept).qty, 1)  # jami 50% = 5, avval 4 -> shu oy 1
        self.assertAlmostEqual(Done.objects.get(item=self.free, period=sept).qty, 1)  # 2 ТН × 50%
        self.assertEqual(Done.objects.filter(period=july).count(), 1)  # avvalgi oy o'zgarmaydi
        self.assertEqual({x["id"] for x in d["sections"]}, {self.work.section_id, self.free.section_id})
        fill([self.work.pk], 30)  # avval 40% bajarilgan — shu oydagisi olib tashlanadi
        self.assertFalse(Done.objects.filter(item=self.work, period=sept).exists())
        fill([self.work.pk], 100)
        self.assertAlmostEqual(Done.objects.get(item=self.work, period=sept).qty, 6)

    def test_other_user_cannot_see_object(self):
        other = User.objects.create_user("b@test.uz", "Kuchli-parol-2026", first_name="B", last_name="B")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("texnazor:object", args=[self.obj.pk])).status_code, 404)
        self.assertEqual(self._post("done", {"item": self.work.pk, "value": "1"}).status_code, 404)
        self.assertNotContains(self.client.get(reverse("texnazor:index")), "CHILONZOR")


def test_money_is_shown_to_tiyin_without_rounding_to_som():
    from texnazor.templatetags.texnazor_tags import pul
    assert pul(2862982.6934) == "2 862 982,69"
    assert pul(1000) == "1 000,00"
    assert pul(-0.001) == "0,00"


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class ZeroVolumeCorrectionTests(TestCase):
    """Ofyortada hajmi 0 bo'lgan ish (resurslarida formula xato): tuzatilsa narxi me'yorlardan tiklanadi."""

    def test_zero_volume_work_gets_unit_cost_from_norms(self):
        user = User.objects.create_user("z@test.uz", "Kuchli-parol-2026", first_name="Z", last_name="Z")
        self.client.force_login(user)
        Path(TMP_MEDIA).mkdir(parents=True, exist_ok=True)
        with open(build_ofyorta(f"{TMP_MEDIA}/z.xlsx", broken_zero_work=True), "rb") as f:
            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(reverse("texnazor:create"), {"ofyorta": SimpleUploadedFile("z.xlsx", f.read())})
        obj = SiteObject.objects.get()
        tasks.run_now(obj.pk)
        zero = Item.objects.get(name="УКЛАДКА ПРОГОНОВ")
        self.assertEqual((zero.qty, zero.unit_cost), (0, 0))
        self.client.post(reverse("texnazor:period_add", args=[obj.pk]), {"year": 2026, "month": 7})
        page = self.client.get(reverse("texnazor:object", args=[obj.pk]))
        self.assertContains(page, "Жадвалда кўрсатиш")
        norms = sum(r.price * r.qty for r in zero.resources.all())
        before = SiteObject.objects.get().contract_totals["direct"]
        r = self.client.post(reverse("texnazor:correct", args=[obj.pk, zero.pk]),
                             json.dumps({"value": "2", "basis": "Смета тузатмаси №1"}), content_type="application/json")
        self.assertTrue(r.json()["ok"])
        zero.refresh_from_db()
        self.assertAlmostEqual(zero.unit_cost, norms)
        self.assertGreater(SiteObject.objects.get().contract_totals["direct"], before)
        self.assertTrue(all(abs(x.total_qty - x.qty * 2) < 1e-9 for x in zero.resources.all()))
        self.assertContains(self.client.get(reverse("texnazor:object", args=[obj.pk])), "Тузатилди: 0 → 2")


def _mini_ofyorta(path, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "ОФЁРТА"
    ws.append(["№ п/п", "Шифр", "Наименование работ", "Единица", "Количество"])
    ws.append([None, None, None, None, "на единицу", "по проектным", "Цена подрятчик", "Цена заказчик", "Сумма", "ЧЕЛ-Ч.", "МАШ-Ч.", "МАТРИАЛ"])
    ws.append(list(range(1, 13)))
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


_WORK = [[1, "Е06-01-001-1", "БЕТОН", "1 М3", 10], ["1.1", "1", "ЗАТРАТЫ ТРУДА", "ЧЕЛ-Ч.", 5, 50, 1000, None, 50000, 50000]]
_FREE = [[2, "30032", "ШВЕЛЛЕР", "ТН", 2, None, 9600000, None, 19200000, None, None, 19200000]]


def test_items_without_any_header_are_not_lost(tmp_path):
    of = parse(_mini_ofyorta(tmp_path / "a.xlsx", _WORK + _FREE))
    assert of.sections == ["Бўлим номи кўрсатилмаган"]
    assert all(r.section == 0 for r in of.rows)
    assert any("сарлавҳаси йўқ" in w["title"] for w in of.warnings)


def test_numbered_header_without_unit_or_price_is_a_section(tmp_path):
    rows = [[1, None, "ДЕМОНТАЖ"], [2, "Е06-01-001-1", "БЕТОН", "1 М3", 10],
            ["2.1", "1", "ЗАТРАТЫ", "ЧЕЛ-Ч.", 5, 50, 1000, None, 50000, 50000]]
    of = parse(_mini_ofyorta(tmp_path / "b.xlsx", rows))
    assert of.sections == ["ДЕМОНТАЖ"]
    assert [r.kind for r in of.rows] == ["section", "work", "res"]


def test_block_titles_with_or_without_na():
    from .ofyorta import BLOCK_RE as _BLOCK
    from .services import block_label
    for t in ("на АС Блок-1", "НА ЭО", "АС БЛОК-2", "БЛОК №3"):
        assert _BLOCK.search(t), t
    for t in ("ДВЕРНЫЕ БЛОКИ (АС ЛИСТ-9)", "ДЕМОНТАЖ 1-ЭТАЖ", "НАРУЖНОЙ ОТДЕЛКИ"):
        assert not _BLOCK.search(t), t
    assert block_label("НА АС БЛОК-1") == "АС БЛОК-1"


@override_settings(MEDIA_ROOT=TMP_MEDIA)
class NegativeVolumeTests(TestCase):
    """Ofyortadagi manfiy hajm («ИСКЛЮЧИТЬ толщ. 10 см»): asosiy ishdan ayiriladi, hajmi ham manfiy yoziladi."""

    def setUp(self):
        self.user = User.objects.create_user("n@test.uz", "Kuchli-parol-2026", first_name="N", last_name="N")
        self.client.force_login(self.user)
        Path(TMP_MEDIA).mkdir(parents=True, exist_ok=True)
        rows = [[None, None, "ПОКРЫТИЯ"],
                [1, "Е27-06-002-17", "БЕТОННЫЕ ПОКРЫТИЯ 20 СМ", "1000 М2", 0.364],
                ["1.1", "1", "ЗАТРАТЫ ТРУДА", "ЧЕЛ-Ч.", 100, 36.4, 1000, None, 36400, 36400],
                [2, "Е27-06-002-18", "ИСКЛЮЧИТЬ ТОЛЩ. 10 СМ", "1000 М2", -0.364],
                ["2.1", "1", "ЗАТРАТЫ ТРУДА", "ЧЕЛ-Ч.", 50, -18.2, 1000, None, -18200, -18200]]
        with open(_mini_ofyorta(f"{TMP_MEDIA}/neg.xlsx", rows), "rb") as f:
            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(reverse("texnazor:create"), {"ofyorta": SimpleUploadedFile("neg.xlsx", f.read())})
        self.obj = SiteObject.objects.get()
        tasks.run_now(self.obj.pk)
        self.obj.refresh_from_db()
        self.base = Item.objects.get(num="1")
        self.minus = Item.objects.get(num="2")
        self.client.post(reverse("texnazor:period_add", args=[self.obj.pk]), {"year": 2026, "month": 7})
        self.july = Period.objects.get()

    def done(self, item, value):
        return self.client.post(reverse("texnazor:done", args=[self.obj.pk]), json.dumps(
            {"period": self.july.pk, "item": item.pk, "value": value}), content_type="application/json").json()

    def test_deduction_row_takes_negative_volume(self):
        self.assertAlmostEqual(self.minus.qty, -0.364)
        self.assertAlmostEqual(self.obj.contract_totals["J"], 36400 - 18200)
        self.assertTrue(self.done(self.minus, "0,364")["ok"])  # musbat yozildi — manfiy saqlanadi
        self.assertAlmostEqual(Done.objects.get(item=self.minus).qty, -0.364)
        self.assertTrue(self.done(self.minus, "50%")["ok"])
        self.assertAlmostEqual(Done.objects.get(item=self.minus).qty, -0.182)
        self.assertIn("Қолдиқ", self.done(self.minus, "-0,5")["error"])  # qoldiqdan oshib bo'lmaydi
        self.assertIn("манфий", self.done(self.base, "-0,1")["error"])  # oddiy ishga manfiy yozib bo'lmaydi
        row = services.item_rows(self.july, [self.minus])[0]
        self.assertEqual(row["base"], self.base)
        self.assertFalse(row["over"])
        full = self.client.get(reverse("texnazor:rows", args=[self.obj.pk]), {"p": self.july.pk}).json()
        self.assertIn("айириш — 1-ишдан", "".join(c["h"] for c in full["chunks"]))

    def test_fill_section_marks_base_and_deduction(self):
        services.fill_section(self.july, self.base.section, 100)
        self.assertAlmostEqual(Done.objects.get(item=self.base).qty, 0.364)
        self.assertAlmostEqual(Done.objects.get(item=self.minus).qty, -0.364)
        t = services.period_totals(self.july)
        self.assertAlmostEqual(t["this"]["J"], 36400 - 18200)
        self.assertAlmostEqual(t["left"]["J"], 0)
        row = services.item_rows(self.july, [self.minus])[0]
        self.assertTrue(row["complete"])

    def test_correction_of_deduction_row(self):
        self.done(self.minus, "0,2")
        self.assertIn("Бажарилган", services_error(lambda: services.correct_item(self.user, self.minus, self.july, "-0,1", "ҚК №1")))
        services.correct_item(self.user, self.minus, self.july, "-0,3", "ҚК №1")
        self.assertAlmostEqual(Item.objects.get(pk=self.minus.pk).qty, -0.3)


def services_error(fn):
    try:
        fn()
    except services.TexnazorError as e:
        return str(e)
    return ""


# ======== «Менинг офёрталарим»: narxlangan lokallardan ofyorta ========
def _priced_local(job, rows):
    """Narxlangan lokal: «_ЛРВ» varag'i (A–F) va ``job.result`` (narx, turi, jami miqdor)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "_ЛРВ"
    res = []
    for r, (cells, info) in enumerate(rows, start=5):
        for j, v in enumerate(cells, start=1):
            if v is not None:
                ws.cell(row=r, column=j, value=v)
        if info:
            res.append({"row": r, "name": cells[2] if len(cells) > 2 else "", "unit": cells[3] if len(cells) > 3 else "", **info})
    job.output_path = "out/local_narxlandi.xlsx"
    job.abs(job.output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(job.abs(job.output_path))
    job.local_sheet = "_ЛРВ"
    job.result = {"rows": res, "start_row": 5, "end_row": 4 + len(rows), "totals": {"s_nds": 1}}
    job.status = job.Status.DONE
    job.save()


@override_settings(MEDIA_ROOT=TMP_MEDIA, NARXLASH_TASKS="worker", TEXNAZOR_MY_OFYORTA=True)
class MyOfyortaTests(TestCase):
    def setUp(self):
        from jobs.models import Batch, Job, Project
        self.user = User.objects.create_user("my@test.uz", "Kuchli-parol-2026", first_name="M", last_name="O")
        self.client.force_login(self.user)
        self.project = Project.objects.create(owner=self.user, name="Нурафшон МТТ")
        jobs = []
        for order, (folder, label, pct) in enumerate((("1. Блок-А", "Блок-А", {}),
                                                      ("2. =№81 склад", "Склад", {"transport_mat": 0.07}))):
            batch = Batch.objects.create(owner=self.user, name=label, project=self.project, folder=folder)
            jobs.append(Job.objects.create(owner=self.user, batch=batch, batch_label=label, batch_order=order, percents=pct,
                                           local_name=f"{label}.xlsx", res_name="r.xlsx", local_path="l", res_path="r"))
        a, b = jobs
        _priced_local(a, [
            ((None, None, "РАЗДЕЛ 1. ФУНДАМЕНТ"), None),
            (("N п.п.", "Шифр", "Наименование работ и затрат", "Ед."), {"g": "ТОПИЛМАДИ", "qty": None}),  # takrorlangan sarlavha
            ((1, "Е06-01", "БЕТОН ФУНДАМЕНТОВ", "М3", 10), {"g": None, "qty": 10, "status": "WORK_HEADER", "category": "I"}),
            (("1.1", "1", "ЗАТРАТЫ ТРУДА", "ЧЕЛ-Ч", 5, 50), {"g": 1000, "qty": 50, "category": "I"}),
            (("1.2", "22006", "БЕТОН М200", "М3", 1.02, 10.2), {"g": 500000, "qty": 10.2, "category": "K"}),
            ((2, None, "ТЕПЛОВОЙ НАСОС", "ШТ", 2, 2), {"g": 1000000, "qty": 2, "category": "M"}),
            ((None, None, "Итого по разделу 1"), None),
        ])
        _priced_local(b, [
            ((1, None, "КАБЕЛЬ ВВГ", "М", 100, 100), {"g": 2000, "qty": 100, "category": "L"}),
            ((2, None, "ДВЕРЬ", "ШТ", 1, 1), {"g": "ТОПИЛМАДИ", "qty": 1, "category": "K"}),
        ])
        self.a, self.b = a, b

    def test_setup_build_and_object_with_exact_blocks(self):
        from .build import setup
        from .models import MyOfyorta
        s = setup(self.project)
        self.assertEqual([p["title"] for p in s["parts"]], ["БЛОК-А", "=№81 СКЛАД"])
        self.assertEqual(len(s["equipment"]), 1)
        self.assertEqual(len(s["differ"]), 1)  # ikkinchi lokalda transport boshqacha
        r = self.client.get(reverse("texnazor:my_new") + f"?project={self.project.pk}")
        self.assertContains(r, "фоизлари бошқача")

        data = {"project": self.project.pk, "name": "Нурафшон МТТ", "object_type": "Таълим",
                f"include_{self.a.pk}": "1", f"include_{self.b.pk}": "1", f"title_{self.a.pk}": "БЛОК-А",
                f"title_{self.b.pk}": "=№81 СКЛАД", "customer": [s["equipment"][0]["key"]],
                **{f"rate_{k}": v for k, v in (("tr_material", "5"), ("tr_cable", "1,5"), ("tr_eq_customer", "2"),
                                               ("tr_eq_contractor", "2"), ("overhead", "19"), ("insurance", "0.32"), ("vat", "12"))}}
        r = self.client.post(reverse("texnazor:my_new"), data)
        doc = MyOfyorta.objects.get()
        self.assertRedirects(r, reverse("texnazor:my_detail", args=[doc.pk]), fetch_redirect_response=False)
        tasks.run_now(doc.pk)
        doc.refresh_from_db()
        self.assertEqual(doc.status, "ready", doc.error)
        self.assertEqual(doc.structure, [{"title": "БЛОК-А", "block": True}, {"title": "РАЗДЕЛ 1. ФУНДАМЕНТ", "block": False},
                                         {"title": "=№81 СКЛАД", "block": True}, {"title": "Склад", "block": False}])
        self.assertEqual(doc.rates["tr_cable"], 0.015)
        sums = {"J": 50000.0, "K": 0.0, "L": 5100000.0, "M": 0.0, "N": 2000000.0, "O": 200000.0, "X": 0.0}
        self.assertAlmostEqual(doc.totals["total"], chain(sums, doc.rates)["total"], places=2)
        self.assertEqual([p["name"] for p in doc.problems], ["ДВЕРЬ"])  # takrorlangan sarlavha muammo emas

        of = parse(doc.file)  # Excel: «=» bilan boshlangan sarlavha formula emas, matn
        self.assertEqual(of.sections, [x["title"] for x in doc.structure])
        self.assertAlmostEqual(of.file_totals["total_vat"], doc.totals["total"], places=2)
        r = self.client.get(reverse("texnazor:my_file", args=[doc.pk]))
        self.assertEqual(r.status_code, 200)

        r = self.client.post(reverse("texnazor:create"), {"source": "my", "my": str(doc.pk), "name": "Нурафшон"})
        obj = SiteObject.objects.get()
        self.assertRedirects(r, reverse("texnazor:object", args=[obj.pk]), fetch_redirect_response=False)
        tasks.run_now(obj.pk)
        obj.refresh_from_db()
        self.assertEqual(obj.status, "ready", obj.error)
        self.assertEqual(obj.made_from, doc)
        self.assertEqual(list(obj.sections.order_by("order").values_list("title", "is_block")),
                         [(x["title"], x["block"]) for x in doc.structure])
        self.assertAlmostEqual(obj.contract_totals["total"], doc.totals["total"], places=2)

    def test_structure_not_applied_to_changed_file(self):
        from .ofyorta import Ofyorta
        of = Ofyorta(sheet="ОФЁРТА", sections=["БЛОК-А", "ФУНДАМЕНТ"], blocks={1})
        self.assertFalse(services.apply_structure(of, [{"title": "БЛОК-А", "block": True}]))
        self.assertEqual(of.blocks, {1})
        self.assertIn("фарқ", of.warnings[0]["text"])
        self.assertTrue(services.apply_structure(of, [{"title": "БЛОК-А", "block": True}, {"title": " ФУНДАМЕНТ", "block": False}]))
        self.assertEqual(of.blocks, {0})

    def test_other_user_cannot_open(self):
        from .models import MyOfyorta
        doc = MyOfyorta.objects.create(owner=self.user, project=self.project, name="x")
        other = User.objects.create_user("o@test.uz", "Kuchli-parol-2026", first_name="O", last_name="O")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("texnazor:my_detail", args=[doc.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("texnazor:my_new") + f"?project={self.project.pk}").status_code, 404)
