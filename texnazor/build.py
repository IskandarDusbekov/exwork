"""Ofyorta yasash: narxlash loyihasidagi narxlangan lokallardan bitta ofyorta (uning 40-maktab ofyortasidagidek).

Ofyorta — lokallarning ketma-ket yig'indisi:

* har lokal — **blok** (sarlavhasi qizil): nomi papkadan («2. Блок-А / 1. Подземный часть» -> «БЛОК-А / ПОДЗЕМНЫЙ ЧАСТЬ»),
  smetachi o'zgartira oladi;
* lokal ichidagi «РАЗДЕЛ …» qatorlari — **bo'limlar**; bo'limsiz boshlansa — lokal nomi bilan bo'lim;
* ishlar va resurslar — lokaldagidek (raqamlash har blokda 1 dan): E — hajm / me'yor, F = E × ish hajmi,
  G — narx (narxlashda topilgan), I = G × F, turi ustuni (J..O) = I;
* pastda «JAMI:» va svodka — Ф-2 dagi bilan bir xil (``f2.svodka``), foizlar lokallardan.

Turlar: lokal I (ish haqi) -> J, J (mashina) -> K, K (material) -> L, L (kabel) -> O, M (uskuna) -> M (pudratchi) yoki
N (buyurtmachi — smetachi belgilaydi; narxi H «Цена заказчик» ga yoziladi).

Qaysi sarlavha blok ekani ``MyOfyorta.structure`` da aniq saqlanadi: texnazor Excel'dagi rang yoki nomga qarab
taxmin qilmaydi.
"""
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional

import xlsxwriter
from django.db.models import Q
from openpyxl import load_workbook

from core.svodka import resolve_percents
from jobs.models import Job, Project

from . import calc
from .f2 import CAT_COL, GREEN_T, LINE, PINK, PURPLE, RED, WIDTHS, _height, _num, svodka
from .models import MyOfyorta

LOCAL_TO_OF = {"I": "J", "J": "K", "K": "L", "L": "O", "M": "M"}
# narxlash svodkasi foizlari -> ofyorta (calc.chain) foizlari
RATES_FROM = {"tr_material": "transport_mat", "tr_cable": "transport_kabel", "tr_eq_customer": "transport_obor",
              "tr_eq_contractor": "transport_obor", "overhead": "prochie", "insurance": "strahovanie", "vat": "s_nds"}
RATE_LABELS = {"tr_material": "Транспорт (материал)", "tr_cable": "Транспорт (кабель)",
               "tr_eq_customer": "Транспорт (ускуна, буюртмачи)", "tr_eq_contractor": "Транспорт (ускуна, пудратчи)",
               "overhead": "Пудратчи бошқа харажатлари", "insurance": "Суғурта", "vat": "ҚҚС"}
_NUM_TOP = re.compile(r"^\d+$")
_NUM_SUB = re.compile(r"^(\d+)\.\d+$")
_SKIP = re.compile(r"^\s*(ИТОГО|ВСЕГО|ЖАМИ|JAMI)\b", re.I)
_LEAD = re.compile(r"^\s*\d+(\.\d+)*[.)]?\s*")
ZERO_OK = ("WORK_HEADER", "MACHINIST")  # narxi 0 bo'lishi to'g'ri: ish sarlavhasi, mashinist mehnati


class BuildError(Exception):
    pass


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _natural(path: str):
    """«10. Благаустройство» «2. Блок-А» dan keyin (raqamlar son sifatida)."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", path)]


def default_title(job: Job) -> str:
    """Papka yo'lidan: «2. Блок-А/1. Подземный часть» -> «БЛОК-А / ПОДЗЕМНЫЙ ЧАСТЬ»; papkasiz — lokal nomi."""
    folder = (job.batch.folder if job.batch_id else "") or ""
    parts = [_LEAD.sub("", p).strip() for p in folder.split("/") if p.strip()]
    title = " / ".join(p for p in parts if p) or job.batch_label or Path(job.local_name).stem
    return title.upper()[:300]


def project_jobs(project: Project) -> List[Job]:
    """Loyihaning narxlangan lokallari — papkalar tartibida (raqamlari bo'yicha), papka ichida — juftlik tartibida."""
    jobs = list(Job.objects.filter(batch__project=project, status=Job.Status.DONE).exclude(output_path="")
                .select_related("batch"))
    return sorted(jobs, key=lambda j: (_natural(j.batch.folder or ""), j.batch_order, j.created))


def default_parts(project: Project) -> List[dict]:
    return [{"job": str(j.pk), "title": default_title(j), "include": True} for j in project_jobs(project)]


# ===================== lokalni o'qish =====================
def read_local(job: Job) -> dict:
    """Narxlangan lokal: bo'limlar va qatorlar (A–E Excel'dan, narx / turi / jami miqdor — narxlash natijasidan).
    ``{"items": [...], "problems": [...], "percents": {...}, "total": s_nds}``; item — dict:
    kind (section | work | res | free), num, code, name, unit, e (hajm / me'yor), f (jami miqdor), price, cat, row."""
    res = job.result or {}
    by_row = {r["row"]: r for r in res.get("rows", [])}
    start, end = res.get("start_row") or 0, res.get("end_row") or 0
    path = job.abs(job.output_path)
    if not path.exists():
        raise BuildError(f"«{job.local_name}»: нархланган файл топилмади — смета қайта нархлансин")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb[job.local_sheet] if job.local_sheet in wb.sheetnames else wb.worksheets[0]
        sheet_rows = {i: r for i, r in enumerate(ws.iter_rows(min_row=start, max_row=end, max_col=6, values_only=True), start)}
    finally:
        wb.close()
    items, problems = [], []
    extra_cols = {c for c, _n in res.get("extra_columns") or []}
    for r in range(start, end + 1):
        row = sheet_rows.get(r) or ()
        cell = lambda k: row[k] if k < len(row) else None  # noqa: E731
        info = by_row.get(r)
        if info is None:
            # narxlanmaydigan qator: «РАЗДЕЛ 1. …» — bo'lim; «Итого по разделу» — tashlanadi
            text = next((_text(cell(k)) for k in range(4) if _text(cell(k))), "")
            if text and not _SKIP.match(text) and not _NUM_TOP.match(_text(cell(0))) and not _NUM_SUB.match(_text(cell(0))):
                items.append({"kind": "section", "name": text, "row": r})
            continue
        num = _text(cell(0))
        unit = _text(cell(3)) or _text(info.get("unit"))
        if _table_head(row):  # sahifa boshida takrorlangan «N п.п. | Наименование…» yoki «1 | 2 | 3 | 4»
            continue
        if not (_NUM_TOP.match(num) or _NUM_SUB.match(num)):
            text = _text(cell(2)) or num or _text(cell(1))
            if not unit:  # «РАЗДЕЛ 2. БЛОК-А», «ДВЕРНЫЕ БЛОКИ (АР-23)» — narxlashga kirgan, lekin sarlavha
                if text and not _SKIP.match(text):
                    items.append({"kind": "section", "name": text, "row": r})
                continue
            problems.append({"row": r, "name": text, "text": f"тартиб рақами йўқ («{num}») — офёртага олинмади, локалда рақам қўйилсин"})
            continue
        status = info.get("status") or ""
        g = info.get("g")
        price = float(g) if isinstance(g, (int, float)) and not isinstance(g, bool) else 0.0
        if not isinstance(g, (int, float)) or isinstance(g, bool):
            if status not in ZERO_OK:
                problems.append({"row": r, "name": info.get("name") or _text(cell(2)),
                                 "text": "нархи топилмаган" + (f" («{_text(g)}»)" if _text(g) else "")})
        cat = info.get("category") or "K"
        if cat in extra_cols or cat not in LOCAL_TO_OF:
            problems.append({"row": r, "name": info.get("name") or "", "text": f"{cat} устуни офёртада йўқ — материал деб олинди"})
            cat = "K"
        kind = "res" if _NUM_SUB.match(num) else "top"
        e = _num_value(cell(4))
        f = _num_value(info.get("qty")) or 0.0
        items.append({"kind": kind, "num": num, "code": _text(cell(1)), "name": _text(cell(2)) or info.get("name") or "",
                      "unit": unit, "e": e if kind == "res" else (e if e is not None else f),
                      "f": f, "price": price, "cat": LOCAL_TO_OF[cat], "row": r, "status": status})
    # ish: ostida «N.k» resurslari bor; qolgan raqamli qatorlar — mustaqil resurs
    for i, it in enumerate(items):
        if it["kind"] == "top":
            nxt = next((x for x in items[i + 1:] if x["kind"] != "section"), None)
            it["kind"] = "work" if nxt and nxt["kind"] == "res" and nxt["num"].split(".")[0] == it["num"] else "free"
    return {"items": items, "problems": problems, "percents": resolve_percents(job.percents),
            "total": (res.get("totals") or {}).get("s_nds")}


def _table_head(row) -> bool:
    cells = [_text(v).lower() for v in row[:4]] + [""] * (4 - len(row[:4]))
    return cells == ["1", "2", "3", "4"] or cells[2].startswith("наименование") or cells[0] in ("n п.п.", "№ п/п", "n п/п", "№ п.п.")


def _num_value(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace("\xa0", "").replace(" ", "").replace(",", "."))
        except ValueError:
            return None
    return None


def default_rates(locals_: List[dict]) -> Dict[str, float]:
    """Lokallardagi eng ko'p uchragan foizlar (odatda hammasida bir xil)."""
    out = {}
    for key, src in RATES_FROM.items():
        vals = Counter(round(float(lc["percents"].get(src, 0.0)), 6) for lc in locals_)
        out[key] = vals.most_common(1)[0][0] if vals else 0.0
    return out


def equipment(locals_: List[dict], jobs: List[Job]) -> List[dict]:
    """Uskunalar (lokalda M ustun) — smetachi qaysilari buyurtmachiniki ekanini belgilaydi."""
    out = []
    for job, lc in zip(jobs, locals_):
        for it in lc["items"]:
            if it["kind"] in ("res", "free") and it["cat"] == "M":
                out.append({"key": f"{job.pk}:{it['row']}", "name": it["name"], "unit": it["unit"], "qty": it["f"],
                            "price": it["price"], "amount": it["price"] * it["f"]})
    return out


# ===================== Excel =====================
def write(doc: MyOfyorta, blocks: List[dict], rates: Dict[str, float], customer: set) -> dict:
    """``blocks`` — [{"title", "job", "items"}]. Fayl ``doc.file`` ga yoziladi; qaytadi: structure, sums, rows."""
    buf = BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("ОФЁРТА")
    cache = {}

    def fmt(bold=False, color=None, bg="#FFFFFF", box=True, align="center", money=False):
        key = (bold, color, bg, box, align, money)
        if key not in cache:
            p = {"font_name": "Times New Roman", "font_size": 10, "bold": bold, "align": align, "valign": "vcenter",
                 "text_wrap": True}
            if color:
                p["font_color"] = color
            if bg:
                p["bg_color"] = bg
            if box:
                p["border"] = 1
            if money:
                p["num_format"] = "#,##0"
            cache[key] = wb.add_format(p)
        return cache[key]

    for col, w in WIDTHS.items():
        ws.set_column(f"{col}:{col}", w)
    ws.write("C1", "OFYORTA", fmt(bold=True, box=False))
    ws.write_string("C3", doc.name, fmt(bold=True, box=False))
    ws.write("C4", "obyekt nomi", fmt(box=False))
    ws.write_string("C5", doc.object_type or "", fmt(bold=True, box=False))
    ws.write("C6", "obyekt turi", fmt(box=False))
    head = fmt(bold=True, bg=PINK)
    for col, text in (("A", "№ п/п"), ("B", "Шифр номера нормативов и  коды ресурсов"),
                      ("C", "Наименование работ и затрат, характеристика оборудования и его масса, расход ресурсов на единицу измерения"),
                      ("D", "Единица измерения")):
        ws.merge_range(f"{col}9:{col}10", text, head)
    ws.merge_range("E9:I9", "Количество", head)
    for col, text in zip("EFGHI", ("на единицу\nизмерения", "по проектным данным", "Цена подрятчик", "Цена заказчик", "Сумма")):
        ws.write(f"{col}10", text, head)
    for col, text in zip("JKLMNO", ("ЧЕЛ-Ч.", "МАШ-Ч.", "МАТРИАЛ", "ЖИХОЗ ПОДРЯТЧИК", "ЖИХОЗ ЗАКАЗЧИК", "КАБЕЛЬ")):
        ws.write(f"{col}10", text, head)
    ws.set_row(9, 38.25)
    for col, n in zip("ABCDEFGHI", (1, 2, 3, 4, 5, 6, None, None, 8)):
        ws.write(f"{col}11", n, head)

    r = first = 12
    structure, rows, notes = [], 0, []
    sums = dict.fromkeys(CAT_COL, 0.0)

    def heading(title, block):
        nonlocal r
        f = fmt(bold=True, color=RED if block else None, bg=PINK)
        for col in "ABDEFGHI":
            ws.write_blank(f"{col}{r}", None, f)
        ws.write_string(f"C{r}", title, f)
        structure.append({"title": " ".join(title.split()), "block": block})
        r += 1

    def priced(rr, cat, price, qty_ref, qty, color):
        pc = "H" if cat == "N" else "G"
        ws.write(f"{pc}{rr}", price, fmt(bold=True, color=color, money=True))
        if not price:
            ws.write_blank(f"I{rr}", None, fmt(bold=True, color=PURPLE, money=True))
            amount = 0.0
        else:
            amount = price * qty
            ws.write_formula(f"I{rr}", f"=+{pc}{rr}*{qty_ref}{rr}", fmt(bold=True, color=PURPLE, money=True), amount)
        ws.write_formula(f"{CAT_COL[cat]}{rr}", f"=+I{rr}", fmt(color=color, money=True), amount)
        sums[cat] += amount

    for b in blocks:
        heading(b["title"], True)
        items = b["items"]
        if items and items[0]["kind"] != "section":  # lokal «РАЗДЕЛ»siz boshlansa — bo'limi lokal nomi bilan
            heading(b.get("local_title") or b["title"], False)
        work_row, work_qty = None, 0.0
        for it in items:
            if it["kind"] == "section":
                heading(it["name"], False)
                continue
            rows += 1
            cat = "N" if it["cat"] == "M" and f"{b['job']}:{it['row']}" in customer else it["cat"]
            u = it["unit"].upper()
            color = PURPLE if "ЧЕЛ" in u else GREEN_T if "МАШ" in u else None
            bold = it["kind"] != "res"
            f, fl = fmt(bold=bold, color=None if bold else color), fmt(bold=bold, color=None if bold else color, align="left")
            num, code = _num(it["num"]), _num(it["code"])
            ws.write_number(f"A{r}", num, f) if isinstance(num, int) else ws.write_string(f"A{r}", num, f)
            ws.write_number(f"B{r}", code, f) if isinstance(code, int) else ws.write_string(f"B{r}", code, f)
            ws.write_string(f"C{r}", it["name"], fl)
            ws.write_string(f"D{r}", it["unit"], f)
            h = _height((it["name"], 62 if bold else 72), (it["unit"], 14))
            if h > LINE:
                ws.set_row(r - 1, h)
            if it["kind"] == "work":
                ws.write(f"E{r}", it["e"], f)
                for col in "FGHI":
                    ws.write_blank(f"{col}{r}", None, f)
                work_row, work_qty = r, it["e"] or 0.0
            elif it["kind"] == "res":
                ws.write(f"E{r}", it["e"], f)
                qty = (it["e"] or 0.0) * work_qty if work_row else it["f"]
                if work_row and abs(qty - it["f"]) <= max(1e-9, abs(it["f"]) * 1e-3):
                    ws.write_formula(f"F{r}", f"=E{r}*E{work_row}", f, qty)
                else:  # lokalda F qo'lda yozilgan (me'yor × hajm emas) — lokaldagidek, smeta summasi o'zgarmasin
                    qty = it["f"]
                    ws.write(f"F{r}", qty, f)
                    if work_row:
                        notes.append({"part": b["title"], "row": it["row"], "name": it["name"], "kind": "note",
                                      "text": f"жами миқдор (F) локалда қўлда ёзилган: {qty:g}, меъёр × ҳажм = {(it['e'] or 0) * work_qty:g}"})
                ws.write_blank(f"H{r}" if cat != "N" else f"G{r}", None, f)
                priced(r, cat, it["price"], "F", qty, color)
            else:  # mustaqil resurs: hajm × narx
                ws.write(f"E{r}", it["e"], f)
                ws.write_blank(f"F{r}", None, f)
                priced(r, cat, it["price"], "E", it["e"] or 0.0, None)
                work_row = None
            r += 1
    last = max(first, r - 1)
    svodka(ws, fmt, r, first, last, sums, rates)
    ws.repeat_rows(8, 10)
    ws.freeze_panes(11, 0)
    ws.fit_to_pages(1, 0)
    ws.set_paper(9)
    wb.close()
    doc.dir.mkdir(parents=True, exist_ok=True)
    doc.file.write_bytes(buf.getvalue())
    return {"structure": structure, "sums": sums, "rows": rows, "notes": notes}


# ===================== yasash (fonda) =====================
def make(doc: MyOfyorta, progress=None):
    """Hamma kiritilgan lokallarni o'qib, ofyorta faylini yozadi; tuzilish, summalar va muammolar saqlanadi."""
    def step(text, done=0, total=0):
        doc.progress = {"step": text, "done": done, "total": total}
        MyOfyorta.objects.filter(pk=doc.pk).update(progress=doc.progress)
        if progress:
            progress()

    try:
        parts = [p for p in doc.parts if p.get("include", True)]
        if not parts:
            raise BuildError("Битта ҳам локал танланмаган")
        jobs = {str(j.pk): j for j in Job.objects.filter(pk__in=[p["job"] for p in parts]).select_related("batch")}
        blocks, problems, locals_ = [], [], []
        for i, p in enumerate(parts, 1):
            job = jobs.get(p["job"])
            if job is None:
                problems.append({"part": p["title"], "row": 0, "name": "", "text": "смета ўчирилган — блок олинмади"})
                continue
            step(f"Локал ўқилмоқда: {p['title']}", i - 1, len(parts))
            lc = read_local(job)
            locals_.append(lc)
            problems += [{"part": p["title"], **x} for x in lc["problems"]]
            blocks.append({"title": p["title"], "job": job.pk, "local_title": job.batch_label or p["title"], "items": lc["items"]})
        rates = doc.rates or default_rates(locals_)
        step("Excel ёзилмоқда", len(parts), len(parts))
        out = write(doc, blocks, rates, set(doc.customer or []))
        doc.rates = rates
        doc.structure = out["structure"]
        doc.rows = out["rows"]
        doc.totals = calc.chain(out["sums"], rates)
        doc.problems = (problems + out["notes"])[:500]
        doc.status, doc.error = MyOfyorta.Status.READY, ""
        doc.progress = {"step": "Тайёр", "done": len(parts), "total": len(parts)}
        doc.save()
    except Exception as e:  # noqa: BLE001 — sabab sahifada ko'rinadi
        doc.status = MyOfyorta.Status.FAILED
        doc.error = str(e) if isinstance(e, BuildError) else f"{type(e).__name__}: {e}"
        doc.save(update_fields=["status", "error"])
    return doc


# ===================== sozlash sahifasi (Excel ochilmaydi — narxlash natijasidan) =====================
def local_rates(job: Job) -> Dict[str, float]:
    p = resolve_percents(job.percents)
    return {key: round(float(p.get(src, 0.0)), 6) for key, src in RATES_FROM.items()}


def setup(project: Project, doc: Optional[MyOfyorta] = None) -> dict:
    """Sozlash sahifasi uchun: lokallar (blok nomi, kiritilsinmi), uskunalar (M ustun — buyurtmachiniki belgilanadi),
    svodka foizlari va lokallarda foizlar har xil bo'lsa — qaysilari qanday."""
    jobs = project_jobs(project)
    saved = {p["job"]: p for p in (doc.parts if doc else [])}
    order = {p["job"]: i for i, p in enumerate(doc.parts)} if doc else {}
    parts = []
    for j in jobs:
        p = saved.get(str(j.pk)) or {"title": default_title(j), "include": not doc}
        res = j.result or {}
        parts.append({"job": str(j.pk), "title": p["title"], "include": p.get("include", True), "local": j.local_name,
                      "folder": j.batch.folder or "", "label": j.batch_label,
                      "total": (res.get("totals") or {}).get("s_nds") or 0, "rates": local_rates(j)})
    parts.sort(key=lambda p: order.get(p["job"], len(order)))
    customer = set(doc.customer or []) if doc else set()
    eq = []
    for j in jobs:
        for r in (j.result or {}).get("rows", []):
            if r.get("category") != "M":
                continue
            g, q = r.get("g"), _num_value(r.get("qty")) or 0.0
            price = float(g) if isinstance(g, (int, float)) and not isinstance(g, bool) else 0.0
            key = f"{j.pk}:{r['row']}"
            eq.append({"key": key, "part": default_title(j), "name": r.get("name") or "", "unit": r.get("unit") or "",
                       "qty": q, "price": price, "amount": price * q, "customer": key in customer})
    sets = Counter(tuple(sorted(p["rates"].items())) for p in parts)
    rates = dict(sets.most_common(1)[0][0]) if sets else {}
    if doc and doc.rates:
        rates = {k: float(doc.rates.get(k, v)) for k, v in rates.items()}
    common = dict(sets.most_common(1)[0][0]) if sets else {}
    differ = [{"title": p["title"], "text": ", ".join(f"{RATE_LABELS[k]} {round(v * 100, 4):g}% (кўпчиликда {round(common[k] * 100, 4):g}%)"
                                                      for k, v in p["rates"].items() if v != common.get(k))}
              for p in parts if p["rates"] != common]
    for p in parts:
        p["odd"] = p["rates"] != common
    return {"parts": parts, "equipment": eq, "rates": rates, "differ": differ,
            "labels": [(k, RATE_LABELS[k], rates.get(k, 0.0)) for k in RATE_LABELS]}


def projects_for(user) -> List[Project]:
    """Ofyorta yasasa bo'ladigan loyihalar: kamida bitta narxlangan lokali bor."""
    ids = Job.objects.filter(batch__project__owner=user, status=Job.Status.DONE).exclude(output_path="") \
        .values_list("batch__project_id", flat=True).distinct()
    return list(Project.objects.filter(Q(pk__in=ids)).order_by("-created"))
